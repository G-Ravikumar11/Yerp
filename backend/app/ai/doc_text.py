"""Any document a vendor sends, as the words (or pictures) a model can read. No AI is involved here.

Every vendor has his own way of sending a quotation: a typed PDF, a scan, a Word file (new .docx or the old .doc), an
Excel sheet (.xlsx or the old .xls), a CSV, an OpenOffice file, a photo of a printed page. A PDF or a picture is passed on
as it is (llm_client reads a PDF's own text, and looks at a scan's pages); everything else is turned into plain text
here, tables kept as rows with the cells separated by tabs so the rates stay beside their items.

The file type is decided by what the bytes are, then the name, then what the browser said - a vendor's file often
arrives as "application/octet-stream" or with the wrong extension.
"""
import csv
import io
import logging
import re
import zipfile
from xml.etree import ElementTree

logger = logging.getLogger("app.ai")

MAX_CHARS = 30000

PICTURES = ("image/jpeg", "image/png", "image/webp", "image/gif")
# what is shown in the file chooser and named in messages
ACCEPT = (".pdf,.doc,.docx,.docm,.dotx,.xls,.xlsx,.xlsm,.csv,.txt,.rtf,.odt,.ods,"
          ".jpg,.jpeg,.png,.webp,.gif,.bmp,.tif,.tiff,.heic,.heif")


class Unreadable(ValueError):
    """The file is not one that can be read, said in words for the person who sent it."""


def kind_of(data: bytes, filename: str = "", content_type: str = "") -> str:
    """pdf | picture | docx | xlsx | odt | ods | doc | xls | rtf | text | heic | unknown"""
    name = (filename or "").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    head = data[:16]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head[:3] == b"\xff\xd8\xff" or head.startswith(b"\x89PNG") or head[:4] == b"GIF8" or (head[:4] == b"RIFF" and data[8:12] == b"WEBP"):
        return "picture"
    if head[:2] == b"BM" or head[:4] in (b"II*\x00", b"MM\x00*"):
        return "picture"                                           # BMP and TIFF are made into a JPEG
    if data[4:12] in (b"ftypheic", b"ftypheix", b"ftypmif1", b"ftypheif", b"ftyphevc"):
        return "heic"
    if head.startswith(b"PK"):                                     # a zip: the modern office formats
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = set(z.namelist())
                if "word/document.xml" in names:
                    return "docx"
                if "xl/workbook.xml" in names:
                    return "xlsx"
                if "content.xml" in names:
                    mime = z.read("mimetype").decode("ascii", "ignore") if "mimetype" in names else ""
                    return "ods" if "spreadsheet" in mime or ext == "ods" else "odt"
        except zipfile.BadZipFile:
            return "unknown"
        return "unknown"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):     # the old Office container
        if ext in ("xls", "xlt") or b"W\x00o\x00r\x00k\x00b\x00o\x00o\x00k" in data[:200000]:
            return "xls"
        return "doc"
    if head.startswith(b"{\\rtf"):
        return "rtf"
    if ext in ("csv", "txt", "tsv") or (content_type or "").startswith("text/"):
        return "text"
    # anything else that is mostly readable characters is read as text
    try:
        sample = data[:4000].decode("utf-8")
    except UnicodeDecodeError:
        return "unknown"
    return "text" if _mostly_printable(sample) else "unknown"


def _mostly_printable(text: str) -> bool:
    if not text:
        return False
    good = sum(1 for ch in text if ch.isprintable() or ch in "\r\n\t")
    return good / len(text) > 0.9


def to_text_or_picture(data: bytes, filename: str = "", content_type: str = ""):
    """(kind, text, picture) for a document: the words of an office file or a text file, or a picture to look at
    (bytes, media type). A PDF comes back as kind "pdf" with neither: the model reads it as it is."""
    kind = kind_of(data, filename, content_type)
    if kind == "pdf":
        return kind, "", None
    if kind == "picture":
        return kind, "", _as_picture(data, content_type)
    if kind == "heic":
        raise Unreadable("This is an iPhone photo (HEIC). Send it as a JPG: on the iPhone, Settings > Camera > Formats > Most Compatible, "
                         "or share the photo by WhatsApp, which sends a JPG.")
    reader = {"docx": _docx, "xlsx": _xlsx, "odt": _odf, "ods": _odf, "doc": _doc, "xls": _xls, "rtf": _rtf, "text": _text}.get(kind)
    if reader is None:
        raise Unreadable("That file could not be read. Send a PDF, a Word or Excel file, or a photo of the quotation.")
    try:
        text = reader(data)
    except Unreadable:
        raise
    except Exception as exc:
        logger.warning("Could not read a %s: %s", kind, exc)
        raise Unreadable("That %s file could not be opened. Save it again (or as a PDF) and send it once more." % kind.upper())
    text = re.sub(r"[ \t]+\n", "\n", re.sub(r"\n{3,}", "\n\n", text or "")).strip()
    if len(text) < 10:
        raise Unreadable("No words were found in that file. If it is a scan, send it as a PDF or a photo.")
    return kind, text[:MAX_CHARS], None


# --- the readers --------------------------------------------------------------------------------------------------

def _docx(data: bytes) -> str:
    """Paragraphs in order, and each table row as one line with its cells tab-separated."""
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        parts = ["word/document.xml"] + sorted(n for n in z.namelist() if re.match(r"word/(header|footer)\d*\.xml$", n))
        out = []
        for part in parts:
            root = ElementTree.fromstring(z.read(part))
            body = root.find(W + "body") if part == "word/document.xml" else root
            if body is None:
                continue
            for block in body:
                if block.tag == W + "p":
                    out.append("".join(t.text or "" for t in block.iter(W + "t")))
                elif block.tag == W + "tbl":
                    for row in block.iter(W + "tr"):
                        cells = ["".join(t.text or "" for t in cell.iter(W + "t")).strip() for cell in row.findall(W + "tc")]
                        out.append("\t".join(cells))
        return "\n".join(out)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        out = []
        for sheet in book.worksheets[:6]:
            out.append("## %s" % sheet.title)
            for row in sheet.iter_rows(values_only=True, max_col=40):
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    out.append("\t".join(cells).rstrip("\t"))
                if len(out) > 3000:
                    break
        return "\n".join(out)
    finally:
        book.close()


def _odf(data: bytes) -> str:
    """OpenOffice / LibreOffice: paragraphs in order, and each table row as one line of tab-separated cells."""
    T = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
    TB = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = ElementTree.fromstring(z.read("content.xml"))
    out = []

    def walk(el):
        if el.tag == TB + "table-row":                  # a row is one line; its paragraphs are not repeated below it
            cells = ["".join(c.itertext()).strip() for c in el.findall(TB + "table-cell")]
            if any(cells):
                out.append("\t".join(cells).rstrip("\t"))
            return
        if el.tag in (T + "p", T + "h"):
            out.append("".join(el.itertext()))
            return
        for child in el:
            walk(child)

    walk(root)
    return "\n".join(out)


def _doc(data: bytes) -> str:
    """An old Word file (.doc). There is no library for it here, so the text is taken from the file itself: Word keeps
    the document's characters either as plain bytes or as UTF-16, in long runs. Good enough for a model to read the
    items and rates; for anything odd the vendor can be asked for a PDF."""
    runs = []
    for m in re.finditer(rb"(?:[\x20-\x7e\r\n\t]\x00){8,}", data):        # UTF-16 text
        runs.append(m.group().decode("utf-16-le", "ignore"))
    if sum(len(r) for r in runs) < 200:
        for m in re.finditer(rb"[\x20-\x7e\r\n\t]{12,}", data):            # 8-bit text
            runs.append(m.group().decode("latin-1"))
    text = "\n".join(r.replace("\r", "\n").replace("\x07", "\t") for r in runs)
    # drop the font tables and style names Word stores beside the text
    keep = [line for line in text.splitlines() if line.strip() and not re.fullmatch(r"[A-Za-z ]{0,30}(Normal|Font|Times New Roman|Calibri|Arial|Heading \d|Table Grid|Default Paragraph Font|Microsoft Word.*)", line.strip())]
    return "\n".join(keep)


def _xls(data: bytes) -> str:
    try:
        import xlrd
    except ImportError:
        raise Unreadable("Old Excel files (.xls) cannot be read on this server yet. Save it as .xlsx or PDF and send it again.")
    book = xlrd.open_workbook(file_contents=data)
    out = []
    for sheet in book.sheets()[:6]:
        out.append("## %s" % sheet.name)
        for r in range(min(sheet.nrows, 3000)):
            cells = [str(v) if v not in ("", None) else "" for v in sheet.row_values(r)[:40]]
            if any(c.strip() for c in cells):
                out.append("\t".join(cells).rstrip("\t"))
    return "\n".join(out)


def _rtf(data: bytes) -> str:
    text = data.decode("latin-1", "ignore")
    text = re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes.fromhex(m.group(1)).decode("cp1252", "ignore"), text)
    text = re.sub(r"\\(par|line|row)\b ?", "\n", text)
    text = re.sub(r"\\cell\b ?", "\t", text)
    text = re.sub(r"\{\\\*[^{}]*\}", "", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    return re.sub(r"[{}]", "", text)


def _text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text.count(",") > text.count("\t") and "\n" in text:            # a CSV: cells tab-separated like the others
        try:
            return "\n".join("\t".join(row) for row in csv.reader(io.StringIO(text)))
        except csv.Error:
            pass
    return text


def _as_picture(data: bytes, content_type: str):
    """A picture the vision model takes: JPEG, PNG, WebP and GIF as they are; BMP or TIFF made into a JPEG."""
    media = (content_type or "").lower()
    if data[:3] == b"\xff\xd8\xff":
        return data, "image/jpeg"
    if data.startswith(b"\x89PNG"):
        return data, "image/png"
    if data[:4] == b"GIF8":
        return data, "image/gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return data, "image/webp"
    from PIL import Image
    img = Image.open(io.BytesIO(data))
    img.seek(0)
    out = io.BytesIO()
    img.convert("RGB").save(out, "JPEG", quality=88)
    return out.getvalue(), "image/jpeg"


def scan_pages(data: bytes, max_pages: int = 4):
    """A scanned PDF's first pages as JPEG pictures, for a model that reads pictures and not PDFs. [] when they cannot
    be drawn (pypdfium2 missing, or a broken file)."""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return []
    try:
        doc = pdfium.PdfDocument(data)
        pages = []
        for i in range(min(len(doc), max_pages)):
            page = doc[i]
            width = page.get_width() or 595
            image = page.render(scale=min(2.2, 1700 / width)).to_pil().convert("RGB")
            out = io.BytesIO()
            image.save(out, "JPEG", quality=80)
            pages.append(out.getvalue())
        return pages
    except Exception as exc:
        logger.warning("Could not draw the pages of a scanned PDF: %s", exc)
        return []
