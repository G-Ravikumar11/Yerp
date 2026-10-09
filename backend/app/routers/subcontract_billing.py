"""The subcontract billing endpoints."""
import io
import json
import os
import re
import types
import uuid
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import Response, StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.validators import import_guard
from app import models
from app.documents import sheet_forms
from app.db import get_db

from app.constants.common import HARDCOPY_TYPES, RETENTION_STAGES, SUB_TRANSITIONS
from app.core.audit import log_audit
from app.services.compliance import free_back_charges_of
from app.core.auth import (
    STAFF_VIEWER,
    get_client_user,
    require_erp_read,
    require_items_access,
    require_owner,
    wo_actor,
)
from app.core.currency import format_money_plain, inr, money, qty_text, sheet_number, unit_rate
from app.core.dates import _parse_date
from app.core.files import _file_or_404, drop_files_of, file_bytes, file_headers, served_type, store_file
from app.core.gst import WORKS_CONTRACT_SAC, our_state, split_gst, supply_state_for_job
from app.core.lifecycle import convert_legacy_holds
from app.core.notifications import notify, notify_employee
from app.core.queries import prime, prime_chains
from app.core.sheets import sheet_response
from app.documents.forms import form_pdf_response, sub_bill_certificate, sub_bill_form_spec
from app.documents.pdf_twins import SHEET_AS_PDF
from app.schemas.subcontract_billing import (
    DimensionIn,
    MeasurementIn,
    RetentionReleaseIn,
    SubBillEditIn,
    SubBillIn,
    SubMeasureBatchIn,
    SubMeasurementIn,
)
from app.services.approvals import owner_label
from app.services.client_billing import (
    billed_qty_to_date,
    dimensions_for,
    measured_to_date,
    refuse_cancel_with_money,
    refuse_unapproved_order,
)
from app.services.crm import settled_amounts, settled_on
from app.services.invoicing import refuse_cancel_with_irn
from app.services.subcontract_billing import (
    _is_owner,
    _live_releases,
    add_sub_measurement,
    apply_material_recovery,
    bill_scan_required,
    dimension_total,
    draw_sub_bill_lines,
    material_recovery_dict,
    mb_item_label,
    mb_match_item,
    past_tense,
    recost_sub_bill,
    release_dict,
    release_material_recovery,
    release_or_404,
    retention_positions,
    sub_advance_recovery,
    sub_bill_chain_rows,
    sub_bill_current_step,
    sub_bill_decide,
    sub_bill_dict,
    sub_bill_or_404,
    sub_bill_start_chain,
    sub_bill_step_name,
    sub_bill_visible_ids,
    sub_billed_to_date,
    sub_claimable_lines,
    sub_gross_measured,
    sub_held_to_date,
    sub_hold_remaining,
    sub_measure_multiplier,
    sub_measured_to_date,
    write_dimensions,
)
from app.services.subcontract_orders import (
    cascade_delete_referrers,
    contractor_state,
    drop_alerts_about,
    item_ceiling,
    run_bulk_delete,
    wo_dict,
    wo_or_404,
    work_order_or_404,
    work_order_to_dict,
)


router = APIRouter()


@router.get("/api/sub-bills/{bill_id}/delete-preview")
def sub_bill_delete_preview(bill_id: int, request: Request, db: Session = Depends(get_db)):
    client = get_client_user(request, db)
    require_owner(request, db)
    bill = sub_bill_or_404(db, client.id, bill_id)
    paid = settled_on(db, client.id, "sub_bill", bill.id)
    warnings = []
    if paid > 0 or (bill.status or "") == "PAID":
        warnings.append("%s is paid%s - the payments are deleted too." % (bill.number, " (%s)" % inr(paid) if paid > 0 else ""))
    return {"numbers": [bill.number], "blockers": [], "warnings": warnings, "can_delete": True, "counts": {
        "lines": db.query(models.DBSubBillLine).filter(models.DBSubBillLine.sub_bill_id == bill.id).count(),
        "measurements": db.query(models.DBSubMeasurement).filter(models.DBSubMeasurement.sub_bill_id == bill.id).count()}}


@router.delete("/api/sub-bills/{bill_id}")
def sub_bill_delete(bill_id: int, request: Request, db: Session = Depends(get_db)):
    """Takes the bill, its approval route and its payments away. What it measured goes back to waiting
    for the next bill - the work was done; only the claim for it is removed."""
    client = get_client_user(request, db)
    require_owner(request, db)
    bill = sub_bill_or_404(db, client.id, bill_id)
    number = bill.number
    drop = lambda q: q.delete(synchronize_session=False)
    db.query(models.DBSubMeasurement).filter(models.DBSubMeasurement.sub_bill_id == bill.id).update(
        {"sub_bill_id": None}, synchronize_session=False)
    release_material_recovery(db, client.id, bill.id)
    free_back_charges_of(db, bill.id)
    drop_files_of(db, "sub_bill_scan", [bill.id])
    drop(db.query(models.DBSubBillLine).filter(models.DBSubBillLine.sub_bill_id == bill.id))
    drop(db.query(models.DBApprovalChain).filter(
        models.DBApprovalChain.entity_type == "sub_bill", models.DBApprovalChain.entity_id == bill.id))
    drop(db.query(models.DBMoneyEntry).filter(
        models.DBMoneyEntry.doc_type == "sub_bill", models.DBMoneyEntry.doc_id == bill.id))
    drop_alerts_about(db, "sub_bill", [bill.id])
    cascade_delete_referrers(db, "sub_bills", [bill.id])
    drop(db.query(models.DBSubBill).filter(models.DBSubBill.id == bill.id))
    log_audit(db, client.id, "sub_bill_deleted", "sub_bill", bill_id, number, "Deleted by the Master", request)
    db.commit()
    return {"ok": True, "message": "%s deleted. What it measured is free to be billed again." % number}


@router.get("/api/mb/{work_order_id}")
def measurement_book(work_order_id: int, request: Request,
                     db: Session = Depends(get_db)):
    """Every ordered line, with what has been measured and what is left.

    The balance is the number somebody on site actually wants: how much of
    this item is still to do.
    """
    client = require_erp_read(request, db)
    wo = work_order_or_404(db, client.id, work_order_id)
    measured = measured_to_date(db, wo.id)
    billed = billed_qty_to_date(db, wo.id)

    lines = []
    for l in db.query(models.DBWorkOrderLine).filter(
            models.DBWorkOrderLine.work_order_id == wo.id).all():
        done = money(measured.get(l.id, 0.0))
        claimed = money(billed.get(l.id, 0.0))
        lines.append({
            "line_id": l.id, "fg_code": l.fg_code or "",
            "description": (l.description or l.item_name or "").split("\n")[0],
            "uom": l.uom or "", "ordered_qty": money(l.qty),
            "rate": unit_rate(l.rate),
            "measured_to_date": done, "billed_to_date": claimed,
            "unbilled": money(done - claimed),
            "balance_to_measure": money(money(l.qty) - done),
            "percent_measured": round(done / l.qty * 100, 1) if l.qty else 0.0,
            "over_measured": money(done - money(l.qty)) if done > money(l.qty) else 0.0,
        })

    rows = db.query(models.DBMeasurement).filter(
        models.DBMeasurement.work_order_id == wo.id).order_by(
            models.DBMeasurement.id.desc()).limit(400).all()
    dims = dimensions_for(db, [m.id for m in rows],
                          models.DBMeasurementDimension.measurement_id)
    entries = [{
        "id": m.id, "line_id": m.line_id, "fg_code": m.fg_code or "",
        "measured_on": m.measured_on or "", "quantity": money(m.quantity),
        "mb_ref": m.mb_ref or "", "location": getattr(m, "location", "") or "", "remarks": m.remarks or "",
        "recorded_by_name": m.recorded_by_name or "",
        "witnessed_by": m.witnessed_by or "",
        "billed": bool(m.ra_bill_id), "created_at": m.created_at or "",
        "dimensions": dims.get(m.id, []),
    } for m in rows]

    return {
        "work_order": work_order_to_dict(db, wo),
        "lines": lines, "entries": entries,
        "summary": {
            "ordered_value": money(wo.total_value),
            "measured_value": money(sum(l["measured_to_date"] * l["rate"] for l in lines)),
            "unbilled_value": money(sum(l["unbilled"] * l["rate"] for l in lines)),
            "lines_over_measured": len([l for l in lines if l["over_measured"] > 0]),
        },
    }


@router.post("/api/mb/{work_order_id}/entries")
def record_measurement(work_order_id: int, body: MeasurementIn, request: Request,
                       db: Session = Depends(get_db)):
    """Write one entry into the book."""
    client, actor_id, actor_name = wo_actor(request, db, "site.record")
    wo = work_order_or_404(db, client.id, work_order_id)
    if (wo.status or "") == "Draft":
        raise HTTPException(
            409, "Nothing is measured against an order that has not been placed.")
    refuse_unapproved_order(wo, "measured")

    line = db.query(models.DBWorkOrderLine).filter(
        models.DBWorkOrderLine.id == body.line_id,
        models.DBWorkOrderLine.work_order_id == wo.id).first()
    if not line:
        raise HTTPException(404, "That line is not on this order")
    # Dimensions, where given, are the measurement; the total is their sum.
    # A bare total is still accepted - a count of manholes has no dimensions.
    quantity = money(body.quantity or 0)
    if body.dimensions:
        quantity = money(dimension_total(body.dimensions))
    if not quantity:
        raise HTTPException(400, "A measurement of nothing is not a measurement")

    entry = models.DBMeasurement(
        client_id=client.id, work_order_id=wo.id, line_id=line.id,
        fg_code=line.fg_code or "", quantity=quantity,
        measured_on=(body.measured_on or datetime.now().strftime("%Y-%m-%d")),
        mb_ref=(body.mb_ref or "").strip(), location=(body.location or "").strip()[:200],
        remarks=(body.remarks or "").strip(),
        witnessed_by=(body.witnessed_by or "").strip(),
        recorded_by=actor_id, recorded_by_name=actor_name)
    db.add(entry)
    db.flush()
    write_dimensions(db, client.id, body.dimensions, measurement_id=entry.id)
    log_audit(db, client.id, "measurement_recorded", "work_order", wo.id, wo.number,
              "%s %s %s" % (line.fg_code, quantity, line.uom or ""), request)
    db.commit()

    measured = money(measured_to_date(db, wo.id).get(line.id, 0.0))
    ordered = money(line.qty)
    return {
        "ok": True, "measured_to_date": measured,
        "balance_to_measure": money(ordered - measured),
        # Not refused: site conditions genuinely differ from the schedule, and
        # a variation is agreed afterwards. Said plainly so nobody bills past
        # the order without knowing they have.
        "over_measured": money(measured - ordered) if measured > ordered else 0.0,
        "message": ("Recorded. %s measured against %s of %s ordered."
                    % (qty_text(measured), line.fg_code, qty_text(ordered))),
    }


@router.delete("/api/mb/entries/{entry_id}")
def delete_measurement(entry_id: int, request: Request, db: Session = Depends(get_db)):
    """Only while it is unbilled. Once claimed it is part of a bill's history."""
    client, actor_id, actor_name = wo_actor(request, db, "site.record")
    entry = db.query(models.DBMeasurement).filter(
        models.DBMeasurement.id == entry_id,
        models.DBMeasurement.client_id == client.id).first()
    if not entry:
        raise HTTPException(404, "Entry not found")
    if entry.ra_bill_id:
        raise HTTPException(
            409, "This measurement has been billed. Record a correcting entry "
                 "instead - a book that can be rubbed out is not a record.")
    db.query(models.DBMeasurementDimension).filter(
        models.DBMeasurementDimension.measurement_id == entry.id).delete()
    db.delete(entry)
    db.commit()
    return {"ok": True, "message": "Entry removed."}


@router.get("/api/sub-mb/template.xlsx")
def sub_mb_template(request: Request, db: Session = Depends(get_db)):
    """A worked example of the measurement book layout the import reads."""
    require_erp_read(request, db)
    if not sheet_forms.XLSX_AVAILABLE:
        raise HTTPException(503, "The workbook library is not installed on this server.")
    data = sheet_forms.build_mb_template()
    return StreamingResponse(io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="measurement_book_template.xlsx"'})


@router.get("/api/sub-mb/{order_id}")
def sub_measurement_book(order_id: int, request: Request, db: Session = Depends(get_db)):
    client = require_erp_read(request, db)
    order = wo_or_404(db, client.id, order_id)
    held_now = sub_held_to_date(db, order.id)
    measured = {k: v + held_now.get(k, 0.0) for k, v in sub_measured_to_date(db, order.id).items()}
    billed = sub_billed_to_date(db, order.id)
    lines = []
    for it in db.query(models.DBSubcontractItem).filter(
            models.DBSubcontractItem.order_id == order.id).order_by(
                models.DBSubcontractItem.display_order, models.DBSubcontractItem.id).all():
        if it.is_header:
            lines.append({"item_id": it.id, "activity_no": it.activity_no or "",
                          "description": (it.item_description or "").split("\n")[0],
                          "is_header": True})
            continue
        done = money(measured.get(it.id, 0.0))
        claimed = money(billed.get(it.id, 0.0))
        held = money(held_now.get(it.id, 0.0))
        ordered = money(it.quantity)
        lines.append({
            "item_id": it.id, "activity_no": it.activity_no or "",
            "item_code": it.item_code or "",
            "description": (it.item_description or "").split("\n")[0],
            "uom": it.uom or "", "ordered_qty": ordered, "rate": unit_rate(it.unit_rate),
            "tolerance_percent": it.tolerance_percent or 0,
            "max_quantity": item_ceiling(it),
            "measured_to_date": done, "billed_to_date": claimed, "held": held,
            "unbilled": money(done - held - claimed),
            "balance_to_measure": money(ordered - done),
            "percent_measured": round(done / ordered * 100, 1) if ordered else 0.0,
            "over_measured": money(done - ordered) if done > ordered else 0.0,
        })
    # The newest entries first. A book imported from Excel runs to hundreds, so the cap is high, and the page
    # is told when there are more than it was sent rather than quietly showing part of the book.
    entries_total = db.query(models.DBSubMeasurement).filter(models.DBSubMeasurement.order_id == order.id).count()
    sub_rows = db.query(models.DBSubMeasurement).filter(
        models.DBSubMeasurement.order_id == order.id).order_by(
            models.DBSubMeasurement.id.desc()).limit(3000).all()
    sub_dims = dimensions_for(db, [m.id for m in sub_rows],
                              models.DBMeasurementDimension.sub_measurement_id)
    entries = [{
        "id": m.id, "code": m.code or "", "item_id": m.item_id, "activity_no": m.activity_no or "",
        "measured_on": m.measured_on or "", "quantity": money(m.quantity),
        "multiplier": getattr(m, "multiplier", None) or 1,
        "mb_ref": m.mb_ref or "", "location": getattr(m, "location", "") or "", "remarks": m.remarks or "",
        "recorded_by_name": m.recorded_by_name or "", "billed": bool(m.sub_bill_id),
        "section": getattr(m, "section", "") or "", "block_label": getattr(m, "block_label", "") or "",
        "group_ref": getattr(m, "group_ref", "") or "",
        "kind": getattr(m, "kind", "") or "", "hold_of": getattr(m, "hold_of", None),
        "held_remaining": sub_hold_remaining(db, m) if getattr(m, "kind", "") == "hold" else 0.0,
        "dimensions": sub_dims.get(m.id, []),
    } for m in sub_rows]
    return {
        "order": wo_dict(db, order), "lines": lines, "entries": entries, "entries_total": entries_total,
        "summary": {
            "ordered_value": money(order.gross_amount),
            "measured_value": money(sum(l["measured_to_date"] * l["rate"]
                                        for l in lines if not l.get("is_header"))),
            "unbilled_value": money(sum(l["unbilled"] * l["rate"]
                                        for l in lines if not l.get("is_header"))),
            "held_value": money(sum(l["held"] * l["rate"] for l in lines if not l.get("is_header"))),
            "lines_over_measured": len([l for l in lines
                                        if not l.get("is_header") and l["over_measured"] > 0]),
        },
    }


@router.post("/api/sub-mb/{order_id}/entries")
def record_sub_measurement(order_id: int, body: SubMeasurementIn, request: Request,
                           db: Session = Depends(get_db)):
    client, actor_id, actor_name = wo_actor(request, db, "site.record")
    order = wo_or_404(db, client.id, order_id)
    if (order.status or "") not in ("APPROVED", "EXECUTED"):
        raise HTTPException(
            409, "Nothing is measured against an order that has not been approved.")
    item = db.query(models.DBSubcontractItem).filter(
        models.DBSubcontractItem.id == body.item_id,
        models.DBSubcontractItem.order_id == order.id).first()
    if not item:
        raise HTTPException(404, "That item is not on this order")
    entry = add_sub_measurement(db, client, order, item, body, actor_id, actor_name)
    quantity = entry.quantity
    log_audit(db, client.id, "sub_measurement_recorded", "subcontract_order", order.id,
              order.wo_number or "", "%s %s %s" % (item.activity_no, quantity,
                                                   item.uom or ""), request)
    db.commit()
    measured = money(sub_measured_to_date(db, order.id).get(item.id, 0.0))
    ordered = money(item.quantity)
    return {"ok": True, "measured_to_date": measured,
            "balance_to_measure": money(ordered - measured),
            "over_measured": money(measured - ordered) if measured > ordered else 0.0,
            "message": "Recorded. %s measured against %s of %s ordered."
                       % (qty_text(measured), ("activity " + item.activity_no) if item.activity_no
                          else "the item", qty_text(ordered))}


@router.post("/api/sub-mb/{order_id}/import")
async def import_sub_measurement_book(order_id: int, request: Request, file: UploadFile = File(...),
                                      commit: str = Form("0"), mapping: str = Form(""),
                                      sheet: str = Form(""), measured_on: str = Form(""), include_dims: str = Form("0"),
                                      entries: str = Form(""), allow_duplicates: str = Form("0"), same_item_ok: str = Form("0"),
                                      db: Session = Depends(get_db)):
    """The measurement book as the site keeps it in Excel - S.No, Description,
    UoM, No's, NoM, Length, Width, Height, Total Quantity - read into the
    gang's book. Each section of the sheet is matched to an item on the order
    by its description; the first call shows what would go in, and a second
    with commit=1 (and any corrected matches) records it. All of it goes in
    or none of it does."""
    client, actor_id, actor_name = wo_actor(request, db, "site.record")
    order = wo_or_404(db, client.id, order_id)
    if (order.status or "") not in ("APPROVED", "EXECUTED"):
        raise HTTPException(409, "Nothing is measured against an order that has not been approved.")
    raw = await file.read()
    try:
        import_guard.check_workbook_bytes(raw, file.filename)
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    try:
        import openpyxl
        values = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, keep_links=False)
        formulas = openpyxl.load_workbook(io.BytesIO(raw), keep_links=False)
    except Exception:
        raise HTTPException(400, "That file could not be read as an Excel workbook. Open it in Excel and save it again as .xlsx.")
    if not values.sheetnames:
        raise HTTPException(400, "That workbook has no sheets.")
    names = values.sheetnames
    pick = sheet if sheet in names else next((n for n in names if n.strip().upper().startswith("MB")), None)
    tried = [pick] if pick else names
    book, errors = None, []
    for n in tried:
        try:
            book = sheet_forms.read_measurement_book(values[n], formulas[n])
            if book["items"]:
                break
        except ValueError as exc:
            errors.append(str(exc))
            book = None
    if not book or not book["items"]:
        raise HTTPException(400, "No measurement book was found in that workbook. " + " ".join(errors[:2]))
    items = [it for it in db.query(models.DBSubcontractItem).filter(
        models.DBSubcontractItem.order_id == order.id).order_by(
            models.DBSubcontractItem.display_order, models.DBSubcontractItem.id).all() if not it.is_header]
    by_id = {it.id: it for it in items}
    try:
        chosen = {int(k): int(v) for k, v in (json.loads(mapping) if mapping else {}).items() if v}
    except (ValueError, TypeError, AttributeError):
        raise HTTPException(400, "The item matches could not be read.")
    # Only some entries of the sheet, when the person ticked them: [[section, entry], ...].
    picked = None
    if entries:
        try:
            picked = {(int(a), int(b)) for a, b in json.loads(entries)}
        except (ValueError, TypeError):
            raise HTTPException(400, "The entries chosen could not be read.")
    # What is already in the book, by item and where it was measured: a block imported twice is billed twice.
    # (The same block measured for another item - flooring, then painting - is not a repeat.)
    plain_place = lambda t: re.sub(r"[^a-z0-9]+", "", (t or "").lower())
    in_book = {}
    for m in db.query(models.DBSubMeasurement).filter(models.DBSubMeasurement.order_id == order.id).all():
        if plain_place(m.location):
            key = (m.item_id, plain_place(m.location))
            in_book[key] = in_book.get(key, 0.0) + (m.quantity or 0.0)
    book_of = lambda item_id, place: (in_book.get((item_id, place)) if item_id else next((q for (i, p), q in in_book.items() if p == place), None))
    # A sheet from another contractor's bill is not stopped, but it is said.
    con = db.query(models.DBContractor).filter(models.DBContractor.id == order.contractor_id).first() if order.contractor_id else None
    sheet_con = (book["meta"].get("contractor") or "").strip()
    if con and sheet_con:
        words = lambda t: {w for w in re.sub(r"[^a-z0-9]+", " ", (t or "").lower()).split() if len(w) >= 3 and w not in ("m/s", "the", "and", "pvt", "ltd", "private", "limited")}
        if not (words(sheet_con) & words(con.company_name)):
            book["warnings"].append("The sheet is for %s, but this order is with %s." % (sheet_con, con.company_name))
    sections, missing = [], []
    # The measure window reads a sheet to fill its own lines, so it asks for them to come back.
    want_dims = (include_dims or "0") in ("1", "true", "yes")
    for i, sec in enumerate(book["items"]):
        item = by_id.get(chosen[i]) if i in chosen else mb_match_item(items, sec["description"], sec.get("sno") or "")
        if item is None:
            missing.append((i, sec["description"]))
        sections.append({"index": i, "description": sec["description"], "sno": sec["sno"],
                         "item_id": item.id if item else None,
                         "item": mb_item_label(item) if item else "",
                         "uom": item.uom if item else "", "quantity": sec["quantity"],
                         "held": sec.get("held", 0.0), "payable": sec.get("payable", sec["quantity"]),
                         "rate": unit_rate(item.unit_rate) if item else 0.0,
                         "holds": [dict(h, held=round(sum(e["held_back"] for e in sec["entries"] if e.get("group") == h["group"]), 3))
                                   for h in book.get("holds", []) if any(e.get("group") == h["group"] for e in sec["entries"])],
                         "entries": [{"location": e["location"], "multiplier": e["multiplier"],
                                      "lines": len([d for d in e["dims"] if not d["is_heading"]]),
                                      "one_block": e["one"], "quantity": e["quantity"],
                                      "stated": e.get("stated_total"), "group": e.get("group"), "letter": e.get("letter") or "",
                                      "held_back": e.get("held_back", 0.0), "payable": e.get("payable", e["quantity"]),
                                      "full_quantity": e["quantity"],
                                      "already_in_book": bool(plain_place(e["location"]) and book_of(item.id if item else None, plain_place(e["location"])) is not None),
                                      "already_quantity": money(book_of(item.id if item else None, plain_place(e["location"]))) if plain_place(e["location"]) and book_of(item.id if item else None, plain_place(e["location"])) is not None else None,
                                      **({"dims": e["dims"]} if want_dims else {})}
                                     for e in sec["entries"]]})
    # Different works in the sheet put onto one item of the order are named, and not recorded until the person
    # says they are one item: three works at one rate is how a bill comes out at three crore.
    squash = lambda t: re.sub(r"[^a-z0-9]+", "", (t or "").lower())
    wanted = {a for a, _ in picked} if picked is not None else None
    by_target = {}
    for row in sections:
        if row["item_id"] and (wanted is None or row["index"] in wanted):
            by_target.setdefault(row["item_id"], []).append(row)
    conflicts = [{"item_id": iid, "item": rows[0]["item"],
                  "sections": [{"index": r["index"], "description": r["description"], "quantity": r["quantity"]} for r in rows]}
                 for iid, rows in by_target.items() if len({squash(r["description"]) for r in rows}) > 1]
    preview = {"sheet": book["sheet"], "meta": book["meta"], "sections": sections,
               "warnings": book["warnings"], "conflicts": conflicts,
               "items": [{"id": it.id, "label": mb_item_label(it), "uom": it.uom or "", "rate": unit_rate(it.unit_rate)} for it in items]}
    if (commit or "0") not in ("1", "true", "yes"):
        return dict(preview, ok=True, committed=False)
    needed = [d for i, d in missing if picked is None or i in {a for a, _ in picked}]
    if needed:
        raise HTTPException(400, "Say which item on the order these are for: " + "; ".join(needed[:5]))
    if conflicts and (same_item_ok or "0") not in ("1", "true", "yes"):
        c = conflicts[0]
        raise HTTPException(409, "%s different works in the sheet are all set to %s: %s. Put each on its own item, or confirm "
                                 "they are all this one item." % (len(c["sections"]), c["item"] or "one item",
                                                                   "; ".join(x["description"] for x in c["sections"][:4])))
    when = (measured_on or book["meta"].get("date") or datetime.now().strftime("%Y-%m-%d"))[:10]
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", when):
        when = datetime.now().strftime("%Y-%m-%d")
    written = 0
    skipped = []
    source = os.path.basename(file.filename or "workbook")
    # Blocks that share one hold are tied together by a reference no other import can repeat.
    batch = uuid.uuid4().hex[:6]
    held_for = {}
    for si, (sec, row) in enumerate(zip(book["items"], sections)):
        if picked is not None and not any(a == si for a, _ in picked):
            continue
        item = by_id[row["item_id"]]
        for ei, e in enumerate(sec["entries"]):
            if picked is not None and (si, ei) not in picked:
                continue
            # Whole-book imports leave out a block that is already in the book; ticking one is a decision.
            if picked is None and (allow_duplicates or "0") not in ("1", "true", "yes") and plain_place(e["location"])                     and (item.id, plain_place(e["location"])) in in_book:
                skipped.append(e["location"] or sec["description"])
                continue
            dims = [DimensionIn(particulars=d["particulars"][:200], is_heading=d["is_heading"],
                                nos=d.get("nos"), nom=d.get("nom"), length=d.get("length"),
                                breadth=d.get("breadth"), depth=d.get("depth"), deduct=d.get("deduct", False))
                    for d in e["dims"]]
            body = SubMeasurementIn(item_id=item.id, dimensions=dims, multiplier=e["multiplier"],
                                    measured_on=when, mb_ref=("%s / %s" % (source, book["sheet"]))[:120],
                                    location=e["location"], remarks="Imported from the measurement book",
                                    section=("%s %s" % (sec["sno"], sec["description"])).strip(), block_label=e.get("letter") or "",
                                    group_ref=("%s-%d-%d" % (batch, si, e["group"])) if e.get("group") else "")
            try:
                # Each entry is flushed as it is written, so the ceiling check
                # on the next one already counts it.
                add_sub_measurement(db, client, order, item, body, actor_id, actor_name)
            except HTTPException as exc:
                db.rollback()
                raise HTTPException(exc.status_code, "%s - nothing was imported. %s" % (
                    e["location"] or sec["description"], exc.detail))
            written += 1
            if e.get("held_back"):
                key = (si, e["group"])
                held_for.setdefault(key, {"item": item, "sec": sec, "held": 0.0})["held"] += e["held_back"]
    # What the sheet holds back is recorded once per group of blocks, as a hold beside the measurement - the
    # measurement itself stays as the sheet wrote it - and the owner releases it onto a later bill.
    reasons = {h["group"]: h["reason"] for h in book.get("holds", [])}
    for (si, g), h in held_for.items():
        qty = money(h["held"])
        if qty <= 0:
            continue
        db.add(models.DBSubMeasurement(
            client_id=client.id, order_id=order.id, item_id=h["item"].id, activity_no=h["item"].activity_no or "",
            measured_on=when, quantity=-qty, kind="hold", mb_ref=("%s / %s" % (source, book["sheet"]))[:120],
            location="Held back", remarks=reasons.get(g, "Held back for finishes and handing over")[:300],
            recorded_by=actor_id, recorded_by_name=actor_name,
            section=("%s %s" % (h["sec"]["sno"], h["sec"]["description"])).strip(), group_ref="%s-%d-%d" % (batch, si, g)))
    db.flush()
    log_audit(db, client.id, "sub_mb_imported", "subcontract_order", order.id, order.wo_number or "",
              "%s: %d entries" % (source, written), request)
    db.commit()
    note = (" %d already in the book %s left out: %s." % (len(skipped), "was" if len(skipped) == 1 else "were", "; ".join(skipped[:3]))) if skipped else ""
    return dict(preview, ok=True, committed=True, entries=written, skipped=skipped,
                message="%d measurement%s recorded from %s.%s" % (written, "" if written == 1 else "s", book["sheet"], note))


@router.delete("/api/sub-mb/entries/{entry_id}")
def delete_sub_measurement(entry_id: int, request: Request, db: Session = Depends(get_db)):
    client, _, _ = wo_actor(request, db, "site.record")
    entry = db.query(models.DBSubMeasurement).filter(
        models.DBSubMeasurement.id == entry_id,
        models.DBSubMeasurement.client_id == client.id).first()
    if not entry:
        raise HTTPException(404, "Entry not found")
    wo_or_404(db, client.id, entry.order_id)        # an order they cannot see is an entry they cannot see
    if (entry.kind or "") == "hold" and sub_hold_remaining(db, entry) < money(-(entry.quantity or 0.0)):
        raise HTTPException(409, "Part of this hold has been released. Delete the release first.")
    if not (entry.kind or ""):
        held = sub_held_to_date(db, entry.order_id).get(entry.item_id, 0.0)
        gross = sub_gross_measured(db, entry.order_id).get(entry.item_id, 0.0)
        if held > 0 and (entry.quantity or 0.0) > 0 and gross - (entry.quantity or 0.0) < held - 0.0001:
            raise HTTPException(409, "%s of this item is held back. Release or delete the hold before taking this "
                                     "measurement away." % qty_text(held))
    bill = None
    if entry.sub_bill_id:
        # The owner may take a measurement off a bill that is still a draft; a sent bill has to be deleted first.
        bill = db.query(models.DBSubBill).filter(models.DBSubBill.id == entry.sub_bill_id).first()
        if not (bill and (bill.status or "") == "DRAFT" and _is_owner(request, db)):
            raise HTTPException(409, "This measurement is on %s, which has been %s. %s" % (
                bill.number if bill else "a bill", ((bill.status if bill else "") or "sent").lower(),
                "Delete that bill first." if _is_owner(request, db) else "Ask the Master, or record a correcting entry instead."))
    db.query(models.DBMeasurementDimension).filter(
        models.DBMeasurementDimension.sub_measurement_id == entry.id).delete()
    drop_files_of(db, "sub_measurement", [entry.id])
    db.delete(entry)
    db.flush()
    if bill is not None:
        order = db.query(models.DBSubcontractOrder).filter(models.DBSubcontractOrder.id == bill.order_id).first()
        draw_sub_bill_lines(db, bill, order)
        recost_sub_bill(db, bill)
    db.commit()
    return {"ok": True, "message": "Entry removed."}


@router.put("/api/sub-mb/entries/{entry_id}")
def update_sub_measurement(entry_id: int, body: SubMeasurementIn, request: Request, db: Session = Depends(get_db)):
    """Change an entry's calculation after it was imported or typed: its lines, blocks, date, place or
    total. The quantity is worked out again and held to the same ceiling as a new entry. An entry on a bill
    that has been sent is not changed under it; one on a draft bill is, and the bill is drawn up again."""
    client, actor_id, actor_name = wo_actor(request, db, "site.record")
    entry = db.query(models.DBSubMeasurement).filter(
        models.DBSubMeasurement.id == entry_id, models.DBSubMeasurement.client_id == client.id).first()
    if not entry:
        raise HTTPException(404, "Entry not found")
    if (entry.kind or ""):
        raise HTTPException(409, "A hold is placed or released from the Hold button, not edited.")
    order = wo_or_404(db, client.id, entry.order_id)
    item = db.query(models.DBSubcontractItem).filter(models.DBSubcontractItem.id == entry.item_id).first()
    bill = None
    if entry.sub_bill_id:
        bill = db.query(models.DBSubBill).filter(models.DBSubBill.id == entry.sub_bill_id).first()
        if not (bill and (bill.status or "") == "DRAFT"):
            raise HTTPException(409, "This measurement is on %s, which has been %s. Delete that bill to change it." % (
                bill.number if bill else "a bill", ((bill.status if bill else "") or "sent").lower()))
    if (order.status or "") not in ("APPROVED", "EXECUTED"):
        raise HTTPException(409, "Nothing is measured against an order that has not been approved.")
    multiplier = sub_measure_multiplier(getattr(body, "multiplier", 1))
    quantity = money(body.quantity or 0)
    if body.dimensions:
        quantity = money(dimension_total(body.dimensions))
    quantity = money(quantity * multiplier)
    if not quantity:
        raise HTTPException(400, "A measurement of nothing is not a measurement")
    if quantity > 0 and item is not None:
        done = money(sub_gross_measured(db, order.id).get(item.id, 0.0) - (entry.quantity or 0.0))
        ceiling = item_ceiling(item)
        if money(done + quantity) > ceiling + 0.0001:
            raise HTTPException(409, "%s: %s already measured elsewhere; %s would make %s against %s ordered%s. Amend the order to measure beyond it." % (
                item.activity_no or "Item", done, quantity, money(done + quantity), money(item.quantity),
                " (+%g%% tolerance = %s)" % (item.tolerance_percent, ceiling) if item.tolerance_percent else ""))
    held_now = sub_held_to_date(db, order.id).get(item.id, 0.0) if item is not None else 0.0
    if held_now > 0:
        gross_after = sub_gross_measured(db, order.id).get(item.id, 0.0) - (entry.quantity or 0.0) + quantity
        if gross_after < held_now - 0.0001:
            raise HTTPException(409, "%s of this item is held back, so its measurement cannot come below that. "
                                     "Release or delete the hold first." % qty_text(held_now))
    entry.quantity, entry.multiplier = quantity, multiplier
    entry.measured_on = body.measured_on or entry.measured_on
    entry.mb_ref = (body.mb_ref or "").strip()
    entry.location = (body.location or "").strip()[:200]
    entry.remarks = (body.remarks or "").strip()
    db.query(models.DBMeasurementDimension).filter(
        models.DBMeasurementDimension.sub_measurement_id == entry.id).delete(synchronize_session=False)
    write_dimensions(db, client.id, body.dimensions, sub_measurement_id=entry.id)
    db.flush()
    if bill is not None:
        draw_sub_bill_lines(db, bill, order)
        recost_sub_bill(db, bill)
    log_audit(db, client.id, "sub_measurement_changed", "subcontract_order", order.id, order.wo_number or "",
              "%s %s" % (entry.activity_no or "", quantity), request)
    db.commit()
    return {"ok": True, "quantity": quantity, "message": "Entry updated."}


@router.post("/api/sub-mb/{order_id}/entries/batch")
def record_sub_measurements_batch(order_id: int, body: SubMeasureBatchIn, request: Request, db: Session = Depends(get_db)):
    """Several blocks checked together in the grid, recorded as one entry each - and what the sheet holds back
    on them, recorded once per group as a hold. All of it goes in, or none of it does."""
    client, actor_id, actor_name = wo_actor(request, db, "site.record")
    order = wo_or_404(db, client.id, order_id)
    if (order.status or "") not in ("APPROVED", "EXECUTED"):
        raise HTTPException(409, "Nothing is measured against an order that has not been approved.")
    item = db.query(models.DBSubcontractItem).filter(
        models.DBSubcontractItem.id == body.item_id, models.DBSubcontractItem.order_id == order.id).first()
    if not item or item.is_header:
        raise HTTPException(404, "That item is not on this order")
    if not body.entries:
        raise HTTPException(400, "There is nothing to record.")
    if len(body.entries) > 500:
        raise HTTPException(400, "Record at most 500 blocks at a time.")
    if any((h.quantity or 0) > 0 for h in body.holds or []):
        wo_actor(request, db, "billing.manage")
    batch = uuid.uuid4().hex[:6]
    ref = lambda g: ("%s-%s" % (batch, g))[:60] if g else ""
    written = 0
    for n, e in enumerate(body.entries, 1):
        one = SubMeasurementIn(item_id=item.id, dimensions=e.dimensions, quantity=e.quantity or 0, multiplier=e.multiplier,
                               measured_on=body.measured_on, mb_ref=body.mb_ref, location=e.location, remarks=body.remarks,
                               section=e.section, block_label=e.block_label, group_ref=ref(e.group))
        try:
            add_sub_measurement(db, client, order, item, one, actor_id, actor_name)
        except HTTPException as exc:
            db.rollback()
            raise HTTPException(exc.status_code, "Block %d (%s) - nothing was recorded. %s" % (
                n, e.location or "no name", exc.detail))
        written += 1
    held_total = 0.0
    for h in body.holds or []:
        qty = money(h.quantity or 0)
        if qty <= 0:
            continue
        if not (h.reason or "").strip():
            db.rollback()
            raise HTTPException(400, "Say why it is held.")
        held_total = money(held_total + qty)
        db.add(models.DBSubMeasurement(
            client_id=client.id, order_id=order.id, item_id=item.id, activity_no=item.activity_no or "",
            measured_on=(body.measured_on or datetime.now().strftime("%Y-%m-%d"))[:10], quantity=-qty, kind="hold",
            mb_ref=(body.mb_ref or "")[:120], location="Held back", remarks=h.reason.strip()[:300],
            recorded_by=actor_id, recorded_by_name=actor_name, group_ref=ref(h.group)))
    if held_total:
        db.flush()
        net = money(sub_measured_to_date(db, order.id).get(item.id, 0.0))
        billed = money(sub_billed_to_date(db, order.id).get(item.id, 0.0))
        if net - billed < -0.0001:
            db.rollback()
            raise HTTPException(409, "More is held back than is measured and not yet billed. Nothing was recorded.")
    log_audit(db, client.id, "sub_measurement_recorded", "subcontract_order", order.id, order.wo_number or "",
              "%d blocks against %s%s" % (written, item.activity_no or "the item", (", %s held" % held_total) if held_total else ""), request)
    db.commit()
    return {"ok": True, "entries": written, "held": held_total,
            "message": "%d block%s recorded%s." % (written, "" if written == 1 else "s",
                                                  (", %s %s held back" % (qty_text(held_total), item.uom or "")) if held_total else "")}


@router.post("/api/sub-mb/convert-old-holds")
def convert_old_holds(request: Request, db: Session = Depends(get_db)):
    client = get_client_user(request, db)
    require_owner(request, db)
    out = convert_legacy_holds(db, client.id)
    db.commit()
    return dict(out, ok=True, message="%d entries put right%s." % (
        out["converted"], ("; left as sent: %s" % ", ".join(out["left_on_sent_bills"])) if out["left_on_sent_bills"] else ""))


@router.post("/api/sub-mb/{order_id}/holds")
def hold_sub_measurement(order_id: int, body: dict, request: Request, db: Session = Depends(get_db)):
    """Hold some of an item's measured work back from billing: a quantity, or a percent of what is
    measured and not yet billed (or of one entry), with the reason."""
    client, actor_id, actor_name = wo_actor(request, db, "billing.manage")
    order = wo_or_404(db, client.id, order_id)
    if (order.status or "") not in ("APPROVED", "EXECUTED"):
        raise HTTPException(409, "Nothing is measured against an order that has not been approved.")
    item = db.query(models.DBSubcontractItem).filter(
        models.DBSubcontractItem.id == int(body.get("item_id") or 0),
        models.DBSubcontractItem.order_id == order.id).first()
    if not item or item.is_header:
        raise HTTPException(404, "That item is not on this order")
    reason = (body.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "Say why it is held - it is read back when it is released.")
    net = money(sub_measured_to_date(db, order.id).get(item.id, 0.0))
    billed = money(sub_billed_to_date(db, order.id).get(item.id, 0.0))
    available = money(net - billed)
    entry = None
    if body.get("entry_id"):
        entry = db.query(models.DBSubMeasurement).filter(
            models.DBSubMeasurement.id == int(body["entry_id"]), models.DBSubMeasurement.order_id == order.id,
            models.DBSubMeasurement.item_id == item.id).first()
        if not entry or (entry.kind or ""):
            raise HTTPException(404, "That measurement is not in this item's book")
    try:
        quantity = float(body.get("quantity") or 0)
        percent = float(body.get("percent") or 0)
    except (TypeError, ValueError):
        raise HTTPException(400, "The quantity and the percent have to be numbers.")
    if percent:
        if not 0 < percent <= 100:
            raise HTTPException(400, "A percent to hold is between 0 and 100.")
        base = (entry.quantity if entry else available) or 0.0
        quantity = money(base * percent / 100.0)
    quantity = money(quantity)
    if quantity <= 0:
        raise HTTPException(400, "Give a quantity or a percent to hold.")
    if quantity > available + 0.0001:
        raise HTTPException(409, "%s has %s measured and not yet billed, so %s cannot be held. Only work not yet "
                                 "billed can be held." % (item.activity_no or "This item", available, quantity))
    held = models.DBSubMeasurement(
        client_id=client.id, order_id=order.id, item_id=item.id, activity_no=item.activity_no or "",
        measured_on=(body.get("measured_on") or datetime.now().strftime("%Y-%m-%d"))[:10],
        quantity=-quantity, kind="hold", mb_ref="HOLD",
        location="Held back" + ((" from %s" % entry.location) if entry and entry.location else ""),
        remarks=reason, recorded_by=actor_id, recorded_by_name=actor_name)
    db.add(held)
    log_audit(db, client.id, "sub_measurement_held", "subcontract_order", order.id, order.wo_number or "",
              "%s %s: %s" % (item.activity_no or "", quantity, reason), request)
    db.commit()
    return {"ok": True, "id": held.id, "held": quantity,
            "message": "%s %s held back from billing." % (qty_text(quantity), item.uom or "")}


@router.post("/api/sub-mb/holds/{hold_id}/release")
def release_sub_hold(hold_id: int, request: Request, body: dict = None, db: Session = Depends(get_db)):
    """Put held work back into what can be billed. The owner's alone."""
    client = get_client_user(request, db)
    require_owner(request, db)
    hold = db.query(models.DBSubMeasurement).filter(
        models.DBSubMeasurement.id == hold_id, models.DBSubMeasurement.client_id == client.id,
        models.DBSubMeasurement.kind == "hold").first()
    if not hold:
        raise HTTPException(404, "Hold not found")
    remaining = sub_hold_remaining(db, hold)
    if remaining <= 0:
        raise HTTPException(409, "That hold has already been released.")
    try:
        quantity = money(float((body or {}).get("quantity") or remaining))
    except (TypeError, ValueError):
        raise HTTPException(400, "The quantity has to be a number.")
    if quantity <= 0 or quantity > remaining + 0.0001:
        raise HTTPException(400, "Release between nothing and %s." % remaining)
    order = wo_or_404(db, client.id, hold.order_id)
    item = db.query(models.DBSubcontractItem).filter(models.DBSubcontractItem.id == hold.item_id).first()
    db.add(models.DBSubMeasurement(
        client_id=client.id, order_id=order.id, item_id=hold.item_id, activity_no=hold.activity_no or "",
        measured_on=datetime.now().strftime("%Y-%m-%d"), quantity=quantity, kind="release", hold_of=hold.id,
        mb_ref="RELEASE", location="Hold released", remarks="Released: %s" % (hold.remarks or ""),
        recorded_by=None, recorded_by_name=owner_label(db, client.id)))
    log_audit(db, client.id, "sub_hold_released", "subcontract_order", order.id, order.wo_number or "",
              "%s %s" % (hold.activity_no or "", quantity), request)
    db.commit()
    return {"ok": True, "released": quantity, "message": "%s %s released - it goes on the next bill." % (
        qty_text(quantity), (item.uom if item else "") or "")}


@router.get("/api/sub-bills")
def list_sub_bills(request: Request, order_id: int = 0, db: Session = Depends(get_db)):
    client = require_erp_read(request, db)
    q = db.query(models.DBSubBill).filter(models.DBSubBill.client_id == client.id)
    if order_id:
        q = q.filter(models.DBSubBill.order_id == order_id)
    viewer = STAFF_VIEWER.get()
    if viewer:
        q = q.filter(models.DBSubBill.id.in_(list(sub_bill_visible_ids(db, client.id, viewer)) or [0]))
    every = q.with_entities(models.DBSubBill.status, models.DBSubBill.this_bill, models.DBSubBill.net_payable,
                            models.DBSubBill.retention_amount, models.DBSubBill.order_id).all()
    bills = q.order_by(models.DBSubBill.id.desc()).limit(300).all()
    prime(db, models.DBSubcontractOrder, [b.order_id for b in bills])
    prime(db, models.DBContractor, [b.contractor_id for b in bills])
    prime(db, models.DBJob, [b.job_id for b in bills])
    prime_chains(db, "sub_bill", [b.id for b in bills])
    prime(db, models.DBEmployee, [r.approver_id for b in bills for r in sub_bill_chain_rows(db, b.id)])
    rows = [sub_bill_dict(db, b) for b in bills]
    seen_orders = {e.order_id for e in every}
    released = money(sum(r.amount or 0 for r in _live_releases(db, client.id, "contractor")
                         if (not order_id or r.sub_order_id == order_id)
                         and (not viewer or r.sub_order_id in seen_orders)))
    return {
        "bills": rows,
        "summary": {
            "claimed": money(sum(e.this_bill or 0 for e in every if (e.status or "DRAFT") != "CANCELLED")),
            "awaiting_certification": len([e for e in every if e.status == "SUBMITTED"]),
            "certified_unpaid": money(sum(e.net_payable or 0 for e in every if e.status == "CERTIFIED")),
            "retention_held": money(sum(e.retention_amount or 0 for e in every
                                        if e.status in ("CERTIFIED", "PAID")) - released),
            "paid": money(sum(e.net_payable or 0 for e in every if e.status == "PAID")),
        },
    }


@router.post("/api/sub-bills")
def create_sub_bill(body: SubBillIn, request: Request, db: Session = Depends(get_db)):
    client, actor_id, actor_name = wo_actor(request, db, "billing.manage")
    order = wo_or_404(db, client.id, body.order_id)
    if (order.status or "") not in ("APPROVED", "EXECUTED"):
        raise HTTPException(409, "Approve the order before billing against it.")
    open_bill = db.query(models.DBSubBill).filter(
        models.DBSubBill.order_id == order.id,
        models.DBSubBill.status.in_(("DRAFT", "SUBMITTED"))).first()
    if open_bill:
        raise HTTPException(
            409, "%s is still open on this order. Finish or cancel it before "
                 "raising another." % open_bill.number)
    if body.entry_ids is not None:
        if not body.entry_ids:
            raise HTTPException(400, "Choose at least one entry of the measurement book to bill.")
    elif not sub_claimable_lines(db, order):
        raise HTTPException(
            409, "Nothing has been measured since the last bill. Record the work "
                 "in the measurement book first.")

    seq = (db.query(func.max(models.DBSubBill.sequence)).filter(
        models.DBSubBill.order_id == order.id).scalar() or 0) + 1
    while db.query(models.DBSubBill.id).filter(
            models.DBSubBill.client_id == client.id,
            models.DBSubBill.number == "%s/RA-%02d" % (order.wo_number or "SC", seq)).first():
        seq += 1
    prior = money(sum(b.this_bill or 0 for b in db.query(models.DBSubBill).filter(
        models.DBSubBill.order_id == order.id,
        models.DBSubBill.status != "CANCELLED").all()))
    con = db.query(models.DBContractor).filter(models.DBContractor.id == order.contractor_id).first() \
        if order.contractor_id else None
    debit = money(body.debit_notes or 0)
    if debit < 0:
        raise HTTPException(400, "Recoveries in debit notes cannot be negative.")
    bill = models.DBSubBill(
        client_id=client.id, order_id=order.id, job_id=order.job_id,
        contractor_id=order.contractor_id,
        number="%s/RA-%02d" % (order.wo_number or "SC", seq), sequence=seq,
        period_from=(body.period_from or ""), period_to=(body.period_to or
                                                          datetime.now().strftime("%Y-%m-%d")),
        bill_date=(body.bill_date or datetime.now().strftime("%Y-%m-%d"))[:10],
        # What the certificate calls the type of work: what the gang was
        # registered to do, else the order's trade.
        work_type=((body.work_type or "").strip() or (con.nature_of_work if con else "") or
                   order.work_type or "")[:200],
        work_name=((body.work_name or "").strip() or (order.subject or "").split("\n")[0])[:300],
        hsn_sac=(body.hsn_sac or "").strip()[:20],
        debit_notes=debit,
        status="DRAFT", previously_billed=prior,
        retention_percent=order.retention_percent or 0,
        advance_recovery=0.0,
        other_deductions=money(body.other_deductions or 0),
        deduction_notes=(body.deduction_notes or "").strip(),
        gst_percent=order.gst_rate or 0, tds_percent=order.tds_rate or 0,
        labour_cess_percent=order.labour_cess_percent or 0)
    db.add(bill)
    db.flush()
    draw_sub_bill_lines(db, bill, order, entry_ids=body.entry_ids)
    recost_sub_bill(db, bill)
    bill.advance_recovery = sub_advance_recovery(db, order, bill, body.advance_recovery)
    recost_sub_bill(db, bill)
    if apply_material_recovery(db, client.id, order, bill):
        recost_sub_bill(db, bill)
    log_audit(db, client.id, "sub_bill_raised", "subcontract_order", order.id,
              order.wo_number or "", "%s %s" % (bill.number, inr(bill.this_bill)), request)
    db.commit()
    db.refresh(bill)
    return {"ok": True, "bill": sub_bill_dict(db, bill, detail=True),
            "message": "%s drawn up from the measurement book - %s of work."
                       % (bill.number, inr(bill.this_bill))}


@router.post("/api/sub-bills/{bill_id}/hardcopy")
def sub_bill_attach_hardcopy(bill_id: int, request: Request, file: UploadFile = File(...), amount: str = Form(""),
                             db: Session = Depends(get_db)):
    """The hard copy of this same bill, on paper and scanned (a PDF, or a photo of it), attached before the bill is sent for
    approval; with the amount of work it shows, so the bill in the app can be checked against it. Attaching again replaces it."""
    client, actor_id, actor_name = wo_actor(request, db, "billing.manage")
    bill = sub_bill_or_404(db, client.id, bill_id)
    if (bill.status or "DRAFT") in ("PAID", "CANCELLED"):
        raise HTTPException(409, "A bill that is %s keeps the hard copy it has." % bill.status.lower())
    name = (file.filename or "bill").strip()
    ctype = (file.content_type or "").lower()
    if ctype not in HARDCOPY_TYPES and os.path.splitext(name.lower())[1] not in (".pdf", ".jpg", ".jpeg", ".png"):
        raise HTTPException(400, "Attach the hard copy as a PDF, or a photo of it (JPG or PNG).")
    if ctype not in HARDCOPY_TYPES:
        ctype = {".pdf": "application/pdf", ".png": "image/png"}.get(os.path.splitext(name.lower())[1], "image/jpeg")
    claimed = None
    if (amount or "").strip():
        claimed = money(sheet_number(amount))
        if claimed < 0:
            raise HTTPException(400, "The amount on the hard copy cannot be negative.")
    data = file.file.read()
    drop_files_of(db, "sub_bill_scan", [bill.id])
    stored = store_file(db, client.id, types.SimpleNamespace(filename=name, content_type=ctype), data, job_id=bill.job_id, attached_type="sub_bill_scan", attached_id=bill.id,
                        kind="document", caption="Hard copy of %s" % bill.number, by=actor_name)
    bill.scan_file_id, bill.scan_name, bill.scan_type, bill.scan_size = stored.id, stored.name, stored.content_type, stored.size
    bill.scan_amount, bill.scan_by_name = claimed, actor_name
    bill.scan_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    log_audit(db, client.id, "sub_bill_hardcopy", "sub_bill", bill.id, bill.number or "", stored.name, request)
    db.commit()
    db.refresh(bill)
    return {"ok": True, "bill": sub_bill_dict(db, bill, detail=True), "message": "The hard copy is attached to %s." % bill.number}


@router.get("/api/sub-bills/{bill_id}/hardcopy")
def sub_bill_hardcopy(bill_id: int, request: Request, db: Session = Depends(get_db)):
    client = require_erp_read(request, db)
    bill = sub_bill_or_404(db, client.id, bill_id)
    if not bill.scan_file_id:
        raise HTTPException(404, "No hard copy has been attached to this bill.")
    f = _file_or_404(db, client.id, bill.scan_file_id, bytes_too=True)
    media = served_type(f.content_type)
    name = re.sub(r'[^A-Za-z0-9._ -]', "_", f.name or "bill")
    return StreamingResponse(io.BytesIO(file_bytes(db, f) or b""), media_type=media,
                             headers=dict(file_headers(media), **{
                                 "Content-Disposition": '%s; filename="%s"' % ("inline" if media != "application/octet-stream" else "attachment", name),
                                 "Cache-Control": "private, max-age=3600"}))


@router.delete("/api/sub-bills/{bill_id}/hardcopy")
def sub_bill_remove_hardcopy(bill_id: int, request: Request, db: Session = Depends(get_db)):
    client, _, _ = wo_actor(request, db, "billing.manage")
    bill = sub_bill_or_404(db, client.id, bill_id)
    if (bill.status or "DRAFT") in ("PAID", "CANCELLED"):
        raise HTTPException(409, "A bill that is %s keeps the hard copy it has." % bill.status.lower())
    if not bill.scan_file_id:
        raise HTTPException(404, "No hard copy has been attached to this bill.")
    log_audit(db, client.id, "sub_bill_hardcopy_removed", "sub_bill", bill.id, bill.number or "", bill.scan_name or "", request)
    drop_files_of(db, "sub_bill_scan", [bill.id])
    bill.scan_file_id, bill.scan_name, bill.scan_type, bill.scan_size = None, "", "", 0
    bill.scan_amount, bill.scan_by_name, bill.scan_at = None, "", ""
    db.commit()
    return {"ok": True, "message": "Removed."}


@router.get("/api/sub-bills/{bill_id}")
def get_sub_bill(bill_id: int, request: Request, db: Session = Depends(get_db)):
    client = require_erp_read(request, db)
    return sub_bill_dict(db, sub_bill_or_404(db, client.id, bill_id), detail=True)


@router.put("/api/sub-bills/{bill_id}")
def edit_sub_bill(bill_id: int, body: SubBillEditIn, request: Request, db: Session = Depends(get_db)):
    """The certificate's boxes - the dates, the type of work, the SAC, the
    recoveries - put right on a draft before it is sent."""
    client, _, _ = wo_actor(request, db, "billing.manage")
    bill = sub_bill_or_404(db, client.id, bill_id)
    if (bill.status or "DRAFT") != "DRAFT":
        raise HTTPException(409, "Only a draft bill can be changed. Send it back to draft first.")
    order = wo_or_404(db, client.id, bill.order_id)
    for key in ("period_from", "period_to", "bill_date"):
        v = getattr(body, key)
        if v is not None:
            setattr(bill, key, (v or "").strip()[:10])
    for key, size in (("work_type", 200), ("work_name", 300), ("hsn_sac", 20), ("deduction_notes", 500)):
        v = getattr(body, key)
        if v is not None:
            setattr(bill, key, (v or "").strip()[:size])
    for key in ("debit_notes", "advance_recovery", "other_deductions"):
        v = getattr(body, key)
        if v is not None and money(v) < 0:
            raise HTTPException(400, "Deductions and recoveries cannot be negative.")
    if body.debit_notes is not None:
        bill.debit_notes = money(body.debit_notes)
        if bill.debit_notes > (bill.this_bill or 0):
            raise HTTPException(400, "Recoveries in debit notes cannot be more than the work in the bill.")
    if body.other_deductions is not None:
        material = money(sum(r.amount or 0 for r in db.query(models.DBMaterialRecovery).filter(
            models.DBMaterialRecovery.client_id == client.id,
            models.DBMaterialRecovery.sub_bill_id == bill.id).all()))
        bill.other_deductions = money(material + money(body.other_deductions))
    recost_sub_bill(db, bill)
    if body.advance_recovery is not None:
        bill.advance_recovery = 0.0
        recost_sub_bill(db, bill)
        bill.advance_recovery = sub_advance_recovery(db, order, bill, body.advance_recovery)
        recost_sub_bill(db, bill)
    if (bill.net_payable or 0) < 0:
        raise HTTPException(400, "Those deductions take the bill below nothing (%s)." % inr(bill.net_payable))
    log_audit(db, client.id, "sub_bill_edited", "sub_bill", bill.id, bill.number or "", "", request)
    db.commit()
    db.refresh(bill)
    return {"ok": True, "bill": sub_bill_dict(db, bill, detail=True), "message": "%s saved." % bill.number}


@router.post("/api/sub-bills/{bill_id}/{action}")
def act_on_sub_bill(bill_id: int, action: str, request: Request, body: dict = None,
                    db: Session = Depends(get_db)):
    client, actor_id, actor_name = wo_actor(
        request, db, ("billing.manage", "subcontracts.approve", "bills.pay"))
    bill = sub_bill_or_404(db, client.id, bill_id)
    move = (action or "").upper()
    body = body or {}
    comments = (body.get("comments") or "").strip()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if move == "ACCEPT":
        # "Accepted for Sub Contractor": the gang signed the certificate -
        # recorded by the office from the signed copy, or by the gang itself
        # from the portal.
        require_items_access(request, db, "billing.manage")
        if (bill.status or "") not in ("SUBMITTED", "CERTIFIED", "PAID"):
            raise HTTPException(409, "A bill is accepted by the sub contractor once it has been sent.")
        con = db.query(models.DBContractor).filter(models.DBContractor.id == bill.contractor_id).first() \
            if bill.contractor_id else None
        bill.accepted_by_name = ((body.get("name") or "").strip() or (con.contact_person if con else "")
                                 or (con.company_name if con else "") or "Sub contractor")[:120]
        bill.accepted_at = ((body.get("date") or "").strip()[:10] or now)
        log_audit(db, client.id, "sub_bill_accepted", "sub_bill", bill.id, bill.number or "",
                  "Accepted for the sub contractor by %s" % bill.accepted_by_name, request)
        db.commit()
        db.refresh(bill)
        return {"ok": True, "bill": sub_bill_dict(db, bill, detail=True),
                "message": "%s marked as accepted by %s." % (bill.number, bill.accepted_by_name)}

    allowed = SUB_TRANSITIONS.get(bill.status or "DRAFT", {})
    if move not in allowed:
        raise HTTPException(409, "A %s bill cannot be %s."
                                 % ((bill.status or "draft").lower(), past_tense(move)))
    order = wo_or_404(db, client.id, bill.order_id)
    if move in ("SUBMIT", "CERTIFY") and (order.status or "") not in ("APPROVED", "EXECUTED"):
        raise HTTPException(409, "%s is %s - nothing more is billed or certified against it."
                                 % (order.wo_number, (order.status or "").lower()))

    # Drawing up and sending is billing's; certifying or sending back is the
    # approver's; paying is the accounts department's.
    require_items_access(request, db, {"CERTIFY": "subcontracts.approve",
                                       "REJECT": "subcontracts.approve",
                                       "PAY": "bills.pay"}.get(move, "billing.manage"))
    if move == "SUBMIT":
        if bill_scan_required() and not bill.scan_file_id:
            raise HTTPException(409, "Attach the hard copy of this bill (as a PDF or a photo) before sending it for "
                                     "approval, so whoever approves it can read the two side by side.")
        # Redrawn from the book on the way out, so anything measured since
        # the draft was opened is on it.
        draw_sub_bill_lines(db, bill, order)
        recost_sub_bill(db, bill)
        if not bill.this_bill:
            raise HTTPException(409, "There is nothing on this bill to submit.")
        # The owner signs as themselves, not as the company.
        bill.submitted_by, bill.submitted_at = actor_id, now
        bill.submitted_by_name = actor_name if actor_id else ((client.contact_name or "").strip() or actor_name)
        bill.approved_by_name, bill.remarks = "", ""
        sub_bill_start_chain(db, client.id, bill, actor_id)
    elif move == "CERTIFY":
        # Certifying a subcontractor's bill is agreeing to pay it - once the
        # last signature on its route is on.
        if not sub_bill_decide(db, client, bill, actor_id, actor_name, True, comments):
            bill.updated_at = now
            db.commit()
            db.refresh(bill)
            step = sub_bill_current_step(db, bill)
            return {"ok": True, "bill": sub_bill_dict(db, bill, detail=True),
                    "message": "Signed. %s now waits with %s." % (bill.number, sub_bill_step_name(db, step))}
        bill.certified_by, bill.certified_by_name, bill.certified_at = actor_id, actor_name, now
        bill.approved_by_name = actor_name if actor_id else (owner_label(db, client.id))
    elif move == "REJECT":
        if not comments:
            raise HTTPException(400, "Say why it is going back.")
        sub_bill_decide(db, client, bill, actor_id, actor_name, False, comments)
        bill.remarks = comments
        if (bill.entry_mode or "") != "chosen":     # a bill of chosen entries goes back with the same entries
            db.query(models.DBSubMeasurement).filter(
                models.DBSubMeasurement.sub_bill_id == bill.id).update(
                    {"sub_bill_id": None}, synchronize_session=False)
    elif move == "PAY":
        bill.paid_at = now
        bill.paid_reference = (body.get("reference") or "").strip()
    elif move == "CANCEL":
        if not comments:
            raise HTTPException(400, "Say why it is being cancelled.")
        refuse_cancel_with_money(db, client.id, "sub_bill", bill, "paid")
        bill.remarks = comments
        db.query(models.DBSubMeasurement).filter(
            models.DBSubMeasurement.sub_bill_id == bill.id).update(
                {"sub_bill_id": None}, synchronize_session=False)
        # Material it recovered goes back to waiting for the next bill.
        release_material_recovery(db, client.id, bill.id)
        free_back_charges_of(db, bill.id)

    was, bill.status = bill.status, allowed[move]
    bill.updated_at = now
    log_audit(db, client.id, "sub_bill_%s" % move.lower(), "sub_bill", bill.id,
              bill.number or "", "%s -> %s %s" % (was, bill.status, comments), request)
    db.commit()
    db.refresh(bill)
    if move in ("SUBMIT", "CERTIFY"):
        con = db.query(models.DBContractor).filter(models.DBContractor.id == bill.contractor_id).first() \
            if bill.contractor_id else None
        gang = (con.company_name if con else "") or "the contractor"
        if move == "SUBMIT":
            step = sub_bill_current_step(db, bill)
            notify(db, client.id, "sub_bill_submitted", "%s from %s is waiting to be certified" % (bill.number, gang),
                   "%s of work billed - with %s." % (inr(bill.this_bill), sub_bill_step_name(db, step) or "the Master"),
                   view="subbills-view", ref_type="sub_bill", ref_id=bill.id, severity="action")
            if step is not None and step.approver_id:
                notify_employee(db, client.id, step.approver_id, "Subcontractor bill awaiting your certification",
                                "%s from %s - %s of work, net %s." % (
                                    bill.number, gang, format_money_plain(bill.this_bill),
                                    format_money_plain(bill.net_payable)),
                                link="/next/approvals")
            db.commit()
        else:
            notify(db, client.id, "sub_bill_certified", "%s certified - %s to pay %s" % (
                bill.number, inr(bill.net_payable), gang), "Pay it from Subcontractor Bills.",
                   view="subbills-view", ref_type="sub_bill", ref_id=bill.id, severity="money")
        db.refresh(bill)
    return {"ok": True, "bill": sub_bill_dict(db, bill, detail=True),
            "message": "%s %s." % (bill.number, bill.status.lower())}


@router.get("/api/sub-bills.xlsx")
def sub_bills_register_xlsx(request: Request, order_id: int = 0, db: Session = Depends(get_db)):
    """Every gang bill as the register the office reconciles - with its PDF
    twin beside it, like every other register. The screen only offered to
    print the page."""
    bills = list_sub_bills(request, order_id, db)["bills"]
    headers = ["Bill No", "Status", "Sub Contractor", "Work Order", "Project", "Period To", "This Bill",
               "Retention", "Advance Recovered", "GST", "TDS", "Labour Cess", "Net Payable",
               "Certified By", "Paid On", "Paid Ref"]
    money_cols = ("this_bill", "retention_amount", "advance_recovery", "gst_amount", "tds_amount",
                  "labour_cess_amount", "net_payable")
    rows = [[b["number"], b["status"], b["contractor"], b["order"], b["project"], b["period_to"]] +
            [b[k] for k in money_cols] +
            [b["certified_by_name"], (b["paid_at"] or "")[:10], b["paid_reference"]] for b in bills]
    live = [b for b in bills if b["status"] != "CANCELLED"]
    rows.append(["TOTAL (excluding cancelled)", "", "", "", "", ""] +
                [money(sum(b[k] or 0 for b in live)) for k in money_cols] + ["", "", ""])
    return sheet_response(headers, rows, "subcontractor_bills.xlsx")


@router.get("/api/sub-bills/{bill_id}/export.xlsx")
def export_sub_bill(bill_id: int, request: Request, db: Session = Depends(get_db)):
    """The bill as the workbook it was always sent as - Top Sheet, AB-1 and
    MB-1 - drawn from the same description as its PDF."""
    client = require_erp_read(request, db)
    bill = sub_bill_or_404(db, client.id, bill_id)
    if SHEET_AS_PDF.get() is not None:
        return form_pdf_response(sub_bill_form_spec(db, client, bill), bill.number)
    cert = sub_bill_certificate(db, client, bill)
    data = sheet_forms.ra_bill_workbook(cert, cert.get("company"))
    return Response(content=data,
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="sub_bill_%s.xlsx"'
                                                    % re.sub(r"[^A-Za-z0-9]+", "-", bill.number or str(bill.id))})


@router.get("/api/material-recoveries")
def list_material_recoveries(request: Request, order_id: int = 0, db: Session = Depends(get_db)):
    client = require_erp_read(request, db)
    q = db.query(models.DBMaterialRecovery).filter(
        models.DBMaterialRecovery.client_id == client.id)
    if order_id:
        q = q.filter(models.DBMaterialRecovery.order_id == order_id)
    rows = [material_recovery_dict(r) for r in q.order_by(models.DBMaterialRecovery.id.desc()).all()]
    return {"recoveries": rows, "summary": {
        "waiting": money(sum(r["amount"] for r in rows if r["state"] == "waiting")),
        "recovered": money(sum(r["amount"] for r in rows if r["state"] == "recovered"))}}


@router.post("/api/sub-bills/bulk-delete")
def sub_bill_bulk_delete(body: dict, request: Request, db: Session = Depends(get_db)):
    return run_bulk_delete(body.get("ids"), lambda i: sub_bill_delete(i, request, db), db)


@router.post("/api/sub-mb/entries/bulk-delete")
def sub_measurement_bulk_delete(body: dict, request: Request, db: Session = Depends(get_db)):
    ids = []
    for raw_id in list(body.get("ids") or [])[:200]:
        try:
            ids.append(int(raw_id))
        except (TypeError, ValueError):
            continue
    kinds = dict(db.query(models.DBSubMeasurement.id, models.DBSubMeasurement.kind).filter(
        models.DBSubMeasurement.id.in_(ids or [0])).all())
    rank = {"release": 0, "hold": 1}
    ids.sort(key=lambda i: rank.get(kinds.get(i) or "", 2))
    return run_bulk_delete(ids, lambda i: delete_sub_measurement(i, request, db), db)


@router.get("/api/retention")
def retention_overview(request: Request, db: Session = Depends(get_db)):
    client = require_items_access(request, db, "bills.view_all")
    positions = retention_positions(db, client.id)
    settled = settled_amounts(db, client.id)
    releases = [release_dict(db, r, settled) for r in db.query(models.DBRetentionRelease).filter(
        models.DBRetentionRelease.client_id == client.id).order_by(
            models.DBRetentionRelease.id.desc()).limit(500).all()]
    ours = [p for p in positions if p["side"] == "client"]
    theirs = [p for p in positions if p["side"] == "contractor"]
    live = [r for r in releases if r["status"] != "CANCELLED"]
    return {
        "positions": positions, "releases": releases, "stages": list(RETENTION_STAGES),
        "summary": {
            "client_held": money(sum(p["balance"] for p in ours)),
            "client_released": money(sum(p["released"] for p in ours)),
            "client_on_finished": money(sum(p["balance"] for p in ours if p["finished"])),
            "client_to_receive": money(sum(r["outstanding"] for r in live if r["side"] == "client")),
            "contractor_held": money(sum(p["balance"] for p in theirs)),
            "contractor_released": money(sum(p["released"] for p in theirs)),
            "contractor_dlp_over": money(sum(p["balance"] for p in theirs if p["dlp_over"])),
            "contractor_to_pay": money(sum(r["outstanding"] for r in live if r["side"] == "contractor")),
        },
    }


@router.post("/api/retention/releases")
def create_release(body: RetentionReleaseIn, request: Request, db: Session = Depends(get_db)):
    """Release retention: claim it from the client, or owe it to a gang.

    Never more than is still held on that order. The tax is worked out the
    way the order's own bills worked theirs out - the site's state against
    ours, or against the gang's for their side."""
    side = (body.side or "").strip().lower()
    if side not in ("client", "contractor"):
        raise HTTPException(400, "Is this retention the client holds, or retention we hold from a contractor?")
    # Owing a gang money is a decision, as certifying their bill is.
    client, actor_id, actor_name = wo_actor(
        request, db, ("billing.manage", "accounts.manage") if side == "client" else "subcontracts.approve")
    stage = (body.stage or "").strip()
    if stage not in RETENTION_STAGES:
        raise HTTPException(400, "Which stage is this - %s?" % ", ".join(RETENTION_STAGES))
    amount = money(body.amount or 0)
    if amount <= 0:
        raise HTTPException(400, "How much is being released?")
    pos = next((p for p in retention_positions(db, client.id)
                if p["side"] == side and p["order_id"] == body.order_id), None)
    if not pos:
        raise HTTPException(404, "No retention is held on that order.")
    if amount > pos["balance"] + 0.009:
        raise HTTPException(400, "Only %s is still held on %s; %s is more than that."
                                 % (inr(pos["balance"]), pos["order_number"] or "that order", inr(amount)))
    on = (body.release_on or date.today().isoformat())[:10]
    if not _parse_date(on):
        raise HTTPException(400, "Release date should be YYYY-MM-DD.")
    rate = pos["gst_percent"] if body.gst_percent is None else max(0.0, min(28.0, float(body.gst_percent)))

    supply = supply_state_for_job(db, pos["job_id"]) if pos["job_id"] else ""
    if side == "client":
        origin = our_state(db, client.id)
    else:
        origin = contractor_state(db, pos.get("contractor_id")) or our_state(db, client.id)
    gst = split_gst(amount, rate, origin, supply)

    stem = "RET" if side == "client" else "RETG"
    n = db.query(models.DBRetentionRelease).filter(
        models.DBRetentionRelease.client_id == client.id,
        models.DBRetentionRelease.side == side).count() + 1
    r = models.DBRetentionRelease(
        client_id=client.id, side=side,
        work_order_id=body.order_id if side == "client" else None,
        sub_order_id=body.order_id if side == "contractor" else None,
        job_id=pos["job_id"], contractor_id=pos.get("contractor_id"),
        number="%s-%04d" % (stem, n), stage=stage, release_on=on, amount=amount,
        gst_percent=rate, gst_amount=gst["total"], cgst_amount=gst["cgst"],
        sgst_amount=gst["sgst"], igst_amount=gst["igst"], place_of_supply=supply,
        net_amount=money(amount + gst["total"]), status="CERTIFIED",
        notes=(body.notes or "").strip()[:500], created_by_name=actor_name or "")
    db.add(r)
    db.flush()
    log_audit(db, client.id, "retention_released", "retention_release", r.id, r.number,
              "%s %s on %s" % (stage, inr(amount), pos["order_number"]), request)
    db.commit()
    db.refresh(r)
    if side == "client":
        notify(db, client.id, "retention_released", "%s - retention to claim from %s" % (r.number, pos["party"] or "the client"),
               "%s released at %s on %s, %s with GST." % (inr(amount), stage.lower(), pos["project"], inr(r.net_amount)),
               view="money-view", ref_type="retention_release", ref_id=r.id, severity="money")
    else:
        notify(db, client.id, "retention_released", "%s - retention due to %s" % (r.number, pos["party"] or "the contractor"),
               "%s released at %s on %s, %s with GST." % (inr(amount), stage.lower(), pos["order_number"], inr(r.net_amount)),
               view="money-view", ref_type="retention_release", ref_id=r.id, severity="action")
    db.refresh(r)
    return {"release": release_dict(db, r),
            "message": "%s raised: %s released, %s with GST." % (r.number, inr(amount), inr(r.net_amount))}


@router.get("/api/retention/releases/{release_id}")
def get_release(release_id: int, request: Request, db: Session = Depends(get_db)):
    client = require_items_access(request, db, "bills.view_all")
    return release_dict(db, release_or_404(db, client.id, release_id))


@router.post("/api/retention/releases/{release_id}/cancel")
def cancel_release(release_id: int, request: Request, body: dict = None,
                   db: Session = Depends(get_db)):
    """A release raised in error goes back to being held. Not once money has
    moved against it - that is voided first, as with any bill."""
    client, actor_id, actor_name = wo_actor(request, db, ("billing.manage", "accounts.manage", "subcontracts.approve"))
    r = release_or_404(db, client.id, release_id)
    reason = ((body or {}).get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "Say why it is being cancelled.")
    if r.status == "CANCELLED":
        raise HTTPException(409, "%s is already cancelled." % r.number)
    refuse_cancel_with_money(db, client.id, "retention_release", r,
                             "received" if r.side == "client" else "paid")
    refuse_cancel_with_irn(db, client.id, "retention_release", r)
    r.status, r.cancel_reason = "CANCELLED", reason[:300]
    log_audit(db, client.id, "retention_release_cancelled", "retention_release", r.id, r.number, reason, request)
    db.commit()
    return {"release": release_dict(db, r), "message": "%s cancelled; the retention is held again." % r.number}


@router.get("/api/retention/releases/{release_id}/export.xlsx")
def export_release(release_id: int, request: Request, db: Session = Depends(get_db)):
    """The release as a bill: the RA bills whose retention it gives back,
    then what was held, released before, released now, and the tax."""
    client = require_items_access(request, db, "bills.view_all")
    r = release_or_404(db, client.id, release_id)
    d = release_dict(db, r)
    if r.side == "client":
        bills = db.query(models.DBRABill).filter(
            models.DBRABill.work_order_id == r.work_order_id,
            models.DBRABill.status.in_(("CERTIFIED", "PAID"))).order_by(models.DBRABill.id).all()
    else:
        bills = db.query(models.DBSubBill).filter(
            models.DBSubBill.order_id == r.sub_order_id,
            models.DBSubBill.status.in_(("CERTIFIED", "PAID"))).order_by(models.DBSubBill.id).all()
    held = money(sum(b.retention_amount or 0 for b in bills))
    before = money(sum(x.amount or 0 for x in _live_releases(db, client.id, r.side)
                       if x.id < r.id and (x.work_order_id, x.sub_order_id) == (r.work_order_id, r.sub_order_id)))
    rows = [(b.number or "", (b.certified_at or "")[:10], money(b.this_bill),
             b.retention_percent or 0, money(b.retention_amount)) for b in bills]
    tax = []
    if r.cgst_amount or r.sgst_amount:
        tax = [("CGST @ %s%%" % (d["gst_percent"] / 2), "", "", "", d["cgst_amount"]),
               ("SGST @ %s%%" % (d["gst_percent"] / 2), "", "", "", d["sgst_amount"])]
    elif r.igst_amount:
        tax = [("IGST @ %s%%" % d["gst_percent"], "", "", "", d["igst_amount"])]
    title = "RETENTION RELEASE" if r.side == "client" else "RETENTION RELEASED TO SUBCONTRACTOR"
    return sheet_response(
        ("Bill", "Certified", "Bill value", "Retention %", "Retention held"),
        rows, "%s.xlsx" % (r.number or "release"),
        preamble=[(title, client.company_name or ""),
                  ("Number", d["number"]), ("Date", d["release_on"]),
                  ("To" if r.side == "client" else "Payable to", d["party"]),
                  ("Project", d["project"]), ("Order", d["order_number"]),
                  ("Stage", d["stage"]),
                  ("Place of supply", d["place_of_supply"]), ("SAC", WORKS_CONTRACT_SAC),
                  ()],
        closing=[(), ("Retention held on these bills", "", "", "", held),
                 ("Released before", "", "", "", -before),
                 ("Released now", "", "", "", d["amount"])] + tax +
                [("Total %s" % ("claimed" if r.side == "client" else "payable"), "", "", "", d["net_amount"])] +
                ([(), ("Note", d["notes"])] if d["notes"] else []))


@router.get("/api/sub-bills/{bill_id}/document.pdf")
def sub_bill_pdf(bill_id: int, request: Request, db: Session = Depends(get_db)):
    client = require_erp_read(request, db)
    bill = sub_bill_or_404(db, client.id, bill_id)
    return form_pdf_response(sub_bill_form_spec(db, client, bill), bill.number)
