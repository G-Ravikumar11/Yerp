"""The office's own schedule sheets, read into the programme.

Two kinds are read, and a sheet this app exported comes back exactly (by its hidden ID and Parent columns):

* the GH layout - SL.NO, V.NO, ITEMES, Start, End, No of working Days, then a column a day - where the tree is told by
  bold rows and the SL.NO (I, A, 1): a bold row with a different kind of number than the ones open goes inside the
  last one; a number of a kind already open closes everything inside it;
* the APTIDCO / finishing-works layout - SL.NO, ITEMES / ACTIVITY, Start, End, No of working Days, U.O.M, Quantity or
  Labour Required - where a heading's own formulas name its rows (=MIN(C8:C45), =+MAX(D8:D28), =SUM(E9:E12), or =+C14
  and =+D17 at its first and last), and a row's Start written as a formula on another row is a dependency.

Dependencies from the formulas: =+D9 starts the day row 9 ends (FS-1: the sheet's End is the day the work finishes,
the day it ends), =+D11+1 the day after (FS+0), =+C23+2 two days after row 23 starts (SS+2). A leaf's Days are the
dates the sheet shows, End - Start + 1, so every date comes back as the sheet has it.

Nothing here touches the database but apply_import; parse_sheet only reads.
"""
import io
import re
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import load_workbook

from app import models
from app.services.programme import _check_loops, _place, add_holiday, job_rows, set_calendar
from app.services.programme_boq import short_name
from app.services.programme_engine import parse_predecessors

MAX_ROWS = 5000
FORMULA_REF = re.compile(r"\$?([A-Z]{1,2})\$?(\d+)(?::\$?([A-Z]{1,2})\$?(\d+))?")
START_FORMULA = re.compile(r"=\s*\+?\(?\s*\$?([A-Z]{1,2})\$?(\d+)\s*\)?\s*(?:([+-])\s*(\d+))?")


def _norm(v):
    return re.sub(r"[^a-z0-9%]", "", str(v or "").lower())


def _iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v.strip()):
        return v.strip()
    return ""


def _number(v):
    if isinstance(v, bool) or v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _columns(ws):
    """The header row and which column is which, found by the words in the header."""
    for r in range(1, min(ws.max_row, 20) + 1):
        found = {}
        for c in range(1, min(ws.max_column, 16) + 1):
            t = _norm(ws.cell(row=r, column=c).value)
            if not t:
                continue
            if t.startswith("slno") or t == "sno":
                found.setdefault("sno", c)
            elif t in ("vno", "boqref", "boqitem"):
                found.setdefault("vno", c)
            elif t == "start":
                found.setdefault("start", c)
            elif t == "end":
                found.setdefault("end", c)
            elif t == "uom":
                found.setdefault("uom", c)
            elif t in ("quantity", "qty"):
                found.setdefault("qty", c)
            elif "labour" in t:
                found.setdefault("labour", c)
            elif t.startswith("predecessor"):
                found.setdefault("pred", c)
            elif t in ("%done", "done"):
                found.setdefault("pct", c)
            elif t == "id":
                found.setdefault("id", c)
            elif t == "parent":
                found.setdefault("parent", c)
            elif "item" in t or "activity" in t:
                found.setdefault("name", c)
            elif "days" in t:
                found.setdefault("days", c)
        if {"start", "end", "name"} <= found.keys():
            return r, found
    return None, None


def _open(raw, data_only):
    try:
        return load_workbook(io.BytesIO(raw), data_only=data_only)
    except Exception:
        raise HTTPException(400, "That file could not be read as an Excel workbook (.xlsx).")


def _kind(sno):
    """The sort of number an SL.NO is: I (roman), A (a letter), 1 (a number), or none."""
    s = (sno or "").strip()
    if not s:
        return ""
    if re.fullmatch(r"[IVX]+", s):
        return "R"
    if re.fullmatch(r"[A-Za-z]", s):
        return "L"
    if re.fullmatch(r"\d+(\.\d+)*", s):
        return "N"
    return ""


def _formula_range(row, fs, fe, fd):
    """The rows a heading's own formulas name, below it: (first, last), or None for a leaf."""
    found = []
    singles = {}
    for tag, f in (("s", fs), ("e", fe), ("d", fd)):
        if not f:
            continue
        up = f.upper()
        for m in FORMULA_REF.finditer(up):
            r1 = int(m.group(2))
            r2 = int(m.group(4)) if m.group(4) else r1
            if m.group(4):
                if r1 > row:
                    found.append((r1, r2))
            elif r1 > row and re.fullmatch(r"=\s*\+?\s*\$?[A-Z]{1,2}\$?\d+\s*", up):
                singles[tag] = r1
    if "s" in singles and "e" in singles:
        found.append((min(singles.values()), max(singles.values())))
    elif "e" in singles and not found:
        found.append((row + 1, singles["e"]))
    if not found:
        return None
    return min(a for a, _ in found), max(b for _, b in found)


def parse_sheet(raw, sheet=""):
    """The programme in a sheet, without saving it: {"sheet", "sheets", "nodes", "calendar", "warnings", "own"}.
    Each node: key (its row), parent (a key or None), name, description, boq_ref, uom, quantity, labour, start, days,
    percent, links [{"pred": key, "kind", "lag"}], heading."""
    values = _open(raw, data_only=True)
    forms = _open(raw, data_only=False)
    candidates = [ws.title for ws in values if _columns(ws)[0]]
    if not candidates:
        raise HTTPException(400, "No programme found: a sheet needs a header row with Start, End and the item / activity.")
    name = sheet if sheet in candidates else candidates[0]
    ws, wf = values[name], forms[name]
    header, col = _columns(ws)
    own = "id" in col and "parent" in col
    letter = lambda key: re.sub(r"\d", "", wf.cell(row=1, column=col[key]).coordinate) if key in col else ""

    def val(r, key):
        return ws.cell(row=r, column=col[key]).value if key in col else None

    def formula(r, key):
        v = wf.cell(row=r, column=col[key]).value if key in col else None
        return v if isinstance(v, str) and v.startswith("=") else None

    nodes, warnings = [], []
    for r in range(header + 1, ws.max_row + 1):
        text = str(val(r, "name") or "").strip()
        if not text or text.startswith("="):
            continue
        if len(nodes) >= MAX_ROWS:
            warnings.append("Only the first %d rows were read." % MAX_ROWS)
            break
        sno_raw = val(r, "sno")
        sno = ("%.2f" % sno_raw).rstrip("0").rstrip(".") if isinstance(sno_raw, float) else str(sno_raw or "").strip()
        cell = wf.cell(row=r, column=col["name"])
        flat = re.sub(r"\s+", " ", text)
        n = {"key": r, "sno": sno, "bold": bool(cell.font and cell.font.bold) or bool(
                 col.get("sno") and wf.cell(row=r, column=col["sno"]).font.bold),
             "indent": int(cell.alignment.indent or 0) if cell.alignment else 0,
             "name": flat if len(flat) <= 100 else short_name(text), "description": flat if len(flat) > 100 else "",
             "boq_ref": str(val(r, "vno") or "").strip()[:40], "uom": str(val(r, "uom") or "").strip()[:20],
             "quantity": _number(val(r, "qty")) or 0.0, "labour": _number(val(r, "labour")) or 0.0,
             "start": _iso(val(r, "start")), "end": _iso(val(r, "end")), "days": None, "percent": None,
             "links": [], "parent": None, "heading": False}
        n["fs"], n["fe"], n["fd"] = formula(r, "start"), formula(r, "end"), formula(r, "days")
        n["range"] = _formula_range(r, n["fs"], n["fe"], n["fd"])
        typed_days = _number(val(r, "days"))
        pct = _number(val(r, "pct"))
        if pct is not None:
            n["percent"] = pct * 100 if pct <= 1 else pct
        if own:
            n["own_id"], n["own_parent"] = val(r, "id"), val(r, "parent")
            n["days"] = int(typed_days) if typed_days is not None else None
        n["typed_days"] = typed_days
        n["pred_text"] = str(val(r, "pred") or "").strip() if "pred" in col else ""
        nodes.append(n)
    if not nodes:
        raise HTTPException(400, "That sheet has a header but no rows under it.")

    _build_tree(nodes, own)
    by_key = {n["key"]: n for n in nodes}
    kids = {}
    for n in nodes:
        kids.setdefault(n["parent"], []).append(n)
    for n in nodes:
        n["heading"] = n["key"] in kids
    if own:
        _own_links(nodes, by_key, warnings)
    else:
        _formula_links(nodes, by_key, letter("start"), letter("end"), warnings)
    for n in nodes:
        if n["heading"]:
            n["start"], n["end"], n["days"], n["links"] = "", "", None, []
            continue
        if not own:
            if n["start"] and n["end"]:
                n["days"] = (date.fromisoformat(n["end"]) - date.fromisoformat(n["start"])).days + 1
            elif n["typed_days"] is not None:
                n["days"] = int(n["typed_days"])
            if n["days"] is not None and n["days"] < 0:
                n["days"] = None
        if n["links"]:
            n["start"] = ""
        if n["percent"] and not n["quantity"]:
            warnings.append("%s: its %% done was not read, because it has no quantity to measure progress in." % n["name"][:40])
            n["percent"] = None
    calendar = None
    if own and "Calendar" in values.sheetnames:
        calendar = _calendar(values["Calendar"])
    for n in nodes:
        for k in ("bold", "indent", "range", "fs", "fe", "fd", "typed_days", "pred_text", "own_id", "own_parent", "end", "sno"):
            n.pop(k, None)
    return {"sheet": name, "sheets": candidates, "nodes": nodes, "calendar": calendar, "warnings": warnings[:30], "own": own}


def _calendar(ws):
    mask = str(ws["B1"].value or "")
    off = [i + 1 for i, c in enumerate(mask[:7]) if c == "1"]
    out = {"calendar_days": str(ws["B2"].value or "").strip().lower() == "yes", "weekly_off": off or [7], "holidays": []}
    r = 5
    while ws.cell(row=r, column=1).value:
        d = _iso(ws.cell(row=r, column=1).value)
        if d:
            out["holidays"].append({"date": d, "name": str(ws.cell(row=r, column=2).value or "")})
        r += 1
    return out


def _build_tree(nodes, own):
    """Each node's parent: from the hidden ID/Parent columns, else the headings' own formulas, else bold and SL.NO."""
    if own:
        keys = {n["own_id"]: n["key"] for n in nodes if n.get("own_id") is not None}
        for n in nodes:
            n["parent"] = keys.get(n.get("own_parent")) if n.get("own_parent") is not None else None
        return
    if any(n["range"] for n in nodes):
        stack = []
        for n in nodes:
            while stack and stack[-1][0] < n["key"]:
                stack.pop()
            n["parent"] = stack[-1][1] if stack else None
            if n["range"]:
                stack.append((n["range"][1], n["key"]))
        return
    if any(n["indent"] for n in nodes):
        stack = []                                           # (indent, key)
        for n in nodes:
            while stack and stack[-1][0] >= n["indent"]:
                stack.pop()
            n["parent"] = stack[-1][1] if stack else None
            stack.append((n["indent"], n["key"]))
        return
    open_bold = []                                           # (kind, key) of the headings still open
    for n in nodes:
        if not n["bold"]:
            n["parent"] = open_bold[-1][1] if open_bold else None
            continue
        kind = _kind(n["sno"])
        if kind and any(k == kind for k, _ in open_bold):
            while open_bold and open_bold[-1][0] != kind:
                open_bold.pop()
            open_bold.pop()
        n["parent"] = open_bold[-1][1] if open_bold else None
        open_bold.append((kind, n["key"]))


def _own_links(nodes, by_key, warnings):
    wbs = {n["sno"]: n["key"] for n in nodes if n["sno"]}
    for n in nodes:
        if not n["pred_text"] or n["heading"]:
            continue
        found, problems = parse_predecessors(n["pred_text"], wbs)
        n["links"] = [{"pred": p, "kind": k, "lag": lag} for p, k, lag in found if p != n["key"]]
        warnings.extend("%s: %s" % (n["name"][:30], p) for p in problems)


def _ancestors(by_key, n):
    out, p = set(), n["parent"]
    while p is not None and p not in out:
        out.add(p)
        p = by_key[p]["parent"]
    return out


def _formula_links(nodes, by_key, start_col, end_col, warnings):
    """A leaf's Start written as a formula on another row is a dependency on it."""
    for n in nodes:
        f = n["fs"]
        if n["heading"] or not f or "TODAY" in f.upper():
            continue
        up = f.upper().replace(" ", "")
        lag_after = re.search(r"\)\s*([+-]\d+)\s*$", up)
        if "MAX(" in up:                                     # =MAX(D9,D12)+1: it waits for the last of them
            refs = [(c, int(r)) for c, r in re.findall(r"\$?([A-Z]{1,2})\$?(\d+)", up)]
            extra = int(lag_after.group(1)) if lag_after else 0
        else:
            m = START_FORMULA.fullmatch(f.upper())
            if not m:
                continue
            refs = [(m.group(1), int(m.group(2)))]
            extra = int(m.group(4)) * (-1 if m.group(3) == "-" else 1) if m.group(4) else 0
        mine = _ancestors(by_key, n)
        for c, row in refs:
            target = by_key.get(row)
            if target is None or target["key"] == n["key"] or target["key"] in mine:
                continue
            if c == end_col:
                n["links"].append({"pred": row, "kind": "FS", "lag": extra - 1})
            elif c == start_col:
                n["links"].append({"pred": row, "kind": "SS", "lag": extra})


def _clean(nodes):
    """Nodes as the screen sends them back: only the fields a node has, of the right kind, parents before children."""
    out, seen = [], set()
    for n in (nodes or [])[:MAX_ROWS]:
        try:
            key = int(n["key"])
        except (KeyError, TypeError, ValueError):
            raise HTTPException(400, "The rows to import are not in the form the preview gave.")
        parent = n.get("parent")
        parent = int(parent) if parent is not None and int(parent) in seen else None
        seen.add(key)
        days = n.get("days")
        out.append({
            "key": key, "parent": parent, "name": (str(n.get("name") or "").strip() or "Activity")[:200],
            "description": str(n.get("description") or "")[:4000], "boq_ref": str(n.get("boq_ref") or "")[:40],
            "uom": str(n.get("uom") or "")[:20], "quantity": max(0.0, _number(n.get("quantity")) or 0.0),
            "labour": max(0.0, _number(n.get("labour")) or 0.0), "start": _iso(n.get("start")),
            "days": max(0, int(days)) if _number(days) is not None else None,
            "percent": _number(n.get("percent")), "heading": bool(n.get("heading")),
            "links": [{"pred": int(x["pred"]), "kind": str(x.get("kind") or "FS").upper()[:2],
                       "lag": int(x.get("lag") or 0)} for x in (n.get("links") or []) if x.get("pred") is not None
                      and str(x.get("kind") or "FS").upper() in ("FS", "SS", "FF", "SF")]})
    return out


def apply_import(db, client, job, nodes, calendar=None):
    """Make the checked rows real, after what is on the programme already. An empty programme takes the sheet's working
    week; one with rows keeps its own. Returns (rows made, links made, what to tell the user)."""
    nodes = _clean(nodes)
    if not nodes:
        raise HTTPException(400, "There is nothing to import.")
    acts = job_rows(db, client.id, job)
    if not acts:
        c = calendar or {"calendar_days": True, "weekly_off": [7]}
        set_calendar(db, client, job, c.get("calendar_days"), c.get("weekly_off"))
        for h in c.get("holidays") or []:
            if _iso(h.get("date")):
                add_holiday(db, client, job, _iso(h["date"]), str(h.get("name") or ""))
    made = {}
    for n in nodes:
        a = models.DBScheduleActivity(
            client_id=client.id, job_id=job.id, name=n["name"], description=n["description"], boq_ref=n["boq_ref"],
            uom=n["uom"], quantity=0.0 if n["heading"] else n["quantity"], labour_required=n["labour"])
        if not n["heading"]:
            a.planned_start = n["start"]
            a.duration_days = n["days"]
            if n["percent"] and n["quantity"]:
                a.qty_done = round(n["quantity"] * min(n["percent"], 100.0) / 100.0, 3)
        db.add(a)
        db.flush()
        acts.append(a)
        parent = made.get(n["parent"])
        _place(acts, a, parent.id if parent is not None else None)
        made[n["key"]] = a
    linked = 0
    for n in nodes:
        if n["heading"]:
            continue
        for ln in n["links"]:
            pred = made.get(ln["pred"])
            if pred is None or pred.id == made[n["key"]].id:
                continue
            db.add(models.DBScheduleLink(client_id=client.id, job_id=job.id, activity_id=made[n["key"]].id,
                                         predecessor_id=pred.id, kind=ln["kind"], lag=ln["lag"]))
            linked += 1
    _check_loops(db, client, job, {a.id for a in made.values()})
    db.commit()
    leaves = sum(1 for n in nodes if not n["heading"])
    return len(nodes), linked, "%d row%s imported (%d activit%s, %d dependenc%s)." % (
        len(nodes), "" if len(nodes) == 1 else "s", leaves, "y" if leaves == 1 else "ies", linked, "y" if linked == 1 else "ies")
