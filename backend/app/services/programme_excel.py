"""The programme as the office's Excel schedule, and the office's schedules read back in.

The export is laid out as the sheet is - SL.NO, V.NO (BOQ Ref), the item, Start, End, working days, U.O.M, quantity,
labour, predecessor, % done, then a column a day forming the bar chart - and it is live: a leaf's End is worked out
from its Start and Days skipping the days off (WORKDAY.INTL), a linked row's Start from its predecessors' cells, a
heading's dates are MIN and MAX of the rows under it, and the bars are conditional formats over the day columns. Edit
a Start or a Days in Excel and everything after it moves, as it does on the screen.

The import reads the office's own sheets, both kinds: the GH layout and the APTIDCO / finishing-works layout. The
tree comes from the sheet's own formulas where it has them (a heading's =MIN(C8:C45) / =+MAX(D8:D28) / =SUM(E9:E12),
or =+C14 / =+D17 pointing at its first and last rows), else from bold rows and the SL.NO (I, A, 1); a Start written as
a formula on another row becomes a dependency (=+D9 starts the day row 9 ends: FS-1; =+D11+1 the day after: FS+0;
=+C23+2 two days after it starts: SS+2). A sheet this export wrote comes back exactly, by its hidden ID and Parent.
"""
import io
import re
from datetime import date, datetime, timedelta

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

FIRST = 7                                 # the first programme row; the headings are on row 6
COLS = ["SL.NO", "V.NO", "ITEMES / ACTIVITY", "Start", "End", "No of working Days", "U.O.M", "Quantity",
        "Labour Required", "Predecessor", "% Done", "ID", "Parent", "Weight"]
# A..N: the columns above; the day columns start at O.
SNO, VNO, ITEM, START, END, DAYS, UOM, QTY, LAB, PRED, PCT, ID, PARENT, WEIGHT = (get_column_letter(i + 1) for i in range(14))
GANTT = 15
MAX_DAY_COLUMNS = 1100                    # past three years a column a week instead

ORANGE, DARK, GREY, RED, HEAD = "F4A261", "B5651D", "D9D9D9", "E63946", "404040"
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def _mask(weekly_off):
    """WORKDAY.INTL's weekend string: Monday first, 1 for a day off ("0000001": Sundays off)."""
    return "".join("1" if d in weekly_off else "0" for d in range(1, 8))


def _date(cell_value):
    if isinstance(cell_value, datetime):
        return cell_value.date()
    if isinstance(cell_value, date):
        return cell_value
    return None


def export_xlsx(view):
    """The programme_view as a live workbook (bytes)."""
    rows = view["rows"]
    cal = view["calendar"]
    calendar_days = bool(cal["calendar_days"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Programme"

    # --- the calendar the formulas skip ---------------------------------------------------------------------------
    cs = wb.create_sheet("Calendar")
    cs["A1"], cs["B1"] = "Days off each week (Mon..Sun, 1 = off)", _mask(cal["weekly_off"])
    cs["A2"], cs["B2"] = "Count calendar days", "Yes" if calendar_days else "No"
    cs["A4"], cs["B4"] = "Holiday", "Name"
    for c in ("A1", "A2", "A4", "B4"):
        cs[c].font = Font(bold=True)
    hols = cal["holidays"] or []
    for i, h in enumerate(hols):
        cs.cell(row=5 + i, column=1, value=date.fromisoformat(h["date"])).number_format = "dd-mmm-yyyy"
        cs.cell(row=5 + i, column=2, value=h.get("name") or "")
    cs.column_dimensions["A"].width, cs.column_dimensions["B"].width = 36, 30
    wb.defined_names["WeekMask"] = DefinedName("WeekMask", attr_text="Calendar!$B$1")
    wb.defined_names["Holidays"] = DefinedName("Holidays", attr_text="Calendar!$A$5:$A$%d" % (5 + max(len(hols), 1) - 1))

    # --- the header block, as the sheet has it ----------------------------------------------------------------------
    n = len(rows)
    last = FIRST + max(n, 1) - 1
    ws["A1"] = view.get("company") or ""
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = " SCHEDULE"
    ws["A2"].font = Font(bold=True, size=13)
    ws["D2"] = "TIME DURATION IN DAYS"
    ws["D2"].font = Font(bold=True)
    _formula(ws, "F2", "=IF(COUNT(%s%d:%s%d)=0,\"\",MAX(%s%d:%s%d)-MIN(%s%d:%s%d)+1)" % (START, FIRST, START, last, END, FIRST, END, last,
                                                                                START, FIRST, START, last))
    ws["F2"].font = Font(bold=True, size=13)
    job = view["job"]
    ws["A3"] = "Name of the work :- %s %s" % (job.get("number") or "", job.get("name") or "")
    ws["A4"] = " Location : %s" % (job.get("location") or "")
    for c in ("A3", "A4"):
        ws[c].font = Font(bold=True)
    if view.get("baseline"):
        ws["D3"] = "Baseline: %s (%s)" % (view["baseline"]["label"], (view["baseline"]["saved_at"] or "")[:10])

    for i, title in enumerate(COLS):
        c = ws.cell(row=6, column=i + 1, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = box
    for col, width in zip("ABCDEFGHIJKLMN", (9, 10, 46, 12, 12, 9, 8, 11, 9, 16, 8, 6, 6, 10)):
        ws.column_dimensions[col].width = width
    for col in (ID, PARENT, WEIGHT):
        ws.column_dimensions[col].hidden = True
    ws.row_dimensions[6].height = 32

    # --- the rows --------------------------------------------------------------------------------------------------
    at = {r["id"]: FIRST + i for i, r in enumerate(rows)}
    last_under = {}
    for i, r in enumerate(rows):                     # a heading's rows run to its last descendant (rows are in tree order)
        j = i
        while j + 1 < n and rows[j + 1]["level"] > r["level"]:
            j += 1
        last_under[r["id"]] = FIRST + j
    incoming = {}
    for ln in view.get("links", []):
        if ln not in view.get("broken_links", []) and ln["predecessor"] in at:
            incoming.setdefault(ln["successor"], []).append(ln)

    def shift(ref, k):
        """The working day k days after ref (k may be 0 or below)."""
        if calendar_days:
            return "%s%+d" % (ref, k) if k else ref
        return "WORKDAY.INTL(%s,%d,WeekMask,Holidays)" % (ref, k)

    def snap(expr):
        """The first working day on or after expr."""
        return expr if calendar_days else "WORKDAY.INTL(%s-1,1,WeekMask,Holidays)" % expr

    rng = lambda col: "$%s$%d:$%s$%d" % (col, FIRST, col, last)
    for r in rows:
        y = at[r["id"]]
        ws["%s%d" % (SNO, y)] = r["wbs"]
        ws["%s%d" % (VNO, y)] = r["boq_ref"] or None
        name = ws["%s%d" % (ITEM, y)]
        name.value = r["name"]
        name.alignment = Alignment(indent=min(r["level"] * 2, 15))
        ws["%s%d" % (UOM, y)] = r["uom"] or None
        ws["%s%d" % (QTY, y)] = r["quantity"] or None
        ws["%s%d" % (LAB, y)] = r["labour_required"] or None
        ws["%s%d" % (PRED, y)] = r["predecessors"] or None
        ws["%s%d" % (ID, y)] = r["id"]
        ws["%s%d" % (PARENT, y)] = r["parent_id"]
        s, e, d, p = ("%s%d" % (c, y) for c in (START, END, DAYS, PCT))
        if r["summary"]:
            first, end_row = y + 1, last_under[r["id"]]
            span = lambda col: "%s%d:%s%d" % (col, first, col, end_row)
            _formula(ws, s, "=IF(COUNT(%s)=0,\"\",MIN(%s))" % (span(START), span(START)))
            _formula(ws, e, "=IF(COUNT(%s)=0,\"\",MAX(%s))" % (span(END), span(END)))
            _formula(ws, d, "=IF(%s=\"\",\"\",%s)" % (s, "%s-%s+1" % (e, s) if calendar_days else "NETWORKDAYS.INTL(%s,%s,WeekMask,Holidays)" % (s, e)))
            # weighted as the screen weighs it: by amount when every row under it has one, else by working days. Only the
            # rows under it are looked at, so a heading never counts itself (no circular reference).
            under = lambda col: "$%s$%d:$%s$%d" % (col, first, col, end_row)
            kids = "%s,$%s%d" % (under(PARENT), ID, y)
            _formula(ws, "%s%d" % (WEIGHT, y), "=SUMIF(%s,%s)" % (kids, under(WEIGHT)))
            _formula(ws, p, "=IF(COUNTIF(%s)=0,0,IF(COUNTIFS(%s,%s,\"<=0\")=0,SUMPRODUCT((%s=$%s%d)*%s*%s)/SUMIF(%s,%s),"
                     "IF(SUMIF(%s,%s)>0,SUMPRODUCT((%s=$%s%d)*IFERROR(%s*1,0)*%s)/SUMIF(%s,%s),AVERAGEIF(%s,%s))))" % (
                         kids, kids, under(WEIGHT), under(PARENT), ID, y, under(WEIGHT), under(PCT), kids, under(WEIGHT),
                         kids, under(DAYS), under(PARENT), ID, y, under(DAYS), under(PCT), kids, under(DAYS), kids, under(PCT)))
            for c in (SNO, VNO, ITEM, START, END, DAYS, PCT):
                ws["%s%d" % (c, y)].font = Font(bold=True)
        else:
            ws[d] = r["days"] if r["days"] is not None else (r["typed"]["days"] if r["typed"]["days"] is not None else None)
            needs = []
            for ln in incoming.get(r["id"], []):
                py = at[ln["predecessor"]]
                pe, ps = "%s%d" % (END, py), "%s%d" % (START, py)
                lag = int(ln["lag"] or 0)
                if ln["kind"] == "FS":
                    t = shift(pe, 1 + lag)
                elif ln["kind"] == "SS":
                    t = shift(ps, lag)
                elif ln["kind"] == "FF":
                    t = _back(shift(pe, lag), d, calendar_days)
                else:                                                         # SF
                    t = _back(shift(ps, lag), d, calendar_days)
                needs.append("IF(%s=\"\",0,%s)" % (pe if ln["kind"] in ("FS", "FF") else ps, t))
            floor = [_date_formula(x) for x in (r["typed"]["start"], r["not_before"]) if x]
            if needs:
                m = "MAX(%s)" % ",".join(needs + floor)
                _formula(ws, s, "=IF(%s=0,\"\",%s)" % (m, snap(m)))
            elif r["start"]:
                ws[s] = date.fromisoformat(r["start"])
            end_of = ("%s+MAX(%s-1,0)" % (s, d)) if calendar_days else "IF(%s<=1,%s,WORKDAY.INTL(%s,%s-1,WeekMask,Holidays))" % (d, s, s, d)
            _formula(ws, e, "=IF(OR(%s=\"\",%s=\"\"),\"\",%s)" % (s, d, end_of))
            ws[p] = round((r["percent"] or 0) / 100.0, 4)
            amount = r["amount"] or (r["quantity"] * r["rate"] if r["quantity"] and r["rate"] else 0)
            ws["%s%d" % (WEIGHT, y)] = amount
        for c in (START, END):
            ws["%s%d" % (c, y)].number_format = "dd-mmm-yy"
        ws[p].number_format = "0%"
        for c in "ABCDEFGHIJK":
            ws["%s%d" % (c, y)].border = box

    # --- the day columns and the bars ----------------------------------------------------------------------------
    begin = date.fromisoformat(view["start"]) if view.get("start") else date.today()
    finish = date.fromisoformat(view["finish"]) if view.get("finish") else begin + timedelta(days=30)
    today = date.today()
    begin, finish = min(begin, today) - timedelta(days=2), max(finish, today) + timedelta(days=7)
    weekly = (finish - begin).days + 1 > MAX_DAY_COLUMNS
    step = 7 if weekly else 1
    if weekly:
        begin -= timedelta(days=begin.weekday())
    count = (finish - begin).days // step + 1
    end_col = get_column_letter(GANTT + count - 1)
    first_col = get_column_letter(GANTT)
    month = None
    for k in range(count):
        day = begin + timedelta(days=k * step)
        col = get_column_letter(GANTT + k)
        h = ws["%s6" % col]
        h.value = day
        h.number_format = "d"
        h.font = Font(bold=True, size=8, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor=HEAD)
        h.alignment = Alignment(horizontal="center", text_rotation=0)
        ws.column_dimensions[col].width = 3.2 if not weekly else 4
        if (day.year, day.month) != month:
            month = (day.year, day.month)
            ws["%s5" % col] = day.strftime("%b %Y")
            ws["%s5" % col].font = Font(bold=True, size=9)
    area = "%s%d:%s%d" % (first_col, FIRST, end_col, last)
    c0, r0 = first_col, FIRST
    within = "AND($%s%d<>\"\",%s$6<=$%s%d,%s$6%s>=$%s%d)" % (START, r0, c0, END, r0, c0, "+6" if weekly else "", START, r0)
    heading = "COUNTIF($%s$%d:$%s$%d,$%s%d)>0" % (PARENT, FIRST, PARENT, last, ID, r0)
    rules = [
        ("%s$6=TODAY()" % c0 if not weekly else "AND(%s$6<=TODAY(),%s$6+6>=TODAY())" % (c0, c0), PatternFill("solid", fgColor=RED, bgColor=RED)),
        ("AND(%s,$%s%d<TODAY(),$%s%d<1,NOT(%s))" % (within, END, r0, PCT, r0, heading), PatternFill("solid", fgColor="F28B82", bgColor="F28B82")),
        ("AND(%s,%s)" % (within, heading), PatternFill("solid", fgColor="595959", bgColor="595959")),
        ("AND(%s,%s$6<$%s%d+($%s%d-$%s%d+1)*$%s%d)" % (within, c0, START, r0, END, r0, START, r0, PCT, r0), PatternFill("solid", fgColor=DARK, bgColor=DARK)),
        (within, PatternFill("solid", fgColor=ORANGE, bgColor=ORANGE)),
    ]
    if not weekly and not calendar_days:
        rules.append(("NETWORKDAYS.INTL(%s$6,%s$6,WeekMask,Holidays)=0" % (c0, c0), PatternFill("solid", fgColor=GREY, bgColor=GREY)))
    for formula, fill in rules:
        ws.conditional_formatting.add(area, FormulaRule(formula=[formula], fill=fill, stopIfTrue=True))

    ws.freeze_panes = "D7"
    ws.print_title_rows = "1:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _formula(ws, ref, text):
    """Write one of this export's own formulas. Every workbook the app writes stores text that looks like a formula as
    text (app/core/excel_guard.py), so a name typed as "=..." can never run; the formulas built here hold nothing but
    cell references, numbers and dates, and are marked as formulas one cell at a time."""
    cell = ws[ref]
    cell.value = text
    cell.data_type = "f"
    return cell


def _back(ref_expr, days_cell, calendar_days):
    """The start that makes a row of `days_cell` days end on ref_expr."""
    if calendar_days:
        return "%s-MAX(%s-1,0)" % (ref_expr, days_cell)
    return "WORKDAY.INTL(%s,-MAX(%s-1,0),WeekMask,Holidays)" % (ref_expr, days_cell)


def _date_formula(iso):
    d = date.fromisoformat(iso[:10])
    return "DATE(%d,%d,%d)" % (d.year, d.month, d.day)
