"""Reading a spreadsheet somebody actually sent, whatever its layout."""
import contextvars
import csv
import io
import re
from datetime import datetime

from fastapi import HTTPException, UploadFile
from fastapi.responses import Response

from app.core import identity
from app.validators import import_guard
from app import models
from app.db import SessionLocal

from app.core.auth import CURRENT_CLIENT_ID
from app.documents.pdf_twins import SHEET_AS_PDF


def sheet_response(headers, sample, filename, fmt="xlsx", preamble=None, closing=None, branded=True):
    """The template, as a real workbook by default.

    A workbook is what people asked for and what they will edit; CSV stays
    available for anything that has to be read by a script.

    A report may pass a preamble - who it is for, when it was printed, which
    order it covers - and a closing set of totals. They are written above and
    below the table rather than into it, so what sits between the headings and
    the totals stays a rectangle that can still be sorted and filtered.
    """
    preamble = list(preamble or [])
    closing = list(closing or [])

    # Asked for at the .pdf twin of this address: the same rows, on paper.
    as_pdf = SHEET_AS_PDF.get()
    if as_pdf is not None:
        return form_pdf_response(sheet_report_spec(headers, sample, filename, preamble, closing,
                                                   as_pdf.get("client")), filename.rsplit(".", 1)[0])

    if fmt == "csv" or filename.endswith(".csv"):
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerows(preamble)
        writer.writerow(headers)
        writer.writerows(sample)
        writer.writerows(closing)
        return Response(
            # utf-8-sig so Excel opens it in the right encoding rather than
            # mangling anything non-ASCII in a description.
            content=buf.getvalue().encode("utf-8-sig"),
            media_type="text/csv",
            headers={"Content-Disposition": 'attachment; filename="%s"' % filename},
        )

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return sheet_response(headers, sample, filename.rsplit(".", 1)[0] + ".csv",
                              "csv", preamble, closing)

    # Whatever the company issues carries its name, address, tax numbers and logo - but a template is
    # read back in, so it stays exactly the shape it is read in.
    company = None
    # branded="logo": a report whose own layout is fixed keeps its rows and takes only the logo.
    if branded and "template" not in filename and CURRENT_CLIENT_ID.get():
        try:
            with SessionLocal() as own:
                owner = own.query(models.DBClient).filter(models.DBClient.id == CURRENT_CLIENT_ID.get()).first()
                company = letterhead(own, owner) if owner else None
        except Exception:
            company = None
    head_rows = identity.sheet_head_rows(company) if company and branded is True else []
    if head_rows:
        preamble = head_rows + [[]] + preamble

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    for row in preamble:
        sheet.append(list(row))
    header_row = len(preamble) + 1
    sheet.append(list(headers))
    for row in sample:
        sheet.append(list(row))
    for row in closing:
        sheet.append(list(row))

    if preamble:
        sheet.cell(row=1, column=1).font = Font(bold=True, size=13)
        if head_rows:
            for n in range(2, len(head_rows) + 1):
                sheet.cell(row=n, column=1).font = Font(size=10)
            if len(preamble) > len(head_rows) + 1:
                sheet.cell(row=len(head_rows) + 2, column=1).font = Font(bold=True, size=12)
    header_font = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4F46E5")
    for cell in sheet[header_row]:
        cell.font = header_font
        cell.fill = fill
    # Width from the longest value in each column, so nothing opens as ####.
    for index, name in enumerate(headers, start=1):
        longest = max([len(str(name))] + [len(str(r[index - 1])) for r in sample
                                          if index - 1 < len(r)])
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = min(40, longest + 4)
    # Freeze under the headings, so the columns stay named while the rows move.
    sheet.freeze_panes = "A%d" % (header_row + 1)
    if company:
        if head_rows:
            for n in range(1, 4):
                sheet.row_dimensions[n].height = 20
        identity.add_logo(sheet, company, "%s1" % openpyxl.utils.get_column_letter(max(len(headers), 4) + 1))

    stream = io.BytesIO()
    book.save(stream)
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="%s"' % filename},
    )


async def parse_sheet(upload: UploadFile, columns, aliases=None, sheet=""):
    """Read an uploaded sheet into dicts keyed by our column names.

    By header where the sheet is headed in words we know, by position only as a
    fallback. Position used to be the whole rule, and it had two ways of being
    wrong at once. It read the workbook we hand out as a template by decoding
    the zip as latin-1 and letting the csv module make rows of the wreckage,
    which failed as a mess of unreadable codes rather than as an error. And on
    the work order template people actually fill in - whose fifth column is a
    price where ours is a unit - it banked the price as the unit of measure and
    the discount as the rate, silently, on a document that prices a contract.
    """
    header, body = await read_sheet_rows(upload, sheet)

    mapping = {}
    if aliases:
        mapping, _ = match_headers(header, aliases)
    # One lucky hit is not recognition. A sheet we cannot read the headings of
    # is still read in our own column order, which is what the template says.
    if len(mapping) < 2:
        mapping = {index: name for index, name in enumerate(columns)}

    return rows_from(header, body, mapping)


#
# The template is a suggestion, not a contract. Suppliers send their own
# layouts, Excel reorders columns, and people type "MTR" where the dropdown
# said "Meters". Rejecting all of that is easy and useless; the work is in
# reading it correctly and saying exactly what was assumed.
# Header text -> our column. Matched on letters and digits only, so spacing,
# case, punctuation and the usual "Sr.No"-style noise all fall away.
HEADER_ALIASES = {
    "item_code": ["itemcode", "code", "materialcode", "materialno", "partno",
                  "partnumber", "sku", "erpcode", "itemno", "productcode"],
    "item_name": ["itemname", "name", "material", "materialname", "particulars",
                  "itemdescription", "product", "productname"],
    "segment": ["segment", "division", "businessunit", "bu"],
    "description": ["description", "desc", "longdescription", "details", "specification"],
    "category": ["category", "itemcategory", "maincategory", "group"],
    "sub_category": ["subcategory", "subcat", "itemsubcategory", "subgroup", "type"],
    "hsn_code": ["hsncode", "hsn", "hsnsac", "sac", "hscode"],
    "item_tax_type": ["itemtaxtype", "tax", "taxrate", "gst", "gstrate", "taxpercent", "vat"],
    "item_type": ["itemtype", "procurementtype", "sourcetype", "naturetype"],
    "units_of_measure": ["unitsofmeasure", "uom", "unit", "units", "measure", "unitofmeasure"],
    "make": ["make", "brand", "manufacturer", "makelist"],
}


def squash(text) -> str:
    """Letters and digits only, lowercased - the form headers are compared in."""
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def header_score(key, field, aliases) -> int:
    """How well a header's squashed text names a field. 0 means it does not."""
    if key == squash(field):
        return 3                # the field's own name
    if key in aliases:
        return 2                # a name we were told to expect
    for alias in aliases:
        if len(alias) >= 4 and (alias in key or key in alias):
            return 1            # a partial, and the last resort
    return 0


def match_headers(header_row, alias_map):
    """Work out which column is which, by name rather than by position.

    Positional parsing breaks silently the moment somebody inserts a column,
    and the data lands one field to the left with no complaint. Matching on the
    header means a reordered - or entirely foreign - sheet still reads right.

    Scored rather than first-match: "Units" is the exact name of a unit column
    and a loose match for a quantity one, and whichever field happened to be
    written first in the vocabulary used to win. Settling the best matches
    first means an exact name always beats a partial one, whatever the order.
    """
    scored = []
    for index, raw in enumerate(header_row):
        key = squash(raw)
        if not key:
            continue
        for field, aliases in alias_map.items():
            points = header_score(key, field, aliases)
            if points:
                scored.append((points, index, list(alias_map).index(field), field))

    # Best score first, then left to right. So where one header names two
    # fields equally well - a budget sheet with two "Product Code" columns, the
    # ordered item and then the material - the leftmost takes the first field
    # the vocabulary declares, and the next takes the one after it.
    scored.sort(key=lambda s: (-s[0], s[1], s[2]))

    mapping, taken = {}, set()
    for _, index, _, field in scored:
        if index in mapping or field in taken:
            continue
        mapping[index] = field
        taken.add(field)

    unmapped = [str(raw).strip() for index, raw in enumerate(header_row)
                if squash(raw) and index not in mapping]
    return mapping, unmapped


def looks_like_xlsx(raw: bytes) -> bool:
    """An .xlsx is a zip archive, so it starts with the zip magic number.

    Sniffing the bytes rather than trusting the filename: people rename files,
    and a .csv that is really a workbook should still open.
    """
    return raw[:4] == b"PK\x03\x04"


# Which sheet of which workbook the request in hand is reading. A ContextVar
# rather than a global because requests are served concurrently, and two
# uploads landing together must not describe each other's file.
_SHEET_CONTEXT = contextvars.ContextVar("sheet_context", default=None)


def workbook_context(raw: bytes, sheet: str):
    """The name of the sheet being read, and the others that were available."""
    try:
        book = open_workbook(raw)
    except HTTPException:
        return None
    try:
        names = list(book.sheetnames)
        _, chosen = pick_sheet(book, sheet)
        return {"sheet": chosen, "sheets": names}
    finally:
        book.close()


def sheet_note() -> str:
    """A sentence naming the tab that was read, when there was a choice.

    Appended to the refusals about missing columns. "No item code column
    found" is a fair complaint about the sheet it read and a baffling one
    about the file, which may well have the codes on the next tab along.
    """
    ctx = _SHEET_CONTEXT.get()
    if not ctx or len(ctx["sheets"]) < 2:
        return ""
    others = [n for n in ctx["sheets"] if n != ctx["sheet"]]
    return (" This was read from the '%s' sheet; the file also has %s. "
            "Open it under Contracts → Open a file to look at the others."
            % (ctx["sheet"], ", ".join("'%s'" % n for n in others)))


def open_workbook(raw: bytes):
    """The workbook, or a plain refusal saying why it would not open."""
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(
            400, "Excel workbooks are not supported on this server yet. "
                 "Save the sheet as CSV and upload that.")
    try:
        return openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise HTTPException(400, "That workbook could not be opened (%s)." % exc)


def cell_text(value):
    """One cell as the text it was written as.

    Everything downstream compares and cleans text, so the numbers Excel hands
    back as floats are rendered the way they were written - 5000 rather than
    5000.0, which would otherwise reach a code column and not match.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def pick_sheet(book, wanted):
    """The worksheet asked for, by name or by position, else the first.

    A workbook off somebody's desk is rarely one sheet. A programme arrives
    with eight - a Gantt chart, two revisions of it, a manpower tab - and
    reading whichever happened to be saved first means the file is judged on
    a sheet nobody meant to send.
    """
    names = book.sheetnames
    wanted = (str(wanted) if wanted is not None else "").strip()
    if not wanted:
        return book[names[0]], names[0]
    for name in names:                          # by name, exactly
        if name == wanted:
            return book[name], name
    for name in names:                          # then case-insensitively
        if name.strip().lower() == wanted.lower():
            return book[name], name
    if re.fullmatch(r"\d+", wanted):            # then by position
        index = int(wanted)
        if 0 <= index < len(names):
            return book[names[index]], names[index]
    raise HTTPException(
        400, "This workbook has no sheet called '%s'. It has: %s."
             % (wanted, ", ".join("'%s'" % n for n in names)))


def read_xlsx_table(raw: bytes, sheet: str = ""):
    """Rows from one worksheet, as strings."""
    book = open_workbook(raw)
    try:
        worksheet, _ = pick_sheet(book, sheet)
        return [[cell_text(c) for c in row] for row in worksheet.iter_rows(values_only=True)]
    finally:
        book.close()


# The templates people actually fill in carry a row of instructions under the
# header - "Without Spaces", "Please Select Option From Dropdown" - and it
# stays there, because it is what tells the typist what to enter. Read as data
# it becomes a row whose item code is "Without Spaces", and it fails every
# check on every upload for ever.
HINT_PHRASES = re.compile(
    r"please\s*select|drop\s*down|dropdown|without\b|only\s*num|mandatory|"
    r"do\s*not\b|choose\s*from|select\s*from|for\s*example|as\s*shown|^e\.?g\.?\b",
    re.IGNORECASE)


def is_hint_row(values) -> bool:
    """Whether a row is guidance for the typist rather than data.

    Every filled cell has to read as an instruction. One real value - a code, a
    quantity, a name - and the row is data again, which is what stops a
    description that happens to say "do not exceed" from taking its row with it.
    """
    filled = [str(v).strip() for v in values if str(v or "").strip()]
    return bool(filled) and all(HINT_PHRASES.search(v) for v in filled)


async def read_sheet_rows(upload: UploadFile, sheet: str = ""):
    """Split an upload into its header row and its data rows, unparsed.

    Takes a real Excel workbook or a CSV; which one is decided by the bytes,
    not the extension. Each data row is returned with the line it came from,
    because blank rows and the hint row are dropped on the way past - and a
    complaint about "line 1945" has to mean the row somebody can scroll to,
    not the row it happened to end up at once we had thrown some away.
    """
    raw = await upload.read()
    if len(raw) > 8_000_000:
        raise HTTPException(400, "That file is too large. Split it and upload in parts.")
    if raw[:8] == import_guard.OLE_SIGNATURE:
        # An old .xls, or a workbook with a password: neither can be read as text, and decoding them as a
        # CSV gave a screen of unreadable codes for an error.
        try:
            import_guard.check_workbook_bytes(raw, upload.filename)
        except ValueError as exc:
            raise HTTPException(400, str(exc))

    if looks_like_xlsx(raw):
        try:
            import_guard.check_workbook_bytes(raw, upload.filename)
        except ValueError as exc:
            raise HTTPException(400, str(exc))
        table = read_xlsx_table(raw, sheet)
        # Remembered so a complaint further down can say which of a workbook's
        # tabs it was actually looking at. Being told a column is missing is
        # no use at all when the file has eight sheets and you never got to
        # say which one you meant.
        _SHEET_CONTEXT.set(workbook_context(raw, sheet))
    else:
        _SHEET_CONTEXT.set(None)
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        try:
            # newline="" because Excel writes CRLF; without it the CR is read
            # as part of the field and the whole file is refused.
            table = list(csv.reader(io.StringIO(text, newline="")))
        except csv.Error as exc:
            raise HTTPException(400, f"That file could not be read as a sheet ({exc}). "
                                     "Save it as CSV or .xlsx and try again.")

    numbered = [(line, row) for line, row in enumerate(table, start=1)
                if any((c or "").strip() for c in row)]
    if len(numbered) < 2:
        raise HTTPException(400, "That file had no rows under its header.")

    header, body = numbered[0][1], numbered[1:]
    if is_hint_row(body[0][1]):
        body = body[1:]
    if not body:
        raise HTTPException(400, "That file had no rows under its header.")
    return header, body


#
# Items already get read, repaired and shown before anything is saved. A work
# order or a budget arriving from a customer or an estimator deserves the same:
# they carry prices, and a sheet that goes straight in is a price nobody
# checked. Both analysers return lines in exactly the shape /build accepts, so
# the review grid commits through the same validated path as on-screen entry.
# "Nos" and "Units" are units, not quantities, and naming them as quantities
# is what made a budget sheet read its unit column as its quantity. The
# misspellings are the ones on the templates in circulation, kept because the
# file people have on their desk is the one that has to open.
WO_HEADER_ALIASES = {
    "fg_code": ["fgcode", "code", "itemcode", "productcode", "sku", "erpcode"],
    "item_name": ["itemname", "name", "productname", "particulars", "scope",
                  "workdescription"],
    "description": ["description", "productdiscription", "discription",
                    "longdescription", "details", "remarks", "notes", "specification"],
    "qty": ["qty", "quantity", "volume", "orderqty", "woqty"],
    "uom": ["uom", "unit", "units", "measure", "unitofmeasure", "unitsofmeasure"],
    "rate": ["rate", "price", "unitrate", "unitprice", "amountperunit"],
}
# Their budget sheet heads the ordered item and the material it consumes with
# the same words - "Product Code" twice over. The scoring settles that by
# position: the leftmost takes fg_code because it is declared first here.
BOM_HEADER_ALIASES = {
    "fg_code": ["fgcode", "finishedgood", "sellingcode", "outputcode", "againstcode",
                "ordereditems", "ordereditem", "productcode", "itemcode"],
    "rm_code": ["rmcode", "materialcode", "rawmaterial", "inputcode", "consumescode",
                "componentcode", "component", "productcode", "itemcode"],
    "rm_name": ["rmname", "materialname", "material", "particulars", "description",
                "productname", "itemname"],
    "qty": ["qty", "quantity", "consumption", "usage", "bomqty", "bomquantity"],
    "uom": ["uom", "unit", "units", "measure", "unitofmeasure", "unitsofmeasure"],
    "rate": ["rate", "price", "cost", "unitrate", "unitcost", "finalrate"],
}


def map_headers_with(header_row, aliases):
    """Same name-matching as the item intake, against a different vocabulary."""
    return match_headers(header_row, aliases)


def mapping_report(header, mapping):
    """Which heading we read as which field, for showing back to the person.

    Keyed by the heading, which is what somebody checking the import wants to
    read. A sheet is allowed to use one heading twice, though - the budget
    sheet heads both the ordered item and the material it consumes "Product
    Code" - and keyed by text alone the second would land on top of the first,
    reporting one column where we read two. The column number separates them,
    because "we read both of these, and differently" is the whole of what is
    being checked.
    """
    seen, report = {}, {}
    for index in sorted(mapping):
        name = (str(header[index]).strip() if index < len(header) else "")
        name = name or "Column %d" % (index + 1)
        seen[name] = seen.get(name, 0) + 1
        if seen[name] > 1:
            name = "%s (column %d)" % (name, index + 1)
        report[name] = mapping[index]
    return report


def rows_from(header, body, mapping):
    """Rows keyed by our field names, with the ones that are not lines left out.

    A sheet off a real desk ends in its own totals: a number alone in a column
    we never read, sitting under fifteen hundred priced lines. Read as a line
    it has no code, no quantity and no rate, so it fails - and because nothing
    is saved unless every line passes, it takes the whole order down with it.
    A row with nothing in any column we mapped is not a line we are failing to
    read; it is a row that was never addressed to us.
    """
    rows = []
    for line, values in body:
        row = {field: (values[i].strip() if i < len(values) else "")
               for i, field in mapping.items()}
        if not any(row.values()):
            continue
        row["_line"] = line
        rows.append(row)
    return rows


# Names this module only needs once it is running. They come from modules that in turn need this one,
# so they are imported last, after everything above has been defined.
from app.documents.forms import form_pdf_response, sheet_report_spec
from app.documents.letterhead import letterhead
