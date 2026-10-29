"""The payslip as the office sets it: employee details in two column pairs, earnings beside deductions, net pay in
figures and in words. One layout feeds the PDF and the Excel download, so the two always carry the same columns.

Every column of the sheet is here. A column the application holds nothing for (father's name, HRA and the other
allowance heads, professional tax) is left empty rather than filled in, for the office to complete.
"""
import io
from datetime import datetime

from app.core.currency import amount_in_words
from app.documents import form_pdf

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

NOTE = "** This is computer generated payslip no need any signature"


def _date(iso):
    try:
        return datetime.strptime((iso or "")[:10], "%Y-%m-%d")
    except ValueError:
        return None


def month_year(iso):
    d = _date(iso)
    return "%s-%s" % (MONTHS[d.month - 1], str(d.year)[2:]) if d else ""


def joined_on(iso):
    d = _date(iso)
    return "%02d-%s-%d" % (d.day, MONTHS[d.month - 1], d.year) if d else ""


def _figure(value):
    """A money figure as written on the slip: whole rupees plainly, paise only when there are some."""
    n = round(float(value or 0), 2)
    return str(int(n)) if n == int(n) else "%.2f" % n


def _blank_or(value):
    return "" if value in (None, "") else value


def layout(p):
    """The payslip's rows from the payslip as the API gives it (see get_payslip)."""
    e = p.get("employee") or {}
    days = p.get("days") or {}
    work = days.get("work") or ""
    info = [
        ("Employee No.", e.get("employee_id", ""), "Location .", e.get("work_location", "")),
        ("Name.", e.get("full_name", ""), "Department.", e.get("department_name", "")),
        ("Bank Name.", e.get("bank_name", ""), "Designation.", e.get("job_title", "")),
        ("Bank Acc.No.", e.get("bank_account", ""), "Total Days.", days.get("total") or ""),
        ("Father Name.", "", "LOP .", ""),
        ("Salary Per Month.", _figure(p.get("gross_pay")), "Balance Leaves.", _blank_or(days.get("balance_leaves"))),
        ("DOJ.", joined_on(e.get("start_date")), "Work Days.", work),
        ("Aadhar No.", e.get("aadhaar_masked", ""), "Effective Work days.", work),
        ("UAN No", e.get("uan_no", ""), "ESI No", e.get("esi_no", "")),
    ]
    earnings = [
        ("Basic", p.get("basic_salary")), ("HRA", None), ("Site Allowance", None), ("Special Allownace", None),
        ("Mobile Allowance", None), ("Transport Allowance", None), ("Other Allowance", p.get("allowances")),
    ]
    # Pay the application records that the sheet has no column for stays on the slip, so the total still adds up.
    if p.get("overtime_pay"):
        hours = p.get("overtime_hours")
        earnings.append(("Overtime%s" % (" (%sh)" % _figure(hours) if hours else ""), p["overtime_pay"]))
    if p.get("bonus"):
        earnings.append(("Bonus", p["bonus"]))
    deductions = [
        ("PF", p.get("retirement")), ("ESI", p.get("insurance")), ("Prof Tax", None), ("Salary Advance", None),
        ("TDS Deduction", p.get("tax_amount")), ("Loss Of Pay", None), ("Other Deductions", p.get("other_deductions")),
    ]
    if (p.get("standing_deduction") or 0) > 0:
        deductions.append(("Standing Deduction", p["standing_deduction"]))
    rows = max(len(earnings), len(deductions))
    earnings += [("", None)] * (rows - len(earnings))
    deductions += [("", None)] * (rows - len(deductions))
    net = p.get("net_pay") or 0
    return {
        "company": (p.get("company") or {}),
        "title": "Payslip For The Month of %s" % month_year(p.get("period_end") or p.get("period_start")),
        "info": info, "earnings": earnings, "deductions": deductions,
        "total_earnings": _figure(p.get("gross_pay")), "total_deductions": _figure(p.get("total_deductions")),
        "net": _figure(net), "words": amount_in_words(net), "note": NOTE, "notes": p.get("notes") or "",
    }


def _cell(value):
    return "" if value is None else _figure(value) if isinstance(value, (int, float)) else str(value)


def pdf_spec(p):
    """The payslip described for the form engine: a six-column ruled sheet, the way the office's own is."""
    lay = layout(p)
    company = lay["company"]
    head = [(company["name"], "banner")] if company.get("name") else []
    if company.get("address"):
        head.append((company["address"], "centrebold"))
    head.append((lay["title"], "band"))
    # label | value | value | label | value | value: the value is two columns wide so earnings can sit beside deductions
    widths = [34, 30, 28, 34, 30, 28]
    rows, spans = [], []
    for l1, v1, l2, v2 in lay["info"]:
        r = len(rows)
        rows.append([l1, v1, "", l2, v2, ""])
        spans += [(1, r, 2, r), (4, r, 5, r)]
    r = len(rows)
    rows.append([{"t": "Earnings", "b": True}, "", "", {"t": "Deductions", "b": True}, "", ""])
    spans += [(0, r, 2, r), (3, r, 5, r)]
    shade = [r]
    for (el, ev), (dl, dv) in zip(lay["earnings"], lay["deductions"]):
        r = len(rows)
        rows.append([el, "", {"t": _cell(ev), "a": "R"}, dl, "", {"t": _cell(dv), "a": "R"}])
        spans += [(0, r, 1, r), (3, r, 4, r)]
    r = len(rows)
    rows.append([{"t": "Total Earnings : Rs.", "b": True}, "", {"t": lay["total_earnings"], "b": True, "a": "R"},
                 {"t": "Total Deductions : Rs.", "b": True}, "", {"t": lay["total_deductions"], "b": True, "a": "R"}])
    spans += [(0, r, 1, r), (3, r, 4, r)]
    blocks = [
        {"type": "banner", "logo": company.get("logo_url") or "", "lines": head},
        {"type": "grid", "widths": widths, "rows": rows, "spans": spans, "shade": shade},
        {"type": "text", "style": "bold", "text": "Net Pay :   Rs.   %s" % lay["net"]},
        {"type": "text", "style": "small", "text": "In Words :  %s" % lay["words"]},
    ]
    if lay["notes"]:
        blocks.append({"type": "text", "text": lay["notes"]})
    blocks.append({"type": "text", "style": "small", "text": lay["note"]})
    return {"title": "Payslip %s" % (p.get("number") or ""), "blocks": blocks,
            "footer": " ".join(x for x in [company.get("name"), company.get("phone")] if x)}


def workbook(p):
    """The same sheet as an .xlsx, cell for cell: figures are numbers, so the office can add to them."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    lay = layout(p)
    company = lay["company"]
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Payslip"
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    grey = PatternFill("solid", fgColor="C0C0C0")
    for col, width in zip("ABCDEF", (24, 22, 16, 24, 22, 16)):
        sheet.column_dimensions[col].width = width

    def line(row, text, bold=False, size=11, align="center"):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        c = sheet.cell(row=row, column=1, value=text)
        c.font = Font(bold=bold, size=size)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)

    row = 1
    if company.get("name"):
        line(row, company["name"], True, 13)
        row += 1
    if company.get("address"):
        line(row, company["address"], True)
        row += 1
    line(row, lay["title"], True)
    row += 2

    def put(r, c, value, bold=False, fill=None, align=None):
        cell = sheet.cell(row=r, column=c, value=value)
        cell.font = Font(bold=bold)
        cell.border = box
        if fill:
            cell.fill = fill
        if align:
            cell.alignment = Alignment(horizontal=align)
        return cell

    def number(value):
        return None if value is None else float(value)

    for l1, v1, l2, v2 in lay["info"]:
        put(row, 1, l1)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        put(row, 2, v1)
        put(row, 3, None)
        put(row, 4, l2)
        sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        put(row, 5, v2)
        put(row, 6, None)
        row += 1
    for c in range(1, 7):
        put(row, c, "Earnings" if c == 1 else "Deductions" if c == 4 else None, True, grey)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
    row += 1
    first = row
    for (el, ev), (dl, dv) in zip(lay["earnings"], lay["deductions"]):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
        put(row, 1, el)
        put(row, 2, None)
        put(row, 3, number(ev), align="right")
        put(row, 4, dl)
        put(row, 5, None)
        put(row, 6, number(dv), align="right")
        row += 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
    put(row, 1, "Total Earnings : Rs.", True)
    put(row, 2, None)
    put(row, 3, float(lay["total_earnings"]), True, align="right")
    put(row, 4, "Total Deductions : Rs.", True)
    put(row, 5, None)
    put(row, 6, float(lay["total_deductions"]), True, align="right")
    for r in range(first, row + 1):
        for c in (3, 6):
            sheet.cell(row=r, column=c).number_format = "#,##0.##"
    row += 2
    sheet.cell(row=row, column=1, value="Net Pay :").font = Font(bold=True)
    sheet.cell(row=row, column=2, value="Rs.").font = Font(bold=True)
    net = sheet.cell(row=row, column=3, value=float(lay["net"]))
    net.font = Font(bold=True)
    net.number_format = "#,##0.##"
    row += 1
    sheet.cell(row=row, column=1, value="In Words :").font = Font(bold=True)
    sheet.cell(row=row, column=2, value=lay["words"]).font = Font(bold=True, size=9)
    row += 2
    if lay["notes"]:
        sheet.cell(row=row, column=1, value=lay["notes"])
        row += 1
    sheet.cell(row=row, column=1, value=lay["note"])
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def pdf_bytes(p):
    return form_pdf.build_form_pdf(pdf_spec(p))


# --- the register: every payslip of a selection, one row each -------------------------------------------------

INFO_COLUMNS = ["Employee No.", "Name.", "Department.", "Designation.", "Location .", "Bank Name.", "Bank Acc.No.",
                "Father Name.", "Aadhar No.", "UAN No", "ESI No", "DOJ.", "Total Days.", "Work Days.",
                "Effective Work days.", "LOP .", "Balance Leaves.", "Salary Per Month."]
EARNING_COLUMNS = ["Basic", "HRA", "Site Allowance", "Special Allownace", "Mobile Allowance", "Transport Allowance",
                   "Other Allowance", "Overtime", "Bonus"]
DEDUCTION_COLUMNS = ["PF", "ESI", "Prof Tax", "Salary Advance", "TDS Deduction", "Loss Of Pay", "Other Deductions",
                     "Standing Deduction"]
LEADING_COLUMNS = ["Payslip No.", "Month", "Pay Date"]
TRAILING_COLUMNS = ["Total Earnings", "Total Deductions", "Net Pay", "Status"]
REGISTER_COLUMNS = LEADING_COLUMNS + INFO_COLUMNS + EARNING_COLUMNS + DEDUCTION_COLUMNS + TRAILING_COLUMNS
# The columns that are sums of money: written as numbers, and totalled at the foot (but not the salary, which is
# the same figure as total earnings and would be counted twice).
MONEY_COLUMNS = {"Salary Per Month."} | set(EARNING_COLUMNS) | set(DEDUCTION_COLUMNS) | {"Total Earnings", "Total Deductions", "Net Pay"}


def register_row(p):
    """One payslip as one row of the register, in REGISTER_COLUMNS order."""
    lay = layout(p)
    cells = {}
    for l1, v1, l2, v2 in lay["info"]:
        cells[l1], cells[l2] = v1, v2
    for label, value in lay["earnings"] + lay["deductions"]:
        if label:
            cells[label] = value
    cells["Overtime"], cells["Bonus"] = p.get("overtime_pay") or None, p.get("bonus") or None
    cells["Standing Deduction"] = p.get("standing_deduction") or None
    cells.update({"Payslip No.": p.get("number") or "", "Month": month_year(p.get("period_end") or p.get("period_start")),
                  "Pay Date": joined_on(p.get("pay_date")), "Status": p.get("status") or "",
                  "Total Earnings": p.get("gross_pay"), "Total Deductions": p.get("total_deductions"),
                  "Net Pay": p.get("net_pay"), "Salary Per Month.": p.get("gross_pay")})
    out = []
    for column in REGISTER_COLUMNS:
        value = cells.get(column)
        if column in MONEY_COLUMNS:
            out.append(None if value in (None, "") else round(float(value), 2))
        else:
            out.append("" if value is None else value)
    return out


def register_workbook(payloads, company="", span=""):
    """Every payslip given, one row each with all the columns, and a total under the money columns."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Payroll register"
    if company:
        sheet.append([company])
    sheet.append(["Payroll register" + (" - %s" % span if span else "")])
    sheet.cell(row=1, column=1).font = Font(bold=True, size=13)
    sheet.append([])
    header_row = sheet.max_row + 1
    sheet.append(REGISTER_COLUMNS)
    grey = PatternFill("solid", fgColor="C0C0C0")
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = grey
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rows = [register_row(p) for p in payloads]
    for row in rows:
        sheet.append(row)
    totals = [""] * len(REGISTER_COLUMNS)
    totals[0] = "Total"
    for i, column in enumerate(REGISTER_COLUMNS):
        if column in MONEY_COLUMNS and column != "Salary Per Month.":
            totals[i] = round(sum(r[i] or 0 for r in rows), 2)
    sheet.append(totals)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for i, column in enumerate(REGISTER_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = max(12, min(30, len(column) + 4))
        if column in MONEY_COLUMNS:
            for r in range(header_row + 1, sheet.max_row + 1):
                sheet.cell(row=r, column=i).number_format = "#,##0.##"
    sheet.column_dimensions[get_column_letter(REGISTER_COLUMNS.index("Name.") + 1)].width = 24
    # The payslip number, month, pay date and employee number stay in view while the rest scrolls.
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=5)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
