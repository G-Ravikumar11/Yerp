"""The office's own workbooks, read and written in their own layout.

Two of the business's papers lived in Excel before the app did, and both are
laid out the way the site office has always set them:

  The Sub Contractor Registration Form - one sheet per vendor, labels down
  column A and what was written against them in column B: the vendor code,
  the personal details, the bank, the documents collected, the declaration.

  The gang's RA bill - three sheets. The Top Sheet is the certificate of
  payment (what was earned, the GST on it, what comes off, what is paid);
  AB-1 is the abstract (each item up to the previous bill, in this bill, and
  up to this bill); MB-1 is the measurement book the abstract's quantities
  are drawn from (No's x NoM x Length x Width x Height, line under line,
  deductions taken away, a block measured once and multiplied).

This module reads those sheets into plain data and writes the same layouts
back out, so what the app downloads is the paper people already sign. The
arithmetic is the app's; nothing here decides a figure.
"""
import copy
import io

import identity
import re
from datetime import date, datetime

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    XLSX_AVAILABLE = True
except ImportError:  # pragma: no cover
    XLSX_AVAILABLE = False


# --- Reading helpers --------------------------------------------------------------

def _norm(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _key(text):
    """A label as a key: lower case, letters and digits only."""
    return re.sub(r"[^a-z0-9]+", "", str(text or "").lower())


def _number(value):
    """A cell as a number, or None. Text that is a number counts; "-" and
    blanks do not."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return _sane(float(value))
    text = str(value).strip()
    # A comma with one or two digits after it and nothing else is a decimal comma ("12,5"): the Indian way of
    # grouping puts three digits after the last comma, so it cannot be thousands.
    if re.match(r"^-?\d+,\d{1,2}$", text):
        text = text.replace(",", ".")
    text = text.replace(",", "")
    if re.match(r"^-?\d+(\.\d+)?$", text):
        return _sane(float(text))
    return None


def _sane(number):
    """A figure no measurement book could hold - infinity, or past ten million - is no figure at all, so
    it cannot carry a quantity out of range into the book or the bill."""
    return number if number == number and abs(number) <= 1e7 else None


def _plain(value):
    """What was written in a box, as text: a long number as its digits (an
    account number is not 5.02E+13), a date as ISO."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value == int(value) else repr(value)
    if isinstance(value, int):
        return str(value)
    return _norm(value)


def date_iso(value):
    """12.07.2024, 1.4.25, 17/02/2026 or a real date, as 2024-07-12. Anything
    else is returned as it was written."""
    if isinstance(value, (datetime, date)):
        return (value.date() if isinstance(value, datetime) else value).isoformat()
    text = _norm(value)
    m = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})$", text)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        y = y + 2000 if y < 100 else y
        try:
            return date(y, mo, d).isoformat()
        except ValueError:
            return text
    if re.match(r"^\d{4}-\d{2}-\d{2}", text):
        return text[:10]
    return text


def _cell_value(values, formulas, row, col):
    """A cell's value: the one Excel last calculated, else a formula that is
    plain arithmetic (=2.1+0.75+0.75) worked out here."""
    v = values.cell(row=row, column=col).value
    if v is not None or formulas is None:
        return v
    f = formulas.cell(row=row, column=col).value
    if isinstance(f, str) and f.startswith("="):
        return _arithmetic(f[1:])
    return v


def _arithmetic(expr):
    """=2.1+0.75+0.75 or =2.7-0.915-0.435, worked out: numbers with + - * /
    and brackets, nothing else - a formula that refers to a cell or calls a
    function is not guessed at."""
    import ast
    import operator
    ops = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv}

    def walk(node):
        if isinstance(node, ast.Expression):
            return walk(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)) and not isinstance(node.value, bool):
            return float(node.value)
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            v = walk(node.operand)
            return v if isinstance(node.op, ast.UAdd) else -v
        if isinstance(node, ast.BinOp) and type(node.op) in ops:
            return ops[type(node.op)](walk(node.left), walk(node.right))
        raise ValueError("not plain arithmetic")

    text = str(expr or "").strip()
    if len(text) > 200 or not re.match(r"^[0-9.+\-*/() ]+$", text) or "**" in text:
        return None
    try:
        return walk(ast.parse(text, mode="eval"))
    except (ValueError, SyntaxError, ZeroDivisionError, RecursionError):
        return None


# --- The registration form ---------------------------------------------------------

# Label on the form -> field. Matched on the label's letters and digits, so
# "E - mail Id :" and "E-mail Id" are the same box.
REGISTRATION_LABELS = {
    "vendorcode": "vendor_code",
    "nameofthesubcontractor": "company_name",
    "residentialaddress": "address",
    "dateofjoining": "joining_date",
    "pincode": "pin_code",
    "city": "city",
    "state": "state",
    "natureofwork": "nature_of_work",
    "telno": "phone_number",
    "emailid": "email",
    "nameofcontactperson": "contact_person",
    "typeofentity": "entity_type",
    "panno": "pan",
    "gstregno": "gst_number",
    "aadharcard": "aadhaar",
    "bankname": "bank_name",
    "accountno": "bank_account",
    "ifsccode": "bank_ifsc",
    "branch": "bank_branch",
}


def read_registration_forms(book):
    """Every sheet that is a registration form, as a dict of its boxes, in the
    workbook's order. A sheet that is not one (a rate list left in the same
    file) is passed over."""
    out = []
    for ws in book.worksheets:
        form = {"sheet": ws.title}
        for row in ws.iter_rows():
            cells = [c for c in row if c.value is not None and _plain(c.value) != ""]
            if not cells:
                continue
            first = _plain(cells[0].value)
            if re.match(r"(?i)^project\s*:", first):
                form.setdefault("registered_project", _norm(first.split(":", 1)[1]))
                continue
            # "C) Aadhar Card" in the documents list is not the Aadhaar box.
            if re.match(r"^[A-Fa-f]\)\s", first):
                continue
            field = REGISTRATION_LABELS.get(_key(first))
            if not field or field in form:
                continue
            value = cells[1].value if len(cells) > 1 else None
            if field == "joining_date":
                form[field] = date_iso(value) if value not in (None, "") else ""
            else:
                form[field] = _plain(value)
        if form.get("vendor_code") or form.get("company_name"):
            if form.get("email") in ("0", "-"):
                form["email"] = ""
            out.append(form)
    return out


def registration_workbook(form, company=None):
    """The registration form as its own workbook, laid out as the office's."""
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", "%s %s" % (form.get("vendor_code") or "", form.get("name") or ""))[:31] or "Form"
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 58
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    r = 1

    def merged(text, size=11, fill=False, align="center"):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(bold=True, size=size)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if fill:
            c.fill = PatternFill("solid", fgColor="D9D9D9")
        for col in (1, 2):
            ws.cell(row=r, column=col).border = box
        r += 1

    def pair(label, value, strong=False):
        nonlocal r
        a = ws.cell(row=r, column=1, value=label)
        b = ws.cell(row=r, column=2, value=value)
        a.font = bold
        if strong:
            b.font = bold
        b.alignment = Alignment(wrap_text=True, vertical="center")
        a.border = b.border = box
        r += 1

    merged(form.get("company") or "", size=14)
    merged("SUB CONTRACTOR REGISTRATION FORM", size=12)
    merged("PROJECT: %s" % (form.get("project") or ""), size=11)
    pair("VENDOR CODE:", form.get("vendor_code") or "", strong=True)
    r += 1
    merged("1. Subcontractor Personal Details", fill=True, align="left")
    for label, value in form.get("personal") or []:
        pair(label, value)
    r += 1
    merged("2. Bank details ", fill=True, align="left")
    for label, value in form.get("bank") or []:
        pair(label, value)
    r += 1
    merged("3.Documents Required", fill=True, align="left")
    for label, received in form.get("documents") or []:
        pair(label, "Received" if received else "")
    r += 1
    merged("Declaration:", align="left")
    for line in form.get("declaration") or []:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(row=r, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30
        r += 1
    if form.get("declaration_signed"):
        ws.cell(row=r, column=1, value="Declaration signed by the sub contractor.").font = Font(italic=True)
        r += 1
    r += 3
    for col, text in ((1, "Authorized Signature "), (2, "Contractor Signature ")):
        c = ws.cell(row=r, column=col, value=text)
        c.font = bold
    identity.add_logo(ws, company, "C1", height=48)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


# --- The measurement book ------------------------------------------------------

MB_HEADINGS = {
    "sno": ("sno", "slno", "sino", "sl", "no"),
    "description": ("description", "particulars", "descriptionofwork", "details"),
    "uom": ("uom", "unit", "units"),
    "nos": ("nos", "nosno", "number", "numbers"),
    "nom": ("nom",),
    "length": ("length", "l", "len"),
    "width": ("width", "breadth", "b", "w"),
    "height": ("height", "depth", "d", "h", "ht"),
    "total": ("totalquantity", "quantity", "qty", "totalqty", "total"),
    "remarks": ("remarks", "remark"),
}


def build_mb_template():
    """A worked example of the measurement book layout this reads: the
    heading row it looks for, one item with a block entry, two dimension
    lines and the closing total - so a blank sheet started from this reads
    back exactly as it prints."""
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "MB"
    bold = Font(bold=True)
    ws.append(["Name of the Work:", "295 KLD STP, Vanya City"])
    ws.append(["Name of the Contractor:", "Rani Labour Contractors"])
    ws.append(["Date:", "01-10-2026"])
    ws.append([])
    header = ["S.No", "Description", "UoM", "No's", "NoM", "Length", "Width", "Height", "Total Quantity", "Remarks"]
    ws.append(header)
    for c in range(1, len(header) + 1):
        ws.cell(row=5, column=c).font = bold
    ws.append(["1", "Shuttering for slabs and beams", "sqm", "", "", "", "", "", "", ""])
    ws.append(["a", "Block C-3, First Floor", "", "", "", "", "", "", "", ""])
    ws.append(["", "Slab", "", 1, 1, 12.5, 8, "", "", ""])
    ws.append(["", "Beam", "", 4, 1, 6, 0.6, "", "", "Deduct openings separately"])
    ws.append(["", "Total Quantity for 1 Block", "", "", "", "", "", "", 114.4, ""])
    for col, width in zip("ABCDEFGHIJ", (7, 34, 8, 6, 6, 8, 8, 8, 14, 24)):
        ws.column_dimensions[col].width = width
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _mb_columns(values, formulas, max_scan=30):
    """Where the book's heading row is and which column is which."""
    for r in range(1, min(values.max_row, max_scan) + 1):
        found = {}
        for c in range(1, min(values.max_column, 30) + 1):
            k = _key(_cell_value(values, formulas, r, c))
            if not k:
                continue
            for field, names in MB_HEADINGS.items():
                if field not in found and k in names:
                    found[field] = c
                    break
        if "description" in found and ("total" in found or "length" in found):
            return r, found
    return None, {}


# Rows that are the sheet's own arithmetic on a total - never measurements. A sheet may call the quantity it is
# billed at "to be paid", "payable", "net quantity", "billable"...; and the part kept back "hold", "held",
# "retention", "withheld". Whatever the words, a payable figure is also recognised by the arithmetic: it is the
# total less the hold.
PAYABLE_WORDS = re.compile(r"(?i)to\s*be\s*paid|payable|net\s+(?:qty|quantity)|billable|for\s+billing|to\s+pay\b")
HOLD_WORDS = re.compile(r"(?i)\bhold\b|\bheld\b|\bretention\b|\bwithh[eo]ld\b")
PERCENT = re.compile(r"(\d+(?:\.\d+)?)\s*%")


def read_measurement_book(values, formulas=None):
    """The measurement book on a sheet, as items, the entries under each and
    their dimension lines.

    It reads the book the way it is written. A numbered row with no figures
    starts an item ("2  BUFFERING WORKS"); a lettered one starts an entry
    under it, usually a block or a floor ("b  430 SFT-Block C-3 FF,SF,TF").
    Words with no figures are a heading inside the entry ("Living Room",
    "Deductions"). A row with figures is a dimension line, a negative count
    being a deduction. "Total Quantity for one Block" closes the entry, and
    "Total Quantity for 4 Blocks" says how many times it is built.

    Some books hold part of the work back: under the totals there are rows such as
    "Hold 5 % for Finishes & Handing over" and "Total Qty To be paid - 45%". Those are
    not measurements. The entries keep the quantity the sheet measured; each group
    of blocks is given the share the sheet holds back (`held_back`, `payable`), and
    the holds are returned beside the items (`holds`), to be recorded as holds - never
    as a line inside a block.
    """
    head, cols = _mb_columns(values, formulas)
    max_col = values.max_column      # asked once: on a big sheet every ask walks all of its cells
    max_row = values.max_row
    if not head:
        raise ValueError("No heading row with Description and Total Quantity was found on '%s'." % values.title)

    def get(r, field):
        c = cols.get(field)
        return _cell_value(values, formulas, r, c) if c else None

    meta = {}
    for r in range(1, head):
        for c in range(1, min(max_col, 20) + 1):
            text = _plain(_cell_value(values, formulas, r, c))
            low = text.lower()
            if not text:
                continue
            if "name of the work" in low:
                meta["work_name"] = _norm(re.split(r":-?|:", text, 1)[-1])
            elif "name of the contractor" in low:
                meta["contractor"] = _norm(re.split(r":-?|:", text, 1)[-1])
            elif "bill period" in low:
                meta["period"] = _norm(text.split(":", 1)[-1])
            elif re.match(r"(?i)^bill\s*no", text):
                meta["bill_no"] = _norm(re.split(r":-?|:", text, 1)[-1])
            elif re.match(r"(?i)^date", text):
                meta["date"] = date_iso(_norm(re.split(r":-?|:", text, 1)[-1]))

    items, warnings = [], []
    item = entry = None
    # The sheet's arithmetic, read ahead: a subtotal followed by two stated figures that add back up to it is a
    # total, the part held and what is payable - whatever the rows are called.
    forced = {}
    seq = []
    for r in range(head + 1, max_row + 1):
        d = _plain(get(r, "description"))
        if any(_number(get(r, f)) is not None for f in ("nos", "nom", "length", "width", "height")):
            seq.append(None)                         # a measured line: nothing stated after it belongs to what came before
            continue
        t = _number(get(r, "total"))
        if t is not None and d:
            seq.append((r, d.lower(), t))
        elif d:
            seq.append(None)                         # a heading or a new block: what follows is not the hold of what came before
    for i, entry_row in enumerate(seq):
        if not entry_row or not entry_row[1].startswith("total") or re.search(r"(?i)for\s+(\d+(?:\.\d+)?\s+)?blocks?|for\s+one\s+block", entry_row[1]):
            continue
        sub = entry_row[2]
        ahead = []
        for nxt in seq[i + 1:i + 4]:
            if nxt is None:
                break
            ahead.append(nxt)
        for a in range(len(ahead)):
            for b in range(a + 1, len(ahead)):
                (ra, la, h), (rb, lb, p) = ahead[a], ahead[b]
                if not (0.011 < h < sub and 0 < p < sub and abs(sub - h - p) <= 0.011):
                    continue
                if la.startswith("total") and not HOLD_WORDS.search(la):
                    continue
                if lb.startswith("total") and not PAYABLE_WORDS.search(lb):
                    continue
                forced.setdefault(ra, "hold")
                forced.setdefault(rb, "payable")
    # Entries not yet covered by a "to be paid" row, and the groups that have been.
    pending, groups, held_label = [], [], []
    # The latest subtotal the sheet states ("Total Quantity"), and the figures stated since it that are not
    # lines: the arithmetic that says which of them is the hold and which the payable.
    last_sub = {"row": 0, "value": None}
    since_sub = []
    since_text = []

    def close():
        nonlocal entry
        if entry is not None and entry["dims"]:
            one = round(sum(d["quantity"] for d in entry["dims"]), 3)
            entry["one"] = one
            entry["quantity"] = round(one * entry["multiplier"], 3)
            stated = entry.get("stated_total")
            if stated is not None and abs(stated - entry["quantity"]) > 0.011:
                warnings.append("Row %d: the sheet says %s, the lines come to %s." % (
                    entry["row"], round(stated, 3), entry["quantity"]))
        entry = None

    def open_entry(r, location="", letter=""):
        nonlocal entry
        entry = {"row": r, "location": location, "letter": letter, "dims": [], "multiplier": 1.0, "stated_total": None}
        item["entries"].append(entry)
        pending.append(entry)

    for r in range(head + 1, max_row + 1):
        desc = _plain(get(r, "description"))
        sno = get(r, "sno")
        sno_text = _plain(sno)
        low = desc.lower()
        figures = {f: _number(get(r, f)) for f in ("nos", "nom", "length", "width", "height")}
        given = {f: v for f, v in figures.items() if v is not None}
        total = _number(get(r, "total"))
        row_text = " ".join(_plain(_cell_value(values, formulas, r, c)) for c in range(1, min(max_col, 14) + 1))

        # Derived rows - the part held back, and what is to be paid - are the sheet's own arithmetic on the
        # totals above them, not measurements: read them for the payable figure, never as lines.
        text_here = (desc + " " + row_text).strip()
        if total is not None and not given:
            sub_value = last_sub["value"]
            is_payable = bool(PAYABLE_WORDS.search(text_here)) or forced.get(r) == "payable"
            is_hold = (bool(HOLD_WORDS.search(text_here)) and bool(PERCENT.search(text_here)) and not is_payable
                       and not low.startswith("total")) or forced.get(r) == "hold"
            if not is_payable and not is_hold and sub_value and not low.startswith("total") and not re.match(r"(?i)^release\b", text_here):
                # No telltale words: it is the payable if it is the subtotal less something stated since.
                is_payable = any(0.011 < h < sub_value and abs(sub_value - h - total) <= 0.011 and 0 < total < sub_value
                                 for _, h in since_sub)
            if is_hold or re.match(r"(?i)^release\b", text_here):
                since_sub.append((r, total))
                words = desc or next((t for t in (_plain(_cell_value(values, formulas, r, c)) for c in range(1, min(max_col, 14) + 1))
                                      if len(t) > 3 and _number(t) is None), "")
                if words:
                    since_text.append(words)
                m = PERCENT.search(text_here)
                if is_hold and m and sub_value:
                    held_label.append((r, float(m.group(1)), sub_value, total))
                continue
            if is_payable and total:
                # It covers every entry since the last such row: one block, or several under one subtotal.
                group = [e for e in pending if e["dims"]]
                if entry is not None and entry["dims"] and entry not in group:
                    group.append(entry)
                pending.clear()
                close()
                if group:
                    groups.append((r, group, total, sub_value if (sub_value and sub_value >= total - 0.011) else None,
                                   (since_text[0] if since_text else "") or desc))
                else:
                    warnings.append("Row %d: a 'to be paid' figure with no entry above it was left out." % r)
                since_sub.clear()
                since_text.clear()
                last_sub["value"] = None
                continue
        if not desc and not given and not sno_text:
            continue
        if low.startswith("total"):
            if total is not None and not re.search(r"(?i)for\s+(\d+(?:\.\d+)?)\s+blocks?|for\s+one\s+block", row_text):
                last_sub["row"], last_sub["value"] = r, total
                since_sub.clear()
                since_text.clear()
            m = re.search(r"(?i)for\s+(\d+(?:\.\d+)?)\s+blocks?", row_text)
            if entry is not None:
                if m:
                    entry["multiplier"] = float(m.group(1)) or 1.0
                    entry["stated_total"] = total
                    b = re.search(r"(?i)block\s*no\.?\s*-?\s*(.+)$", desc)
                    if b and not entry["location"]:
                        entry["location"] = _norm(b.group(1))
                    close()
                else:
                    entry["one_stated"] = total
            continue
        if not given:
            is_item = isinstance(sno, (int, float)) and not isinstance(sno, bool) or re.match(r"^\d+(\.\d+)?$", sno_text)
            # Capital roman numerals of two letters or more (II, III, IV...) number the sections of a book,
            # as in "II  Internal Painting Work"; one letter is a block ("A", "B").
            is_item = is_item or bool(re.match(r"^[IVXL]{2,}$", sno_text))
            is_entry = bool(re.match(r"^[A-Za-z]{1,3}\)?$|^\([A-Za-z]{1,3}\)$|^[ivx]+$", sno_text))
            if desc and (is_item or item is None):
                close()
                item = {"row": r, "sno": sno_text, "description": desc, "entries": []}
                items.append(item)
                continue
            if desc and is_entry:
                close()
                open_entry(r, desc, sno_text)
                continue
            if desc and total and abs(total) > 0.0001:
                # Words and a quantity but no dimensions: a lump sum line.
                given = {"nos": total}
            elif desc:
                if entry is None:
                    open_entry(r)
                entry["dims"].append({"particulars": desc, "is_heading": True, "quantity": 0.0})
                continue
            else:
                continue
        if item is None:
            warnings.append("Row %d has figures before any item heading, so it was left out." % r)
            continue
        if entry is None:
            open_entry(r)
        sign = 1
        clean = {}
        for f, v in given.items():
            if v < 0:
                sign = -sign
            clean[f] = abs(v)
        qty = 1.0
        for v in clean.values():
            qty *= v
        qty = round(qty, 3)
        shown = get(r, "total")
        if total is None and shown is not None and _plain(shown) not in ("", "-"):
            # The sheet's own total for this line is not a figure ("`", "#REF!"): its subtotal leaves the line out,
            # while the dimensions say what it comes to.
            warnings.append("Row %d (%s): the total cell holds '%s', not a figure. The lines come to %s%s, which "
                            "the sheet's own totals leave out." % (r, desc or "no description", _plain(shown), "-" if sign < 0 else "", qty))
        if not qty:
            warnings.append("Row %d (%s) comes to nothing and was left out." % (r, desc or "no description"))
            continue
        last_sub["value"] = None
        since_sub.clear()
        since_text.clear()
        entry["dims"].append({
            "particulars": desc, "is_heading": False,
            "nos": clean.get("nos"), "nom": clean.get("nom"), "length": clean.get("length"),
            "breadth": clean.get("width"), "depth": clean.get("height"),
            "deduct": sign < 0, "quantity": -qty if sign < 0 else qty,
            "remarks": _plain(get(r, "remarks")),
        })
    close()
    for it in items:
        it["entries"] = [e for e in it["entries"] if any(not d["is_heading"] for d in e["dims"])]
        for e in it["entries"]:
            if "quantity" not in e:
                one = round(sum(d["quantity"] for d in e["dims"]), 3)
                e["one"], e["quantity"] = one, round(one * e["multiplier"], 3)
    for it in items:
        for e in it["entries"]:
            if not e["location"]:
                # An entry the sheet gives no place to is named by the first heading inside it ("365 SFT-Block ( B24, B21 & B19)").
                e["location"] = _norm(next((d["particulars"] for d in e["dims"] if d["is_heading"] and d["particulars"]), ""))
    holds = []
    for number, (row, group, pay, stated, said) in enumerate(groups, 1):
        for e in group:
            e["group"] = number         # blocks under one subtotal share one hold-back
        full = round(sum(e["quantity"] for e in group), 3)
        if stated and 0 < pay <= stated + 0.011:
            # The hold is the sheet's own rule - what it keeps back of what it totals - and it is that share
            # of the lines that is held, even when the sheet's total and its lines differ.
            share = min(1.0, pay / stated)
            if abs(stated - full) > 0.011:
                warnings.append("Row %d: the sheet totals %s but its lines come to %s. The sheet's hold (%g%%) was applied to the lines." % (
                    row, round(stated, 3), full, round((1 - share) * 100, 2)))
        else:
            if pay > full + 0.011:
                warnings.append("Row %d: the sheet says %s is to be paid but the lines come to %s, so the lines are used." % (row, round(pay, 3), full))
                continue
            share = pay / full if full else 1.0          # the same share of every entry under that subtotal
        if not full or share >= 1 - 0.00001:
            continue
        percent = round((1 - share) * 100, 2)
        for e in group:
            e["held_back"] = round(e["quantity"] * (1 - share), 3)
            e["payable"] = round(e["quantity"] - e["held_back"], 3)
        reason = _norm(said) or "Held back for finishes and handing over"
        if re.match(r"(?i)^release\b", reason):
            # A release names the stages paid for ("Putty 40% & Primer 10%"); what is held is said beside it.
            reason = "%s - %g%% held" % (reason, percent)
        elif not any(abs(float(p) - percent) < 0.01 for p in PERCENT.findall(reason)):
            # The row's own percent is not what its figures hold ("Hold 5 %" over a 10% calculation): the figure is
            # what is held, so that is what the reason says.
            reason = "%s - %g%% held" % (_norm(PERCENT.sub("", reason)).replace("  ", " "), percent)
        holds.append({"group": number, "row": row, "reason": reason[:300], "percent": percent, "full": full,
                      "held": round(sum(e["held_back"] for e in group), 3), "payable": round(sum(e["payable"] for e in group), 3)})
    for it in items:
        for e in it["entries"]:
            e.setdefault("held_back", 0.0)
            e.setdefault("payable", e["quantity"])
    # A hold's label can say one percent while its figures hold another ("Hold 5 %" over a 10% calculation).
    for row, said, sub_value, held in held_label:
        actual = held / sub_value * 100 if sub_value else None
        if actual is not None and abs(actual - said) > 0.51:
            warnings.append("Row %d says %g%% is held but its figure is %g%% of %s. The figure was used." % (
                row, said, round(actual, 2), round(sub_value, 3)))
    for it in items:
        it["quantity"] = round(sum(e["quantity"] for e in it["entries"]), 3)
        it["held"] = round(sum(e["held_back"] for e in it["entries"]), 3)
        it["payable"] = round(sum(e["payable"] for e in it["entries"]), 3)
    return {"sheet": values.title, "meta": meta, "items": [i for i in items if i["entries"]], "warnings": warnings,
            "holds": holds}


# --- The RA bill, three sheets --------------------------------------------------

def _styles():
    thin = Side(style="thin", color="000000")
    return {"box": Border(left=thin, right=thin, top=thin, bottom=thin),
            "bold": Font(bold=True), "title": Font(bold=True, size=14), "sub": Font(bold=True, size=12),
            "shade": PatternFill("solid", fgColor="D9D9D9"),
            "wrap": Alignment(wrap_text=True, vertical="center"),
            "centre": Alignment(horizontal="center", vertical="center", wrap_text=True),
            "right": Alignment(horizontal="right", vertical="center")}


def _put(ws, st, row, col, value, bold=False, align=None, fmt=None, span=None, shade=False, box=True):
    c = ws.cell(row=row, column=col, value=value)
    # Setting a font, border and fill on a cell costs a lookup in the workbook's style tables each, and a
    # measurement book is tens of thousands of cells in a handful of looks. A cell that looks like one
    # already made takes its style as it is.
    key = None
    if not span and (align is None or isinstance(align, str)):
        key = (id(ws.parent), bold, align, fmt if (fmt and isinstance(value, (int, float))) else None, shade, box)
        known = st.setdefault("_looks", {}).get(key)
        if known is not None:
            c._style = copy.copy(known)
            return c
    if bold:
        c.font = st["bold"]
    c.alignment = st.get(align or "wrap") if isinstance(align or "wrap", str) else align
    if fmt and isinstance(value, (int, float)):
        c.number_format = fmt
    if span:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=span)
    last = span or col
    for cc in range(col, last + 1):
        cell = ws.cell(row=row, column=cc)
        if box:
            cell.border = st["box"]
        if shade:
            cell.fill = st["shade"]
    if key is not None:
        st["_looks"][key] = copy.copy(c._style)
    return c


MONEY = "#,##0.00"
QTY = "#,##0.000"


def ra_bill_workbook(cert, company=None):
    """The bill as the three sheets it was always sent as: Top Sheet, AB-1, MB-1."""
    book = openpyxl.Workbook()
    st = _styles()
    _top_sheet(book.active, st, cert)
    _abstract_sheet(book.create_sheet("AB-1"), st, cert)
    _mb_sheet(book.create_sheet("MB-1"), st, cert)
    # The company's mark beside each sheet's head, outside the ruled table.
    for ws, column in zip(book.worksheets, ("I", "N", "L")):
        identity.add_logo(ws, company, "%s1" % column, height=48)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _banner_rows(ws, st, lines, last_col):
    r = 1
    for text, kind in lines:
        _put(ws, st, r, 1, text, bold=True, align="centre", span=last_col, shade=(kind == "band"), box=True)
        ws.cell(row=r, column=1).font = st["title"] if kind == "banner" else (st["sub"] if kind == "band" else st["bold"])
        r += 1
    return r


def _top_sheet(ws, st, cert):
    ws.title = "Top Sheet"
    for col, width in zip("ABCDEFG", (7, 44, 16, 17, 17, 9, 12)):
        ws.column_dimensions[col].width = width
    r = _banner_rows(ws, st, cert["top"]["banner"], 7)
    for sl, label, cells in cert["top"]["info"]:
        _put(ws, st, r, 1, sl, align="centre")
        _put(ws, st, r, 2, label)
        c, d, e, f = (list(cells) + [None, None, None, None])[:4]
        if e is None and f is None:
            _put(ws, st, r, 3, c, span=7)
        else:
            if d is None:
                _put(ws, st, r, 3, c, span=4)
            else:
                _put(ws, st, r, 3, c)
                _put(ws, st, r, 4, d)
            _put(ws, st, r, 5, e)
            _put(ws, st, r, 6, f, span=7)
        r += 1
    for i, text in enumerate(("Sl\nNo", "Description", "Reference", "Upto This\nBill Amount",
                              "Upto Previous\nBill Amount", "For This\nBill Amount")):
        _put(ws, st, r, i + 1, text, bold=True, align="centre", span=7 if i == 5 else None, shade=True)
    ws.row_dimensions[r].height = 30
    r += 1
    for row in cert["top"]["money"]:
        if row.get("section"):
            _put(ws, st, r, 1, row["section"], bold=True, span=7)
            r += 1
            continue
        bold = row.get("bold", False)
        _put(ws, st, r, 1, row.get("sl") or "", bold=bold, align="centre", span=(2 if row.get("span_label") else None))
        if not row.get("span_label"):
            _put(ws, st, r, 2, row["label"], bold=bold)
        else:
            ws.cell(row=r, column=1, value=row["label"])
        _put(ws, st, r, 3, row.get("ref") or "")
        _put(ws, st, r, 4, row["upto"], bold=bold, align="right", fmt=MONEY)
        _put(ws, st, r, 5, row["prev"], bold=bold, align="right", fmt=MONEY)
        _put(ws, st, r, 6, row["this"], bold=bold, align="right", fmt=MONEY, span=7)
        r += 1
    _put(ws, st, r, 1, "AMOUNT IN WORDS: ", bold=True, span=2)
    _put(ws, st, r, 3, cert["top"]["words"], bold=True, span=7)
    r += 1
    _signature_rows(ws, st, r, cert["top"]["signatures"], (1, 3, 4, 5, 6), (2, 3, 4, 5, 7))


def _signature_rows(ws, st, r, boxes, starts, ends):
    for (role, name, caption), c0, c1 in zip(boxes, starts, ends):
        _put(ws, st, r, c0, role, bold=True, align="centre", span=c1 if c1 > c0 else None)
        _put(ws, st, r + 1, c0, name or "", align="centre", span=c1 if c1 > c0 else None)
        ws.row_dimensions[r + 1].height = 42
        _put(ws, st, r + 2, c0, caption or "", bold=True, align="centre", span=c1 if c1 > c0 else None)


def _abstract_sheet(ws, st, cert):
    a = cert["abstract"]
    for col, width in zip("ABCDEFGHIJKL", (7, 34, 8, 11, 10, 14, 11, 10, 14, 11, 14, 14)):
        ws.column_dimensions[col].width = width
    r = _banner_rows(ws, st, a["banner"], 12)
    for left, mid, right in a["meta"]:
        _put(ws, st, r, 1, left, span=5)
        _put(ws, st, r, 6, mid, span=8)
        _put(ws, st, r, 9, right, span=12)
        r += 1
    for c0, c1, text in ((1, 1, "SI.No"), (2, 2, "Description"), (3, 3, "Unit"), (4, 6, "Up To Previous Bill"),
                         (7, 9, "In This Bill Claimed"), (10, 11, "Up to This Bill"), (12, 12, "Remarks")):
        _put(ws, st, r, c0, text, bold=True, align="centre", span=c1 if c1 > c0 else None, shade=True)
    r += 1
    for i, text in enumerate(("", "", "", "Qty", "Rate", "Amount", "Qty", "Rate", "Amount", "Qty", "Amount", "")):
        _put(ws, st, r, i + 1, text, bold=True, align="centre", shade=True)
    r += 1
    for row in a["rows"]:
        if row.get("header"):
            _put(ws, st, r, 1, "")
            _put(ws, st, r, 2, row["description"], bold=True)
            for c in range(3, 13):
                _put(ws, st, r, c, "")
            r += 1
            continue
        vals = (row["sl"], row["description"], row["unit"], row["prev_qty"], row["prev_rate"], row["prev_amount"],
                row["this_qty"], row["this_rate"], row["this_amount"], row["upto_qty"], row["upto_amount"],
                row.get("remarks") or "")
        fmts = (None, None, None, QTY, MONEY, MONEY, QTY, MONEY, MONEY, QTY, MONEY, None)
        for i, (v, f) in enumerate(zip(vals, fmts)):
            _put(ws, st, r, i + 1, v, align="right" if f else ("centre" if i in (0, 2) else None), fmt=f)
        r += 1
    t = a["totals"]
    _put(ws, st, r, 1, "A)", bold=True, align="centre")
    _put(ws, st, r, 2, "Total Invoice Amount", bold=True)
    _put(ws, st, r, 3, "Up to Previous Bill Amount :-", bold=True, span=5)
    _put(ws, st, r, 6, t["prev"], bold=True, align="right", fmt=MONEY)
    _put(ws, st, r, 7, "In this Bill", bold=True, span=8)
    _put(ws, st, r, 9, t["this"], bold=True, align="right", fmt=MONEY)
    _put(ws, st, r, 10, "")
    _put(ws, st, r, 11, t["upto"], bold=True, align="right", fmt=MONEY)
    _put(ws, st, r, 12, "")
    r += 3
    _signature_rows(ws, st, r, a["signatures"], (1, 3, 6, 10), (2, 5, 9, 12))


def _mb_sheet(ws, st, cert):
    m = cert["mb"]
    ws.title = "MB-1"
    for col, width in zip("ABCDEFGHIJ", (7, 44, 8, 8, 8, 10, 10, 10, 15, 16)):
        ws.column_dimensions[col].width = width
    r = _banner_rows(ws, st, m["banner"], 10)
    for left, right in m["meta"]:
        _put(ws, st, r, 1, left, span=6)
        _put(ws, st, r, 7, right, span=10)
        r += 1
    for i, text in enumerate(("S.No", "Description", "UoM", "No's", "NoM", "Length", "Width", "Height",
                              "Total Quantity", "Remarks")):
        _put(ws, st, r, i + 1, text, bold=True, align="centre", shade=True)
    r += 1
    for row in m["rows"]:
        kind = row.get("kind")
        if kind in ("item", "entry", "heading"):
            _put(ws, st, r, 1, row.get("sno") or "", bold=kind != "heading", align="centre")
            _put(ws, st, r, 2, row.get("description") or "", bold=True)
            if kind == "heading":
                ws.cell(row=r, column=2).font = Font(bold=True, italic=True)
            for c in range(3, 11):
                _put(ws, st, r, c, "")
        elif kind in ("subtotal", "total"):
            _put(ws, st, r, 1, "")
            _put(ws, st, r, 2, row.get("description") or "", bold=True)
            _put(ws, st, r, 3, "")
            _put(ws, st, r, 4, "")
            _put(ws, st, r, 5, "")
            _put(ws, st, r, 6, row.get("label") or "", bold=True, span=8)
            _put(ws, st, r, 9, row.get("quantity"), bold=True, align="right", fmt=QTY)
            _put(ws, st, r, 10, row.get("uom") or "", bold=True)
        else:
            vals = ("", row.get("description") or "", row.get("uom") or "", row.get("nos"), row.get("nom"),
                    row.get("length"), row.get("width"), row.get("height"), row.get("quantity"),
                    row.get("remarks") or "")
            for i, v in enumerate(vals):
                _put(ws, st, r, i + 1, v, align="right" if i >= 3 and i <= 8 else None,
                     fmt=QTY if i == 8 else ("0.###" if i >= 3 else None))
        r += 1
    r += 2
    _signature_rows(ws, st, r, m["signatures"], (1, 3, 7), (2, 6, 10))
