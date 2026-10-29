"""The salary structure: HR keeps the components, a payslip gets its own lines when it is made, and nothing is
recalculated when it is viewed or downloaded."""
import io
import uuid

import pytest
from fastapi.testclient import TestClient

import main
from conftest import make_employee

PASSWORD = "Crew1234"
SEP = {"period_start": "2026-09-01", "period_end": "2026-09-30", "pay_date": "2026-09-30"}


def structure(tenant):
    return tenant.get("/api/salary-structure").json()["components"]


def named(components, name):
    return next(c for c in components if c["name"] == name)


def lines_of(tenant, pid):
    return {l["name"]: l for l in tenant.get("/api/payslips/%d" % pid).json()["lines"]}


def make_slip(tenant, salary=65000.0, **kw):
    emp = make_employee(tenant, salary=salary, tax_rate=0.0, **kw)
    res = tenant.post("/api/payslips", json={"employee_id": emp["id"], **SEP})
    assert res.status_code == 200, res.text
    return emp, res.json()["id"]


def test_a_company_starts_with_the_offices_own_list(tenant):
    got = structure(tenant)
    assert [(c["name"], c["kind"], c["basis"], c["percent"]) for c in got] == [
        ("Basic", "earning", "percent", 40.0), ("HRA", "earning", "manual", 0.0), ("Conveyance", "earning", "percent", 10.0),
        ("Medical allowance", "earning", "percent", 20.0), ("Spcl allowance", "earning", "manual", 0.0), ("PF", "deduction", "percent", 12.0)]


def test_a_new_payslip_gets_its_lines_worked_out_from_the_monthly_salary(tenant):
    _, pid = make_slip(tenant, 65000)
    lines = lines_of(tenant, pid)
    assert lines["Basic"]["amount"] == 26000 and lines["Conveyance"]["amount"] == 6500
    assert lines["Medical allowance"]["amount"] == 13000 and lines["PF"]["amount"] == 7800
    assert lines["HRA"]["amount"] == 0 and lines["Spcl allowance"]["amount"] == 0          # blank, for HR to type
    slip = tenant.get("/api/payslips/%d" % pid).json()
    assert slip["monthly_salary"] == 65000
    assert slip["basic_salary"] == 26000 and slip["retirement"] == 7800                     # the payslip's own totals follow
    assert slip["gross_pay"] == 45500 and slip["net_pay"] == 45500 - 7800


def test_hr_types_the_blank_ones_and_the_totals_follow(tenant):
    _, pid = make_slip(tenant, 65000)
    lines = lines_of(tenant, pid)
    res = tenant.put("/api/payslips/%d" % pid, json={"lines": [{"id": lines["HRA"]["id"], "amount": 13000},
                                                               {"id": lines["Spcl allowance"]["id"], "amount": 6500}]})
    assert res.status_code == 200, res.text
    slip = tenant.get("/api/payslips/%d" % pid).json()
    assert slip["gross_pay"] == 65000 and slip["net_pay"] == 65000 - 7800
    assert tenant.put("/api/payslips/%d" % pid, json={"lines": [{"id": lines["HRA"]["id"], "amount": -1}]}).status_code == 400
    assert tenant.put("/api/payslips/%d" % pid, json={"lines": [{"id": 99999999, "amount": 1}]}).status_code == 400


def test_changing_the_structure_leaves_payslips_already_made_alone(tenant):
    _, pid = make_slip(tenant, 65000)
    basic = named(structure(tenant), "Basic")
    assert tenant.put("/api/salary-structure/%d" % basic["id"], json={"percent": 50}).status_code == 200
    assert lines_of(tenant, pid)["Basic"]["amount"] == 26000                      # still 40%
    _, newer = make_slip(tenant, 65000)
    assert lines_of(tenant, newer)["Basic"]["amount"] == 32500                    # 50% now
    # and a payslip can be brought up to date on request
    assert tenant.put("/api/payslips/%d" % pid, json={"reapply_structure": True}).status_code == 200
    assert lines_of(tenant, pid)["Basic"]["amount"] == 32500


def test_nothing_is_recalculated_when_a_payslip_is_viewed_or_downloaded(tenant):
    import openpyxl
    _, pid = make_slip(tenant, 65000)
    before = tenant.get("/api/payslips/%d" % pid).json()
    for c in structure(tenant):
        tenant.put("/api/salary-structure/%d" % c["id"], json={"percent": 1 if c["basis"] == "percent" else 0})
    after = tenant.get("/api/payslips/%d" % pid).json()
    assert (after["gross_pay"], after["net_pay"], after["lines"]) == (before["gross_pay"], before["net_pay"], before["lines"])
    sheet = openpyxl.load_workbook(io.BytesIO(tenant.get("/api/payslips/%d/xlsx" % pid).content)).active
    figures = [c.value for row in sheet.iter_rows() for c in row]
    assert 26000.0 in figures and 7800.0 in figures                                # the amounts the payslip was made with


def test_hr_can_add_change_reorder_and_remove_components(tenant):
    made = tenant.post("/api/salary-structure", json={"name": "Site allowance", "kind": "earning", "basis": "percent", "percent": 5})
    assert made.status_code == 200, made.text
    comp = made.json()["component"]
    assert tenant.put("/api/salary-structure/%d" % comp["id"], json={"percent": 7.5, "is_active": False}).json()["component"]["is_active"] is False
    ids = [c["id"] for c in structure(tenant)]
    assert tenant.post("/api/salary-structure/reorder", json={"ids": list(reversed(ids))}).status_code == 200
    assert [c["id"] for c in structure(tenant)] == list(reversed(ids))
    _, pid = make_slip(tenant, 10000)
    assert "Site allowance" not in lines_of(tenant, pid)                           # switched off
    assert tenant.delete("/api/salary-structure/%d" % comp["id"]).status_code == 200
    assert comp["id"] not in [c["id"] for c in structure(tenant)]
    assert tenant.post("/api/salary-structure", json={"name": "", "kind": "earning"}).status_code == 400
    assert tenant.post("/api/salary-structure", json={"name": "X", "kind": "earning", "basis": "percent", "percent": 150}).status_code == 400
    assert tenant.post("/api/salary-structure", json={"name": "X", "kind": "gift"}).status_code == 400


def test_removing_everything_is_respected_and_payslips_fall_back_to_the_salary(tenant):
    for c in structure(tenant):
        assert tenant.delete("/api/salary-structure/%d" % c["id"]).status_code == 200
    assert structure(tenant) == []                                                  # the starting list does not come back
    emp = make_employee(tenant, salary=3000.0, tax_rate=0.0)
    pid = tenant.post("/api/payslips", json={"employee_id": emp["id"], **SEP}).json()["id"]
    slip = tenant.get("/api/payslips/%d" % pid).json()
    assert slip["lines"] == [] and slip["basic_salary"] == 3000


def test_somebody_paid_by_the_hour_is_paid_the_way_they_were(tenant):
    emp = make_employee(tenant, salary=0.0, hourly_rate=100.0, pay_frequency="hourly")
    pid = tenant.post("/api/payslips", json={"employee_id": emp["id"], **SEP, "hours_worked": 10}).json()["id"]
    assert tenant.get("/api/payslips/%d" % pid).json()["lines"] == []


def test_a_payroll_run_uses_the_structure_too(tenant):
    emp = make_employee(tenant, salary=65000.0, tax_rate=0.0)
    res = tenant.post("/api/payroll/run", json={**SEP, "include_attendance_hours": False, "skip_existing": True, "employee_ids": [emp["id"]]})
    assert res.status_code == 200, res.text
    slip = next(p for p in tenant.get("/api/payslips").json() if p["employee_id"] == emp["id"])
    assert lines_of(tenant, slip["id"])["Conveyance"]["amount"] == 6500


def test_the_payslip_and_register_show_the_structures_own_components(tenant):
    import openpyxl
    from pypdf import PdfReader
    _, pid = make_slip(tenant, 65000)
    sheet = openpyxl.load_workbook(io.BytesIO(tenant.get("/api/payslips/%d/xlsx" % pid).content)).active
    text = [str(c.value) for row in sheet.iter_rows() for c in row if c.value not in (None, "")]
    for name in ("Basic", "HRA", "Conveyance", "Medical allowance", "Spcl allowance", "PF"):
        assert name in text, name
    assert "Site Allowance" not in text                                              # not the older fixed sheet
    assert "65000" in text                                                           # salary per month
    pdf = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(tenant.get("/api/payslips/%d/pdf" % pid).content)).pages)
    assert "Medical allowance" in pdf and "Conveyance" in pdf
    reg = openpyxl.load_workbook(io.BytesIO(tenant.get("/api/payslips/register.xlsx", params={"ids": str(pid)}).content)).active
    header = next(r for r in reg.iter_rows(values_only=True) if r and r[0] == "Payslip No.")
    for name in ("Basic", "Conveyance", "Medical allowance", "PF", "Net Pay"):
        assert name in header, name
    assert "Site Allowance" not in header


@pytest.fixture
def hr_login(tenant):
    email = "hr-%s@example.com" % uuid.uuid4().hex[:6]
    make_employee(tenant, email=email, password=PASSWORD, permission_role="hr", status="active")
    c = TestClient(main.app)
    c.__enter__()
    assert c.post("/api/employee/auth/login", json={"email": email, "password": PASSWORD}).status_code == 200
    yield c
    c.__exit__(None, None, None)


def test_hr_and_the_master_keep_the_structure_nobody_else_does(tenant, hr_login):
    assert hr_login.get("/api/salary-structure").status_code == 200
    made = hr_login.post("/api/salary-structure", json={"name": "Night allowance", "kind": "earning", "basis": "manual"})
    assert made.status_code == 200
    cid = made.json()["component"]["id"]
    assert hr_login.put("/api/salary-structure/%d" % cid, json={"name": "Night shift"}).status_code == 200
    assert hr_login.delete("/api/salary-structure/%d" % cid).status_code == 200
    email = "st-%s@example.com" % uuid.uuid4().hex[:6]
    make_employee(tenant, email=email, password=PASSWORD, permission_role="staff", status="active")
    staff = TestClient(main.app)
    staff.__enter__()
    try:
        assert staff.post("/api/employee/auth/login", json={"email": email, "password": PASSWORD}).status_code == 200
        assert staff.get("/api/salary-structure").status_code == 403
        assert staff.post("/api/salary-structure", json={"name": "X"}).status_code == 403
        assert staff.put("/api/salary-structure/1", json={"percent": 1}).status_code == 403
        assert staff.delete("/api/salary-structure/1").status_code == 403
    finally:
        staff.__exit__(None, None, None)


def test_a_structure_belongs_to_one_company(tenant, second_tenant):
    mine = named(structure(tenant), "Basic")
    assert named(structure(second_tenant), "Basic")["id"] != mine["id"]
    assert second_tenant.put("/api/salary-structure/%d" % mine["id"], json={"percent": 1}).status_code == 404
    assert second_tenant.delete("/api/salary-structure/%d" % mine["id"]).status_code == 404


def test_deleting_a_payslip_or_an_employee_takes_the_lines_with_it(tenant):
    from app import models
    from app.db import SessionLocal
    emp, pid = make_slip(tenant, 65000)
    _, other = make_slip(tenant, 65000)

    def count(**by):
        with SessionLocal() as db:
            return db.query(models.DBPayslipLine).filter_by(**by).count()
    assert count(payslip_id=pid) == 6 and count(employee_id=emp["id"]) == 6
    assert tenant.delete("/api/payslips/%d" % pid).status_code == 200
    assert count(payslip_id=pid) == 0 and count(payslip_id=other) == 6
    _, third = make_slip(tenant, 65000)
    emp3 = tenant.get("/api/payslips/%d" % third).json()["employee_id"]
    assert tenant.delete("/api/employees/%d" % emp3).status_code == 200
    assert count(employee_id=emp3) == 0
