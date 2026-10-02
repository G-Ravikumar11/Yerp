"""The sub contractor, from registration form to certificate of payment.

The office kept two workbooks: a Sub Contractor Registration Form per gang,
numbered IV0001 onwards, and for every bill three sheets - the Top Sheet
(certificate of payment), AB-1 (abstract) and MB-1 (measurement book). These
tests are the rules those papers carry: who is taken on and by whom, how the
book is written, how the certificate adds up, and who signs it in what order.
"""
import io

import openpyxl

from conftest import as_owner
from test_subcontract_orders import staff, sign_in, draft, fund_order
from test_subcontractor_bills import live_order, book, measure, raise_bill


def registered(tenant, **form):
    body = {"company_name": "M/s Arathi Malik %s" % len(form), "pan": "AFVPF9080M"}
    body.update(form)
    res = tenant.post("/api/wo/contractors", json=body)
    assert res.status_code == 200, res.text
    return res.json()


def contractor(tenant, con_id):
    return next(c for c in tenant.get("/api/wo/contractors").json()["contractors"] if c["id"] == con_id)


# --- The registration form ----------------------------------------------------

def test_vendor_codes_run_in_the_iv_series_and_continue_from_the_last(tenant):
    assert registered(tenant, company_name="One")["vendor_code"] == "IV0001"
    registered(tenant, company_name="Imported", vendor_code="IV0197")
    assert registered(tenant, company_name="Next")["vendor_code"] == "IV0198"


def test_a_vendor_code_is_not_given_to_two_sub_contractors(tenant):
    registered(tenant, company_name="First", vendor_code="IV0037")
    res = tenant.post("/api/wo/contractors", json={"company_name": "Second", "vendor_code": "iv0037"})
    assert res.status_code == 409


def test_the_form_holds_every_box_on_the_paper(tenant):
    con = registered(tenant, company_name="Vishwa Sankalp Enterprises", registered_project="ICONICA CAPITAL and CROWN",
                     joining_date="2025-04-01", pin_code="761208", city="Gajapati", state="Odisha",
                     nature_of_work="Labour shed fabrication works", contact_person="Mr.Balaram",
                     aadhaar="6713 0091 0542", bank_name="HDFC Bank", bank_account="5020 0099 710320",
                     bank_ifsc="HDFC0000042", bank_branch="Kondapur",
                     documents=["pan", "cheque", "not-a-document"], declaration_signed=True)
    c = contractor(tenant, con["id"])
    assert c["nature_of_work"] == "Labour shed fabrication works"
    assert c["aadhaar"] == "671300910542"
    assert c["bank_account"] == "50200099710320"
    assert c["documents"] == ["pan", "cheque"]
    assert c["declaration_signed"] is True
    assert c["registration_status"] == "APPROVED", "registered by the owner, it is signed off as made"


def test_an_aadhaar_is_twelve_digits(tenant):
    res = tenant.post("/api/wo/contractors", json={"company_name": "Short", "aadhaar": "1234"})
    assert res.status_code == 400


def test_a_gang_registered_on_site_waits_for_approval(tenant, portal):
    qs = staff(tenant, "planning_billing")
    pm = staff(tenant, "project_manager")
    sign_in(portal, qs)
    con = portal.post("/api/wo/contractors", json={"company_name": "Najir Hossain", "nature_of_work": "Concrete"}).json()
    assert con["registration_status"] == "PENDING"
    # Not by the person who filled it in.
    assert portal.post("/api/wo/contractors/%d/approve" % con["id"], json={}).status_code == 403

    sign_in(portal, pm)
    inbox = portal.get("/api/approvals/inbox").json()["items"]
    item = next(i for i in inbox if i["kind"] == "contractor" and i["id"] == con["id"])
    assert item["pdf"].endswith("/registration.pdf")
    res = portal.post("/api/approvals/decide", json={"kind": "contractor", "id": con["id"], "decision": "approve"})
    assert res.status_code == 200, res.text
    assert contractor(tenant, con["id"])["registration_status"] == "APPROVED"


def test_a_form_sent_back_says_why_and_goes_back_when_put_right(tenant, portal):
    qs = staff(tenant, "planning_billing")
    sign_in(portal, qs)
    con = portal.post("/api/wo/contractors", json={"company_name": "Rahul"}).json()
    assert tenant.post("/api/wo/contractors/%d/reject" % con["id"], json={}).status_code == 400
    res = tenant.post("/api/wo/contractors/%d/reject" % con["id"], json={"comments": "No PAN"})
    assert res.status_code == 200
    assert contractor(tenant, con["id"])["rejection_reason"] == "No PAN"
    res = portal.put("/api/wo/contractors/%d" % con["id"], json={"company_name": "Rahul", "pan": "BNAPH8617N"})
    assert res.status_code == 200, res.text
    assert contractor(tenant, con["id"])["registration_status"] == "PENDING"


def test_no_order_is_approved_for_a_gang_not_yet_taken_on(tenant, portal):
    qs = staff(tenant, "planning_billing")
    sign_in(portal, qs)
    con = portal.post("/api/wo/contractors", json={"company_name": "Pending Gang"}).json()
    order = draft(tenant, contractor_id=con["id"])
    res = tenant.put("/api/wo/orders/%d/boq" % order["id"], json={"lines": [
        {"activity_no": "1", "item_description": "Plaster", "uom": "Sqm", "quantity": 10, "unit_rate": 100}]})
    fund_order(tenant, res.json()["order"])
    assert tenant.post("/api/wo/orders/%d/submit" % order["id"], json={}).status_code == 200
    res = tenant.post("/api/wo/orders/%d/approve" % order["id"], json={})
    assert res.status_code == 409
    assert "not yet a registered sub contractor" in res.json()["detail"]


def test_the_registration_form_prints_and_downloads(tenant):
    con = registered(tenant, company_name="Mahabub Alam", nature_of_work="Water Proofing",
                     documents=["gst", "pan"])
    pdf = tenant.get("/api/wo/contractors/%d/registration.pdf" % con["id"])
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    xlsx = tenant.get("/api/wo/contractors/%d/registration.xlsx" % con["id"])
    ws = openpyxl.load_workbook(io.BytesIO(xlsx.content)).active
    cells = {c.value for row in ws.iter_rows() for c in row if c.value}
    assert "SUB CONTRACTOR REGISTRATION FORM" in cells
    assert con["vendor_code"] in cells and "Water Proofing" in cells
    register = tenant.get("/api/wo/contractors.xlsx")
    assert register.status_code == 200
    assert tenant.get("/api/wo/contractors.pdf").content[:4] == b"%PDF"


def _form(ws, code, name, col=1, **boxes):
    """A registration form as the office's sheets lay it out."""
    rows = [("YALAVARTI PROJECTS PVT LTD", None), ("SUB CONTRACTOR REGISTRATION FORM", None),
            ("PROJECT: CONSTRUCTION OF APTIDCO EWS HOUSES", None), ("VENDOR CODE:", code), (None, None),
            ("1. Subcontractor Personal Details", None), ("Name of the Sub Contractor:", name),
            ("Residential Address:", boxes.get("address")), ("Date of Joining :", boxes.get("joining")),
            ("Pin Code :", boxes.get("pin")), ("City :", "Katihar"), ("State :", "Bihar"),
            ("Nature of Work :", "Water Proofing"), ("Tel no. :", 8178967569), ("E - mail Id :", 0),
            ("Name of Contact Person :", name), ("Type of Entity : ", None), ("PAN no. :", boxes.get("pan")),
            ("GST Reg No :", boxes.get("gst")), ("Aadhar Card :", boxes.get("aadhaar")), (None, None),
            ("2. Bank details ", None), ("Bank Name :", "AXIS BANK LTD"), ("Account no :", 924020072127415),
            ("IFSC Code :", boxes.get("ifsc", "UTIB0000767")), ("Branch :", None), (None, None),
            ("3.Documents Required", None), ("A) GST Certificate", None), ("C) Aadhar Card", None)]
    for r, (label, value) in enumerate(rows, 1):
        if label is not None:
            ws.cell(row=r + (1 if col == 2 else 0), column=col, value=label)
        if value is not None:
            ws.cell(row=r + (1 if col == 2 else 0), column=col + 1, value=value)


def test_the_vendor_codes_workbook_is_brought_in_once(tenant):
    wb = openpyxl.Workbook()
    _form(wb.active, "IV0001", "NAIR ALAM", col=2, pan="FXTPS0149P", aadhaar="4100 5030 7735", joining="12.07.2024")
    wb.active.title = "tg"
    _form(wb.create_sheet("IV0184= Biswajit Tarai"), "IV0184", "M/s Biswajit tarai", pan="CSHPBA391N")
    # A copied form that kept its original's code: its sheet's name carries the right one.
    _form(wb.create_sheet("IV0185= Habibur Rehaman"), "IV0184", "M/s HABIBUR REHAMAN", gst="318192889417",
          ifsc=":HDFC0003297")
    rates = wb.create_sheet("Sheet1")
    rates["A1"], rates["B1"] = "S.No", "Description of the item"
    buf = io.BytesIO()
    wb.save(buf)

    res = tenant.post("/api/wo/contractors/import",
                      files={"file": ("INFRA_VENDOR_CODES.xlsx", buf.getvalue())})
    assert res.status_code == 200, res.text
    out = res.json()
    assert out["created"] == 3 and not out["skipped"]
    cons = {c["vendor_code"]: c for c in tenant.get("/api/wo/contractors").json()["contractors"]}
    assert set(cons) == {"IV0001", "IV0184", "IV0185"}
    nair = cons["IV0001"]
    assert nair["aadhaar"] == "410050307735" and nair["joining_date"] == "2024-07-12"
    assert nair["bank_account"] == "924020072127415", "an account number is its digits, not 9.24E+14"
    assert nair["email"] == "" and nair["registration_status"] == "APPROVED"
    # What is not the right shape is left blank and said.
    assert cons["IV0184"]["pan"] == ""
    assert cons["IV0185"]["gst_number"] == "" and cons["IV0185"]["bank_ifsc"] == ""
    said = " ".join(out["warnings"])
    assert "CSHPBA391N" in said and "318192889417" in said and "IV0185, the code in the sheet's name" in said
    # Brought in again, nothing is doubled.
    again = tenant.post("/api/wo/contractors/import", files={"file": ("again.xlsx", buf.getvalue())}).json()
    assert again["created"] == 0 and again["updated"] == 3
    assert registered(tenant, company_name="After import")["vendor_code"] == "IV0186"


# --- The measurement book -----------------------------------------------------------

def painting_order(tenant, **over):
    """An order shaped like the painting gang's: a heading and two items,
    priced as the abstract prices them."""
    order = draft(tenant, gst_rate=0, tds_rate=2, retention_percent=5, **over)
    res = tenant.put("/api/wo/orders/%d/boq" % order["id"], json={"lines": [
        {"activity_no": "", "item_description": "Painting Works", "is_header": True},
        {"activity_no": "1", "item_description": "Buffering works", "uom": "Sqm", "quantity": 6000, "unit_rate": 6},
        {"activity_no": "2", "item_description": "Hole Packing", "uom": "Sqm", "quantity": 20000, "unit_rate": 8}]})
    assert res.status_code == 200, res.text
    order = fund_order(tenant, res.json()["order"])
    tenant.post("/api/wo/orders/%d/submit" % order["id"], json={})
    res = tenant.post("/api/wo/orders/%d/approve" % order["id"], json={})
    assert res.status_code == 200, res.text
    return res.json()["order"]


def items_of(tenant, order_id):
    return [l for l in book(tenant, order_id)["lines"] if not l.get("is_header")]


def test_a_line_is_nos_times_nom_times_its_dimensions(tenant):
    order = painting_order(tenant)
    buff = items_of(tenant, order["id"])[0]["item_id"]
    # Ceiling - Hall: 6 flats x 3 floors x 3.25 x 2.5 = 146.25
    out = measure(tenant, order["id"], buff, 0, dimensions=[
        {"particulars": "Ceiling - Hall", "nos": 6, "nom": 3, "length": 3.25, "breadth": 2.5}])
    assert out["measured_to_date"] == 146.25


def test_headings_group_the_lines_and_measure_nothing(tenant):
    order = painting_order(tenant)
    buff = items_of(tenant, order["id"])[0]["item_id"]
    measure(tenant, order["id"], buff, 0, dimensions=[
        {"particulars": "Living Room", "is_heading": True},
        {"particulars": "Long Walls", "nos": 6, "nom": 6, "length": 4.15, "depth": 2.7},
        {"particulars": "Deductions", "is_heading": True},
        {"particulars": "Main Door Deduct", "nos": 6, "nom": 3, "length": 0.9, "depth": 1.9, "deduct": True}])
    entry = book(tenant, order["id"])["entries"][0]
    assert entry["quantity"] == round(403.38 - 30.78, 2)
    assert [d["is_heading"] for d in entry["dimensions"]] == [True, False, True, False]


def test_a_block_measured_once_counts_for_every_block_built(tenant):
    order = painting_order(tenant)
    hole = items_of(tenant, order["id"])[1]["item_id"]
    measure(tenant, order["id"], hole, 0, multiplier=4, location="300 SFT-Block A-25,24,16,18", dimensions=[
        {"particulars": "Long Walls", "nos": 6, "nom": 8, "length": 3.45, "depth": 2.75}])
    entry = book(tenant, order["id"])["entries"][0]
    assert entry["multiplier"] == 4
    assert entry["quantity"] == round(6 * 8 * 3.45 * 2.75 * 4, 2)


def mb_workbook():
    """The MB-1 sheet as the site keeps it - the painting gang's bill in small."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MB-1 (2)"
    ws["A1"] = "YALAVARTI INFRA PROJECTS"
    ws["A2"] = "Name of the Work :- PAINTING WORKS"
    ws["A4"] = "Bill no :- 01 (Measurements)"
    ws["I4"] = "Date :- 28-02-2026"
    for c, h in enumerate(("S.No", "Description", "UoM", "No's", "NoM", "Length", "Width", "Height",
                           "Total Quantity", "Remarks"), 1):
        ws.cell(row=5, column=c, value=h)
    rows = [
        (2, "BUFFERING WORKS"), ("b", "430 SFT-Block C-3 FF,SF,TF"), (None, "Type-1 Flats - 6 No's"),
        (None, "Ceiling"), (None, "Ceiling - Hall", "Sqm", 6, 3, 3.25, 2.5, None, 146.25),
        (None, "Living Room"), (None, "Long Walls", "Sqm", 6, 6, 4.15, None, 2.7, 403.38),
        (None, "Deductions"), (None, "Main Door Deduct", "Sqm", -6, 3, 0.9, None, 1.9, -30.78),
        (None, "Dado Deduct", "Sqm", -6, 3, "=2.1+0.75+0.75", None, 1.2, -77.76),
        (None, "Total Quantity for one Block", None, None, None, "Total Quantity", None, None, 441.09),
        (None, "Total Quantity for Block No. - C3 FF,SF,TF", None, None, None, "Total Quantity for 1 Blocks", None, None, 441.09),
        (3, "Hole packing"), ("a", "300 SFT-Block A-25,24,16,18"),
        (None, "Long Walls", "Sqm", 6, 8, 3.45, None, 2.75, 455.4),
        (None, "Window Deduct", "Sqm", -1, 20, 0.9, None, 1.2, -21.6),
        (None, "Total Quantity for one Block", None, None, None, "Total Quantity", None, None, 433.8),
        (None, "Total Quantity for Block No. - A-25,24,16,18", None, None, None, "Total Quantity for 4 Blocks", None, None, 1735.2),
    ]
    for r, row in enumerate(rows, 6):
        for c, v in enumerate(row, 1):
            if v is not None:
                ws.cell(row=r, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_the_measurement_book_is_read_from_its_sheet(tenant):
    order = painting_order(tenant)
    raw = mb_workbook()
    res = tenant.post("/api/sub-mb/%d/import" % order["id"], files={"file": ("mb.xlsx", raw)})
    assert res.status_code == 200, res.text
    pv = res.json()
    assert pv["committed"] is False
    assert [s["item"] for s in pv["sections"]] == ["1 Buffering works", "2 Hole Packing"]
    assert pv["sections"][0]["quantity"] == 441.09
    assert pv["sections"][1]["entries"][0]["multiplier"] == 4
    assert pv["sections"][1]["quantity"] == 1735.2
    assert not pv["warnings"], "the lines come to what the sheet says"
    assert book(tenant, order["id"])["entries"] == [], "a preview records nothing"

    res = tenant.post("/api/sub-mb/%d/import" % order["id"], files={"file": ("mb.xlsx", raw)},
                      data={"commit": "1"})
    assert res.status_code == 200, res.text
    lines = {l["description"]: l["measured_to_date"] for l in items_of(tenant, order["id"])}
    assert lines == {"Buffering works": 441.09, "Hole Packing": 1735.2}
    entry = [e for e in book(tenant, order["id"])["entries"] if e["multiplier"] == 4][0]
    assert entry["location"] == "300 SFT-Block A-25,24,16,18"
    assert entry["measured_on"] == "2026-02-28"


def test_the_preview_returns_the_lines_when_the_measure_window_asks_for_them(tenant):
    """The measure window fills its own grid from a sheet, so it needs each entry's
    dimension lines back - and only when it asks, as the import screen does not."""
    order = painting_order(tenant)
    raw = mb_workbook()
    plain = tenant.post("/api/sub-mb/%d/import" % order["id"], files={"file": ("mb.xlsx", raw)}).json()
    assert "dims" not in plain["sections"][0]["entries"][0]

    res = tenant.post("/api/sub-mb/%d/import" % order["id"], files={"file": ("mb.xlsx", raw)},
                      data={"include_dims": "1"})
    assert res.status_code == 200, res.text
    entry = res.json()["sections"][1]["entries"][0]
    assert entry["multiplier"] == 4
    lines = [d for d in entry["dims"] if not d["is_heading"]]
    assert [d["particulars"] for d in lines][-1] == "Window Deduct"
    assert lines[0]["nos"] is not None and lines[0]["length"] is not None
    assert book(tenant, order["id"])["entries"] == [], "asking for the lines records nothing"


def test_an_import_past_the_order_records_none_of_it(tenant):
    order = draft(tenant, gst_rate=0)
    res = tenant.put("/api/wo/orders/%d/boq" % order["id"], json={"lines": [
        {"activity_no": "1", "item_description": "Buffering works", "uom": "Sqm", "quantity": 6000, "unit_rate": 6},
        {"activity_no": "2", "item_description": "Hole Packing", "uom": "Sqm", "quantity": 100, "unit_rate": 8}]})
    order = fund_order(tenant, res.json()["order"])
    tenant.post("/api/wo/orders/%d/submit" % order["id"], json={})
    tenant.post("/api/wo/orders/%d/approve" % order["id"], json={})
    res = tenant.post("/api/sub-mb/%d/import" % order["id"], files={"file": ("mb.xlsx", mb_workbook())},
                      data={"commit": "1"})
    assert res.status_code == 409
    assert "nothing was imported" in res.json()["detail"]
    assert book(tenant, order["id"])["entries"] == []


# --- The certificate of payment ---------------------------------------------------

def measured_painting_bill(tenant):
    order = painting_order(tenant)
    buff, hole = [l["item_id"] for l in items_of(tenant, order["id"])]
    measure(tenant, order["id"], buff, 5964.42)
    measure(tenant, order["id"], hole, 18424.8)
    res = raise_bill(tenant, order["id"], period_from="2026-01-02", period_to="2026-02-28",
                     bill_date="2026-02-17", work_type="Putty and Painting works")
    assert res.status_code == 200, res.text
    return order, res.json()["bill"]


def test_the_certificate_adds_up_as_the_top_sheet_does(tenant):
    """Arti Malik's first bill, figure for figure: 1,83,184.92 of work, 5%
    retention and 2% TDS on it, whole rupees, 1,70,362 to pay."""
    order, bill = measured_painting_bill(tenant)
    assert bill["this_bill"] == 183184.92
    assert bill["gross_value"] == 183185
    assert bill["retention_amount"] == 9159
    assert bill["tds_amount"] == 3664
    assert bill["net_payable"] == 170362
    assert bill["work_type"] == "Putty and Painting works"
    assert bill["bill_date"] == "2026-02-17"


def test_gst_is_charged_on_the_measured_value_not_the_gross_after_debit_notes(tenant):
    order = live_order(tenant, retention_percent=5)          # 18% GST, 1% TDS
    item = book(tenant, order["id"])["lines"][0]["item_id"]
    measure(tenant, order["id"], item, 10)                   # 68,000 of work
    bill = raise_bill(tenant, order["id"]).json()["bill"]
    res = tenant.put("/api/sub-bills/%d" % bill["id"], json={"debit_notes": 8000, "hsn_sac": "995414"})
    assert res.status_code == 200, res.text
    bill = res.json()["bill"]
    assert bill["gross_value"] == 60000
    assert bill["gst_amount"] == 12240                       # 18% of the 68,000 measured
    assert bill["retention_amount"] == 3400                  # 5% of the 68,000 of work
    assert bill["tds_amount"] == 680
    assert bill["net_payable"] == 60000 + 12240 - 3400 - 680


def test_a_sent_bill_s_boxes_are_not_changed(tenant):
    order = live_order(tenant)
    item = book(tenant, order["id"])["lines"][0]["item_id"]
    measure(tenant, order["id"], item, 10)
    bill = raise_bill(tenant, order["id"]).json()["bill"]
    tenant.post("/api/sub-bills/%d/submit" % bill["id"], json={})
    assert tenant.put("/api/sub-bills/%d" % bill["id"], json={"hsn_sac": "9954"}).status_code == 409


def test_the_bill_downloads_as_its_three_sheets(tenant):
    order, bill = measured_painting_bill(tenant)
    res = tenant.get("/api/sub-bills/%d/export.xlsx" % bill["id"])
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Top Sheet", "AB-1", "MB-1"]
    top = {c.value for row in wb["Top Sheet"].iter_rows() for c in row if c.value is not None}
    assert {"CERTIFICATE OF PAYMENT", "Gross Total Value ", "Net Amount for Payment ", 170362, 9159} <= top
    ab = {c.value for row in wb["AB-1"].iter_rows() for c in row if c.value is not None}
    assert {"ABSTRACT SHEET", "Painting Works", 35786.52, 147398.4, 183184.92} <= ab
    for url in ("/api/sub-bills/%d/document.pdf", "/api/sub-bills/%d/export.pdf"):
        pdf = tenant.get(url % bill["id"])
        assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"


def test_the_second_certificate_carries_the_first_as_upto_previous(tenant):
    order = painting_order(tenant)
    buff = items_of(tenant, order["id"])[0]["item_id"]
    measure(tenant, order["id"], buff, 1000)
    first = raise_bill(tenant, order["id"]).json()["bill"]
    tenant.post("/api/sub-bills/%d/submit" % first["id"], json={})
    tenant.post("/api/sub-bills/%d/certify" % first["id"], json={})
    measure(tenant, order["id"], buff, 500)
    second = raise_bill(tenant, order["id"]).json()["bill"]
    wb = openpyxl.load_workbook(io.BytesIO(tenant.get("/api/sub-bills/%d/export.xlsx" % second["id"]).content))
    rows = {r[1]: r for r in wb["Top Sheet"].iter_rows(values_only=True) if r[0] == "4.01"}
    work = next(iter(rows.values()))
    assert work[3:6] == (9000, 6000, 3000), "up to this = up to previous + this bill"


# --- Who signs it, in turn ---------------------------------------------------------

def test_a_bill_climbs_its_route_before_it_is_certified(tenant, portal):
    """Prepared by the QS, certified by the Head QS, approved by the site
    incharge - one signature at a time, each told when it is their turn."""
    qs = staff(tenant, "planning_billing")
    head_qs = staff(tenant, "project_manager")
    incharge = staff(tenant, "head_projects")
    order = live_order(tenant)
    item = book(tenant, order["id"])["lines"][0]["item_id"]
    measure(tenant, order["id"], item, 10)

    sign_in(portal, qs)
    bill = raise_bill(portal, order["id"]).json()["bill"]
    sent = portal.post("/api/sub-bills/%d/submit" % bill["id"], json={}).json()["bill"]
    assert [s["name"] for s in sent["route"]] == [
        "%s %s" % (head_qs["first_name"], head_qs["last_name"]),
        "%s %s" % (incharge["first_name"], incharge["last_name"])]
    assert portal.post("/api/sub-bills/%d/certify" % bill["id"], json={}).status_code == 403

    sign_in(portal, incharge)
    assert not [i for i in portal.get("/api/approvals/inbox").json()["items"] if i["kind"] == "sub_bill"]
    early = portal.post("/api/sub-bills/%d/certify" % bill["id"], json={})
    assert early.status_code == 403 and "comes to you after that" in early.json()["detail"]

    sign_in(portal, head_qs)
    item = next(i for i in portal.get("/api/approvals/inbox").json()["items"] if i["kind"] == "sub_bill")
    assert item["approve_label"] == "Sign and pass on"
    res = portal.post("/api/approvals/decide", json={"kind": "sub_bill", "id": bill["id"], "decision": "approve"})
    assert res.status_code == 200, res.text
    assert res.json()["bill"]["status"] == "SUBMITTED"

    sign_in(portal, incharge)
    res = portal.post("/api/sub-bills/%d/certify" % bill["id"], json={})
    assert res.status_code == 200, res.text
    done = res.json()["bill"]
    assert done["status"] == "CERTIFIED"
    assert done["approved_by_name"] == "%s %s" % (incharge["first_name"], incharge["last_name"])
    assert done["submitted_by_name"] == "%s %s" % (qs["first_name"], qs["last_name"])
    assert [s["status"] for s in done["route"]] == ["approved", "approved"]


def test_sending_a_bill_back_at_any_step_returns_it_to_draft(tenant, portal):
    qs = staff(tenant, "planning_billing")
    head_qs = staff(tenant, "project_manager")
    order = live_order(tenant)
    item = book(tenant, order["id"])["lines"][0]["item_id"]
    measure(tenant, order["id"], item, 10)
    sign_in(portal, qs)
    bill = raise_bill(portal, order["id"]).json()["bill"]
    portal.post("/api/sub-bills/%d/submit" % bill["id"], json={})
    sign_in(portal, head_qs)
    res = portal.post("/api/sub-bills/%d/reject" % bill["id"], json={"comments": "Recheck the deductions"})
    assert res.status_code == 200, res.text
    assert res.json()["bill"]["status"] == "DRAFT"
    assert res.json()["bill"]["remarks"] == "Recheck the deductions"


def test_the_owner_s_signature_is_the_last_word(tenant, portal):
    qs = staff(tenant, "planning_billing")
    staff(tenant, "project_manager")
    staff(tenant, "head_projects")
    order = live_order(tenant)
    item = book(tenant, order["id"])["lines"][0]["item_id"]
    measure(tenant, order["id"], item, 10)
    sign_in(portal, qs)
    bill = raise_bill(portal, order["id"]).json()["bill"]
    portal.post("/api/sub-bills/%d/submit" % bill["id"], json={})
    as_owner(tenant)
    res = tenant.post("/api/sub-bills/%d/certify" % bill["id"], json={})
    assert res.json()["bill"]["status"] == "CERTIFIED"
    assert [s["status"] for s in res.json()["bill"]["route"]][:2] == ["skipped", "skipped"]


def test_the_sub_contractor_accepts_the_certificate(tenant):
    order, bill = measured_painting_bill(tenant)
    assert tenant.post("/api/sub-bills/%d/accept" % bill["id"], json={}).status_code == 409, "a draft is not accepted"
    tenant.post("/api/sub-bills/%d/submit" % bill["id"], json={})
    res = tenant.post("/api/sub-bills/%d/accept" % bill["id"], json={"name": "Arti Malik"})
    assert res.status_code == 200, res.text
    assert res.json()["bill"]["accepted_by_name"] == "Arti Malik"
    assert res.json()["bill"]["status"] == "SUBMITTED", "accepting is a signature, not a step"


def test_the_gang_accepts_the_certificate_from_its_own_login(tenant, portal):
    from test_partner_portal import gang_with_a_certified_bill, invite, join
    order, bill = gang_with_a_certified_bill(tenant)
    join(portal, invite(tenant, "contractor", order["contractor_id"], "babu@gang.in")["invite_url"])
    mine = portal.get("/api/portal/bills").json()["bills"][0]
    assert mine["can_accept"] is True
    res = portal.post("/api/portal/bills/%d/accept" % bill["id"])
    assert res.status_code == 200, res.text
    assert portal.post("/api/portal/bills/%d/accept" % bill["id"]).status_code == 409, "a signature is given once"
    assert tenant.get("/api/sub-bills/%d" % bill["id"]).json()["accepted_by_name"] == "Babu"


def test_only_plain_arithmetic_in_a_sheet_is_worked_out():
    """=2.1+0.75+0.75 is a length written the way the book writes it; a
    formula that is anything more is not guessed at, nor run."""
    import sheet_forms
    assert sheet_forms._arithmetic("2.1+0.75+0.75") == 3.6
    assert sheet_forms._arithmetic("+2.7-0.915-0.435") == 1.35
    for hostile in ("9**9**9**9", "A1+2", "__import__('os')", "1/0", "SUM(1,2)"):
        assert sheet_forms._arithmetic(hostile) is None


def test_a_section_is_matched_by_the_item_code_written_in_its_heading(tenant):
    order = draft(tenant, gst_rate=0)
    res = tenant.put("/api/wo/orders/%d/boq" % order["id"], json={"lines": [
        {"activity_no": "1", "item_code": "STR001", "item_description": "Wall finishing, level one", "uom": "Sqm", "quantity": 6000, "unit_rate": 6},
        {"activity_no": "2", "item_code": "STR002", "item_description": "Packing of holes", "uom": "Sqm", "quantity": 6000, "unit_rate": 8}]})
    order = fund_order(tenant, res.json()["order"])
    tenant.post("/api/wo/orders/%d/submit" % order["id"], json={})
    tenant.post("/api/wo/orders/%d/approve" % order["id"], json={})
    raw = mb_workbook().replace(b"x", b"x")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb.active
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "BUFFERING WORKS":
                cell.value = "STR002 BUFFERING WORKS"          # words that match no item, but the code does
            if cell.value == "Hole packing":
                cell.value = "STR001 Hole packing"
    buf = io.BytesIO()
    wb.save(buf)
    pv = tenant.post("/api/sub-mb/%d/import" % order["id"], files={"file": ("mb.xlsx", buf.getvalue())}).json()
    assert [s["item"] for s in pv["sections"]] == ["STR002 2 Packing of holes", "STR001 1 Wall finishing, level one"]


def test_a_serial_number_is_not_taken_for_an_activity_number(tenant):
    """Section 2 of the sheet is not the order's line 2: only words or a code decide it."""
    order = painting_order(tenant)
    pv = tenant.post("/api/sub-mb/%d/import" % order["id"], files={"file": ("mb.xlsx", mb_workbook())}).json()
    assert [s["item"] for s in pv["sections"]] == ["1 Buffering works", "2 Hole Packing"]
