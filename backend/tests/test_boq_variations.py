"""Variations to the BOQ: extra items and quantities past it, priced, approved on a route, and applied as the next revision."""
from test_boq import CLIENT_SHEET, import_lines, make_boq, saved, gang_order
from test_subcontract_orders import draft, fund_order
from test_subcontract_orders import staff, sign_in


def ready(tenant):
    """A BOQ with its lines saved and coded, and R0 issued so the working revision is R1."""
    _, boq = make_boq(tenant)
    out = saved(tenant, boq)
    lines = {l["sno"]: l for l in out["lines"]}
    sections = [l for l in out["lines"] if l["kind"] == "section"]
    return boq, lines, sections


def variation(tenant, boq, lines, reason="Site conditions differ"):
    return tenant.post("/api/boq-variations", json={"boq_id": boq["id"], "reason": reason, "lines": lines})


def current_lines(tenant, boq):
    return {l["sno"]: l for l in tenant.get("/api/boqs/%d" % boq["id"]).json()["lines"]}


# --- raising one ---------------------------------------------------------------------------------------------------

def test_a_variation_has_quantity_changes_and_extra_items_each_priced(tenant):
    boq, lines, sections = ready(tenant)
    res = variation(tenant, boq, [
        {"kind": "quantity", "boq_key": lines["1.1"]["key"], "change_qty": 80},                       # rate defaults to the BOQ's 180
        {"kind": "extra", "description": "Dewatering during excavation", "uom": "day", "change_qty": 20, "rate": 4500,
         "section_key": sections[0]["key"]}])
    assert res.status_code == 200, res.text
    v = res.json()["variation"]
    assert v["number"].endswith("/BV-01") and v["status"] == "DRAFT"
    assert v["value"] == 80 * 180 + 20 * 4500
    q, e = v["lines"]
    assert (q["old_qty"], q["new_qty"], q["rate"], q["description"]) == (500, 580, 180, "Excavation in all types of soil")
    assert e["kind"] == "extra" and e["amount"] == 90000


def test_a_variation_is_refused_in_words_when_a_line_is_wrong(tenant):
    boq, lines, _ = ready(tenant)
    key = lines["1.1"]["key"]
    cases = [
        ([{"kind": "quantity", "boq_key": "nope", "change_qty": 5}], "choose the BOQ line"),
        ([{"kind": "quantity", "boq_key": key, "change_qty": 0}], "say how much"),
        ([{"kind": "quantity", "boq_key": key, "change_qty": -900}], "below nothing"),
        ([{"kind": "quantity", "boq_key": key, "change_qty": 5}, {"kind": "quantity", "boq_key": key, "change_qty": 5}], "twice"),
        ([{"kind": "extra", "description": "", "change_qty": 5, "rate": 10}], "description"),
        ([{"kind": "extra", "description": "X", "change_qty": 5, "rate": 0}], "rate"),
        ([{"kind": "extra", "description": "X", "change_qty": 0, "rate": 10}], "quantity"),
        ([], "at least one line"),
    ]
    for rows, word in cases:
        res = variation(tenant, boq, rows)
        assert res.status_code == 400 and word in res.json()["detail"], (rows, res.text)


def test_a_draft_can_be_changed_and_deleted_but_not_once_sent(tenant):
    boq, lines, _ = ready(tenant)
    v = variation(tenant, boq, [{"kind": "quantity", "boq_key": lines["1.1"]["key"], "change_qty": 10}]).json()["variation"]
    res = tenant.put("/api/boq-variations/%d" % v["id"], json={"reason": "more", "lines": [{"kind": "quantity", "boq_key": lines["1.1"]["key"], "change_qty": 25, "rate": 200}]})
    assert res.status_code == 200 and res.json()["variation"]["value"] == 5000
    assert tenant.post("/api/boq-variations/%d/submit" % v["id"]).status_code == 200
    assert tenant.put("/api/boq-variations/%d" % v["id"], json={"lines": []}).status_code == 409
    assert tenant.delete("/api/boq-variations/%d" % v["id"]).status_code == 409
    other = variation(tenant, boq, [{"kind": "quantity", "boq_key": lines["1.2"]["key"], "change_qty": 1}]).json()["variation"]
    assert tenant.delete("/api/boq-variations/%d" % other["id"]).status_code == 200


# --- approving: the next revision -------------------------------------------------------------------------------------

def test_approval_moves_the_boq_to_its_next_revision_with_the_changes_in_it(tenant):
    boq, lines, sections = ready(tenant)
    v = variation(tenant, boq, [
        {"kind": "quantity", "boq_key": lines["1.1"]["key"], "change_qty": 80},
        {"kind": "extra", "description": "Dewatering during excavation", "uom": "day", "change_qty": 20, "rate": 4500, "section_key": sections[0]["key"], "sno": "1.9"}]).json()["variation"]
    assert tenant.post("/api/boq-variations/%d/submit" % v["id"]).status_code == 200
    done = tenant.post("/api/boq-variations/%d/approve" % v["id"], json={})
    assert done.status_code == 200, done.text
    assert done.json()["variation"]["status"] == "APPROVED" and done.json()["variation"]["applied_rev"].startswith("R1")
    book = tenant.get("/api/boqs/%d" % boq["id"]).json()
    assert book["boq"]["current_rev"] == 1 and [r["status"] for r in book["revisions"]] == ["ISSUED", "OPEN"]
    now = {l["sno"]: l for l in book["lines"]}
    assert now["1.1"]["quantity"] == 580 and now["1.1"]["amount"] == 580 * 180
    extra = now["1.9"]
    assert extra["item_code"] and extra["quantity"] == 20 and extra["rate"] == 4500
    order = [l["sno"] for l in book["lines"]]
    assert order.index("1.9") < order.index("B"), "the extra item sits in the section it was put under"
    assert book["total"] == 1472800 + 80 * 180 + 90000
    old = tenant.get("/api/boqs/%d" % boq["id"], params={"rev": 0}).json()
    assert old["total"] == 1472800, "the issued revision is untouched"
    ch = tenant.get("/api/boqs/%d/changes" % boq["id"], params={"frm": 0}).json()
    assert {c["change"] for c in ch["changes"]} == {"added", "changed"} and ch["difference"] == 80 * 180 + 90000


def test_the_clients_work_order_follows_the_boq(tenant):
    boq, lines, sections = ready(tenant)
    order = tenant.post("/api/boqs/%d/client-order" % boq["id"], json={}).json()["work_order"]
    v = variation(tenant, boq, [
        {"kind": "quantity", "boq_key": lines["1.2"]["key"], "change_qty": 30},
        {"kind": "extra", "description": "Anti-termite treatment", "uom": "sqm", "change_qty": 200, "rate": 60}]).json()["variation"]
    tenant.post("/api/boq-variations/%d/submit" % v["id"])
    assert tenant.post("/api/boq-variations/%d/approve" % v["id"], json={}).status_code == 200
    wo = tenant.get("/api/erp/work-orders/%d" % order["id"]).json()
    by = {l["fg_code"]: l for l in wo["lines"]}
    assert by[lines["1.2"]["item_code"]]["qty"] == 150 and len(wo["lines"]) == 5
    assert wo["total_value"] == 1472800 + 30 * 8200 + 200 * 60


def test_a_variation_whose_line_was_removed_meanwhile_cannot_be_applied(tenant):
    boq, lines, _ = ready(tenant)
    v = variation(tenant, boq, [{"kind": "quantity", "boq_key": lines["1.1"]["key"], "change_qty": 5}]).json()["variation"]
    tenant.post("/api/boq-variations/%d/submit" % v["id"])
    kept = [l for l in tenant.get("/api/boqs/%d" % boq["id"]).json()["lines"] if l["sno"] != "1.1"]
    tenant.put("/api/boqs/%d/lines" % boq["id"], json={"lines": kept})
    res = tenant.post("/api/boq-variations/%d/approve" % v["id"], json={})
    assert res.status_code == 409 and "no longer in the BOQ" in res.json()["detail"]
    assert tenant.get("/api/boq-variations/%d" % v["id"]).json()["status"] == "SUBMITTED"


# --- the route ----------------------------------------------------------------------------------------------------------

def test_sending_it_back_needs_a_reason_and_returns_it_to_draft(tenant):
    boq, lines, _ = ready(tenant)
    v = variation(tenant, boq, [{"kind": "quantity", "boq_key": lines["1.1"]["key"], "change_qty": 5}]).json()["variation"]
    tenant.post("/api/boq-variations/%d/submit" % v["id"])
    assert tenant.post("/api/boq-variations/%d/reject" % v["id"], json={}).status_code == 400
    back = tenant.post("/api/boq-variations/%d/reject" % v["id"], json={"comments": "Rate too high"}).json()["variation"]
    assert back["status"] == "DRAFT" and back["rejection_reason"] == "Rate too high" and back["editable"]
    assert tenant.post("/api/boq-variations/%d/submit" % v["id"]).status_code == 200


def test_it_waits_in_the_approvals_inbox_and_is_decided_there(tenant):
    boq, lines, _ = ready(tenant)
    v = variation(tenant, boq, [{"kind": "quantity", "boq_key": lines["1.1"]["key"], "change_qty": 5}]).json()["variation"]
    tenant.post("/api/boq-variations/%d/submit" % v["id"])
    inbox = tenant.get("/api/approvals/inbox").json()["items"]
    mine = [i for i in inbox if i["kind"] == "boq_variation"]
    assert len(mine) == 1 and mine[0]["number"] == v["number"] and mine[0]["amount"] == 900
    res = tenant.post("/api/approvals/decide", json={"kind": "boq_variation", "id": v["id"], "decision": "approve"})
    assert res.status_code == 200, res.text
    assert tenant.get("/api/boq-variations/%d" % v["id"]).json()["status"] == "APPROVED"


def test_staff_who_may_not_approve_cannot_and_the_master_can(tenant, portal):
    boq, lines, _ = ready(tenant)
    site = staff(tenant, "staff")
    sign_in(portal, site)
    v = variation(tenant, boq, [{"kind": "quantity", "boq_key": lines["1.1"]["key"], "change_qty": 5}]).json()["variation"]
    tenant.post("/api/boq-variations/%d/submit" % v["id"])
    refused = portal.post("/api/boq-variations/%d/approve" % v["id"], json={})
    assert refused.status_code in (403, 404)
    assert tenant.get("/api/boq-variations/%d" % v["id"]).json()["status"] == "SUBMITTED"
    assert tenant.post("/api/boq-variations/%d/approve" % v["id"], json={}).status_code == 200


# --- what the work already done calls for ------------------------------------------------------------------------------

def test_the_lines_executed_past_the_boq_are_suggested(tenant):
    boq, lines, _ = ready(tenant)
    order = draft(tenant)
    pulled = tenant.post("/api/wo/orders/%d/boq-lines" % order["id"], json={"lines": [{"key": lines["1.1"]["key"], "quantity": 500, "rate": 150}]}).json()
    sched = dict(pulled["lines"][0], tolerance_percent=50)                          # the gang may measure 50% over
    order = tenant.put("/api/wo/orders/%d/boq" % order["id"], json={"lines": [sched]}).json()["order"]
    fund_order(tenant, order)
    assert tenant.post("/api/wo/orders/%d/submit" % order["id"], json={}).status_code == 200
    assert tenant.post("/api/wo/orders/%d/approve" % order["id"], json={}).status_code == 200
    item = tenant.get("/api/sub-mb/%d" % order["id"]).json()["lines"][0]["item_id"]
    assert tenant.post("/api/sub-mb/%d/entries" % order["id"], json={"item_id": item, "quantity": 620}).status_code == 200
    s = tenant.get("/api/boq-variations/suggest/%d" % boq["id"]).json()
    assert len(s["lines"]) == 1 and s["lines"][0]["boq_key"] == lines["1.1"]["key"]
    assert s["lines"][0]["change_qty"] == 120 and s["value"] == 120 * 180


def test_an_extra_item_with_no_number_is_numbered_after_its_variation(tenant):
    boq, lines, _ = ready(tenant)
    v = variation(tenant, boq, [{"kind": "extra", "description": "Rock breaking", "uom": "cum", "change_qty": 5, "rate": 1800},
                                {"kind": "extra", "description": "Dewatering", "uom": "day", "change_qty": 3, "rate": 4500}]).json()["variation"]
    tenant.post("/api/boq-variations/%d/submit" % v["id"])
    tenant.post("/api/boq-variations/%d/approve" % v["id"], json={})
    got = current_lines(tenant, boq)
    assert got["V1/1"]["description"] == "Rock breaking" and got["V1/2"]["description"] == "Dewatering"


# --- on the client's printed bill --------------------------------------------------------------------------------------

def test_the_clients_printed_bill_lists_the_variations_agreed_on_its_order(tenant):
    import io
    import re
    import openpyxl
    import pypdf
    from conftest import owner_places
    from test_measurement_and_ra_bills import measure, line_of, raise_bill
    boq, lines, _ = ready(tenant)
    order = tenant.post("/api/boqs/%d/client-order" % boq["id"], json={}).json()["work_order"]
    # a client order is approved against a material budget
    rm = tenant.post("/api/erp/items/bulk", json={"items": [{"kind": "RM", "item_name": "Cement for the works", "units_of_measure": "Nos"}]}).json()["codes"][0]
    first = tenant.get("/api/erp/work-orders/%d" % order["id"]).json()["lines"][0]
    assert tenant.post("/api/erp/bom/build", json={"work_order_id": order["id"], "lines": [
        {"fg_code": first["fg_code"], "rm_code": rm, "qty": 10, "rate": 50}]}).status_code == 200
    owner_places(tenant, order["id"])
    measure(tenant, order["id"], line_of(tenant, order["id"]), 10)
    before = raise_bill(tenant, order["id"]).json()["bill"]
    plain = pypdf.PdfReader(io.BytesIO(tenant.get("/api/ra-bills/%d/document.pdf" % before["id"]).content))
    assert "VARIATIONS AGREED" not in " ".join(p.extract_text() or "" for p in plain.pages), "no variation, no section"

    v = variation(tenant, boq, [{"kind": "quantity", "boq_key": lines["1.2"]["key"], "change_qty": 30},
                                {"kind": "extra", "description": "Anti-termite treatment", "uom": "sqm", "change_qty": 200, "rate": 60}],
                  reason="Client asked for treatment").json()["variation"]
    tenant.post("/api/boq-variations/%d/submit" % v["id"])
    assert tenant.post("/api/boq-variations/%d/approve" % v["id"], json={}).status_code == 200

    pdf = tenant.get("/api/ra-bills/%d/document.pdf" % before["id"])
    flat = re.sub(r"\s+", " ", " ".join(p.extract_text() or "" for p in pypdf.PdfReader(io.BytesIO(pdf.content)).pages))
    assert "VARIATIONS AGREED ON THIS ORDER" in flat and "BV-01" in flat and "Client asked for treatment" in flat
    assert "Anti-termite treatment" in flat
    assert "Order value as first placed" in flat and "14,72,800.00" in flat      # 1,472,800 as first placed
    assert "Order value as varied" in flat and "17,30,800.00" in flat            # + 30 x 8,200 + 200 x 60 = 2,58,000
    x = tenant.get("/api/ra-bills/%d/export.xlsx" % before["id"])
    cells = [str(c.value) for ws in openpyxl.load_workbook(io.BytesIO(x.content)).worksheets for row in ws.iter_rows() for c in row if c.value]
    assert any("variations agreed" in c and "BV-01" in c for c in cells)


def test_the_clients_work_order_prints_its_variations_too(tenant):
    import io
    import re
    import openpyxl
    import pypdf
    boq, lines, _ = ready(tenant)
    order = tenant.post("/api/boqs/%d/client-order" % boq["id"], json={}).json()["work_order"]

    def printed():
        pdf = tenant.get("/api/erp/work-orders/%d/export.pdf" % order["id"])
        assert pdf.status_code == 200, pdf.text[:200]
        return re.sub(r"\s+", " ", " ".join(p.extract_text() or "" for p in pypdf.PdfReader(io.BytesIO(pdf.content)).pages))

    assert "VARIATIONS AGREED" not in printed()
    v = variation(tenant, boq, [{"kind": "quantity", "boq_key": lines["1.2"]["key"], "change_qty": 30},
                                {"kind": "extra", "description": "Anti-termite treatment", "uom": "sqm", "change_qty": 200, "rate": 60}],
                  reason="Client asked for treatment").json()["variation"]
    tenant.post("/api/boq-variations/%d/submit" % v["id"])
    tenant.post("/api/boq-variations/%d/approve" % v["id"], json={})
    text = printed()
    assert "VARIATIONS AGREED ON THIS ORDER" in text and "BV-01" in text and "Client asked for treatment" in text
    assert "Order value as first placed" in text and "14,72,800" in text and "17,30,800" in text
    x = tenant.get("/api/erp/work-orders/%d/export.xlsx" % order["id"])
    cells = [str(c.value) for ws in openpyxl.load_workbook(io.BytesIO(x.content)).worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    numbers = [float(c) for c in cells if re.fullmatch(r"-?\d+(\.\d+)?", c)]
    assert "VARIATIONS AGREED ON THIS ORDER" in cells and "Order value as varied" in cells and 1730800.0 in numbers
