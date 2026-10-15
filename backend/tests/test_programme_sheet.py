"""Excel in and out of the programme: the export is the office's layout with live formulas; the import reads the GH layout
(bold rows and SL.NO), the APTIDCO / finishing-works layout (headings' own formulas, Start formulas as dependencies)
and a sheet this app wrote, which comes back exactly."""
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from test_programme_edit import _monday, ok, rows, unit_with_painting, url
from test_schedule import job


def upload(tenant, j, book, sheet="", tail="/import"):
    buf = io.BytesIO()
    book.save(buf)
    return tenant.post(url(j, tail), files={"file": ("schedule.xlsx", buf.getvalue())}, data={"sheet": sheet})


def confirm(tenant, j, preview):
    return tenant.post(url(j, "/import/confirm"), json={"nodes": preview["nodes"], "calendar": preview["calendar"]})


def aptidco():
    """The slabs sheet's shape, with its formulas as the office writes them (no cached values, as from a script)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Slabs"
    ws.append([" SCHEDULE"])
    ws.append([None, None, "TIME DURATION IN DAYS", None, "=+(D6-C6)+1"])
    ws.append(["YALAVARTI INFRA PROJECTS Ltd."])
    ws.append(["PROJECT - CONSTRUCTION OF APTIDCO"])
    ws.append(["   SL.NO", "                   ITEMES / ACTIVITY", "Start", "End", "No of working Days", "U.O.M", "Quantity"])
    ws.append([None, None, "=+TODAY()", "=+MAX(D7:D14)", "=+(D6-C6)+1"])
    ws.append([1, "300 SFT Block Type - A", "=MIN(C8:C14)", "=+MAX(D8:D14)", "=D7-C7"])
    ws.append([1.01, "A-16 Third floor slab", "=+C9", "=+D12", "=SUM(E9:E12)"])
    ws.append([1.12, "Reinforcement", date(2026, 4, 20), "=+C9+E9", 5, "MT", 10.62])
    ws.append([None, "Shuttering", "=+D9", "=+C10+E10", 6, "Sqm", 1873.61])
    ws.append([None, "Concreting", "=+D10", "=+C11+E11", 1, "Cum", 118])
    ws.append([None, "De- shuttering", "=+D11+1", "=+C12+E12", 2])
    ws.append([1.02, "A-16 Terrace floor slab", "=+C14", "=+D15", "=SUM(E14:E15)"])
    ws.append([None, "Reinforcement", "=+D12+1", "=+C14+E14", 4, "MT", 10.62])
    ws.append([None, "Shuttering", "=+C14+2", "=+C15+E15", 5, "Sqm", 1873.61])
    for r in (7, 8, 13):
        for c in "ABCDE":
            ws["%s%d" % (c, r)].font = Font(bold=True)
    return wb


def gh():
    """The guest-house sheet's shape: no formulas to say what is under what, only bold rows and the SL.NO."""
    wb = Workbook()
    ws = wb.active
    ws.title = "GH"
    ws.append([])
    ws.append([" SCHEDULE"])
    ws.append([None, None, None, "TIME DURATION IN DAYS"])
    ws.append([" Location :"])
    ws.append(["Name of the work :- INTERIOR"])
    ws.append(["   SL.NO", "   V.NO", "                   ITEMES", "Start", "End", "No of working Days"])
    ws.append([])
    ws.append([None, None, "PHASE-II PART-B", date(2026, 3, 19), date(2026, 10, 18)])
    ws.append([None, None, "GUEST HOUSE"])
    ws.append(["I", None, "Flat No. 201"])
    ws.append(["A", None, "Guest room"])
    ws.append([None, None, "Wall finishing"])
    ws.append([1, None, "Putty & Primer"])
    ws.append([None, None, "Cement based putty", date(2026, 3, 19), date(2026, 3, 19), 1])
    ws.append([None, None, "Primer", date(2026, 3, 20), date(2026, 3, 20), 1])
    ws.append([2, None, "Ply panelling"])
    ws.append([None, None, "Ply panelling for wardrobe backing", date(2026, 3, 21), date(2026, 3, 22), 2])
    ws.append(["B", None, "Guest room 2"])
    ws.append([None, None, "Wall finishing"])
    for r in (8, 9, 10, 11, 12, 13, 16, 18, 19):
        for c in "ABC":
            ws["%s%d" % (c, r)].font = Font(bold=True)
    return wb


def test_export_is_the_offices_layout_with_live_formulas(tenant):
    mon = _monday()
    j = job(tenant)
    unit_with_painting(tenant, j, mon)
    res = tenant.get(url(j, "/export"))
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    ws = load_workbook(io.BytesIO(res.content))["Programme"]
    heads = [ws.cell(row=6, column=c).value for c in range(1, 12)]
    assert heads[:5] == ["SL.NO", "V.NO", "ITEMES / ACTIVITY", "Start", "End"] and "days" in heads[5].lower()
    assert heads[6:] == ["U.O.M", "Quantity", "Labour Required", "Predecessor", "% Done"]
    names = {ws.cell(row=r, column=3).value: r for r in range(7, 12)}
    putty, primer, unit = names["Putty"], names["Primer"], names["A-16"]
    assert ws.cell(row=putty, column=4).value.date() == mon                      # a leaf's typed start is a date
    assert str(ws.cell(row=putty, column=5).value).startswith("=")               # its End is a formula on Start and Days
    assert str(ws.cell(row=primer, column=4).value).startswith("=")              # a linked leaf's Start follows its predecessor
    assert "MIN(" in ws.cell(row=unit, column=4).value                           # a heading's Start is MIN of the rows under it
    assert "MAX(" in ws.cell(row=unit, column=5).value
    assert "MAX(" in ws["F2"].value or "MIN(" in ws["F2"].value                  # TIME DURATION IN DAYS
    assert ws.conditional_formatting                                              # the bars
    assert ws.page_setup.orientation == "landscape" and int(ws.page_setup.paperSize) == int(ws.PAPERSIZE_A3)


def test_a_sheet_this_app_wrote_comes_back_exactly(tenant):
    mon = _monday()
    a, b = job(tenant), job(tenant)
    unit_with_painting(tenant, a, mon)
    out = ok(tenant.get(url(a)))
    book = load_workbook(io.BytesIO(tenant.get(url(a, "/export")).content))
    preview = upload(tenant, b, book).json()
    assert preview["own"] is True
    assert [n["name"] for n in preview["nodes"]] == [r["name"] for r in out["rows"]]
    made = ok(confirm(tenant, b, preview))
    again = ok(tenant.get(url(b)))
    pick = lambda v: [(r["wbs"], r["name"], r["start"], r["end"], r["days"], r["predecessors"]) for r in v["rows"]]
    assert pick(again) == pick(out)
    assert made["links"] == 2


def test_the_aptidco_layout_gives_its_tree_and_its_dependencies(tenant):
    j = job(tenant)
    res = upload(tenant, j, aptidco())
    assert res.status_code == 200, res.text
    preview = res.json()
    by = {n["key"]: n for n in preview["nodes"]}
    assert by[7]["parent"] is None and by[7]["heading"]
    assert by[8]["parent"] == 7 and by[13]["parent"] == 7
    assert by[9]["parent"] == 8 and by[14]["parent"] == 13 and not by[9]["heading"]
    assert by[9]["uom"] == "MT" and abs(by[9]["quantity"] - 10.62) < 1e-9
    link = lambda key: [(l["pred"], l["kind"], l["lag"]) for l in by[key]["links"]]
    assert link(10) == [(9, "FS", -1)]          # =+D9: starts the day it ends
    assert link(12) == [(11, "FS", 0)]          # =+D11+1: the day after
    assert link(15) == [(14, "SS", 2)]          # =+C14+2: two days after it starts
    assert by[9]["start"] == "2026-04-20" and by[10]["start"] == ""
    ok(confirm(tenant, j, preview))
    v = ok(tenant.get(url(j)))
    r = rows(v)
    assert r["300 SFT Block Type - A"]["summary"] and r["A-16 Third floor slab"]["summary"]
    shuttering = sorted(x["predecessors"] for x in v["rows"] if x["name"] == "Shuttering")
    assert shuttering == ["1.1.1FS-1", "1.2.1SS+2"]          # the third-floor one waits on its Reinforcement; the terrace one on its own
    assert v["calendar"]["calendar_days"] is True         # an office sheet counts every day, Sundays included


def test_the_gh_layout_is_told_by_bold_rows_and_numbers(tenant):
    j = job(tenant)
    preview = upload(tenant, j, gh()).json()
    by = {n["key"]: n for n in preview["nodes"]}
    assert by[8]["parent"] is None and by[9]["parent"] == 8 and by[10]["parent"] == 9      # phase > house > flat
    assert by[11]["parent"] == 10 and by[12]["parent"] == 11 and by[13]["parent"] == 12    # room > wall finishing > 1
    assert by[14]["parent"] == 13 and by[15]["parent"] == 13 and not by[14]["heading"]
    assert by[16]["parent"] == 12                                                           # 2 is beside 1
    assert by[18]["parent"] == 10                                                           # B is beside A
    assert by[14]["days"] == 1 and by[17]["days"] == 2 and by[17]["start"] == "2026-03-21"


def test_a_sheet_with_a_loop_is_refused_and_nothing_is_kept(tenant):
    j = job(tenant)
    wb = Workbook()
    ws = wb.active
    ws.append(["   SL.NO", "ITEMES / ACTIVITY", "Start", "End", "No of working Days"])
    ws.append([1, "One", "=+D3", "=+C2+E2", 2])
    ws.append([2, "Two", "=+D2", "=+C3+E3", 2])
    preview = upload(tenant, j, wb).json()
    res = confirm(tenant, j, preview)
    assert res.status_code == 400 and "Circular" in res.json()["detail"]
    assert ok(tenant.get(url(j)))["rows"] == []


def test_a_file_that_is_not_a_schedule_is_refused(tenant):
    j = job(tenant)
    res = tenant.post(url(j, "/import"), files={"file": ("x.xlsx", b"not a workbook")})
    assert res.status_code == 400
    wb = Workbook()
    wb.active.append(["nothing", "here"])
    assert upload(tenant, j, wb).status_code == 400


def test_importing_adds_after_what_is_already_there(tenant):
    j = job(tenant)
    unit_with_painting(tenant, j, _monday())
    before = len(ok(tenant.get(url(j)))["rows"])
    ok(confirm(tenant, j, upload(tenant, j, gh()).json()))
    after = ok(tenant.get(url(j)))
    assert len(after["rows"]) > before
    assert after["rows"][0]["name"] == "A-16"                       # what was there stays first

def test_the_pdf_is_a_landscape_a3_with_every_row(tenant):
    j = job(tenant)
    unit_with_painting(tenant, j, _monday())
    res = tenant.get(url(j, "/pdf"))
    assert res.status_code == 200 and res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF")
    from pypdf import PdfReader
    pdf = PdfReader(io.BytesIO(res.content))
    page = pdf.pages[0]
    assert round(float(page.mediabox.width)) == 1191 and round(float(page.mediabox.height)) == 842      # A3 landscape
    text = page.extract_text()
    for word in ("SCHEDULE", "TIME DURATION IN DAYS", "Internal Painting", "Putty", "Predecessor"):
        assert word in text


def test_a_long_programme_repeats_its_headings_on_every_page(tenant):
    j = job(tenant)
    mon = _monday()
    for i in range(120):
        ok(tenant.post(url(j, "/rows"), json={"anchor_id": None, "position": "end",
                                              "rows": [{"name": "Row %d" % i, "start": mon.isoformat(), "days": 2}]}))
    from pypdf import PdfReader
    pdf = PdfReader(io.BytesIO(tenant.get(url(j, "/pdf")).content))
    assert len(pdf.pages) >= 2
    for page in pdf.pages:
        assert "Predecessor" in page.extract_text()