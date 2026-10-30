"""Client RA bills, retention, credit notes and the Owed page, held to the contract.

A bill does not claim past the order, its deductions are real deductions, an advance recovered is one the client
paid, retention released is not reversed again by a credit note, the last of it waits for the defects period, and the
Owed page shows a client in credit and money sitting on account.
"""
import io
from datetime import date, timedelta

import openpyxl

from app.core.dates import financial_year_label
from test_measurement_and_ra_bills import placed_order, measure, line_of, raise_bill


def certified(tenant, qty=1000, rate=100, measured=None, **rates):
    wo = placed_order(tenant, qty=qty, rate=rate)
    measure(tenant, wo["id"], line_of(tenant, wo["id"]), measured or qty)
    body = dict(retention_percent=5, tax_percent=18, tds_percent=2)
    body.update(rates)
    b = raise_bill(tenant, wo["id"], **body).json()["bill"]
    assert tenant.post("/api/ra-bills/%d/submit" % b["id"], json={}).status_code == 200
    res = tenant.post("/api/ra-bills/%d/certify" % b["id"], json={})
    assert res.status_code == 200, res.text
    return wo, res.json()["bill"]


def release(tenant, order_id, amount, stage="Practical completion", **over):
    body = {"side": "client", "order_id": order_id, "amount": amount, "stage": stage}
    body.update(over)
    return tenant.post("/api/retention/releases", json=body)


def on_account(tenant, amount, party="L&T", direction="IN", party_type="client"):
    res = tenant.post("/api/money/entries", json={"doc_type": "on_account", "direction": direction,
                                                  "party_type": party_type, "party_name": party,
                                                  "amount": amount, "mode": "NEFT", "reference": "UTR"})
    assert res.status_code == 200, res.text
    return res.json()["entry"]


# --- the bill does not claim past the order ---------------------------------------------------------------------

def test_a_bill_past_the_ordered_quantity_is_refused_at_submit(tenant):
    wo = placed_order(tenant, qty=1000, rate=60)
    measure(tenant, wo["id"], line_of(tenant, wo["id"]), 1200)
    b = raise_bill(tenant, wo["id"]).json()["bill"]          # a draft may show what was measured
    detail = tenant.get("/api/ra-bills/%d" % b["id"]).json()["bill"]
    assert detail["lines"][0]["over_order"] is True
    res = tenant.post("/api/ra-bills/%d/submit" % b["id"], json={})
    assert res.status_code == 409 and "past the order" in res.json()["detail"]
    assert "variation" in res.json()["detail"]


def test_an_approved_variation_lets_the_extra_be_billed(tenant):
    wo = placed_order(tenant, qty=1000, rate=60)
    measure(tenant, wo["id"], line_of(tenant, wo["id"]), 1200)
    vo = tenant.post("/api/variations", json={"work_order_id": wo["id"]}).json()["variation"]
    tenant.post("/api/variations/%d/submit" % vo["id"])
    tenant.post("/api/variations/%d/approve" % vo["id"])
    b = raise_bill(tenant, wo["id"]).json()["bill"]
    assert tenant.post("/api/ra-bills/%d/submit" % b["id"], json={}).status_code == 200


# --- TDS and retention defaults ---------------------------------------------------------------------------------

def test_tds_defaults_to_two_per_cent_and_the_company_can_set_its_own(tenant):
    wo = placed_order(tenant, qty=100, rate=100)
    measure(tenant, wo["id"], line_of(tenant, wo["id"]), 100)
    assert raise_bill(tenant, wo["id"]).json()["bill"]["tds_percent"] == 2
    assert tenant.put("/api/billing/ra-defaults", json={"tds_percent": 150}).status_code == 400
    assert tenant.put("/api/billing/ra-defaults", json={"tds_percent": 1}).status_code == 200
    assert tenant.get("/api/billing/ra-defaults").json()["tds_percent"] == 1
    wo2 = placed_order(tenant, qty=100, rate=100)
    measure(tenant, wo2["id"], line_of(tenant, wo2["id"]), 100)
    assert raise_bill(tenant, wo2["id"]).json()["bill"]["tds_percent"] == 1
    # A rate given on the bill still wins over the company's.
    wo3 = placed_order(tenant, qty=100, rate=100)
    measure(tenant, wo3["id"], line_of(tenant, wo3["id"]), 100)
    assert raise_bill(tenant, wo3["id"], tds_percent=0).json()["bill"]["tds_percent"] == 0


def test_a_project_that_states_no_retention_is_billed_none(tenant):
    wo = placed_order(tenant, qty=1000, rate=100)
    tenant.put("/api/jobs/%d" % wo["job_id"], json={"name": "Vizag", "retention_percent": 0})
    measure(tenant, wo["id"], line_of(tenant, wo["id"]), 1000)
    bill = raise_bill(tenant, wo["id"]).json()["bill"]
    assert bill["retention_percent"] == 0 and bill["retention_amount"] == 0
    # Saving the project without the field keeps what it states.
    job = tenant.put("/api/jobs/%d" % wo["job_id"], json={"name": "Vizag again"}).json()
    assert job["retention_stated"] is True and job["retention_percent"] == 0


def test_a_project_left_blank_keeps_the_usual_five(tenant):
    job = tenant.post("/api/jobs", json={"name": "Blank", "customer_name": "C"}).json()
    assert job["retention_stated"] is False
    wo = placed_order(tenant, qty=100, rate=100)
    measure(tenant, wo["id"], line_of(tenant, wo["id"]), 100)
    assert raise_bill(tenant, wo["id"]).json()["bill"]["retention_percent"] == 5


# --- the deductions -----------------------------------------------------------------------------------------------

def test_deductions_are_checked(tenant):
    wo = placed_order(tenant, qty=100, rate=100)
    measure(tenant, wo["id"], line_of(tenant, wo["id"]), 100)
    assert raise_bill(tenant, wo["id"], tax_percent=7).status_code == 400
    assert raise_bill(tenant, wo["id"], retention_percent=-1).status_code == 400
    assert raise_bill(tenant, wo["id"], tds_percent=101).status_code == 400
    assert raise_bill(tenant, wo["id"], other_deductions=-5).status_code == 400
    assert raise_bill(tenant, wo["id"], advance_recovery=-5).status_code == 400
    res = raise_bill(tenant, wo["id"], other_deductions=50000)
    assert res.status_code == 400 and "more than the bill" in res.json()["detail"]
    b = raise_bill(tenant, wo["id"]).json()["bill"]
    res = tenant.put("/api/ra-bills/%d" % b["id"], json={"work_order_id": wo["id"], "tax_percent": 13})
    assert res.status_code == 400


def test_an_advance_recovered_has_to_have_been_received(tenant):
    wo = placed_order(tenant, qty=1000, rate=100)
    measure(tenant, wo["id"], line_of(tenant, wo["id"]), 1000)
    b = raise_bill(tenant, wo["id"], advance_recovery=20000).json()["bill"]
    res = tenant.post("/api/ra-bills/%d/submit" % b["id"], json={})
    assert res.status_code == 400 and "on account" in res.json()["detail"]
    on_account(tenant, 15000)
    assert tenant.post("/api/ra-bills/%d/submit" % b["id"], json={}).status_code == 400
    on_account(tenant, 5000)
    assert tenant.post("/api/ra-bills/%d/submit" % b["id"], json={}).status_code == 200


# --- credit notes and retention -----------------------------------------------------------------------------------

def test_credit_notes_run_in_one_series_a_year(tenant):
    wo, bill = certified(tenant)
    line = tenant.get("/api/ra-bills/%d" % bill["id"]).json()["bill"]["lines"][0]
    url = "/api/ra-bills/%d/credit-notes" % bill["id"]
    first = tenant.post(url, json={"lines": [{"bill_line_id": line["id"], "qty": 10}], "reason": "x"}).json()
    wo2, bill2 = certified(tenant)
    line2 = tenant.get("/api/ra-bills/%d" % bill2["id"]).json()["bill"]["lines"][0]
    second = tenant.post("/api/ra-bills/%d/credit-notes" % bill2["id"],
                         json={"lines": [{"bill_line_id": line2["id"], "qty": 10}], "reason": "y"}).json()
    year = financial_year_label(date.today().isoformat())
    assert first["credit_note"]["number"] == "CN/%s/0001" % year
    assert second["credit_note"]["number"] == "CN/%s/0002" % year


def test_a_credit_note_cannot_reverse_retention_already_released(tenant):
    wo, bill = certified(tenant)                         # 5,000 held
    assert release(tenant, wo["id"], 4800).status_code == 200
    line = tenant.get("/api/ra-bills/%d" % bill["id"]).json()["bill"]["lines"][0]
    res = tenant.post("/api/ra-bills/%d/credit-notes" % bill["id"], json={
        "lines": [{"bill_line_id": line["id"], "qty": 100}], "reason": "Measurement reduced"})    # 500 of retention
    assert res.status_code == 409 and "released" in res.json()["detail"]


def test_a_release_cannot_exceed_what_credit_notes_left_held(tenant):
    wo, bill = certified(tenant)
    line = tenant.get("/api/ra-bills/%d" % bill["id"]).json()["bill"]["lines"][0]
    tenant.post("/api/ra-bills/%d/credit-notes" % bill["id"], json={
        "lines": [{"bill_line_id": line["id"], "qty": 100}], "reason": "Measurement reduced"})
    res = release(tenant, wo["id"], 4900, dlp_ends="2020-01-01")
    assert res.status_code == 400 and "4,500" in res.json()["detail"]


def test_the_ra_bill_register_holds_what_the_retention_page_holds(tenant):
    wo, bill = certified(tenant)
    release(tenant, wo["id"], 2500)
    s = tenant.get("/api/ra-bills?work_order_id=%d" % wo["id"]).json()["summary"]
    assert s["retention_held"] == 2500


def test_the_release_sheet_counts_credit_notes(tenant):
    wo, bill = certified(tenant)
    line = tenant.get("/api/ra-bills/%d" % bill["id"]).json()["bill"]["lines"][0]
    tenant.post("/api/ra-bills/%d/credit-notes" % bill["id"], json={
        "lines": [{"bill_line_id": line["id"], "qty": 100}], "reason": "Measurement reduced"})
    r = release(tenant, wo["id"], 1000).json()["release"]
    res = tenant.get("/api/retention/releases/%d/export.xlsx" % r["id"])
    ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
    cells = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    assert any(row and row[0] == "Taken back by credit notes" for row in cells)


# --- the last of the client's retention waits for the defects period ----------------------------------------------

def test_the_last_of_the_retention_waits_for_the_defects_period(tenant):
    wo, bill = certified(tenant)
    res = release(tenant, wo["id"], 5000)
    assert res.status_code == 409 and "defects period" in res.json()["detail"]
    later = (date.today() + timedelta(days=200)).isoformat()
    res = release(tenant, wo["id"], 2500, stage="End of defects period", dlp_ends=later)
    assert res.status_code == 409 and later in res.json()["detail"]
    # The Master may, saying why.
    assert release(tenant, wo["id"], 2500, stage="End of defects period", dlp_ends=later,
                   override=True).status_code == 400
    res = release(tenant, wo["id"], 2500, stage="End of defects period", dlp_ends=later, override=True,
                  notes="Client agreed against a letter")
    assert res.status_code == 200, res.text
    assert "defects period" in res.json()["release"]["override_note"]
    assert res.json()["release"]["dlp_ends"] == later


def test_once_the_defects_period_is_over_it_is_claimed(tenant):
    wo, bill = certified(tenant)
    past = (date.today() - timedelta(days=1)).isoformat()
    res = release(tenant, wo["id"], 5000, stage="End of defects period", dlp_ends=past)
    assert res.status_code == 200, res.text
    p = next(p for p in tenant.get("/api/retention").json()["positions"]
             if p["side"] == "client" and p["order_id"] == wo["id"])
    assert p["dlp_ends"] == past and p["balance"] == 0


def test_release_numbers_follow_the_highest_not_the_count(tenant):
    wo, bill = certified(tenant)
    first = release(tenant, wo["id"], 1000).json()["release"]
    second = release(tenant, wo["id"], 1000).json()["release"]
    assert (first["number"], second["number"]) == ("RET-0001", "RET-0002")


# --- the Owed page ------------------------------------------------------------------------------------------------

def test_a_client_in_credit_after_a_credit_note_is_shown(tenant):
    wo, bill = certified(tenant)
    got = tenant.post("/api/money/entries", json={"doc_type": "ra_bill", "doc_id": bill["id"],
                                                  "amount": bill["net_payable"], "mode": "NEFT", "reference": "U1"})
    assert got.status_code == 200, got.text
    line = tenant.get("/api/ra-bills/%d" % bill["id"]).json()["bill"]["lines"][0]
    note = tenant.post("/api/ra-bills/%d/credit-notes" % bill["id"], json={
        "lines": [{"bill_line_id": line["id"], "qty": 100}], "reason": "Measurement reduced"}).json()["credit_note"]
    rec = tenant.get("/api/money/receivables").json()
    mine = [c for c in rec["credits"] if c["id"] == bill["id"]]
    assert mine and mine[0]["amount"] == note["net_amount"]
    party = [p for p in rec["parties"] if p["party"] == "L&T"][0]
    assert party["credit"] == note["net_amount"]
    assert rec["summary"]["credit_balances"] == note["net_amount"]


def test_money_on_account_is_netted_per_party(tenant):
    wo, bill = certified(tenant)
    on_account(tenant, 10000)
    rec = tenant.get("/api/money/receivables").json()
    party = [p for p in rec["parties"] if p["party"] == "L&T"][0]
    assert party["on_account"] == 10000
    assert party["net"] == round(party["owed"] - 10000, 2)
    assert rec["summary"]["net_owed"] == round(rec["summary"]["owed"] - 10000, 2)
    on_account(tenant, 3000, party="Sri Balaji Constructions", direction="OUT", party_type="contractor")
    pay = tenant.get("/api/money/payables").json()
    assert [p for p in pay["parties"] if p["party"] == "Sri Balaji Constructions"][0]["on_account"] == 3000


def test_the_receivables_sheet_says_what_each_row_is(tenant):
    certified(tenant)
    res = tenant.get("/api/money/receivables.xlsx")
    assert res.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
    rows = list(ws.iter_rows(values_only=True))
    header = next(r for r in rows if r and r[0] == "Document")
    assert "Kind" in header and "GSTIN" in header


def test_an_unknown_document_is_a_clear_answer_not_a_crash(tenant):
    assert tenant.get("/api/money/outstanding/nonsense/1").status_code == 400
    assert tenant.get("/api/money/outstanding/supplier_bill/999999").status_code == 404
