"""The hindrance register: what held the work up, whose risk it was, what it cost in days, when the client was told.

Raised by hand or from a day's diary (its "held up by" words and its rain), listed per site with the days added up by whose risk they
were, exported as a workbook, and shown on the activity it held up.
"""
import io

from openpyxl import load_workbook

from test_site_scope import staff_on, two_projects


def site(tenant, name="Hindrance Towers"):
    job = tenant.post("/api/jobs", json={"name": name, "customer_name": "Client Ltd"}).json()
    act = tenant.post("/api/jobs/%d/schedule/activities" % job["id"], json={"name": "Raft", "planned_start": "2026-09-01", "planned_finish": "2026-09-20"})
    assert act.status_code == 200, act.text
    return job["id"], act.json()["schedule"]["activities"][0]["id"]


def diary(tenant, jid, on, **kw):
    res = tenant.post("/api/diary", json={"job_id": jid, "diary_date": on, "weather": kw.pop("weather", "Clear"), "working_hours": 8, "work_done": "Shuttering", **kw})
    assert res.status_code == 200, res.text
    return res.json()["diary"]["id"]


def add(tenant, jid, **kw):
    body = {"cause": "Drawings or details awaited", "description": "Raft reinforcement drawing not issued", "from_date": "2026-09-05",
            "to_date": "2026-09-09", "responsibility": "Client", **kw}
    return tenant.post("/api/jobs/%d/hindrances" % jid, json=body)


def test_a_hindrance_is_numbered_per_site_and_its_days_counted_from_its_dates(tenant):
    jid, aid = site(tenant)
    first = add(tenant, jid, activity_id=aid)
    assert first.status_code == 200, first.text
    h = first.json()["hindrance"]
    assert h["number"] == "H-001" and h["impact_days"] == 5 and h["running_days"] == 5      # 5th to the 9th, both counted
    assert h["activity"].endswith("Raft") and h["status"] == "Open"
    assert h["notice_due"] is True                                                           # the client's risk, not yet told
    second = add(tenant, jid, cause="Rain / weather", responsibility="Neither (force majeure)", to_date="", impact_days=None, description="Monsoon")
    assert second.json()["hindrance"]["number"] == "H-002"
    assert second.json()["hindrance"]["ongoing"] is True and second.json()["hindrance"]["impact_days"] == 0
    # Another site starts its own numbers.
    other, _ = site(tenant, "Second site")
    assert add(tenant, other).json()["hindrance"]["number"] == "H-001"

    reg = tenant.get("/api/jobs/%d/hindrances" % jid).json()
    assert [r["number"] for r in reg["hindrances"]] == ["H-002", "H-001"]
    assert reg["summary"]["client_days"] == 5 and reg["summary"]["open"] == 2 and reg["summary"]["ongoing"] == 1 and reg["summary"]["notice_due"] == 1


def test_notice_status_and_closing(tenant):
    jid, _ = site(tenant)
    hid = add(tenant, jid, notice_date="2026-09-06").json()["hindrance"]
    assert hid["status"] == "Notified" and hid["notice_due"] is False                       # told the client: notified at once
    hid = hid["id"]
    res = tenant.put("/api/hindrances/%d" % hid, json={"status": "Agreed", "impact_days": 4})
    assert res.status_code == 200 and res.json()["hindrance"]["impact_days"] == 4 and res.json()["hindrance"]["status"] == "Agreed"
    ongoing = add(tenant, jid, to_date="").json()["hindrance"]["id"]
    assert tenant.put("/api/hindrances/%d" % ongoing, json={"status": "Closed"}).status_code == 400       # no end date yet
    assert tenant.put("/api/hindrances/%d" % ongoing, json={"status": "Closed", "to_date": "2026-09-12"}).status_code == 200
    # A rejected one is left out of the days.
    tenant.put("/api/hindrances/%d" % ongoing, json={"status": "Rejected"})
    assert tenant.get("/api/jobs/%d/hindrances" % jid).json()["summary"]["impact_days"] == 4


def test_what_is_refused(tenant):
    jid, _ = site(tenant)
    other, other_act = site(tenant, "Elsewhere")
    assert add(tenant, jid, cause="Bad luck").status_code == 400
    assert add(tenant, jid, responsibility="Nobody").status_code == 400
    assert add(tenant, jid, from_date="").status_code == 400
    assert add(tenant, jid, to_date="2026-09-01").status_code == 400                       # ends before it starts
    assert add(tenant, jid, notice_date="2026-09-01").status_code == 400                   # told before it started
    assert add(tenant, jid, from_date="2099-01-01", to_date="").status_code == 400         # in the future
    assert add(tenant, jid, activity_id=other_act).status_code == 400                      # another site's activity
    assert add(tenant, jid, impact_days=-1).status_code == 400


def test_raised_from_the_diary_starts_from_what_held_it_up_or_the_rain(tenant):
    jid, aid = site(tenant)
    wet = diary(tenant, jid, "2026-09-10", weather="Rain", rain_hours=4, holdups="")
    d = tenant.get("/api/diary/%d/hindrance-draft" % wet).json()
    assert d["has_cause"] is True and d["already"] == []
    assert d["draft"]["cause"] == "Rain / weather" and d["draft"]["impact_days"] == 0.5 and d["draft"]["responsibility"] == "Neither (force majeure)"
    assert d["draft"]["from_date"] == d["draft"]["to_date"] == "2026-09-10" and "4 hours of rain" in d["draft"]["description"]

    held = diary(tenant, jid, "2026-09-11", holdups="Waiting for GFC drawings of the raft from the architect")
    d = tenant.get("/api/diary/%d/hindrance-draft" % held).json()["draft"]
    assert d["cause"] == "Drawings or details awaited" and d["responsibility"] == "Client" and d["impact_days"] == 1
    assert d["description"].startswith("Waiting for GFC")

    # Saved with the diary it came from; the next draft from that day says one is already raised.
    saved = tenant.post("/api/jobs/%d/hindrances" % jid, json={**d, "activity_id": aid})
    assert saved.status_code == 200, saved.text
    assert saved.json()["hindrance"]["diary_id"] == held
    assert [h["number"] for h in tenant.get("/api/diary/%d/hindrance-draft" % held).json()["already"]] == ["H-001"]

    quiet = diary(tenant, jid, "2026-09-12")
    assert tenant.get("/api/diary/%d/hindrance-draft" % quiet).json()["has_cause"] is False

    # Deleting the draft diary keeps the hindrance (the hold-up still happened).
    assert tenant.delete("/api/diary/%d" % held).status_code == 200
    kept = tenant.get("/api/jobs/%d/hindrances" % jid).json()["hindrances"]
    assert len(kept) == 1 and kept[0]["diary_id"] is None


def test_the_activity_shows_what_held_it_up(tenant):
    jid, aid = site(tenant)
    add(tenant, jid, activity_id=aid)
    add(tenant, jid, cause="Labour shortage", responsibility="Contractor")          # on the site, not this activity
    detail = tenant.get("/api/schedule/activities/%d" % aid).json()
    assert [h["number"] for h in detail["hindrances"]] == ["H-001"]
    assert [h["number"] for h in tenant.get("/api/jobs/%d/hindrances?activity_id=%d" % (jid, aid)).json()["hindrances"]] == ["H-001"]
    # The activity deleted: the register keeps the hindrance, as one of the site's.
    tenant.delete("/api/schedule/activities/%d" % aid)
    assert {h["activity_id"] for h in tenant.get("/api/jobs/%d/hindrances" % jid).json()["hindrances"]} == {None}


def test_the_register_as_a_workbook(tenant):
    jid, aid = site(tenant)
    add(tenant, jid, activity_id=aid)
    add(tenant, jid, cause="Rain / weather", responsibility="Neither (force majeure)", from_date="2026-09-14", to_date="2026-09-14")
    res = tenant.get("/api/jobs/%d/hindrances.xlsx" % jid)
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content)).active
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "HINDRANCE REGISTER" in cells and "H-001" in cells and "H-002" in cells
    assert "Drawings or details awaited" in cells and "NOT YET" in cells


def test_deleting_one(tenant):
    jid, _ = site(tenant)
    hid = add(tenant, jid).json()["hindrance"]["id"]
    assert tenant.delete("/api/hindrances/%d" % hid).status_code == 200
    assert tenant.get("/api/jobs/%d/hindrances" % jid).json()["hindrances"] == []


def test_staff_on_one_site_see_and_change_only_that_sites_register(tenant, portal):
    a, b = two_projects(tenant)
    ha = add(tenant, a["id"]).json()["hindrance"]["id"]
    hb = add(tenant, b["id"]).json()["hindrance"]["id"]
    staff_on(tenant, portal, [a["id"]])
    assert [h["id"] for h in portal.get("/api/jobs/%d/hindrances" % a["id"]).json()["hindrances"]] == [ha]
    assert portal.get("/api/jobs/%d/hindrances" % b["id"]).status_code == 403
    assert portal.get("/api/jobs/%d/hindrances.xlsx" % b["id"]).status_code == 403
    assert portal.put("/api/hindrances/%d" % hb, json={"status": "Closed", "to_date": "2026-09-10"}).status_code == 404
    assert portal.delete("/api/hindrances/%d" % hb).status_code == 404
    assert tenant.get("/api/jobs/%d/hindrances" % b["id"]).json()["hindrances"][0]["status"] == "Open"
