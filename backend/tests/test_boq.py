"""The project BOQ: read as the client sends it, coded from the item master, revised, and tracked against the gangs."""
import io

import openpyxl

import main
from test_subcontract_orders import masters, draft, fund_order


def sheet(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Off a client's desk: a title block above the headings, section rows, items, sub-items and its own grand total.
CLIENT_SHEET = [
    ["M/s Kokapet Towers LLP"], ["BILL OF QUANTITIES - TOWER C"], [],
    ["Group code", "Description of work", "Unit", "Qty", "Rate (Rs)", "Amount (Rs)"],
    [None, "SUBSTRUCTURE", None, None, None, None],
    ["EXC", "Excavation in all types of soil", "cum", 500, 180, 90000],
    ["RCC", "M25 grade RCC in raft", "cum", 120, 8200, 984000],
    ["FWK", "Shuttering to raft edges", "sqm", 80, 410, 32800],
    [None, "SUPERSTRUCTURE", None, None, None, None],
    ["MAS", "Brick masonry in CM 1:6", "cum", 60, 6100, 366000],
    ["", "Grand Total", None, None, None, 1472800],
]


def make_boq(tenant, job_name="Kokapet Towers"):
    job = tenant.post("/api/jobs", json={"name": job_name, "customer_name": "Kokapet LLP"}).json()
    jid = job.get("id") or job["job"]["id"]
    res = tenant.post("/api/boqs", json={"job_id": jid})
    assert res.status_code == 200, res.text
    return jid, res.json()["boq"]


def import_lines(tenant, boq):
    res = tenant.post("/api/boqs/%d/import" % boq["id"], files={"file": ("boq.xlsx", sheet(CLIENT_SHEET))})
    assert res.status_code == 200, res.text
    return res.json()


def saved(tenant, boq, issue_codes=True):
    lines = import_lines(tenant, boq)["lines"]
    res = tenant.put("/api/boqs/%d/lines" % boq["id"], json={"lines": lines, "issue_codes": issue_codes})
    assert res.status_code == 200, res.text
    return res.json()


# --- reading the client's sheet --------------------------------------------------------------------------------------

def test_a_clients_sheet_is_read_as_it_comes(tenant):
    _, boq = make_boq(tenant)
    out = import_lines(tenant, boq)
    kinds = [(l["kind"], l["group_code"]) for l in out["lines"]]
    assert kinds == [("section", ""), ("item", "EXC"), ("item", "RCC"), ("item", "FWK"), ("section", ""), ("item", "MAS")]
    assert out["total"] == 1472800 and not out["warnings"]
    assert out["skipped_rows"] == 1                       # the sheet's own grand total, checked and left out


def test_a_sheets_wrong_total_or_amount_is_called_out(tenant):
    _, boq = make_boq(tenant)
    rows = [list(r) for r in CLIENT_SHEET]
    rows[-1][5] = 999999                                   # the grand total
    rows[5][5] = 95000                                     # an amount that is not qty x rate
    res = tenant.post("/api/boqs/%d/import" % boq["id"], files={"file": ("b.xlsx", sheet(rows))}).json()
    text = " ".join(res["warnings"])
    assert "grand total is" in text and "180" in text


def test_a_sheet_that_is_not_a_boq_is_refused_in_words(tenant):
    _, boq = make_boq(tenant)
    res = tenant.post("/api/boqs/%d/import" % boq["id"], files={"file": ("b.xlsx", sheet([["Name", "Phone"], ["A", "1"]]))})
    assert res.status_code == 400 and "description and quantity" in res.json()["detail"]


# --- saving, totals and item codes ---------------------------------------------------------------------------------------

def test_saved_lines_add_up_by_section_and_get_item_codes_from_the_master(tenant):
    _, boq = make_boq(tenant)
    out = saved(tenant, boq)
    assert out["total"] == 1472800
    sections = [l for l in out["lines"] if l["kind"] == "section"]
    assert [s["subtotal"] for s in sections] == [1106800, 366000]
    priced = [l for l in out["lines"] if l["kind"] in ("item", "sub")]
    codes = [l["item_code"] for l in priced]
    assert all(codes) and len(set(codes)) == len(codes)
    assert "4 item codes given" in out["message"]
    items = tenant.get("/api/erp/items", params={"kind": "FG"}).json()
    names = {i["item_code"]: i["item_name"] for i in (items["items"] if isinstance(items, dict) else items)}
    assert names[codes[0]] == "Excavation in all types of soil"


def test_saving_again_keeps_each_lines_key_and_code_and_issues_nothing_twice(tenant):
    _, boq = make_boq(tenant)
    first = saved(tenant, boq)
    again = tenant.put("/api/boqs/%d/lines" % boq["id"], json={"lines": first["lines"], "issue_codes": True}).json()
    assert [l["key"] for l in again["lines"]] == [l["key"] for l in first["lines"]]
    assert [l["item_code"] for l in again["lines"]] == [l["item_code"] for l in first["lines"]]
    assert "code" not in again["message"].split("saved,")[-1]


def test_a_line_pasted_with_the_name_of_an_existing_item_takes_that_items_code(tenant):
    _, boq = make_boq(tenant)
    first = saved(tenant, boq)
    code = first["lines"][1]["item_code"]
    line = {"kind": "item", "group_code": "NEW", "description": "Excavation in all types of soil", "uom": "cum", "quantity": 10, "rate": 200}
    # the first line holding that code is removed, so the code is free for the pasted line to take
    rest = [l for l in first["lines"] if l["item_code"] != code]
    res = tenant.put("/api/boqs/%d/lines" % boq["id"], json={"lines": rest + [line], "issue_codes": True}).json()
    assert [l for l in res["lines"] if l["group_code"] == "NEW"][0]["item_code"] == code


def test_one_boq_to_a_project_and_negative_figures_are_refused(tenant):
    jid, boq = make_boq(tenant)
    assert tenant.post("/api/boqs", json={"job_id": jid}).status_code == 409
    bad = {"lines": [{"kind": "item", "description": "x", "quantity": -1, "rate": 5}]}
    assert tenant.put("/api/boqs/%d/lines" % boq["id"], json=bad).status_code == 400


# --- revisions ---------------------------------------------------------------------------------------------------------

def test_a_new_revision_locks_the_old_one_and_the_changes_are_listed(tenant):
    _, boq = make_boq(tenant)
    first = saved(tenant, boq)
    rev = tenant.post("/api/boqs/%d/revisions" % boq["id"], json={"label": "R1 - Award", "note": "after negotiation"})
    assert rev.status_code == 200, rev.text
    r1 = tenant.get("/api/boqs/%d" % boq["id"]).json()
    assert r1["boq"]["current_rev"] == 1 and r1["editable"] and [r["status"] for r in r1["revisions"]] == ["ISSUED", "OPEN"]
    assert tenant.get("/api/boqs/%d" % boq["id"], params={"rev": 0}).json()["editable"] is False
    lines = r1["lines"]
    lines[1]["quantity"] = 650                                     # excavation 500 -> 650
    lines = [l for l in lines if l["group_code"] != "MAS"]                # masonry removed
    lines.append({"kind": "item", "group_code": "PLS", "description": "Plastering", "uom": "sqm", "quantity": 100, "rate": 90})
    assert tenant.put("/api/boqs/%d/lines" % boq["id"], json={"lines": lines}).status_code == 200
    ch = tenant.get("/api/boqs/%d/changes" % boq["id"], params={"frm": 0}).json()
    by = {c["change"]: c for c in ch["changes"]}
    assert set(by) == {"added", "removed", "changed"} and "quantity 500.0 -> 650.0" in by["changed"]["detail"]
    assert ch["difference"] == by["added"]["amount"] + by["removed"]["amount"] + by["changed"]["amount"]


def test_an_issued_revision_cannot_be_edited_by_saving(tenant):
    _, boq = make_boq(tenant)
    saved(tenant, boq)
    tenant.post("/api/boqs/%d/revisions" % boq["id"], json={})
    # revision 0 is closed; saving goes to the open one, and the old one is untouched
    assert tenant.get("/api/boqs/%d" % boq["id"], params={"rev": 0}).json()["total"] == 1472800


# --- gang orders against the BOQ ---------------------------------------------------------------------------------------

def gang_order(tenant, boq_lines, picks, order=None):
    """A draft gang order whose schedule is the BOQ lines picked: [(boq line, quantity, gang rate)]."""
    order = order or draft(tenant)
    pulled = tenant.post("/api/wo/orders/%d/boq-lines" % order["id"],
                         json={"lines": [{"key": l["key"], "quantity": q, "rate": r} for l, q, r in picks]}).json()
    res = tenant.put("/api/wo/orders/%d/boq" % order["id"], json={"lines": pulled["lines"]})
    assert res.status_code == 200, res.text
    return order, pulled, res.json()


def test_lines_are_pulled_from_the_boq_into_a_gang_order_and_keep_their_link(tenant):
    _, boq = make_boq(tenant)
    lines = {l["group_code"]: l for l in saved(tenant, boq)["lines"]}
    order, pulled, res = gang_order(tenant, lines, [(lines["EXC"], 300, 150)])
    assert pulled["lines"][0]["item_description"] == "Excavation in all types of soil" and pulled["warnings"] == []
    assert res["order"]["items"][0]["boq_key"] == lines["EXC"]["key"] and res["warnings"] == []


def test_giving_more_than_the_boq_allows_is_warned_not_refused(tenant):
    _, boq = make_boq(tenant)
    lines = {l["group_code"]: l for l in saved(tenant, boq)["lines"]}
    gang_order(tenant, lines, [(lines["EXC"], 300, 150)])
    pulled = tenant.post("/api/wo/orders/%d/boq-lines" % draft(tenant)["id"],
                         json={"lines": [{"key": lines["EXC"]["key"], "quantity": 300, "rate": 150}]}).json()
    assert "600" in pulled["warnings"][0] and "500" in pulled["warnings"][0]
    _, _, res = gang_order(tenant, lines, [(lines["EXC"], 300, 150)])
    assert res["warnings"] and "Over the BOQ" in res["message"]


def test_the_tracker_adds_up_what_the_gangs_have_and_flags_what_is_wrong(tenant):
    _, boq = make_boq(tenant)
    lines = {l["group_code"]: l for l in saved(tenant, boq)["lines"]}
    gang_order(tenant, lines, [(lines["EXC"], 300, 150), (lines["RCC"], 120, 9000)])      # the second is at a loss
    t = tenant.get("/api/boqs/%d/tracker" % boq["id"]).json()
    rows = {r["group_code"]: r for r in t["rows"] if r["kind"] != "section"}
    assert rows["EXC"]["given"] == 300 and rows["EXC"]["gang_rate"] == 150 and rows["EXC"]["left_to_give"] == 200
    assert rows["EXC"]["margin_percent"] == round((180 - 150) / 180 * 100, 1) and rows["EXC"]["flags"] == []
    assert "loss" in rows["RCC"]["flags"] and "no_gang" in rows["MAS"]["flags"] and "no_gang" in rows["FWK"]["flags"]
    assert t["flags"]["loss"] == 1 and t["flags"]["no_gang"] == 2


def test_a_cancelled_gang_order_no_longer_counts_against_the_boq(tenant):
    _, boq = make_boq(tenant)
    lines = {l["group_code"]: l for l in saved(tenant, boq)["lines"]}
    order, _, _ = gang_order(tenant, lines, [(lines["EXC"], 300, 150)])
    tenant.post("/api/wo/orders/%d/cancel" % order["id"], json={"comments": "x"})
    rows = {r["group_code"]: r for r in tenant.get("/api/boqs/%d/tracker" % boq["id"]).json()["rows"] if r["kind"] != "section"}
    assert rows["EXC"]["given"] == 0


# --- the client's order, export, tender -------------------------------------------------------------------------------

def test_the_clients_work_order_is_drawn_from_the_boq_once(tenant):
    jid, boq = make_boq(tenant)
    saved(tenant, boq, issue_codes=False)
    refused = tenant.post("/api/boqs/%d/client-order" % boq["id"], json={})
    assert refused.status_code == 409 and "item code" in refused.json()["detail"]
    saved(tenant, boq, issue_codes=True)
    made = tenant.post("/api/boqs/%d/client-order" % boq["id"], json={"reference": "PO-77"})
    assert made.status_code == 200, made.text
    wo = made.json()["work_order"]
    assert len(wo["lines"]) == 4 and wo["total_value"] == 1472800
    assert tenant.post("/api/boqs/%d/client-order" % boq["id"], json={}).status_code == 409


def test_the_boq_goes_to_excel_and_to_paper(tenant):
    _, boq = make_boq(tenant)
    saved(tenant, boq)
    x = tenant.get("/api/boqs/%d/export.xlsx" % boq["id"])
    assert x.status_code == 200
    cells = [c.value for ws in openpyxl.load_workbook(io.BytesIO(x.content)).worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    assert "SUBSTRUCTURE" in cells and 1472800 in cells
    assert tenant.get("/api/boqs/%d/export.pdf" % boq["id"]).status_code == 200


def test_a_won_tender_becomes_the_boq(tenant):
    job = tenant.post("/api/jobs", json={"name": "Tendered job"}).json()
    jid = job.get("id") or job["job"]["id"]
    est = tenant.post("/api/estimates", json={"title": "Tender", "job_id": jid}).json()
    eid = est.get("id") or est["estimate"]["id"]
    with main.SessionLocal() as db:
        db.add(main.models.DBEstimateItem(estimate_id=eid, item_no="1.1", description="Piling", uom="m", quantity=100, quoted_rate=900, display_order=0))
        db.commit()
    res = tenant.post("/api/boqs/from-estimate/%d" % eid)
    assert res.status_code == 200, res.text
    got = tenant.get("/api/boqs/%d" % res.json()["boq"]["id"]).json()
    assert got["total"] == 90000 and got["lines"][0]["description"] == "Piling"


def test_the_kinds_of_row_are_told_apart():
    kind = main.boq_kind_of
    assert kind("Excavation", 5, 10) == "item" and kind("Edges", 0, 10) == "item"
    assert kind("SUBSTRUCTURE", 0, 0) == "section" and kind("FINISHING WORKS", 0, 0) == "section"
    assert kind("All rates include supply, fixing and curing unless stated otherwise in the specification.", 0, 0) == "note"


# --- the contractor's own layout: material or grade, rate-only items, a summary with GST ------------------------------

CLUB_HOUSE_SHEET = [
    ["Proposed Construction of Residential Apartment - Club House"], [],
    ["GROUP CODE", "PARTICULARS", "MATERIAL", "UNIT", "QUANTITY", "RATE", "AMOUNT"],
    ["A", "PLAIN CEMENT CONCRETE", None, None, None, None, None],
    ["a", "Footings", None, "Cum.", 277, 5230, 1448710],
    ["B", "REINFORCED CEMENT CONCRETE", None, None, None, None, None],
    ["2.01a", "Tower Columns (Found to Plinth)", "M60", "Cum", 61, 7800, 475800],
    ["2.04a", "Podium drops", "M40", "Cum", 269, 6592, 1773248],
    ["C", "REINFORCEMENT STEEL", None, None, None, None, None],
    ["3.01a", "Footings", "Fe550/550D", "MT", 133.05, 73900, 9832395],
    ["E", "MASONRY", None, None, None, None, None],
    ["6", "Solid block masonry 400 x 100 x 200 mm in CM 1:4", None, "Sqm", "R/o", 1200, None],
]


def club_house(tenant):
    _, boq = make_boq(tenant, "Club House")
    res = tenant.post("/api/boqs/%d/import" % boq["id"], files={"file": ("club.xlsx", sheet(CLUB_HOUSE_SHEET))})
    assert res.status_code == 200, res.text
    out = res.json()
    saved = tenant.put("/api/boqs/%d/lines" % boq["id"], json={"lines": out["lines"], "issue_codes": True})
    assert saved.status_code == 200, saved.text
    return boq, out


def test_material_and_rate_only_items_are_read_from_the_sheet(tenant):
    _, out = club_house(tenant)
    by_text = {l["description"]: l for l in out["lines"]}
    assert by_text["Tower Columns (Found to Plinth)"]["material"] == "M60"
    assert by_text["Footings"]["material"] in ("", "Fe550/550D")
    r_o = by_text["Solid block masonry 400 x 100 x 200 mm in CM 1:4"]
    assert r_o["rate_only"] is True and r_o["quantity"] == 0 and r_o["amount"] == 0 and r_o["rate"] == 1200
    assert not out["warnings"]


def test_a_rate_only_item_is_in_no_total(tenant):
    boq, _ = club_house(tenant)
    body = tenant.get("/api/boqs/%d" % boq["id"]).json()
    assert body["total"] == 1448710 + 475800 + 1773248 + 9832395
    assert [l for l in body["lines"] if l["rate_only"]][0]["amount"] == 0


def test_the_summary_adds_sections_and_puts_gst_on_the_total(tenant):
    boq, _ = club_house(tenant)
    s = tenant.get("/api/boqs/%d/summary" % boq["id"]).json()
    amounts = {x["title"]: x["amount"] for x in s["sections"]}
    assert amounts["PLAIN CEMENT CONCRETE"] == 1448710
    assert amounts["REINFORCED CEMENT CONCRETE"] == 475800 + 1773248
    assert amounts["MASONRY"] == 0
    assert s["total"] == sum(amounts.values())
    assert s["gst_percent"] == 18.0
    assert s["gst"] == round(s["total"] * 0.18, 2)
    assert s["total_with_gst"] == round(s["total"] + s["gst"], 2)
    assert s["rate_only_items"] == 1
    # The same money by grade.
    grades = {m["material"]: m["amount"] for m in s["by_material"]}
    assert grades["M40"] == 1773248 and grades["Fe550/550D"] == 9832395


def test_the_gst_rate_can_be_changed_on_the_summary(tenant):
    boq, _ = club_house(tenant)
    s = tenant.get("/api/boqs/%d/summary?gst=12" % boq["id"]).json()
    assert s["gst_percent"] == 12.0 and s["gst"] == round(s["total"] * 0.12, 2)


def test_material_and_rate_only_survive_a_new_revision(tenant):
    boq, _ = club_house(tenant)
    assert tenant.post("/api/boqs/%d/revisions" % boq["id"], json={}).status_code == 200
    lines = tenant.get("/api/boqs/%d" % boq["id"]).json()["lines"]
    assert any(l["material"] == "M60" for l in lines) and any(l["rate_only"] for l in lines)


# --- group code and a typed item code ----------------------------------------------------------------------------------

def save_with(tenant, boq, patch, issue_codes=True):
    lines = import_lines(tenant, boq)["lines"]
    for line in lines:
        if line["description"].startswith("Excavation"):
            line.update(patch)
    return tenant.put("/api/boqs/%d/lines" % boq["id"], json={"lines": lines, "issue_codes": issue_codes})


def test_a_group_code_is_kept_with_the_line(tenant):
    _, boq = make_boq(tenant)
    res = save_with(tenant, boq, {"group_code": "Earthwork"})
    assert res.status_code == 200, res.text
    assert [l["group_code"] for l in res.json()["lines"] if l["description"].startswith("Excavation")] == ["Earthwork"]


def test_a_typed_item_code_is_kept_and_added_to_the_master(tenant):
    _, boq = make_boq(tenant)
    res = save_with(tenant, boq, {"item_code": "ZK4M7X"})
    assert res.status_code == 200, res.text
    assert [l["item_code"] for l in res.json()["lines"] if l["description"].startswith("Excavation")] == ["ZK4M7X"]
    items = tenant.get("/api/erp/items").json()
    assert "ZK4M7X" in str(items)


def test_a_blank_item_code_is_still_given_one(tenant):
    _, boq = make_boq(tenant)
    res = save_with(tenant, boq, {"item_code": ""})
    assert all(l["item_code"] for l in res.json()["lines"] if l["kind"] != "section")


def test_one_item_code_cannot_be_on_two_lines(tenant):
    _, boq = make_boq(tenant)
    lines = import_lines(tenant, boq)["lines"]
    priced = [l for l in lines if l["kind"] != "section"]
    priced[0]["item_code"] = priced[1]["item_code"] = "ZK4M7X"
    res = tenant.put("/api/boqs/%d/lines" % boq["id"], json={"lines": lines, "issue_codes": True})
    assert res.status_code == 400 and "two lines" in res.json()["detail"]


def test_an_item_code_of_any_pattern_is_kept_and_added_to_the_master(tenant):
    _, boq = make_boq(tenant)
    res = save_with(tenant, boq, {"item_code": "civ-rcc-25"})
    assert res.status_code == 200, res.text
    line = [l for l in res.json()["lines"] if l["description"].startswith("Excavation")][0]
    assert line["item_code"] == "CIV-RCC-25", "kept as typed, in capitals"
    items = tenant.get("/api/erp/items").json()
    row = [i for k in ("FG",) for i in items.get(k, items.get("items", [])) if i["item_code"] == "CIV-RCC-25"] if isinstance(items, dict) else []
    assert row and row[0]["item_name"].startswith("Excavation"), "the code is a finished good in the master under the line's name"


def test_a_typed_item_code_with_spaces_or_odd_marks_is_refused(tenant):
    _, boq = make_boq(tenant)
    for bad in ("CIV RCC", "RCC*25", "A" * 31):
        res = save_with(tenant, boq, {"item_code": bad})
        assert res.status_code == 400 and "letters, numbers" in res.json()["detail"], bad


def test_the_code_typed_on_the_boq_is_the_code_on_every_document_made_from_it(tenant):
    jid, boq = make_boq(tenant)
    lines = import_lines(tenant, boq)["lines"]
    typed = {"Excavation in all types of soil": "CIV-EXC-01", "M25 grade RCC in raft": "CIV-RCC-25",
             "Shuttering to raft edges": "CIV-FWK-03", "Brick masonry in CM 1:6": "CIV-MAS-06"}
    for l in lines:
        l["item_code"] = typed.get(l["description"], "")
    saved_boq = tenant.put("/api/boqs/%d/lines" % boq["id"], json={"lines": lines, "issue_codes": True}).json()
    assert {l["description"]: l["item_code"] for l in saved_boq["lines"] if l["kind"] != "section"} == typed
    # the client's work order carries the same codes
    wo = tenant.post("/api/boqs/%d/client-order" % boq["id"], json={}).json()["work_order"]
    full = tenant.get("/api/erp/work-orders/%d" % wo["id"]).json()
    assert sorted(l["fg_code"] for l in full["lines"]) == sorted(typed.values())
    # a gang's order pulls the line with the same code, and the tracker shows it
    key = [l for l in saved_boq["lines"] if l["item_code"] == "CIV-RCC-25"][0]["key"]
    pulled = tenant.post("/api/wo/orders/%d/boq-lines" % draft(tenant)["id"], json={"lines": [{"key": key, "quantity": 10, "rate": 100}]}).json()
    assert pulled["lines"][0]["item_code"] == "CIV-RCC-25"
    rows = tenant.get("/api/boqs/%d/tracker" % boq["id"]).json()["rows"]
    assert "CIV-RCC-25" in [r.get("item_code") for r in rows]


def test_a_blank_item_code_is_still_given_one_beside_typed_ones(tenant):
    _, boq = make_boq(tenant)
    res = save_with(tenant, boq, {"item_code": "CIV-EXC-01"})
    priced = [l for l in res.json()["lines"] if l["kind"] != "section"]
    assert priced[0]["item_code"] == "CIV-EXC-01" and all(l["item_code"] for l in priced)



def test_the_sample_sheet_imports_as_it_is(tenant):
    _, boq = make_boq(tenant)
    sample = tenant.get("/api/boqs/sample.xlsx")
    assert sample.status_code == 200
    out = tenant.post("/api/boqs/%d/import" % boq["id"], files={"file": ("sample_boq.xlsx", sample.content)})
    assert out.status_code == 200, out.text
    lines = out.json()["lines"]
    assert [l["description"] for l in lines if l["kind"] == "section"][0] == "PLAIN CEMENT CONCRETE"
    assert any(l["rate_only"] for l in lines) and any(l["group_code"] == "SUB STRUCTURE WORKS" for l in lines)
    assert not out.json()["warnings"]


def test_a_material_added_once_is_offered_on_every_boq(tenant):
    _, boq = make_boq(tenant)
    res = save_with(tenant, boq, {"material": "Anti-termite chemical"})
    assert res.status_code == 200, res.text
    got = tenant.get("/api/boqs/materials")
    assert got.status_code == 200 and "Anti-termite chemical" in got.json()["materials"]
    # A second project's BOQ sees it too.
    _, other = make_boq(tenant, "Another site")
    assert "Anti-termite chemical" in tenant.get("/api/boqs/materials").json()["materials"]


def test_remarks_are_kept_and_go_to_the_excel(tenant):
    _, boq = make_boq(tenant)
    res = save_with(tenant, boq, {"remarks": "Rate includes pumping"})
    assert res.status_code == 200, res.text
    assert "Rate includes pumping" in [l["remarks"] for l in res.json()["lines"]]
    x = tenant.get("/api/boqs/%d/export.xlsx" % boq["id"])
    cells = [c.value for ws in openpyxl.load_workbook(io.BytesIO(x.content)).worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Remarks" in cells and "Rate includes pumping" in cells
    assert tenant.get("/api/boqs/%d/export.pdf" % boq["id"]).status_code == 200


# --- the company logo on the BOQ's PDF ------------------------------------------------------------------------------

def logo_uri():
    import base64
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (240, 80), (200, 40, 40)).save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def set_logo(account, uri):
    from app import models
    from app.db.session import SessionLocal
    db = SessionLocal()
    db.query(models.DBClient).filter(models.DBClient.email == account["email"]).update({"logo_url": uri})
    db.commit()
    db.close()


def test_the_boq_pdf_carries_the_company_logo_in_the_heading_and_as_a_watermark(tenant, account, monkeypatch):
    from app.documents import form_pdf
    _, boq = make_boq(tenant)
    saved(tenant, boq)
    plain = tenant.get("/api/boqs/%d/export.pdf" % boq["id"])
    assert plain.status_code == 200 and b"/Subtype /Image" not in plain.content

    set_logo(account, logo_uri())
    drawn = []
    real = form_pdf._draw_logo_watermark
    monkeypatch.setattr(form_pdf, "_draw_logo_watermark", lambda *a: (drawn.append(1), real(*a))[1])
    branded = tenant.get("/api/boqs/%d/export.pdf" % boq["id"])
    assert branded.status_code == 200 and b"/Subtype /Image" in branded.content
    assert len(drawn) >= 1, "the watermark is drawn on the page"
    assert len(branded.content) > len(plain.content)


def test_the_watermark_is_on_every_page_of_a_long_boq(tenant, account, monkeypatch):
    from pypdf import PdfReader
    from app.documents import form_pdf
    _, boq = make_boq(tenant)
    lines = [{"kind": "item", "description": "Item %d of the schedule" % n, "uom": "cum", "quantity": 1, "rate": 10} for n in range(160)]
    assert tenant.put("/api/boqs/%d/lines" % boq["id"], json={"lines": lines, "issue_codes": False}).status_code == 200
    set_logo(account, logo_uri())
    drawn = []
    real = form_pdf._draw_logo_watermark
    monkeypatch.setattr(form_pdf, "_draw_logo_watermark", lambda *a: (drawn.append(1), real(*a))[1])
    pdf = tenant.get("/api/boqs/%d/export.pdf" % boq["id"]).content
    pages = len(PdfReader(io.BytesIO(pdf)).pages)
    assert pages >= 2 and len(drawn) == pages


def test_a_logo_that_is_only_an_address_is_left_off_the_pdf_rather_than_fetched(tenant, account):
    _, boq = make_boq(tenant)
    saved(tenant, boq)
    set_logo(account, "http://example.invalid/logo.png")
    res = tenant.get("/api/boqs/%d/export.pdf" % boq["id"])
    assert res.status_code == 200 and b"/Subtype /Image" not in res.content
