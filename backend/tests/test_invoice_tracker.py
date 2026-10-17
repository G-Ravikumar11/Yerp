"""The invoice tracker: the purchase register built from posted receipts, with the supplier's own bill checked against it."""
from test_goods_receipt import order, post, receipt, set_lines
from test_site_pricing import make_item


def delivered(tenant, code, name, category, qty, price, supplier="Sudhakar PVC", **receipt_fields):
    """One posted delivery of one material. Returns the receipt as it stood before posting (its id is what matters)."""
    tenant.post("/api/item-categories", json={"name": category})          # a repeat is simply refused
    if not any(i["item_code"] == code for i in tenant.get("/api/items").json()["items"]):
        make_item(tenant, code, name, item_category=category, specification="", units_of_measure="Nos", item_type="Purchased")
    po = order(tenant, lines=[{"description": name, "item_code": code, "uom": "Nos", "qty": qty, "price": price}],
               supplier=supplier)
    grn = receipt(tenant, po, **receipt_fields)
    set_lines(tenant, grn, [{"received_qty": qty, "accepted_qty": qty}])
    post(tenant, grn)
    return grn


def test_the_summary_adds_gst_and_groups_purchases_by_trade(tenant):
    delivered(tenant, "PL-1", "75MM PVC pipe", "Plumbing", 100, 100)       # 10,000 + 18% = 11,800
    delivered(tenant, "EL-1", "6A switch", "Electrical", 50, 20)           # 1,000 + 18% = 1,180
    t = tenant.get("/api/invoice-tracker").json()
    rows = {r["category"]: r for r in t["summary"]["rows"]}
    assert rows["Plumbing"]["purchase_amount"] == 11800 and rows["Electrical"]["purchase_amount"] == 1180
    assert t["summary"]["purchase_amount"] == 12980
    assert t["totals"]["invoices"] == 2 and t["totals"]["gst"] == 1980


def test_a_draft_receipt_is_not_a_purchase_yet(tenant):
    tenant.post("/api/item-categories", json={"name": "Plumbing"})
    make_item(tenant, "PL-2", "Elbow", item_category="Plumbing", specification="", units_of_measure="Nos", item_type="Purchased")
    po = order(tenant, lines=[{"description": "Elbow", "item_code": "PL-2", "uom": "Nos", "qty": 5, "price": 10}])
    receipt(tenant, po)
    assert tenant.get("/api/invoice-tracker").json()["totals"]["invoices"] == 0


def test_the_suppliers_printed_total_is_checked_against_what_was_received(tenant):
    grn = delivered(tenant, "PL-3", "110MM pipe", "Plumbing", 10, 100, invoice_number="SSMA-2718")   # 1,000 + 180
    t = tenant.get("/api/invoice-tracker").json()
    assert t["invoices"][0]["difference"] is None and t["totals"]["unchecked"] == 1, "nothing keyed in yet"
    assert tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"],
                      json={"invoice_date": "2026-01-02", "invoice_amount": 1180}).status_code == 200
    t = tenant.get("/api/invoice-tracker").json()
    assert t["invoices"][0]["difference"] == 0 and t["totals"]["mismatched"] == 0
    assert t["invoices"][0]["invoice_date"] == "2026-01-02" and t["invoices"][0]["invoice_number"] == "SSMA-2718"
    tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"invoice_amount": 1100, "freight": 50})
    t = tenant.get("/api/invoice-tracker").json()
    assert t["invoices"][0]["total"] == 1230 and t["invoices"][0]["difference"] == -130 and t["totals"]["mismatched"] == 1
    tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"clear_amount": True})
    assert tenant.get("/api/invoice-tracker").json()["invoices"][0]["invoice_amount"] is None


def test_each_material_lists_what_came_on_each_invoice(tenant):
    delivered(tenant, "PL-4", "Clamp 75MM", "Plumbing", 40, 10)
    delivered(tenant, "PL-4", "Clamp 75MM", "Plumbing", 60, 12)
    plumbing = [c for c in tenant.get("/api/invoice-tracker").json()["categories"] if c["name"] == "Plumbing"][0]
    assert len(plumbing["invoice_ids"]) == 2 and len(plumbing["materials"]) == 1
    mat = plumbing["materials"][0]
    assert mat["total_qty"] == 100 and mat["total_amount"] == 1120 and [c["qty"] for c in mat["cells"]] == [40, 60]


def test_ready_mix_gets_a_register_by_lorry_and_block(tenant):
    a = delivered(tenant, "RMC-15", "RMC M15 concrete", "RMC", 7, 3500, supplier="Techno Infra", vehicle_number="AP39UZ4477")
    b = delivered(tenant, "RMC-15", "RMC M15 concrete", "RMC", 3, 3500, supplier="Techno Infra", vehicle_number="AP39WM2377")
    rows = tenant.get("/api/invoice-tracker").json()["rmc"]["rows"]
    assert [r["vehicle"] for r in rows] == ["AP39UZ4477", "AP39WM2377"] and {r["grade"] for r in rows} == {"M15"}
    for row, block in zip(rows, ("B-21", "B-21")):
        assert tenant.put("/api/invoice-tracker/blocks/%d" % row["line_id"], json={"block": block}).status_code == 200
    rmc = tenant.get("/api/invoice-tracker").json()["rmc"]
    assert rmc["qty"] == 10 and rmc["amount"] == 35000 and rmc["gst"] == 6300 and rmc["total"] == 41300
    assert rmc["blocks"] == [{"block": "B-21", "loads": 2, "qty": 10, "amount": 35000}]
    assert a["id"] != b["id"]


def test_the_work_order_value_and_remark_give_the_difference(tenant):
    delivered(tenant, "CV-1", "Cement OPC", "Civil", 100, 100)             # 11,800
    assert tenant.put("/api/invoice-tracker/categories", json={"category": "Civil", "wo_amount": 10000, "remarks": "WO pending"}).status_code == 200
    row = [r for r in tenant.get("/api/invoice-tracker").json()["summary"]["rows"] if r["category"] == "Civil"][0]
    assert row["wo_amount"] == 10000 and row["difference"] == -1800 and row["remarks"] == "WO pending"


def test_bad_input_is_refused_and_one_business_cannot_see_anothers(tenant, second_tenant):
    grn = delivered(tenant, "CV-2", "Sand", "Civil", 10, 10)
    assert tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"invoice_amount": -5}).status_code == 400
    assert tenant.put("/api/invoice-tracker/categories", json={"category": "  "}).status_code == 400
    assert second_tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"freight": 5}).status_code == 404
    assert second_tenant.get("/api/invoice-tracker").json()["totals"]["invoices"] == 0


def test_it_can_be_downloaded_as_a_sheet(tenant):
    delivered(tenant, "CV-3", "Bricks", "Civil", 10, 10)
    res = tenant.get("/api/invoice-tracker/export.xlsx")
    assert res.status_code == 200 and "spreadsheet" in res.headers["content-type"]


def client_of(grn):
    import main
    from app import models
    with main.SessionLocal() as db:
        return db.get(models.DBGoodsReceipt, grn["id"]).client_id


def test_insurance_and_hamali_are_part_of_what_the_bill_should_come_to(tenant):
    grn = delivered(tenant, "CV-4", "Cement", "Civil", 10, 100)               # 1,000 + 180
    res = tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"],
                     json={"freight": 20, "insurance": 30, "hamali": 40, "invoice_amount": 1270})
    assert res.status_code == 200
    t = tenant.get("/api/invoice-tracker").json()
    inv = t["invoices"][0]
    assert (inv["freight"], inv["insurance"], inv["hamali"], inv["total"], inv["difference"]) == (20, 30, 40, 1270, 0)
    assert t["totals"]["insurance"] == 30 and t["totals"]["hamali"] == 40 and t["totals"]["mismatched"] == 0
    assert [r["purchase_amount"] for r in t["summary"]["rows"]] == [1270], "the extra charges belong to the trade's purchases"
    assert tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"hamali": -1}).status_code == 400


def test_each_trade_shows_totals_under_every_invoice_and_the_markup_on_the_basic_rate(tenant):
    grn = delivered(tenant, "CV-5", "Sand", "Civil", 10, 100)
    tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"freight": 50, "invoice_amount": 1200})
    civil = tenant.get("/api/invoice-tracker?markup=20").json()["categories"][0]
    f = civil["footer"][0]
    assert (f["amount"], f["gst"], f["freight"], f["total"], f["invoice_amount"], f["difference"]) == (1000, 180, 50, 1230, 1200, -30)
    assert civil["markup"] == 200 and civil["with_markup"] == 1200 and civil["with_gst"] == 1180
    assert tenant.get("/api/invoice-tracker?markup=50").json()["categories"][0]["markup"] == 500


def test_a_single_trade_invoice_shows_its_charges_in_that_trades_footer(tenant):
    delivered(tenant, "SP-1", "Pipe", "Plumbing", 10, 100)
    t = tenant.get("/api/invoice-tracker").json()
    assert t["invoices"][0]["trades"] == ["Plumbing"] and t["categories"][0]["footer"][0]["whole"] is True


def test_the_work_order_value_comes_from_committed_work_orders_unless_typed_over(tenant):
    import main
    from app import models
    grn = delivered(tenant, "PL-9", "Pipe", "Plumbing", 100, 100)             # 11,800 bought
    with main.SessionLocal() as db:
        for status, amount in (("APPROVED", 9000), ("EXECUTED", 3000), ("DRAFT", 99999)):
            db.add(models.DBSubcontractOrder(client_id=client_of(grn), wo_number="WO-%s" % status, status=status,
                                             work_type="Plumbing", gross_amount=amount))
        db.add(models.DBSubcontractOrder(client_id=client_of(grn), wo_number="WO-E", status="APPROVED",
                                         work_type="Electrical", gross_amount=500))
        db.commit()
    rows = {r["category"]: r for r in tenant.get("/api/invoice-tracker").json()["summary"]["rows"]}
    assert rows["Plumbing"]["wo_amount"] == 12000 and rows["Plumbing"]["wo_source"] == "work orders", "a draft is not awarded yet"
    assert rows["Plumbing"]["difference"] == 200
    assert rows["Electrical"]["wo_amount"] == 500 and rows["Electrical"]["purchase_amount"] == 0, "awarded but nothing bought yet"
    tenant.put("/api/invoice-tracker/categories", json={"category": "Plumbing", "wo_amount": 10000})
    row = [r for r in tenant.get("/api/invoice-tracker").json()["summary"]["rows"] if r["category"] == "Plumbing"][0]
    assert row["wo_amount"] == 10000 and row["wo_source"] == "typed" and row["wo_suggested"] == 12000


def test_the_physical_stock_sheet_lists_every_purchased_material_with_what_the_store_holds(tenant):
    delivered(tenant, "ST-1", "Switch 6A", "Electrical", 40, 20)
    make_item(tenant, "ST-2", "Fan regulator", item_category="Electrical", specification="", units_of_measure="Nos", item_type="Purchased")
    sheet = {c["category"]: c["items"] for c in tenant.get("/api/invoice-tracker").json()["stock"]}["Electrical"]
    by_code = {i["item_code"]: i for i in sheet}
    assert by_code["ST-1"]["on_hand"] == 40 and by_code["ST-1"]["counted"] is True
    assert by_code["ST-2"]["on_hand"] == 0 and by_code["ST-2"]["counted"] is False, "listed even though nothing has been received"


def test_an_invoice_across_two_trades_splits_the_goods_and_keeps_the_charges_on_the_invoice(tenant):
    for cat, code, name in (("Plumbing", "MX-1", "Pipe"), ("Electrical", "MX-2", "Switch")):
        tenant.post("/api/item-categories", json={"name": cat})
        make_item(tenant, code, name, item_category=cat, specification="", units_of_measure="Nos", item_type="Purchased")
    po = order(tenant, lines=[{"description": "Pipe", "item_code": "MX-1", "uom": "Nos", "qty": 10, "price": 100},
                              {"description": "Switch", "item_code": "MX-2", "uom": "Nos", "qty": 5, "price": 40}])
    grn = receipt(tenant, po)
    set_lines(tenant, grn, [{"received_qty": 10, "accepted_qty": 10}, {"received_qty": 5, "accepted_qty": 5}])
    post(tenant, grn)
    tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"freight": 60})
    t = tenant.get("/api/invoice-tracker").json()
    assert t["invoices"][0]["trades"] == ["Plumbing", "Electrical"] and t["invoices"][0]["amount"] == 1200
    feet = {c["name"]: c["footer"][0] for c in t["categories"]}
    assert feet["Plumbing"]["amount"] == 1000 and feet["Electrical"]["amount"] == 200
    assert feet["Plumbing"]["whole"] is False and feet["Plumbing"]["freight"] is None and feet["Plumbing"]["invoice_amount"] is None
    rows = {r["category"]: r["purchase_amount"] for r in t["summary"]["rows"]}
    assert round(rows["Plumbing"] + rows["Electrical"], 2) == t["invoices"][0]["total"] == 1476, "freight is shared by value, nothing lost"


# --- the workbook --------------------------------------------------------------------------------------------------

def workbook_of(tenant, url="/api/invoice-tracker/export.xlsx"):
    import io
    import openpyxl
    res = tenant.get(url)
    assert res.status_code == 200, res.text
    return openpyxl.load_workbook(io.BytesIO(res.content))


def test_the_workbook_has_a_sheet_for_each_part_and_starts_with_the_figures(tenant):
    delivered(tenant, "XL-1", "75MM PVC pipe", "Plumbing", 100, 100)
    delivered(tenant, "XL-2", "RMC M15 concrete", "RMC", 7, 3500, vehicle_number="AP39UZ4477")
    book = workbook_of(tenant)
    names = book.sheetnames
    assert names == ["Summary", "Invoices", "Plumbing", "Ready-mix"], "a sheet per trade; ready-mix has its register instead of a list"
    summary = book["Summary"]
    assert [c.value for c in summary[1]][:5] == ["S.no", "Description", "Purchase amount", "WO amount", "Difference"], "the table opens the sheet"
    assert summary["B2"].value == "Plumbing" and summary["C2"].value == 11800
    assert summary.freeze_panes == "A2" and summary.sheet_view.showGridLines is False


def test_who_issued_the_sheet_and_for_what_comes_after_the_data_not_before(tenant, account):
    from test_boq import logo_uri, set_logo
    delivered(tenant, "XL-3", "Cement", "Civil", 10, 100)
    set_logo(account, logo_uri())
    ws = workbook_of(tenant)["Summary"]
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    last_figure = max(r for (r, c), v in cells.items() if c == 2 and v == "Total amount")
    label_rows = {v: r for (r, c), v in cells.items() if c == 1 and isinstance(v, str)}
    assert label_rows["REPORT DETAILS"] > last_figure, "the details follow the totals"
    assert min(label_rows[k] for k in ("REPORT DETAILS", "Company", "Report", "GST", "Generated") if k in label_rows) > last_figure
    assert "Acme Ltd" in cells.values() and "Invoice Tracker - Summary" in cells.values()
    assert not any(r < last_figure and isinstance(v, str) and v in ("Acme Ltd", "REPORT DETAILS") for (r, c), v in cells.items())


def test_the_logo_is_a_faded_watermark_on_every_sheet(tenant, account):
    import io
    from PIL import Image
    from test_boq import logo_uri, set_logo
    from app.documents.tracker_xlsx import faded_logo
    delivered(tenant, "XL-4", "Cement", "Civil", 10, 100)
    set_logo(account, logo_uri())
    book = workbook_of(tenant)
    assert all(len(ws._images) == 1 for ws in book.worksheets), "one logo on each sheet"
    pixel = Image.open(io.BytesIO(book["Summary"]._images[0]._data())).convert("RGBA").getpixel((10, 10))
    assert pixel[3] <= 0.12 * 255 + 1 and pixel[0] > 150, "the red logo is drawn back to a faint tint"
    import base64
    raw = base64.b64decode(logo_uri().partition(",")[2])
    assert faded_logo(b"not an image") is None and faded_logo(raw)


def test_a_sheet_without_a_logo_still_opens_with_no_picture(tenant):
    delivered(tenant, "XL-5", "Cement", "Civil", 10, 100)
    assert all(len(ws._images) == 0 for ws in workbook_of(tenant).worksheets)


def test_a_mismatched_bill_is_marked_in_the_workbook(tenant):
    grn = delivered(tenant, "XL-6", "Sand", "Civil", 10, 100)
    tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"invoice_amount": 1000})
    ws = workbook_of(tenant)["Invoices"]
    header = [c.value for c in ws[1]]
    row = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(header)}
    assert row["Status"] == "Mismatch" and row["Difference"] == -180 and row["Total"] == 1180


def test_the_trade_sheet_carries_the_totals_and_markup_and_the_pdf_twin_still_works(tenant):
    grn = delivered(tenant, "XL-7", "Bricks", "Civil", 10, 100)
    tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"freight": 50, "invoice_amount": 1230})
    ws = workbook_of(tenant)["Civil"]
    labels = [ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)]
    labels += [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    for want in ("AMOUNT", "GST 18%", "FREIGHT CHARGES", "INSURANCE", "HAMALI", "TOTAL AMOUNT", "INVOICE AMOUNT", "DIFFERENT AMOUNT", "REMARKS"):
        assert want in labels, want
    assert any(isinstance(v, str) and "@ 20%" in v for v in labels) and any(isinstance(v, str) and "total with markup" in v for v in labels)
    assert tenant.get("/api/invoice-tracker/export.pdf").status_code == 200


# --- what the sheet of materials carries: every material, total and working rate, remarks, the stock ---------------------

def test_every_material_of_the_trade_is_listed_even_before_any_has_arrived(tenant):
    delivered(tenant, "LM-1", "75MM PVC pipe", "Plumbing", 100, 100)
    make_item(tenant, "LM-2", "Gully trap", item_category="Plumbing", specification="", units_of_measure="Nos", item_type="Purchased")
    mats = {m["item_code"]: m for m in tenant.get("/api/invoice-tracker").json()["categories"][0]["materials"]}
    assert mats["LM-1"]["total_qty"] == 100 and mats["LM-2"]["cells"] == [] and mats["LM-2"]["total_qty"] == 0
    assert [m["sl"] for m in tenant.get("/api/invoice-tracker").json()["categories"][0]["materials"]] == [1, 2]


def test_total_rate_is_the_rates_added_up_and_the_working_rate_is_one_you_set(tenant):
    delivered(tenant, "LM-3", "Clamp", "Plumbing", 40, 10)
    delivered(tenant, "LM-3", "Clamp", "Plumbing", 60, 12)
    mat = tenant.get("/api/invoice-tracker").json()["categories"][0]["materials"][0]
    assert mat["total_rate"] == 22 and mat["avg_rate"] == 11.2 and mat["working_rate"] is None and mat["working_amount"] is None
    assert tenant.put("/api/invoice-tracker/working-rates", json={"item_key": "lm-3", "rate": 11}).status_code == 200
    cat = tenant.get("/api/invoice-tracker").json()["categories"][0]
    assert cat["materials"][0]["working_rate"] == 11 and cat["materials"][0]["working_amount"] == 1100 and cat["working_total"] == 1100
    assert tenant.put("/api/invoice-tracker/working-rates", json={"item_key": "LM-3", "rate": None}).status_code == 200
    assert tenant.get("/api/invoice-tracker").json()["categories"][0]["materials"][0]["working_rate"] is None
    assert tenant.put("/api/invoice-tracker/working-rates", json={"item_key": "LM-3", "rate": -1}).status_code == 400
    assert tenant.put("/api/invoice-tracker/working-rates", json={"item_key": " "}).status_code == 400


def test_a_remark_stays_with_its_invoice(tenant):
    grn = delivered(tenant, "LM-4", "Wire", "Electrical", 10, 100)
    assert tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"remark": "SOME COILS NOT RECEIVED"}).status_code == 200
    t = tenant.get("/api/invoice-tracker").json()
    assert t["invoices"][0]["remark"] == "SOME COILS NOT RECEIVED" and t["categories"][0]["footer"][0]["remark"] == "SOME COILS NOT RECEIVED"
    assert workbook_of(tenant)["Invoices"]["Q2"].value == "SOME COILS NOT RECEIVED"


def test_what_a_trade_has_bought_is_what_its_suppliers_billed(tenant):
    grn = delivered(tenant, "LM-5", "Cable", "Electrical", 10, 100)           # receipts come to 1,180
    assert tenant.get("/api/invoice-tracker").json()["summary"]["rows"][0]["purchase_amount"] == 1180
    tenant.put("/api/invoice-tracker/invoices/%d" % grn["id"], json={"invoice_amount": 1200})
    t = tenant.get("/api/invoice-tracker").json()
    assert t["summary"]["rows"][0]["purchase_amount"] == 1200, "the printed bill is what was paid"
    assert t["summary"]["purchase_amount"] == 1200


def test_the_trade_sheet_follows_the_layout_of_the_list_people_keep(tenant, account):
    from test_boq import logo_uri, set_logo
    grn = delivered(tenant, "LM-6", "6A 1-way switch", "Electrical", 100, 15, supplier="SRIVEN ELECTRICAL SOLUTIONS", invoice_number="SES/4222")
    make_item(tenant, "LM-7", "Fan regulator", item_category="Electrical", specification="", units_of_measure="Nos", item_type="Purchased")
    tenant.put("/api/invoice-tracker/working-rates", json={"item_key": "LM-6", "rate": 14})
    set_logo(account, logo_uri())
    ws = workbook_of(tenant)["Electrical"]
    assert [ws.cell(row=1, column=c).value for c in (1, 2, 3)] == ["SL NO", "MATERIAL DESCRIPTION", "UOM"]
    assert ws["D1"].value == "SRIVEN ELECTRICAL SOLUTIONS" and ws["D3"].value == "SES/4222", "supplier, date, number, PO down the heading"
    assert [ws.cell(row=5, column=c).value for c in (4, 5, 6)] == ["Qty", "Rate", "Amount"]
    heads = [ws.cell(row=1, column=c).value for c in range(7, 13)]
    assert heads == ["Total qty", "Total rate", "Total amount", "Working rate", "Working amount", "Physical store stock"]
    row = {h: ws.cell(row=6, column=7 + i).value for i, h in enumerate(heads)}
    assert row["Total qty"] == 100 and row["Total amount"] == 1500 and row["Working rate"] == 14 and row["Working amount"] == 1400
    assert ws["A7"].value == 2 and ws["G7"].value is None, "a material still to arrive is listed, with nothing against it"
    assert row["Physical store stock"] == 100 and len(ws._images) == 1 and grn["id"]


def test_ready_mix_is_not_given_a_materials_sheet(tenant):
    delivered(tenant, "LM-8", "RMC M15 concrete", "RMC", 7, 3500)
    names = workbook_of(tenant).sheetnames
    assert "RMC" not in names and "Ready-mix" in names
