"""UAN, Aadhaar and ESI numbers: given when an employee is added, edited by HR, and printed on the payslip."""
import io

from conftest import make_employee

UAN, AADHAAR, ESI = "101185316896", "801874264479", "3100123456"


def test_the_numbers_are_saved_when_the_employee_is_added(tenant):
    emp = make_employee(tenant, uan_no=UAN, aadhaar_no=AADHAAR, esi_no=ESI)
    got = tenant.get("/api/employees/%d" % emp["id"]).json()
    assert got["uan_no"] == UAN and got["esi_no"] == ESI
    assert got["aadhaar_masked"] == "XXXX XXXX 4479"          # the full number is never sent back to the screen


def test_spaces_are_allowed_and_a_wrong_number_is_refused(tenant):
    emp = make_employee(tenant, uan_no="1011 8531 6896", esi_no="3100 123 456")
    got = tenant.get("/api/employees/%d" % emp["id"]).json()
    assert got["uan_no"] == UAN and got["esi_no"] == ESI
    for bad in ({"uan_no": "123"}, {"esi_no": "12345"}, {"aadhaar_no": "123456789012"}):
        res = tenant.put("/api/employees/%d" % emp["id"], json=bad)
        assert res.status_code == 400, bad


def test_the_numbers_can_be_changed_and_cleared(tenant):
    emp = make_employee(tenant)
    assert tenant.put("/api/employees/%d" % emp["id"], json={"esi_no": ESI, "uan_no": UAN}).status_code == 200
    assert tenant.get("/api/employees/%d" % emp["id"]).json()["esi_no"] == ESI
    assert tenant.put("/api/employees/%d" % emp["id"], json={"esi_no": ""}).status_code == 200
    assert tenant.get("/api/employees/%d" % emp["id"]).json()["esi_no"] == ""


def test_the_payslip_carries_the_numbers(tenant):
    import openpyxl
    emp = make_employee(tenant, uan_no=UAN, aadhaar_no=AADHAAR, esi_no=ESI)
    pid = tenant.post("/api/payslips", json={"employee_id": emp["id"], "period_start": "2026-09-01", "period_end": "2026-09-30",
                                             "pay_date": "2026-09-30", "basic_salary": 1000, "use_structure": False}).json()["id"]
    sheet = openpyxl.load_workbook(io.BytesIO(tenant.get("/api/payslips/%d/xlsx" % pid).content)).active
    cells = [str(c.value) for row in sheet.iter_rows() for c in row if c.value not in (None, "")]
    assert UAN in cells and ESI in cells and "XXXX XXXX 4479" in cells
    assert AADHAAR not in cells                                 # masked on the slip
    reg = openpyxl.load_workbook(io.BytesIO(tenant.get("/api/payslips/register.xlsx", params={"ids": str(pid)}).content)).active
    row = [str(c.value) for r in reg.iter_rows() for c in r if c.value not in (None, "")]
    assert UAN in row and ESI in row


def test_the_employee_gives_their_esi_number_in_their_own_portal(tenant):
    from fastapi.testclient import TestClient
    import main
    email = "esi-%s@example.com" % __import__("uuid").uuid4().hex[:6]
    make_employee(tenant, email=email, password="Crew1234", status="active")
    c = TestClient(main.app)
    c.__enter__()
    try:
        assert c.post("/api/employee/auth/login", json={"email": email, "password": "Crew1234"}).status_code == 200
        assert c.put("/api/employee/profile/details", json={"esi_no": ESI}).status_code == 200
        assert c.put("/api/employee/profile/details", json={"esi_no": "12"}).status_code == 400
        assert c.get("/api/employee/profile").json()["esi_no"] == ESI
    finally:
        c.__exit__(None, None, None)


# --- the PAN number, and the Aadhaar, PAN and ESI documents ----------------------------------------------------

import base64
import uuid

PAN = "ABCDE1234F"


def _file(name="aadhaar.pdf", kind="application/pdf", data=b"%PDF-1.4 hello"):
    return {"file_name": name, "file_type": kind, "file_data": base64.b64encode(data).decode()}


def test_the_pan_number_is_saved_with_the_new_employee(tenant):
    emp = make_employee(tenant, pan_no="abcde1234f")
    assert tenant.get("/api/employees/%d" % emp["id"]).json()["pan_no"] == PAN
    assert tenant.put("/api/employees/%d" % emp["id"], json={"pan_no": "12345"}).status_code == 400


def test_documents_can_be_kept_for_aadhaar_pan_and_esi_in_any_accepted_format(tenant):
    emp = make_employee(tenant)
    base = "/api/employees/%d/documents" % emp["id"]
    for title, f in (("Aadhaar card", _file("a.pdf")), ("PAN card", _file("p.jpeg", "image/jpeg", b"\xff\xd8\xff")),
                     ("ESI card", _file("e.png", "image/png", b"\x89PNG"))):
        res = tenant.post(base, json={"title": title, "doc_type": "identity", **f})
        assert res.status_code == 200, res.text
    listed = tenant.get(base).json()
    assert {d["title"] for d in listed} == {"Aadhaar card", "PAN card", "ESI card"}
    first = listed[0]
    got = tenant.get("%s/%d/download" % (base, first["id"])).json()
    assert got["file_name"] == first["file_name"] and base64.b64decode(got["file_data"])


def test_a_document_that_is_the_wrong_type_or_too_big_is_refused(tenant):
    emp = make_employee(tenant)
    base = "/api/employees/%d/documents" % emp["id"]
    assert tenant.post(base, json={"title": "x", **_file("virus.exe", "application/octet-stream")}).status_code == 400
    assert tenant.post(base, json={"title": "x", **_file("empty.pdf", data=b"")}).status_code == 400
    assert tenant.post(base, json={"title": "x", **_file("big.pdf", data=b"0" * (5 * 1024 * 1024 + 10))}).status_code == 413
    assert tenant.get(base).json() == []


def test_a_document_can_be_removed_and_belongs_to_one_company(tenant, second_tenant):
    emp = make_employee(tenant)
    base = "/api/employees/%d/documents" % emp["id"]
    tenant.post(base, json={"title": "PAN card", **_file("p.pdf")})
    doc = tenant.get(base).json()[0]
    assert second_tenant.get("%s/%d/download" % (base, doc["id"])).status_code in (401, 403, 404)
    assert second_tenant.delete("%s/%d" % (base, doc["id"])).status_code in (401, 403, 404)
    assert tenant.delete("%s/%d" % (base, doc["id"])).status_code == 200
    assert tenant.get(base).json() == []


# --- gender is compulsory when an employee is added; address is kept -----------------------------------------------

def test_an_employee_cannot_be_added_without_a_gender(tenant, monkeypatch):
    monkeypatch.setenv("REQUIRE_GENDER", "1")
    body = {"first_name": "Asha", "last_name": "Rao", "email": "g-%s@example.com" % uuid.uuid4().hex[:6]}
    refused = tenant.post("/api/employees", json=body)
    assert refused.status_code == 400 and "Gender" in refused.json()["detail"]
    assert tenant.post("/api/employees", json={**body, "gender": "Female"}).status_code == 200
    body["email"] = "g2-%s@example.com" % uuid.uuid4().hex[:6]
    assert tenant.post("/api/employees", json={**body, "gender": "Prefer not to say"}).status_code == 200      # an answer
    body["email"] = "g3-%s@example.com" % uuid.uuid4().hex[:6]
    assert tenant.post("/api/employees", json={**body, "gender": "Robot"}).status_code == 400


def test_the_address_is_saved_and_can_be_changed(tenant):
    emp = make_employee(tenant, address="Plot 1, Kondapur, Hyderabad")
    assert tenant.get("/api/employees/%d" % emp["id"]).json()["address"] == "Plot 1, Kondapur, Hyderabad"
    assert tenant.put("/api/employees/%d" % emp["id"], json={"address": "Flat 2, Madhapur"}).status_code == 200
    assert tenant.get("/api/employees/%d" % emp["id"]).json()["address"] == "Flat 2, Madhapur"
