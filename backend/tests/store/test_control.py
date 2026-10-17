"""P4 - control: transfers with in-transit and short receipt, blind stock counts that freeze the store, adjustments, scrap, rights."""
import uuid
from datetime import date, timedelta
from decimal import Decimal

import pytest
from conftest import make_employee

from app import db as database, models
from app.models.store_control import (
    DBStoreAdjustment, DBStoreAdjustmentLine, DBStoreControlLog, DBStoreCount, DBStoreCountLine, DBStoreScrap, DBStoreScrapLine,
    DBStoreTransfer, DBStoreTransferLine,
)
from app.models.store_core import (
    DBStore, DBStoreBin, DBStoreDocSequence, DBStoreFreeze, DBStoreItemSettings, DBStoreStockBalance, DBStoreStockLedger,
    DBStoreUomConversion, DBStoreUserAssignment, MovementType, StoreFrozenError,
)
from app.models.store_outward import DBStoreWriteoffRequest
from app.services.stock_ledger import get_balance, post_movement
from app.services.store_control import get_count_variances, get_in_transit, get_scrap_realisation
from app.services.store_outward import list_writeoff_requests

D = Decimal
C = "/api/store/control"
VIEW = ["store.%s.view" % s for s in ("masters", "enquiry", "control")]


@pytest.fixture(autouse=True)
def _clean(account):
    yield
    with database.SessionLocal() as s:
        cid = s.query(models.DBClient).filter(models.DBClient.email == account["email"]).first().id
        stores = [i for (i,) in s.query(DBStore.id).filter(DBStore.client_id == cid).all()]
        items = [i for (i,) in s.query(models.DBItem.id).filter(models.DBItem.client_id == cid).all()]
        for model, line, fk, t in ((DBStoreTransfer, DBStoreTransferLine, DBStoreTransferLine.transfer_id, "TRF"), (DBStoreCount, DBStoreCountLine, DBStoreCountLine.count_id, "CNT"),
                                   (DBStoreAdjustment, DBStoreAdjustmentLine, DBStoreAdjustmentLine.adjustment_id, "ADJ"), (DBStoreScrap, DBStoreScrapLine, DBStoreScrapLine.scrap_id, "SCR")):
            ids = [i for (i,) in s.query(model.id).filter(model.client_id == cid).all()]
            s.query(DBStoreControlLog).filter(DBStoreControlLog.doc_type == t, DBStoreControlLog.doc_id.in_(ids)).delete(synchronize_session=False)
            s.query(line).filter(fk.in_(ids)).delete(synchronize_session=False)
            s.query(model).filter(model.id.in_(ids)).delete(synchronize_session=False)
        s.query(DBStoreWriteoffRequest).filter(DBStoreWriteoffRequest.client_id == cid).delete(synchronize_session=False)
        s.query(DBStoreStockLedger).filter(DBStoreStockLedger.store_id.in_(stores)).delete(synchronize_session=False)
        for model in (DBStoreStockBalance, DBStoreFreeze, DBStoreUserAssignment, DBStoreBin, DBStoreDocSequence):
            s.query(model).filter(model.store_id.in_(stores)).delete(synchronize_session=False)
        for model in (DBStoreUomConversion, DBStoreItemSettings):
            s.query(model).filter(model.item_id.in_(items)).delete(synchronize_session=False)
        s.query(DBStore).filter(DBStore.id.in_(stores)).delete(synchronize_session=False)
        s.query(models.DBJob).filter(models.DBJob.client_id == cid).delete(synchronize_session=False)
        s.commit()


class W:
    pass


def seed(account):
    """Two stores on two projects; cement (kg, 7/kg) and steel (kg, 62/kg) held at store A."""
    o = account["client"]
    w = W()
    tag = uuid.uuid4().hex[:6]
    with database.SessionLocal() as s:
        w.cid = s.query(models.DBClient).filter(models.DBClient.email == account["email"]).first().id
        ja, jb = models.DBJob(client_id=w.cid, number="PA-" + tag, name="Tower A"), models.DBJob(client_id=w.cid, number="PB-" + tag, name="Tower B")
        cement = models.DBItem(client_id=w.cid, kind="RM", item_code="CEM-" + tag, item_name="Cement", units_of_measure="kg")
        steel = models.DBItem(client_id=w.cid, kind="RM", item_code="STL-" + tag, item_name="Steel", units_of_measure="kg")
        s.add_all([ja, jb, cement, steel])
        s.flush()
        s.add_all([DBStoreItemSettings(item_id=cement.id, base_uom="kg"), DBStoreItemSettings(item_id=steel.id, base_uom="kg")])
        s.commit()
        w.ja, w.jb, w.cement, w.steel = ja.id, jb.id, cement.id, steel.id
    mk = lambda code, name, pid: o.post("/api/store/masters/stores", json={"code": code + uuid.uuid4().hex[:3].upper(), "name": name, "type": "SITE", "project_id": pid}).json()["id"]   # noqa: E731
    w.a, w.b = mk("A", "Store A", w.ja), mk("B", "Store B", w.jb)
    with database.SessionLocal() as s:
        for item, qty, rate in ((w.cement, "1000", "7"), (w.steel, "500", "62")):
            post_movement(s, movement_type=MovementType.OPENING_IN, store_id=w.a, item_id=item, qty=D(qty), rate=D(rate), txn_date=date.today() - timedelta(days=5),
                          project_id=None, ref_doc_type="OPENING", ref_doc_id=1, ref_line_id=None, user_id=None)
        s.commit()
    return w


def bal(store, item):
    with database.SessionLocal() as s:
        b = get_balance(s, store, item)
        return b.qty, b.value


def go(o, kind, doc_id, *actions):
    res = None
    for a in actions:
        res = o.post("%s/%s/%d/%s" % (C, kind, doc_id, a), json={"reason": "because"} if a == "cancel" else {})
        assert res.status_code == 200, (a, res.text)
    return res.json()


def trf(o, w, qty="100", item=None, **kw):
    res = o.post(C + "/transfers", json={"store_id": w.a, "to_store_id": w.b, "vehicle_no": "KA01", "lines": [{"item_id": item or w.cement, "qty": qty}], **kw})
    assert res.status_code == 201, res.text
    return res.json()


def test_transfer_goes_in_transit_then_is_received(account):
    o = account["client"]
    w = seed(account)
    t = trf(o, w)
    assert t["number"].startswith("TRF/") and t["inter_project"] and t["status"] == "DRAFT"
    go(o, "transfers", t["id"], "submit", "approve")
    assert bal(w.a, w.cement)[0] == D("1000")
    posted = go(o, "transfers", t["id"], "post")
    assert posted["transit"] == "IN_TRANSIT" and posted["lines"][0]["rate"] == "7.0000" and posted["total_value"] == "700.00"
    assert bal(w.a, w.cement) == (D("900"), D("6300")) and bal(w.b, w.cement) == (0, 0)                    # out of A, not yet in B
    with database.SessionLocal() as s:
        rows = get_in_transit(s, project_id=w.jb)
        assert len(rows) == 1 and (rows[0].qty, rows[0].value, rows[0].inter_project) == (D("100"), D("700.00"), True)
        assert get_in_transit(s, project_id=999999) == []
    assert len(o.get(C + "/in-transit").json()["rows"]) == 1

    done = o.post("%s/transfers/%d/receive" % (C, t["id"]), json={}).json()
    assert done["transit"] == "RECEIVED" and done["lines"][0]["received_qty"] == "100.000"
    assert bal(w.b, w.cement) == (D("100"), D("700"))                                                       # in at the rate it left at
    with database.SessionLocal() as s:
        assert get_in_transit(s) == []
    assert o.post("%s/transfers/%d/receive" % (C, t["id"]), json={}).status_code == 409                     # not twice
    pdf = o.get("%s/transfers/%d/pdf" % (C, t["id"]))
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"


def test_dispatch_details_and_a_transfer_cannot_be_received_twice(account):
    o = account["client"]
    w = seed(account)
    t = trf(o, w, challan_no="CH-88", challan_date=str(date.today()), driver_name="Ramesh")
    assert (t["challan_no"], t["driver_name"], t["vehicle_no"]) == ("CH-88", "Ramesh", "KA01")
    assert o.put("%s/transfers/%d" % (C, t["id"]), json={"challan_no": "CH-89"}).json()["challan_no"] == "CH-89"
    go(o, "transfers", t["id"], "submit", "approve", "post")
    assert o.post("%s/transfers/%d/receive" % (C, t["id"]), json={}).status_code == 200
    assert o.post("%s/transfers/%d/receive" % (C, t["id"]), json={}).status_code == 409
    assert o.post("%s/transfers/%d/approve-loss" % (C, t["id"])).status_code == 409                          # nothing was lost
    assert "CH-89" in o.get("%s/transfers/%d/pdf" % (C, t["id"])).text or o.get("%s/transfers/%d/pdf" % (C, t["id"])).content[:4] == b"%PDF"


def test_loss_in_transit_needs_a_reason_and_becomes_wastage_when_approved(account):
    o = account["client"]
    w = seed(account)
    t = trf(o, w, "100")
    go(o, "transfers", t["id"], "submit", "approve", "post")
    line = o.get("%s/transfers/%d" % (C, t["id"])).json()["lines"][0]["id"]
    url = "%s/transfers/%d/receive" % (C, t["id"])
    assert o.post(url, json={"received": {str(line): "101"}}).status_code == 400                              # more than was sent
    assert o.post(url, json={"received": {str(line): "90"}}).status_code == 400                               # a loss needs a reason
    assert o.post(url, json={"received": {str(line): "90"}, "losses": {str(line): {"qty": "5", "reason": "BREAKAGE"}}}).status_code == 400   # 90 + 5 is not 100
    assert o.post(url, json={"received": {str(line): "90"}, "losses": {str(line): {"reason": "MELTED"}}}).status_code == 400                   # not a known reason
    res = o.post(url, json={"received": {str(line): "90"}, "losses": {str(line): {"reason": "BREAKAGE", "remarks": "dropped at the gate"}}}).json()
    assert res["transit"] == "RECEIVED" and res["loss_status"] == "PENDING" and res["lines"][0]["loss_qty"] == "10.000" and res["lines"][0]["loss_reason"] == "BREAKAGE"
    assert bal(w.b, w.cement)[0] == D("90") and bal(w.a, w.cement)[0] == D("900")                            # the 10 lost are in neither store yet
    with database.SessionLocal() as s:
        pend = get_in_transit(s)
        assert len(pend) == 1 and (pend[0].qty, pend[0].value) == (D("10"), D("70.00"))                    # still on the register until approved
    assert len(o.get(C + "/in-transit").json()["rows"]) == 1

    done = o.post("%s/transfers/%d/approve-loss" % (C, t["id"])).json()
    assert done["loss_status"] == "APPROVED" and bal(w.b, w.cement)[0] == D("90")                            # destination stock unchanged by the pass-through
    with database.SessionLocal() as s:
        assert get_in_transit(s) == []
        rows = s.query(DBStoreStockLedger).filter(DBStoreStockLedger.ref_doc_type == "TRANSFER", DBStoreStockLedger.ref_doc_id == t["id"], DBStoreStockLedger.store_id == w.b).all()
        waste = [r for r in rows if r.movement_type == "WASTAGE_OUT"]
        assert len(waste) == 1 and waste[0].qty_out == D("10") and waste[0].value == D("70.00") and waste[0].project_id == w.jb
    assert o.post("%s/transfers/%d/approve-loss" % (C, t["id"])).status_code == 409                          # once only
    assert any("wastage" in x["note"] for x in o.get("%s/transfers/%d" % (C, t["id"])).json()["log"])


def test_a_big_loss_needs_the_high_value_right(account, portal):
    o = account["client"]
    w = seed(account)
    with database.SessionLocal() as s:
        post_movement(s, movement_type=MovementType.OPENING_IN, store_id=w.a, item_id=w.steel, qty=D(5000), rate=D(62), txn_date=date.today(), project_id=None,
                      ref_doc_type="OPENING", ref_doc_id=7, ref_line_id=None, user_id=None)
        s.commit()
    t = trf(o, w, "2000", item=w.steel)
    go(o, "transfers", t["id"], "submit", "approve", "post")
    line = o.get("%s/transfers/%d" % (C, t["id"])).json()["lines"][0]["id"]
    o.post("%s/transfers/%d/receive" % (C, t["id"]), json={"received": {str(line): "0"}, "losses": {str(line): {"reason": "PILFERAGE"}}})      # the whole load, Rs 1,24,000
    _login(account, portal, VIEW + ["store.control.approve"], [w.a, w.b])
    assert portal.post("%s/transfers/%d/approve-loss" % (C, t["id"])).status_code == 403
    _login(account, portal, VIEW + ["store.control.approve", "store.control.approve_high_value"], [w.a, w.b])
    assert portal.post("%s/transfers/%d/approve-loss" % (C, t["id"])).status_code == 200


def test_transfer_stock_checks_and_cancel(account):
    o = account["client"]
    w = seed(account)
    big = trf(o, w, "1001")
    go(o, "transfers", big["id"], "submit", "approve")
    assert o.post("%s/transfers/%d/post" % (C, big["id"])).status_code == 409                               # insufficient stock
    assert bal(w.a, w.cement)[0] == D("1000")

    t = trf(o, w, "100")
    go(o, "transfers", t["id"], "submit", "approve", "post")
    assert go(o, "transfers", t["id"], "cancel")["status"] == "CANCELLED"                                    # in transit: stock returns to A
    assert bal(w.a, w.cement)[0] == D("1000")

    r = trf(o, w, "100")
    go(o, "transfers", r["id"], "submit", "approve", "post")
    o.post("%s/transfers/%d/receive" % (C, r["id"]), json={})
    with database.SessionLocal() as s:                                                                      # B used some of it
        post_movement(s, movement_type=MovementType.ISSUE_OUT, store_id=w.b, item_id=w.cement, qty=D(20), rate=None, txn_date=date.today(), project_id=w.jb,
                      ref_doc_type="ISSUE", ref_doc_id=5, ref_line_id=1, user_id=None)
        s.commit()
    assert o.post("%s/transfers/%d/cancel" % (C, r["id"]), json={"reason": "x"}).status_code == 409
    assert o.get("%s/transfers/%d" % (C, r["id"])).json()["status"] == "POSTED" and bal(w.b, w.cement)[0] == D("80")
    assert o.post(C + "/transfers", json={"store_id": w.a, "to_store_id": w.a, "lines": [{"item_id": w.cement, "qty": "1"}]}).status_code == 400
    assert o.post("%s/transfers/%d/cancel" % (C, t["id"]), json={"reason": "x"}).status_code == 409        # already cancelled


def _login(account, portal, rights, stores):
    email = "staff-%s@example.com" % uuid.uuid4().hex[:8]
    emp = make_employee(account["client"], email=email, password="Crew1234", permission_role="staff", status="active")
    with database.SessionLocal() as s:
        row = s.get(models.DBEmployee, emp["id"])
        row.extra_permissions = ",".join(rights)
        s.add_all([DBStoreUserAssignment(store_id=i, employee_id=row.id) for i in stores])
        s.commit()
    assert portal.post("/api/employee/auth/login", json={"email": email, "password": "Crew1234"}).status_code == 200


def test_inter_project_transfer_needs_the_high_value_right(account, portal):
    o = account["client"]
    w = seed(account)
    t = trf(o, w)
    go(o, "transfers", t["id"], "submit")
    _login(account, portal, VIEW + ["store.control.approve", "store.control.post"], [w.a, w.b])
    res = portal.post("%s/transfers/%d/approve" % (C, t["id"]))
    assert res.status_code == 403 and "high-value" in res.text
    _login(account, portal, VIEW + ["store.control.approve", "store.control.post", "store.control.approve_high_value"], [w.a, w.b])
    assert portal.post("%s/transfers/%d/approve" % (C, t["id"])).status_code == 200
    assert portal.post("%s/transfers/%d/post" % (C, t["id"])).status_code == 200

    _login(account, portal, VIEW + ["store.control.post"], [w.a])                                           # sender only: cannot receive
    assert portal.post("%s/transfers/%d/receive" % (C, t["id"]), json={}).status_code == 403
    _login(account, portal, VIEW + ["store.control.post"], [w.b])                                           # the receiving store can
    assert portal.get("%s/transfers/%d" % (C, t["id"])).status_code == 200
    assert portal.post("%s/transfers/%d/receive" % (C, t["id"]), json={}).status_code == 200
    assert bal(w.b, w.cement)[0] == D("100")


def test_count_freezes_counts_blind_and_posts_the_differences(account, portal):
    o = account["client"]
    w = seed(account)
    c = o.post(C + "/counts", json={"store_id": w.a}).json()
    assert c["number"].startswith("CNT/") and len(c["lines"]) == 2 and c["lines"][0]["system_qty"] is not None                # the account holder sees the books
    with database.SessionLocal() as s:
        with pytest.raises(StoreFrozenError):
            post_movement(s, movement_type=MovementType.ISSUE_OUT, store_id=w.a, item_id=w.cement, qty=D(1), rate=None, txn_date=date.today(), project_id=None,
                          ref_doc_type="ISSUE", ref_doc_id=9, ref_line_id=1, user_id=None)
    assert o.post(C + "/counts", json={"store_id": w.a}).status_code == 400                                  # already frozen
    by = {l["item_id"]: l["id"] for l in c["lines"]}

    _login(account, portal, VIEW + ["store.control.create"], [w.a])                                          # a counter: no right to see the books
    blind = portal.get("%s/counts/%d" % (C, c["id"])).json()
    assert blind["revealed"] is False and all(l["system_qty"] is None and l["variance_qty"] is None for l in blind["lines"])
    assert portal.get("%s/counts/%d/pdf" % (C, c["id"])).status_code == 200

    assert o.post("%s/counts/%d/submit" % (C, c["id"])).status_code == 400                                   # nothing counted
    upd = o.put("%s/counts/%d" % (C, c["id"]), json={"lines": [{"item_id": w.cement, "counted_qty": "985"}, {"item_id": w.steel, "counted_qty": "503"}]})
    assert upd.status_code == 200 and upd.json()["counted_count"] == 2
    go(o, "counts", c["id"], "submit", "approve")
    done = go(o, "counts", c["id"], "post")
    assert done["status"] == "POSTED" and done["variance_value"] == "81.00"                                  # -15 x 7 + 3 x 62
    assert bal(w.a, w.cement) == (D("985"), D("6895")) and bal(w.a, w.steel)[0] == D("503")
    with database.SessionLocal() as s:
        v = {x.item_id: x for x in get_count_variances(s, w.a, date.today(), date.today())}
        assert (v[w.cement].system_qty, v[w.cement].counted_qty, v[w.cement].variance_qty, v[w.cement].variance_value) == (D("1000"), D("985"), D("-15"), D("-105.00"))
        assert get_count_variances(s, w.a, date.today() + timedelta(days=1), None) == []
        post_movement(s, movement_type=MovementType.ISSUE_OUT, store_id=w.a, item_id=w.cement, qty=D(1), rate=None, txn_date=date.today(), project_id=None,
                      ref_doc_type="ISSUE", ref_doc_id=10, ref_line_id=1, user_id=None)              # unfrozen again
        s.commit()
    assert go(o, "counts", c["id"], "cancel")["status"] == "CANCELLED"                                       # the adjustments are reversed
    assert bal(w.a, w.cement)[0] == D("1000") - 1 and bal(w.a, w.steel)[0] == D("500")


def test_cancelling_a_draft_count_unfreezes_the_store(account):
    o = account["client"]
    w = seed(account)
    c = o.post(C + "/counts", json={"store_id": w.a, "item_ids": [w.cement]}).json()
    assert len(c["lines"]) == 1
    go(o, "counts", c["id"], "cancel")
    with database.SessionLocal() as s:
        post_movement(s, movement_type=MovementType.ISSUE_OUT, store_id=w.a, item_id=w.cement, qty=D(1), rate=None, txn_date=date.today(), project_id=None,
                      ref_doc_type="ISSUE", ref_doc_id=11, ref_line_id=1, user_id=None)
        s.commit()


def test_count_finding_unrecorded_item_and_no_difference(account):
    o = account["client"]
    w = seed(account)
    c = o.post(C + "/counts", json={"store_id": w.b, "item_ids": [w.cement]}).json()                         # B holds none
    o.put("%s/counts/%d" % (C, c["id"]), json={"lines": [{"item_id": w.cement, "counted_qty": "0"}]})
    go(o, "counts", c["id"], "submit", "approve", "post")                                                    # nothing differs: nothing posts
    assert bal(w.b, w.cement) == (0, 0) and o.get("%s/counts/%d" % (C, c["id"])).json()["status"] == "POSTED"


def test_adjustments_stock_wastage_and_high_value(account, portal):
    o = account["client"]
    w = seed(account)
    base = {"store_id": w.a, "reason": "Damaged by rain", "reason_code": "LOST", "attachment_ref": "SURVEY-12", "lines": [{"item_id": w.cement, "direction": "OUT", "qty": "40"}]}
    assert o.post(C + "/adjustments", json={**base, "reason": ""}).status_code == 400                      # the remark is mandatory
    assert o.post(C + "/adjustments", json={**base, "reason_code": ""}).status_code == 400                  # so is the reason code ...
    assert o.post(C + "/adjustments", json={**base, "reason_code": "MISPLACED"}).status_code == 400         # ... one of the known codes
    assert o.post(C + "/adjustments", json={**base, "reason_code": "CEMENT_SET"}).status_code == 400        # a wastage code does not fit a stock correction
    assert o.post(C + "/adjustments", json={**base, "attachment_ref": ""}).status_code == 400               # and the supporting paper
    a = o.post(C + "/adjustments", json=base).json()
    assert a["number"].startswith("ADJ/") and a["kind"] == "STOCK"
    done = go(o, "adjustments", a["id"], "submit", "approve", "post")
    assert done["total_value"] == "280.00" and bal(w.a, w.cement) == (D("960"), D("6720"))
    up = o.post(C + "/adjustments", json={"store_id": w.a, "reason": "Found in yard", "reason_code": "FOUND", "attachment_ref": "PHOTO-3", "lines": [{"item_id": w.cement, "direction": "IN", "qty": "10", "rate": "8"}]}).json()
    go(o, "adjustments", up["id"], "submit", "approve", "post")
    assert bal(w.a, w.cement) == (D("970"), D("6800"))
    assert go(o, "adjustments", up["id"], "cancel")["status"] == "CANCELLED" and bal(w.a, w.cement)[0] == D("960")

    assert o.post(C + "/adjustments", json={"store_id": w.a, "reason": "bad", "reason_code": "LOST", "attachment_ref": "X-1", "lines": [{"item_id": w.cement, "direction": "SIDEWAYS", "qty": "1"}]}).status_code == 400
    plain = o.post("/api/store/masters/stores", json={"code": "NP" + uuid.uuid4().hex[:3].upper(), "name": "No project"}).json()["id"]
    assert o.post(C + "/adjustments", json={"store_id": plain, "kind": "WASTAGE", "reason": "spill", "reason_code": "BREAKAGE", "attachment_ref": "X-2", "lines": [{"item_id": w.cement, "qty": "1"}]}).status_code == 400   # charged to nothing
    wa = o.post(C + "/adjustments", json={"store_id": w.a, "kind": "WASTAGE", "reason": "Spillage", "reason_code": "CEMENT_SET", "attachment_ref": "WASTE-4", "cost_head": "OVERHEAD", "lines": [{"item_id": w.cement, "qty": "5"}]}).json()
    go(o, "adjustments", wa["id"], "submit", "approve", "post")
    with database.SessionLocal() as s:
        e = s.query(DBStoreStockLedger).filter(DBStoreStockLedger.ref_doc_type == "ADJUST", DBStoreStockLedger.ref_doc_id == wa["id"]).one()
        assert (e.movement_type, e.cost_head, e.project_id) == ("WASTAGE_OUT", "OVERHEAD", w.ja)

    with database.SessionLocal() as s:                                                                      # 20,000 kg of steel at 62 is Rs 12.4 lakh
        post_movement(s, movement_type=MovementType.OPENING_IN, store_id=w.a, item_id=w.steel, qty=D(20000), rate=D(62), txn_date=date.today(), project_id=None,
                      ref_doc_type="OPENING", ref_doc_id=2, ref_line_id=None, user_id=None)
        s.commit()
    hv = o.post(C + "/adjustments", json={"store_id": w.a, "reason": "Theft", "reason_code": "THEFT", "attachment_ref": "FIR-9", "lines": [{"item_id": w.steel, "direction": "OUT", "qty": "5000"}]}).json()
    go(o, "adjustments", hv["id"], "submit")
    _login(account, portal, VIEW + ["store.control.approve"], [w.a])
    assert portal.post("%s/adjustments/%d/approve" % (C, hv["id"])).status_code == 403
    _login(account, portal, VIEW + ["store.control.approve", "store.control.approve_high_value"], [w.a])
    assert portal.post("%s/adjustments/%d/approve" % (C, hv["id"])).json()["high_value"] is True


def test_writeoff_adjustment_closes_outward_requests(account):
    o = account["client"]
    w = seed(account)
    with database.SessionLocal() as s:
        wr = DBStoreWriteoffRequest(client_id=w.cid, store_id=w.a, item_id=w.steel, qty=D(3), reason="bent beyond use", status="OPEN")
        s.add(wr)
        s.commit()
        rid = wr.id
    assert o.get(C + "/writeoff-requests").json()["rows"][0]["request_id"] == rid
    assert o.post(C + "/adjustments", json={"store_id": w.a, "kind": "WRITEOFF", "reason": "Lost", "reason_code": "PILFERAGE", "attachment_ref": "R-1", "lines": [{"writeoff_request_id": 999999}]}).status_code == 400
    a = o.post(C + "/adjustments", json={"store_id": w.a, "kind": "WRITEOFF", "reason": "Lost on site", "reason_code": "PILFERAGE", "attachment_ref": "R-2", "lines": [{"writeoff_request_id": rid}]}).json()
    assert a["lines"][0]["qty"] == "3.000"
    before = bal(w.a, w.steel)
    done = go(o, "adjustments", a["id"], "submit", "approve", "post")
    assert bal(w.a, w.steel) == before and done["lines"][0]["value"] == "186.00"                             # no stock moves; valued at the average
    with database.SessionLocal() as s:
        assert list_writeoff_requests(s) == [] and list_writeoff_requests(s, "CLOSED")[0].request_id == rid
    assert o.post("%s/adjustments/%d/cancel" % (C, a["id"]), json={"reason": "x"}).status_code == 409


def test_scrap_records_book_value_and_realisation(account):
    o = account["client"]
    w = seed(account)
    assert o.post(C + "/scrap", json={"store_id": w.a, "lines": [{"item_id": w.steel, "qty": "10"}]}).status_code == 400          # a reason
    assert o.post(C + "/scrap", json={"store_id": w.a, "reason": "x", "lines": [{"item_id": w.steel, "qty": "10", "scrap_class": "GOLD"}]}).status_code == 400   # a known class
    s_ = o.post(C + "/scrap", json={"store_id": w.a, "reason": "Cut-offs", "lines": [{"item_id": w.steel, "qty": "50", "realised_value": "1500", "scrap_class": "STEEL_CUT_PIECES"}]}).json()
    assert s_["number"].startswith("SCR/") and s_["project_id"] == w.ja and s_["lines"][0]["scrap_class"] == "STEEL_CUT_PIECES"
    assert o.post("%s/scrap/%d/submit" % (C, s_["id"])).status_code == 400                                                         # no buyer yet
    o.put("%s/scrap/%d" % (C, s_["id"]), json={"buyer": "Kabadi Traders"})
    assert o.post("%s/scrap/%d/submit" % (C, s_["id"])).status_code == 400                                                         # no gate pass yet
    o.put("%s/scrap/%d" % (C, s_["id"]), json={"gate_pass_no": "GP-2041"})
    done = go(o, "scrap", s_["id"], "submit", "approve", "post")
    assert done["book_value"] == "3100.00" and done["realised_value"] == "1500.00" and done["difference"] == "-1600.00" and done["gate_pass_no"] == "GP-2041"
    assert done["lines"][0]["difference"] == "-1600.00" and bal(w.a, w.steel)[0] == D("450")
    with database.SessionLocal() as s:
        rows = get_scrap_realisation(s, w.ja, date.today(), date.today())
        assert len(rows) == 1 and (rows[0].qty, rows[0].book_value, rows[0].realised_value) == (D("50"), D("3100.00"), D("1500.00"))
        assert get_scrap_realisation(s, w.jb, None, None) == []
    assert o.get("%s/scrap/%d/pdf" % (C, s_["id"])).status_code == 200
    assert go(o, "scrap", s_["id"], "cancel")["status"] == "CANCELLED" and bal(w.a, w.steel)[0] == D("500")


def test_lifecycle_rights_and_wrong_store(account, portal):
    o = account["client"]
    w = seed(account)
    t = trf(o, w)
    for step in ("approve", "post"):
        assert o.post("%s/transfers/%d/%s" % (C, t["id"], step)).status_code == 409
    assert o.put("%s/transfers/%d" % (C, t["id"]), json={"vehicle_no": "KA02"}).json()["vehicle_no"] == "KA02"
    go(o, "transfers", t["id"], "submit")
    assert o.put("%s/transfers/%d" % (C, t["id"]), json={"vehicle_no": "x"}).status_code == 409
    assert o.get(C + "/transfers", params={"status": "BOGUS"}).status_code == 400
    assert o.post(C + "/transfers/%d/bogus" % t["id"]).status_code == 404

    _login(account, portal, VIEW + ["store.control.create"], [w.b])                                          # assigned to B only: A's documents are out of reach
    assert portal.post(C + "/adjustments", json={"store_id": w.a, "reason": "x", "reason_code": "LOST", "attachment_ref": "X", "lines": [{"item_id": w.cement, "direction": "OUT", "qty": "1"}]}).status_code == 403
    assert portal.post(C + "/counts", json={"store_id": w.a}).status_code == 403
    assert portal.post("%s/transfers/%d/cancel" % (C, t["id"]), json={"reason": "x"}).status_code == 403    # no cancel right
    assert portal.get(C + "/scrap").json()["total"] == 0

    _login(account, portal, ["store.inward.view"], [w.a])
    assert portal.get(C + "/transfers").status_code == 403


def test_count_can_be_limited_to_a_category_and_a_bin(account):
    o = account["client"]
    w = seed(account)
    with database.SessionLocal() as s:
        s.query(DBStoreItemSettings).filter(DBStoreItemSettings.item_id == w.steel).update({"category": "TOOL"})
        s.commit()
    assert o.post(C + "/counts", json={"store_id": w.a, "categories": ["WIDGETS"]}).status_code == 400
    c = o.post(C + "/counts", json={"store_id": w.a, "categories": ["TOOL"]}).json()
    assert [l["item_id"] for l in c["lines"]] == [w.steel] and c["categories"] == ["TOOL"]
    go(o, "counts", c["id"], "cancel")
    foreign = o.post("/api/store/masters/stores/%d/bins" % w.b, json={"code": "Z-1"}).json()["id"]
    assert o.post(C + "/counts", json={"store_id": w.a, "bin_id": foreign}).status_code == 400                  # a bin of another store
    mine = o.post("/api/store/masters/stores/%d/bins" % w.a, json={"code": "A-7"}).json()["id"]
    c2 = o.post(C + "/counts", json={"store_id": w.a, "bin_id": mine, "item_ids": [w.cement]}).json()
    assert c2["bin_code"] == "A-7" and len(c2["lines"]) == 1
    assert o.post(C + "/counts", json={"store_id": w.a, "categories": ["CONSUMABLE"]}).status_code == 400    # the store is frozen by the count above


def test_recount_keeps_the_first_count_and_returns_the_sheet_to_the_counters(account):
    o = account["client"]
    w = seed(account)
    c = o.post(C + "/counts", json={"store_id": w.a}).json()
    ids = {l["item_id"]: l["id"] for l in c["lines"]}
    o.put("%s/counts/%d" % (C, c["id"]), json={"lines": [{"item_id": w.cement, "counted_qty": "700"}, {"item_id": w.steel, "counted_qty": "500"}]})
    assert o.post("%s/counts/%d/recount" % (C, c["id"]), json={"line_ids": [ids[w.cement]], "reason": "far too low"}).status_code == 409        # not submitted yet
    go(o, "counts", c["id"], "submit")
    assert o.post("%s/counts/%d/recount" % (C, c["id"]), json={"line_ids": [], "reason": "x"}).status_code == 400
    assert o.post("%s/counts/%d/recount" % (C, c["id"]), json={"line_ids": [ids[w.cement]]}).status_code == 400                               # a reason
    assert o.post("%s/counts/%d/recount" % (C, c["id"]), json={"line_ids": [999999], "reason": "x"}).status_code == 400
    back = o.post("%s/counts/%d/recount" % (C, c["id"]), json={"line_ids": [ids[w.cement]], "reason": "far too low"}).json()
    assert back["status"] == "DRAFT" and back["recount_rounds"] == 1
    cem = [l for l in back["lines"] if l["item_id"] == w.cement][0]
    assert (cem["counted_qty"], cem["first_count_qty"], cem["recount"]) == (None, "700.000", True)
    assert o.post("%s/counts/%d/submit" % (C, c["id"])).status_code == 400                                       # the recount is still to be done
    o.put("%s/counts/%d" % (C, c["id"]), json={"lines": [{"item_id": w.cement, "counted_qty": "995"}]})
    go(o, "counts", c["id"], "submit", "approve", "post")
    assert bal(w.a, w.cement)[0] == D("995")                                                                     # the recount, not the first count, was posted
    final = o.get("%s/counts/%d" % (C, c["id"])).json()
    assert [l for l in final["lines"] if l["item_id"] == w.cement][0]["first_count_qty"] == "700.000" and final["lines"][0]["recount"] is False
    var = o.get("%s/counts/%d/pdf" % (C, c["id"]), params={"form": "variance"})
    assert var.status_code == 200 and var.content[:4] == b"%PDF"
    assert o.get("%s/counts/%d/pdf" % (C, c["id"]), params={"form": "nope"}).status_code == 400


def test_the_variance_statement_waits_for_approval_for_a_blind_counter(account, portal):
    o = account["client"]
    w = seed(account)
    c = o.post(C + "/counts", json={"store_id": w.a}).json()
    _login(account, portal, VIEW, [w.a])
    assert portal.get("%s/counts/%d/pdf" % (C, c["id"]), params={"form": "variance"}).status_code == 403
    assert portal.get("%s/counts/%d/pdf" % (C, c["id"])).status_code == 200                                        # the blind sheet is fine
    _login(account, portal, VIEW + ["store.control.reveal_count"], [w.a])
    assert portal.get("%s/counts/%d/pdf" % (C, c["id"]), params={"form": "variance"}).status_code == 200


def test_write_off_note_and_voucher_forms_and_meta(account):
    o = account["client"]
    w = seed(account)
    wa = o.post(C + "/adjustments", json={"store_id": w.a, "kind": "WASTAGE", "reason": "Rusted", "reason_code": "RUSTING", "attachment_ref": "W-1", "cost_head": "OVERHEAD",
                                           "lines": [{"item_id": w.steel, "qty": "5"}]}).json()
    assert wa["reason_code"] == "RUSTING" and wa["attachment_ref"] == "W-1"
    for form in ("", "writeoff", "voucher"):
        r = o.get("%s/adjustments/%d/pdf" % (C, wa["id"]), params={"form": form} if form else None)
        assert r.status_code == 200 and r.content[:4] == b"%PDF", form
    assert o.get("%s/adjustments/%d/pdf" % (C, wa["id"]), params={"form": "x"}).status_code == 400
    meta = o.get(C + "/meta").json()
    assert meta["adjustment_reasons"] == ["DATA_CORRECTION", "FOUND", "LOST", "THEFT", "CONVERSION"]
    assert meta["wastage_reasons"] == ["CEMENT_SET", "BREAKAGE", "RUSTING", "EXPIRY", "PILFERAGE"]
    assert meta["scrap_classes"] == ["STEEL_CUT_PIECES", "EMPTY_CEMENT_BAGS", "DAMAGED_SHUTTERING", "OTHER"]


def test_reorder_alerts_subtract_what_is_already_on_order_and_export(account):
    o = account["client"]
    w = seed(account)
    tag = uuid.uuid4().hex[:6]
    with database.SessionLocal() as s:
        s.query(DBStoreItemSettings).filter(DBStoreItemSettings.item_id == w.cement).update({"reorder_level": 1200, "min_qty": 300, "max_qty": 5000})
        po = models.DBPurchaseOrder(client_id=w.cid, job_id=None, number="PO-" + tag, supplier_name="Shree", status="Approved", issue_date=str(date.today()))
        s.add(po)
        s.flush()
        cem = s.get(models.DBItem, w.cement)
        s.add(models.DBPurchaseOrderLineItem(order_id=po.id, description="Cement", item_code=cem.item_code, uom="kg", qty=1500, price=7.0))
        s.commit()
        po_id = po.id
    rows = o.get(C + "/reorder-alerts").json()["rows"]
    assert len(rows) == 1 and rows[0]["item_id"] == w.cement and rows[0]["store_id"] == w.a
    r = rows[0]
    assert (r["on_hand"], r["pending_po"], r["suggested"], r["below_min"]) == ("1000.000", "1500.000", "2500.000", False)      # 5,000 - (1,000 + 1,500)
    assert o.get(C + "/reorder-alerts", params={"store_id": w.b}).json()["rows"] == []                                              # the site holds none
    x = o.get(C + "/reorder-alerts", params={"format": "xlsx"})
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    from openpyxl import load_workbook
    import io
    ws = load_workbook(io.BytesIO(x.content)).active
    assert [c.value for c in ws[4]][-1] == "Suggested order qty" and ws["J5"].value == 2500 and ws["I5"].value == 1500
    assert o.get(C + "/reorder-alerts", params={"format": "docx"}).status_code == 400
    with database.SessionLocal() as s:                                                                                               # nothing was ordered
        assert s.query(models.DBPurchaseOrder).filter(models.DBPurchaseOrder.client_id == w.cid).count() == 1
        s.query(models.DBPurchaseOrderLineItem).filter(models.DBPurchaseOrderLineItem.order_id == po_id).delete()
        s.query(models.DBPurchaseOrder).filter(models.DBPurchaseOrder.id == po_id).delete()
        s.commit()
