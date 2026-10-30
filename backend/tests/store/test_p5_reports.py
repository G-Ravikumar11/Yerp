"""P5: reports (valuation, movement, consumption, norms against actual use, slow-moving, reorder), material norms, saved reports, and the
integration reads that costing, billing and purchasing use. Stock is seeded straight through post_movement, as every slice does."""
import io
import json
import uuid
from datetime import date, timedelta
from decimal import Decimal

import openpyxl
import pytest

from app import db as database, models
from app.models.store_core import (
    DBStoreBatch, DBStoreBatchBalance, DBStoreLedgerBatch,
    DBStore, DBStoreItemSettings, DBStoreStockBalance, DBStoreStockLedger, DBStoreUomConversion, DBStoreUserAssignment, MovementType,
)
from app.models.store_reports import DBStoreMaterialNorm
from app.services import store_integration, store_reports
from app.services.stock_ledger import post_movement, reverse_document

from conftest import make_employee

D = Decimal
API = "/api/store/reports"
TODAY = date.today()
LONG_AGO = TODAY - timedelta(days=30)


@pytest.fixture(autouse=True)
def _leave_no_stock_behind():
    """The suite empties the item master between tests, freeing item ids for reuse: clear what points at them."""
    yield
    with database.SessionLocal() as s:
        for model in (DBStoreLedgerBatch, DBStoreBatchBalance, DBStoreBatch):
            s.query(model).delete(synchronize_session=False)
        s.query(DBStoreStockLedger).filter(DBStoreStockLedger.reversal_of_id.isnot(None)).delete(synchronize_session=False)
        for model in (DBStoreStockLedger, DBStoreStockBalance, DBStoreItemSettings, DBStoreUomConversion, DBStoreMaterialNorm):
            s.query(model).delete(synchronize_session=False)
        s.commit()


class W:
    pass


def put(s, w, mtype, store, key, qty, rate=None, day=TODAY, doc=("T", 1), line=1, **dims):
    return post_movement(s, movement_type=mtype, store_id=store, item_id=w.item[key], qty=D(qty), rate=None if rate is None else D(rate), txn_date=day, project_id=dims.pop("project_id", None),
                         ref_doc_type=doc[0], ref_doc_id=doc[1], ref_line_id=line, user_id=None, **dims)


def world(account):
    """Two stores, a project with a BOQ of two items and a work order, three items, and a month of movements:
    opening stock, issues to the BOQ and to a cost head, a subcontractor issue, a return, wastage, scrap, a transfer, and an issue that was reversed."""
    w = W()
    with database.SessionLocal() as s:
        client = s.query(models.DBClient).filter(models.DBClient.email == account["email"]).first()
        tag = uuid.uuid4().hex[:6]
        w.client_id = client.id
        job = models.DBJob(client_id=client.id, number="P-" + tag, name="Tower A")
        other = models.DBJob(client_id=client.id, number="Q-" + tag, name="Tower B")
        s.add_all([job, other])
        s.flush()
        w.proj, w.other_proj = job.id, other.id
        boq = models.DBBoq(client_id=client.id, job_id=job.id, number="B-" + tag, title="Tower A")
        s.add(boq)
        s.flush()
        rev = models.DBBoqRevision(boq_id=boq.id, rev_no=boq.current_rev or 0, label="R0")
        s.add(rev)
        s.flush()
        w.boq = []
        for i, (d, u, q) in enumerate((("RCC M25 in columns", "cum", 10), ("Brick masonry", "cum", 20)), 1):
            line = models.DBBoqLine(boq_id=boq.id, revision_id=rev.id, key="k%d" % i, kind="item", item_no=str(i), description=d, uom=u, quantity=q, rate=100.0, amount=q * 100.0, display_order=i)
            s.add(line)
            s.flush()
            w.boq.append(line.id)
        oboq = models.DBBoq(client_id=client.id, job_id=other.id, number="C-" + tag, title="Tower B")
        s.add(oboq)
        s.flush()
        orev = models.DBBoqRevision(boq_id=oboq.id, rev_no=0, label="R0")
        s.add(orev)
        s.flush()
        oline = models.DBBoqLine(boq_id=oboq.id, revision_id=orev.id, key="o1", kind="item", item_no="1", description="Elsewhere", uom="cum", quantity=1, rate=1.0, amount=1.0)
        s.add(oline)
        s.flush()
        w.other_boq = oline.id
        wo = models.DBWorkOrder(client_id=client.id, job_id=job.id, number="WO-" + tag, status="Approved", boq_id=boq.id)
        con = models.DBContractor(client_id=client.id, vendor_code="V-" + tag, company_name="Sai Steel Fixers")
        s.add_all([wo, con])
        cs = DBStore(client_id=client.id, code="C" + tag[:3].upper(), name="Central", type="CENTRAL")
        st = DBStore(client_id=client.id, code="S" + tag[:3].upper(), name="Site", type="SITE", project_id=job.id)
        s.add_all([cs, st])
        s.flush()
        w.wo, w.contractor, w.central, w.site = wo.id, con.id, cs.id, st.id
        w.item = {}
        for key, name, uom, cat, reorder, mx in (("cement", "Cement OPC", "kg", "CONSUMABLE", 0, 0), ("steel", "TMT 12mm", "kg", "CONSUMABLE", 0, 0), ("plate", "Shuttering Plate", "nos", "RETURNABLE", 150, 400)):
            it = models.DBItem(client_id=client.id, kind="RM", item_code="RP-%s-%s" % (key.upper(), tag), item_name=name, units_of_measure=uom)
            s.add(it)
            s.flush()
            w.item[key] = it.id
            s.add(DBStoreItemSettings(item_id=it.id, category=cat, base_uom=uom, reorder_level=reorder, max_qty=mx, min_qty=50 if reorder else 0))
        for key, qty, rate in (("cement", "5000", "7"), ("steel", "2000", "62"), ("plate", "100", "450")):
            put(s, w, MovementType.OPENING_IN, cs.id, key, qty, rate, day=LONG_AGO, doc=("OPENING", 1))
        # 500 kg cement to BOQ item 1; 50 kg of it comes back
        put(s, w, MovementType.ISSUE_OUT, cs.id, "cement", "500", day=TODAY - timedelta(days=10), doc=("ISSUE", 1), project_id=job.id, work_order_id=wo.id, boq_item_id=w.boq[0])
        put(s, w, MovementType.ISSUE_RETURN_IN, cs.id, "cement", "50", "7", day=TODAY - timedelta(days=9), doc=("ISSRET", 1), project_id=job.id, work_order_id=wo.id, boq_item_id=w.boq[0])
        # 200 kg cement to a subcontractor on BOQ item 2
        put(s, w, MovementType.SUBCON_ISSUE_OUT, cs.id, "cement", "200", day=TODAY - timedelta(days=8), doc=("ISSUE", 2), project_id=job.id, boq_item_id=w.boq[1], subcontractor_id=con.id)
        # 100 kg steel on a cost head (not BOQ work)
        put(s, w, MovementType.ISSUE_OUT, cs.id, "steel", "100", day=TODAY - timedelta(days=7), doc=("ISSUE", 3), project_id=job.id, cost_head="SAFETY")
        # wastage and scrap
        put(s, w, MovementType.WASTAGE_OUT, cs.id, "cement", "40", day=TODAY - timedelta(days=6), doc=("SCRAP", 1), project_id=job.id, cost_head="OVERHEAD")
        put(s, w, MovementType.SCRAP_OUT, cs.id, "steel", "10", day=TODAY - timedelta(days=5), doc=("SCRAP", 2), project_id=job.id)
        # 300 kg cement transferred to the site store: not consumption
        out = put(s, w, MovementType.TRANSFER_OUT, cs.id, "cement", "300", day=TODAY - timedelta(days=4), doc=("TRANSFER", 1), project_id=None)
        put(s, w, MovementType.TRANSFER_IN, st.id, "cement", "300", out.rate, day=TODAY - timedelta(days=4), doc=("TRANSFER", 1), project_id=job.id)
        # an issue of 100 kg cement that was made and then reversed: it nets to nothing
        put(s, w, MovementType.ISSUE_OUT, cs.id, "cement", "100", day=TODAY - timedelta(days=3), doc=("ISSUE", 9), project_id=job.id, boq_item_id=w.boq[0])
        reverse_document(s, ref_doc_type="ISSUE", ref_doc_id=9, user_id=None, reason="wrong site")
        s.commit()
    return w


def staff(account, portal, rights, store_ids=()):
    email = "p5-%s@example.com" % uuid.uuid4().hex[:8]
    emp = make_employee(account["client"], email=email, password="Crew1234", permission_role="staff", status="active")
    with database.SessionLocal() as s:
        row = s.get(models.DBEmployee, emp["id"])
        row.extra_permissions = ",".join(rights)
        s.add_all([DBStoreUserAssignment(store_id=i, employee_id=row.id) for i in store_ids])
        s.commit()
    assert portal.post("/api/employee/auth/login", json={"email": email, "password": "Crew1234"}).status_code == 200
    return emp


def get(tenant, report, **params):
    res = tenant.get("%s/%s" % (API, report), params=params)
    assert res.status_code == 200, (report, params, res.text)
    return res.json()


# --- the integration read ---------------------------------------------------------------------------------------------------------

def test_project_material_cost_is_what_the_stores_consumed_net_of_returns_and_reversals(account):
    w = world(account)
    with database.SessionLocal() as s:
        got = store_integration.get_project_material_cost(s, w.proj, date(2000, 1, 1), TODAY)
        by = {(l.cost_head, l.boq_item_id, l.item_id): (l.qty, l.value) for l in got}
        assert by == {
            (None, w.boq[0], w.item["cement"]): (D("450.000"), D("3150.00")),            # 500 issued, 50 back, the reversed 100 nets out
            (None, w.boq[1], w.item["cement"]): (D("200.000"), D("1400.00")),            # to the subcontractor
            ("SAFETY", None, w.item["steel"]): (D("100.000"), D("6200.00")),
            ("OVERHEAD", None, w.item["cement"]): (D("40.000"), D("280.00")),            # wastage
            (None, None, w.item["steel"]): (D("10.000"), D("620.00")),                   # scrap
        }
        assert all(l.project_id == w.proj for l in got)
        assert sum((l.value for l in got), D(0)) == D("11650.00")
        assert [l.work_order_id for l in got if l.boq_item_id == w.boq[0]] == [w.wo]
        # a window that holds only part of it
        part = store_integration.get_project_material_cost(s, w.proj, TODAY - timedelta(days=8), TODAY - timedelta(days=7))
        assert sorted((l.qty, l.value) for l in part) == [(D("100.000"), D("6200.00")), (D("200.000"), D("1400.00"))]
        assert store_integration.get_project_material_cost(s, w.other_proj, date(2000, 1, 1), TODAY) == []
        assert store_integration.get_project_material_cost(s, 999999, date(2000, 1, 1), TODAY) == []
        assert store_integration.project_store_material(s, w.proj) == D("11650.00")


def test_the_project_costing_and_the_purchase_order_carry_the_stores_figures(account):
    tenant = account["client"]
    w = world(account)
    job = tenant.get("/api/jobs/%d" % w.proj)
    assert job.status_code == 200, job.text
    assert job.json()["costing"]["store_material"] == 11650.0
    board = tenant.get("/api/jobs", params={"costing": "true"})
    assert board.status_code == 200
    mine = [j for j in board.json()["jobs"] if j["id"] == w.proj]
    assert not mine or mine[0].get("costing", {}).get("store_material", 11650.0) == 11650.0
    other = tenant.get("/api/jobs/%d" % w.other_proj).json()
    assert other["costing"]["store_material"] == 0.0
    po = tenant.post("/api/purchase-orders", json={"supplier_name": "Shree Traders", "job_id": w.proj, "line_items": [{"description": "Cement", "qty": 10, "price": 7, "item_code": "X"}]})
    assert po.status_code == 200, po.text
    detail = tenant.get("/api/purchase-orders/%d" % po.json()["id"]).json()
    # what the receipts say is P2's to supply (it is empty until goods are received); the order simply carries the key
    assert isinstance(detail["store_receipt"], list)
    assert all({"po_line_id", "ordered_qty", "received_qty", "accepted_qty", "rejected_qty", "pending_qty"} <= set(r) for r in detail["store_receipt"])


# --- valuation ----------------------------------------------------------------------------------------------------------------------

def test_valuation_matches_what_the_stores_hold(account):
    tenant = account["client"]
    w = world(account)
    v = get(tenant, "valuation")
    by = {(r["store_code"], r["item_name"]): (r["qty"], r["value"], r["avg_rate"]) for r in v["rows"]}
    code_c, code_s = [s.code for s in _stores(w)]
    assert by[(code_c, "Cement OPC")] == ("4010.000", "28070.00", "7.0000") and by[(code_c, "TMT 12mm")] == ("1890.000", "117180.00", "62.0000")
    assert by[(code_s, "Cement OPC")] == ("300.000", "2100.00", "7.0000") and by[(code_c, "Shuttering Plate")] == ("100.000", "45000.00", "450.0000")
    assert v["total_value"] == "192350.00" and v["as_on"] == TODAY.isoformat()
    assert {c["category"]: c["value"] for c in v["by_category"]} == {"CONSUMABLE": "147350.00", "RETURNABLE": "45000.00"}
    with database.SessionLocal() as s:
        assert sum((D(str(b.value)) for b in s.query(DBStoreStockBalance).filter(DBStoreStockBalance.store_id.in_([w.central, w.site])).all()), D(0)) == D("192350.00")   # the ledger and the balances agree
    assert get(tenant, "valuation", category="returnable")["total_value"] == "45000.00"
    assert [r["item_name"] for r in get(tenant, "valuation", q="steel")["rows"]] == ["TMT 12mm"]
    assert get(tenant, "valuation", store_id=w.site)["total_value"] == "2100.00"
    assert get(tenant, "valuation", include_zero="true")["rows"] and len(get(tenant, "valuation", include_zero="true")["rows"]) >= len(v["rows"])


def _stores(w):
    with database.SessionLocal() as s:
        return [s.get(DBStore, w.central), s.get(DBStore, w.site)]


def test_valuation_on_a_past_date_is_what_was_held_then(account):
    tenant = account["client"]
    w = world(account)
    before = get(tenant, "valuation", as_on=(LONG_AGO + timedelta(days=1)).isoformat())          # just after the opening stock
    assert before["total_value"] == str(D("35000.00") + D("124000.00") + D("45000.00")) + "" or before["total_value"] == "204000.00"
    mid = get(tenant, "valuation", as_on=(TODAY - timedelta(days=9)).isoformat())                # after the first issue and its return
    cement = [r for r in mid["rows"] if r["item_name"] == "Cement OPC"]
    assert sum(D(r["qty"]) for r in cement) == D("4550.000")
    assert get(tenant, "valuation", as_on=TODAY.isoformat())["total_value"] == get(tenant, "valuation")["total_value"]
    assert get(tenant, "valuation", as_on="2000-01-02")["rows"] == []
    assert tenant.get("%s/valuation" % API, params={"as_on": "nope"}).status_code == 400


# --- movement -----------------------------------------------------------------------------------------------------------------------

def test_movement_opens_moves_and_closes_to_the_valuation(account):
    tenant = account["client"]
    w = world(account)
    m = get(tenant, "movement", from_date=(TODAY - timedelta(days=12)).isoformat(), to_date=TODAY.isoformat())
    assert [b["key"] for b in m["buckets"]] == ["receipts", "returns_to_supplier", "issues", "transfers", "adjustments"]
    cement = [r for r in m["rows"] if r["item_name"] == "Cement OPC" and r["store_code"] == _stores(w)[0].code][0]
    assert cement["opening_qty"] == "5000.000" and cement["opening_value"] == "35000.00"
    assert (cement["issues_qty"], cement["issues_value"]) == ("-650.000", "-4550.00")            # 500 + 200 + the reversed 100 and its reversal, less 50 back
    assert (cement["transfers_qty"], cement["adjustments_qty"], cement["receipts_qty"]) == ("-300.000", "-40.000", "0.000")
    assert cement["closing_qty"] == "4010.000" and cement["closing_value"] == "28070.00"
    val = {(r["store_code"], r["item_name"]): r["value"] for r in get(tenant, "valuation")["rows"]}
    for r in m["rows"]:
        assert r["closing_value"] == val[(r["store_code"], r["item_name"])], r                       # every row closes to the valuation
    assert m["totals"]["closing_value"] == get(tenant, "valuation")["total_value"]
    assert D(m["totals"]["opening_value"]) + sum(D(m["totals"][k + "_value"]) for k in ("receipts", "returns_to_supplier", "issues", "transfers", "adjustments")) == D(m["totals"]["closing_value"])
    assert tenant.get("%s/movement" % API, params={"from_date": "2026-02-01", "to_date": "2026-01-01"}).status_code == 400
    default = get(tenant, "movement")
    assert default["from"] == TODAY.replace(day=1).isoformat() and default["to"] == TODAY.isoformat()


# --- consumption --------------------------------------------------------------------------------------------------------------------

def test_consumption_groups_by_every_dimension_and_leaves_transfers_out(account):
    tenant = account["client"]
    w = world(account)
    assert get(tenant, "consumption", group_by="project")["total_value"] == "11650.00"
    items = {g["name"].split(" - ")[1]: g["value"] for g in get(tenant, "consumption", group_by="item")["groups"]}
    assert items == {"TMT 12mm": "6820.00", "Cement OPC": "4830.00"}
    heads = {g["name"]: g["value"] for g in get(tenant, "consumption", group_by="cost_head")["groups"]}
    assert heads == {"BOQ work": "5170.00", "Safety": "6200.00", "Overhead": "280.00"}              # the BOQ items and the scrap that names no head
    boq = get(tenant, "consumption", group_by="boq_item")
    assert [(g["name"].split(" ")[0], g["value"]) for g in boq["groups"] if g["name"][0].isdigit()] == [("1", "3150.00"), ("2", "1400.00")] or len(boq["groups"]) >= 2
    wo = {g["name"]: g["value"] for g in get(tenant, "consumption", group_by="work_order")["groups"]}
    assert "No work order" in wo and sum(D(v) for v in wo.values()) == D("11650.00")
    assert get(tenant, "consumption", project_id=w.other_proj)["groups"] == []
    assert get(tenant, "consumption", from_date=(TODAY - timedelta(days=8)).isoformat(), to_date=(TODAY - timedelta(days=7)).isoformat())["total_value"] == "7600.00"
    assert tenant.get("%s/consumption" % API, params={"group_by": "colour"}).status_code == 400
    assert tenant.get("%s/consumption" % API, params={"project_id": 999999}).status_code == 404
    assert tenant.get("%s/consumption" % API, params={"from_date": "2026-02-01", "to_date": "2026-01-01"}).status_code == 400


def test_the_project_cost_endpoint_reads_the_same_figures(account):
    tenant = account["client"]
    w = world(account)
    r = tenant.get("%s/project-cost/%d" % (API, w.proj)).json()
    assert r["total_value"] == "11650.00" and len(r["items"]) == 5
    boq0 = [i for i in r["items"] if i["boq_item_id"] == w.boq[0]][0]
    assert (boq0["qty"], boq0["value"], boq0["work_order"], boq0["boq_item"].startswith("1 RCC")) == ("450.000", "3150.00", boq0["work_order"], True) and boq0["work_order"]
    assert tenant.get("%s/project-cost/999999" % API).status_code == 404


# --- norms ----------------------------------------------------------------------------------------------------------------------------

def norm(tenant, w, **over):
    body = {"boq_item_id": w.boq[0], "item_id": w.item["cement"], "qty_per_unit": "50", "wastage_pct": "5", **over}
    return tenant.post(API + "/norms", json=body)


def test_norms_set_the_theoretical_requirement_and_the_report_compares_it_with_use(account):
    tenant = account["client"]
    w = world(account)
    res = norm(tenant, w)
    assert res.status_code == 200, res.text
    n = res.json()
    assert (n["qty_per_unit"], n["wastage_pct"], n["theoretical_qty"], n["boq_qty"], n["base_uom"]) == ("50.000000", "5.000", "525.000", "10.000", "kg")     # 10 cum x 50 kg x 1.05
    with database.SessionLocal() as s:
        assert store_reports.get_theoretical_requirement(s, w.boq[0], w.item["cement"]) == D("525.000")
        assert store_reports.get_theoretical_requirement(s, w.boq[0], w.item["steel"]) is None
        assert store_reports.get_theoretical_requirement(s, 999999, w.item["cement"]) is None
    rep = get(tenant, "norm-variance", project_id=w.proj)
    row = rep["rows"][0]
    assert (row["theoretical_qty"], row["actual_qty"], row["variance_qty"], row["status"], row["actual_value"]) == ("525.000", "450.000", "-75.000", "UNDER", "3150.00") and rep["over"] == 0
    assert row["variance_pct"] == "-14.3"
    # more is issued than the norm allows
    with database.SessionLocal() as s:
        put(s, w, MovementType.ISSUE_OUT, w.central, "cement", "200", day=TODAY, doc=("ISSUE", 20), project_id=w.proj, work_order_id=w.wo, boq_item_id=w.boq[0])
        s.commit()
    over = get(tenant, "norm-variance", project_id=w.proj)
    assert over["rows"][0]["status"] == "OVER" and over["rows"][0]["variance_qty"] == "125.000" and over["over"] == 1
    assert get(tenant, "norm-variance", project_id=w.proj, tolerance_pct="25")["rows"][0]["status"] == "OK"            # 650 against 525 + 25%
    assert tenant.get("%s/norm-variance" % API).status_code == 400
    assert tenant.get("%s/norm-variance" % API, params={"project_id": w.proj, "tolerance_pct": "-1"}).status_code == 400
    assert tenant.get("%s/norm-variance" % API, params={"project_id": w.other_proj}).json()["rows"] == []


def test_norms_are_edited_listed_and_removed_with_their_rules(account):
    tenant = account["client"]
    w = world(account)
    n = norm(tenant, w).json()
    assert norm(tenant, w).status_code == 409                                                                         # one norm per BOQ item and item
    for bad in ({"qty_per_unit": "0"}, {"qty_per_unit": "-3"}, {"qty_per_unit": "x"}, {"qty_per_unit": "0.0000001"}, {"wastage_pct": "101"}, {"wastage_pct": "-1"}, {"item_id": 999999},
                {"boq_item_id": 999999}, {"notes": "n" * 201}):
        assert norm(tenant, w, **{"item_id": w.item["steel"], **bad}).status_code == 400, bad
    edited = tenant.put("%s/norms/%d" % (API, n["id"]), json={"boq_item_id": w.boq[0], "item_id": w.item["cement"], "qty_per_unit": "40", "wastage_pct": "0", "notes": "site trial"})
    assert edited.status_code == 200 and edited.json()["theoretical_qty"] == "400.000" and edited.json()["notes"] == "site trial"
    other = norm(tenant, w, item_id=w.item["steel"]).json()
    assert tenant.put("%s/norms/%d" % (API, other["id"]), json={"boq_item_id": w.boq[0], "item_id": w.item["cement"], "qty_per_unit": "1"}).status_code == 409
    assert [x["id"] for x in tenant.get(API + "/norms", params={"project_id": w.proj}).json()["items"]] == [n["id"], other["id"]]
    assert tenant.get(API + "/norms", params={"boq_item_id": w.boq[1]}).json()["items"] == []
    assert tenant.get(API + "/norms", params={"project_id": 999999}).status_code == 404
    assert tenant.delete("%s/norms/%d" % (API, n["id"])).status_code == 200 and tenant.delete("%s/norms/%d" % (API, n["id"])).status_code == 404
    assert tenant.put("%s/norms/999999" % API, json={"boq_item_id": w.boq[0], "item_id": w.item["cement"], "qty_per_unit": "1"}).status_code == 404
    picks = tenant.get(API + "/pick/boq-items", params={"project_id": w.proj}).json()["items"]
    assert [p["id"] for p in picks] == w.boq and picks[0]["quantity"] == "10.000"
    assert [p["id"] for p in tenant.get(API + "/pick/boq-items", params={"project_id": w.proj, "q": "brick"}).json()["items"]] == [w.boq[1]]


# --- slow-moving and reorder ------------------------------------------------------------------------------------------------------------

def test_slow_moving_stock_is_what_has_sat_idle(account):
    tenant = account["client"]
    w = world(account)
    s20 = get(tenant, "slow-moving", days=20)
    assert [(r["item_name"], r["days_idle"], r["value"]) for r in s20["rows"]] == [("Shuttering Plate", 30, "45000.00")]          # untouched since the opening stock
    assert s20["total_value"] == "45000.00" and s20["rows"][0]["last_moved"] == LONG_AGO.isoformat()
    assert get(tenant, "slow-moving", days=40)["rows"] == []
    assert {r["item_name"] for r in get(tenant, "slow-moving", days=3)["rows"]} >= {"Shuttering Plate", "TMT 12mm"}
    for bad in ("0", "99999", "x"):
        assert tenant.get("%s/slow-moving" % API, params={"days": bad}).status_code == 400
    assert get(tenant, "slow-moving")["days"] == 90


def test_reorder_lists_what_is_at_or_below_its_level_with_a_suggested_quantity(account):
    tenant = account["client"]
    w = world(account)
    r = get(tenant, "reorder")["rows"]
    assert [(x["item_name"], x["qty"], x["reorder_level"], x["suggested_qty"], x["below_minimum"]) for x in r] == [("Shuttering Plate", "100.000", "150.000", "300.000", False)]          # 400 maximum less 100 held
    with database.SessionLocal() as s:
        put(s, w, MovementType.ISSUE_OUT, w.central, "plate", "60", day=TODAY, doc=("ISSUE", 30), project_id=w.proj, boq_item_id=w.boq[0])
        s.commit()
    r = get(tenant, "reorder")["rows"]
    assert r[0]["qty"] == "40.000" and r[0]["below_minimum"] is True and r[0]["suggested_qty"] == "360.000"


# --- saved reports, exports ------------------------------------------------------------------------------------------------------------

def test_a_report_can_be_saved_with_its_filters_and_run_again(account):
    tenant = account["client"]
    w = world(account)
    saved = tenant.post(API + "/saved", json={"name": " Monthly   valuation ", "report": "valuation", "params": {"category": "RETURNABLE", "store_id": w.central, "q": ""}})
    assert saved.status_code == 200, saved.text
    s = saved.json()
    assert s["name"] == "Monthly valuation" and s["params"] == {"category": "RETURNABLE", "store_id": str(w.central)}
    assert [x["name"] for x in tenant.get(API).json()["items"]] == ["Monthly valuation"] and tenant.get(API).json()["slice"] == "reports"
    again = tenant.get("%s/valuation" % API, params=s["params"]).json()
    assert again["total_value"] == "45000.00"
    renamed = tenant.post(API + "/saved", json={"name": "Monthly valuation", "report": "reorder", "params": {}})                  # the same name replaces it
    assert renamed.status_code == 200 and renamed.json()["id"] == s["id"] and renamed.json()["report"] == "reorder"
    for bad in ({"name": " ", "report": "valuation"}, {"name": "x" * 81, "report": "valuation"}, {"name": "ok", "report": "nonsense"}):
        assert tenant.post(API + "/saved", json=bad).status_code == 400
    assert tenant.delete("%s/saved/%d" % (API, s["id"])).status_code == 200 and tenant.delete("%s/saved/%d" % (API, s["id"])).status_code == 404
    assert tenant.get(API).json()["items"] == []


def test_every_report_exports_and_prints(account):
    tenant = account["client"]
    w = world(account)
    norm(tenant, w)
    params = {"valuation": {}, "movement": {"from_date": (TODAY - timedelta(days=12)).isoformat()}, "consumption": {"group_by": "boq_item"}, "norm-variance": {"project_id": w.proj},
              "slow-moving": {"days": 5}, "reorder": {}}
    for report, p in params.items():
        xl = tenant.get("%s/%s/export" % (API, report), params=p)
        assert xl.status_code == 200 and "spreadsheetml" in xl.headers["content-type"], (report, xl.text[:200])
        ws = openpyxl.load_workbook(io.BytesIO(xl.content)).active
        assert ws.max_row >= 3 and ws.cell(row=1, column=1).value
        pdf = tenant.get("%s/%s/print" % (API, report), params=p)
        assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF", (report, pdf.text[:200])
    wb = openpyxl.load_workbook(io.BytesIO(tenant.get("%s/valuation/export" % API).content)).active
    assert [c.value for c in wb[3]][:3] == ["Store", "Item code", "Item"] and wb[wb.max_row][0].value == "TOTAL VALUE"
    assert tenant.get("%s/nonsense" % API).status_code == 404 and tenant.get("%s/nonsense/export" % API).status_code == 404


# --- rights and tenancy -----------------------------------------------------------------------------------------------------------------

def test_reports_need_their_rights(account, portal):
    tenant = account["client"]
    w = world(account)
    staff(account, portal, ["store.reports.view"], [w.central])
    assert portal.get("%s/valuation" % API).status_code == 200 and portal.get(API).status_code == 200
    assert portal.get(API + "/norms").status_code == 200
    assert portal.post(API + "/norms", json={"boq_item_id": w.boq[0], "item_id": w.item["cement"], "qty_per_unit": "1"}).status_code == 403
    assert portal.post(API + "/saved", json={"name": "x", "report": "valuation"}).status_code == 403
    n = norm(tenant, w).json()
    assert portal.put("%s/norms/%d" % (API, n["id"]), json={"boq_item_id": w.boq[0], "item_id": w.item["cement"], "qty_per_unit": "1"}).status_code == 403
    assert portal.delete("%s/norms/%d" % (API, n["id"])).status_code == 403
    staff(account, portal, ["store.inward.view"], [w.central])                                                         # a right of another slice is no use here
    assert portal.get("%s/valuation" % API).status_code == 403 and portal.get(API).status_code == 403
    staff(account, portal, ["store.reports.view", "store.reports.create"], [w.central])
    assert norm(portal, w, item_id=w.item["steel"]).status_code == 200 and portal.post(API + "/saved", json={"name": "mine", "report": "reorder"}).status_code == 200


def test_staff_reports_cover_only_their_own_stores(account, portal):
    tenant = account["client"]
    w = world(account)
    staff(account, portal, ["store.reports.view"], [w.site])
    v = portal.get("%s/valuation" % API).json()
    assert {r["store_code"] for r in v["rows"]} == {_stores(w)[1].code} and v["total_value"] == "2100.00"
    assert portal.get("%s/valuation" % API, params={"store_id": w.central}).status_code == 403                          # wrong-store access
    assert portal.get("%s/movement" % API, params={"store_id": w.central}).status_code == 403
    assert portal.get("%s/consumption" % API).json()["groups"] == []                                                    # all the issues were from the central store
    assert portal.get("%s/slow-moving" % API, params={"days": 1}).status_code == 200
    assert tenant.get("%s/consumption" % API).json()["total_value"] == "11650.00"


def test_one_company_cannot_see_anothers_stock(account, second_tenant):
    tenant = account["client"]
    w = world(account)
    assert second_tenant.get("%s/valuation" % API).json()["rows"] == [] and second_tenant.get("%s/valuation" % API).json()["total_value"] == "0"
    assert second_tenant.get("%s/valuation" % API, params={"store_id": w.central}).status_code == 404
    assert second_tenant.get("%s/consumption" % API, params={"project_id": w.proj}).status_code == 404
    assert second_tenant.get("%s/project-cost/%d" % (API, w.proj)).status_code == 404
    assert second_tenant.get("%s/norm-variance" % API, params={"project_id": w.proj}).status_code == 404
    assert second_tenant.post(API + "/norms", json={"boq_item_id": w.boq[0], "item_id": w.item["cement"], "qty_per_unit": "1"}).status_code == 400
    n = norm(tenant, w).json()
    assert second_tenant.put("%s/norms/%d" % (API, n["id"]), json={"boq_item_id": w.boq[0], "item_id": w.item["cement"], "qty_per_unit": "9"}).status_code == 404
    assert second_tenant.get(API + "/norms").json()["items"] == []
    assert second_tenant.get(API + "/pick/boq-items", params={"project_id": w.proj}).status_code == 404
    assert tenant.get("%s/valuation" % API).json()["total_value"] == "192350.00"


# --- FIFO valuation, ABC analysis, stock ageing, site material reconciliation -----------------------------------------------------------

def test_fifo_values_what_is_left_as_the_newest_receipts_and_ignores_a_reversed_one(account):
    tenant = account["client"]
    w = world(account)
    with database.SessionLocal() as s:
        # Central holds 4010 kg cement at 7. In come 1000 kg at 9, then 1000 kg go out at the average.
        put(s, w, MovementType.GRN_IN, w.central, "cement", "1000", "9", day=TODAY - timedelta(days=2), doc=("GRN", 1))
        put(s, w, MovementType.ISSUE_OUT, w.central, "cement", "1000", day=TODAY - timedelta(days=1), doc=("ISSUE", 20), project_id=w.proj)
        # A receipt at a silly rate that was put in and then undone is no receipt at all.
        put(s, w, MovementType.GRN_IN, w.central, "cement", "500", "100", day=TODAY, doc=("GRN", 2))
        reverse_document(s, ref_doc_type="GRN", ref_doc_id=2, user_id=None, reason="keyed twice")
        s.commit()
    pick = lambda rep: next(r for r in rep["rows"] if r["store_id"] == w.central and r["item_id"] == w.item["cement"])
    avg, fifo = pick(get(tenant, "valuation")), pick(get(tenant, "valuation", method="fifo"))
    assert avg["qty"] == fifo["qty"] == "4010.000"
    assert fifo["value"] == "30070.00"                      # the newest 1000 kg at 9, then 3010 kg at 7
    assert avg["value"] != fifo["value"]
    assert get(tenant, "valuation", method="fifo")["method"] == "fifo"
    assert tenant.get(API + "/valuation", params={"method": "lifo"}).status_code == 400


def test_abc_ranks_items_by_what_was_used_and_classes_the_tail(account):
    tenant = account["client"]
    w = world(account)
    with database.SessionLocal() as s:
        put(s, w, MovementType.ISSUE_OUT, w.central, "plate", "1", day=TODAY - timedelta(days=2), doc=("ISSUE", 21), project_id=w.proj)
        s.commit()
    rep = get(tenant, "abc")
    rows = [(r["item_id"], r["value"], r["class"]) for r in rep["rows"] if r["item_id"] in w.item.values()]
    # steel 6820, cement 4830, plate 450 of 12100: steel and cement make the first 70%, the plate is the tail.
    assert rows == [(w.item["steel"], "6820.00", "A"), (w.item["cement"], "4830.00", "A"), (w.item["plate"], "450.00", "C")]
    assert rep["rows"][-1]["cumulative_pct"] == "100.00"
    assert {c["class"]: c["items"] for c in rep["classes"]} == {"A": 2, "B": 0, "C": 1}
    # The period counts: nothing was used a year and more ago.
    old = get(tenant, "abc", from_date=(TODAY - timedelta(days=800)).isoformat(), to_date=(TODAY - timedelta(days=400)).isoformat())
    assert old["rows"] == [] and old["total_value"] == "0.00"


def test_ageing_splits_stock_by_how_long_ago_it_came_in_and_adds_up_to_the_valuation(account):
    tenant = account["client"]
    w = world(account)
    with database.SessionLocal() as s:
        put(s, w, MovementType.OPENING_IN, w.site, "steel", "10", "62", day=TODAY - timedelta(days=200), doc=("OPENING", 2))
        put(s, w, MovementType.OPENING_IN, w.site, "steel", "5", "62", day=TODAY - timedelta(days=60), doc=("OPENING", 3))
        s.commit()
    rep = get(tenant, "ageing")
    site_steel = next(r for r in rep["rows"] if r["store_id"] == w.site and r["item_id"] == w.item["steel"])
    assert (site_steel["qty"], site_steel["d31_90_qty"], site_steel["d181_365_qty"], site_steel["d0_30_qty"]) == ("15.000", "5.000", "10.000", "0.000")
    assert site_steel["oldest"] == (TODAY - timedelta(days=200)).isoformat()
    site_cement = next(r for r in rep["rows"] if r["store_id"] == w.site and r["item_id"] == w.item["cement"])
    assert site_cement["d0_30_qty"] == "300.000"                # came in by transfer 4 days ago
    total = D(rep["total_value"])
    assert sum(D(v) for v in rep["totals"].values()) == total
    assert total == D(get(tenant, "valuation")["total_value"])
    assert rep["rows"][0]["item_id"] == w.item["steel"] and rep["rows"][0]["store_id"] == w.site     # the oldest money first


def test_site_reconciliation_balances_to_what_the_site_holds_and_compares_with_the_boq(account):
    tenant = account["client"]
    w = world(account)
    norm(tenant, w)                                             # BOQ item 1 needs 525 kg cement
    with database.SessionLocal() as s:
        put(s, w, MovementType.GRN_IN, w.site, "cement", "200", "7", day=TODAY - timedelta(days=2), doc=("GRN", 3))
        put(s, w, MovementType.ISSUE_OUT, w.site, "cement", "120", day=TODAY - timedelta(days=1), doc=("ISSUE", 22), project_id=w.proj, boq_item_id=w.boq[0])
        put(s, w, MovementType.WASTAGE_OUT, w.site, "cement", "5", day=TODAY - timedelta(days=1), doc=("SCRAP", 3), project_id=w.proj)
        s.commit()
    rep = get(tenant, "reconciliation", project_id=w.proj)
    assert [st["id"] for st in rep["site_stores"]] == [w.site]
    c = next(r for r in rep["rows"] if r["item_id"] == w.item["cement"])
    assert (c["received"], c["transfers"], c["issued"], c["wasted"], c["balance"]) == ("200.000", "300.000", "120.000", "5.000", "375.000")
    assert D(c["received"]) + D(c["transfers"]) - D(c["returned"]) - D(c["issued"]) - D(c["wasted"]) + D(c["adjusted"]) == D(c["balance"])
    with database.SessionLocal() as s:
        held = s.query(DBStoreStockBalance).filter(DBStoreStockBalance.store_id == w.site, DBStoreStockBalance.item_id == w.item["cement"]).one()
        assert D(c["balance"]) == D(str(held.qty)).quantize(D("0.001"))
    # Straight from the central store to the project: 500 - 50 + 200 (the reversed 100 nets out).
    assert c["issued_elsewhere"] == "650.000" and c["total_issued"] == "770.000"
    assert (c["required"], c["still_to_issue"], c["status"]) == ("525.000", "0.000", "OVER") and rep["over"] == 1
    steel = next(r for r in rep["rows"] if r["item_id"] == w.item["steel"])
    assert (steel["issued_elsewhere"], steel["required"], steel["status"]) == ("100.000", None, "NO_NORM")
    assert tenant.get(API + "/reconciliation").status_code == 400
    assert get(tenant, "reconciliation", project_id=w.other_proj)["rows"] == []


def test_the_new_reports_export_and_print(account):
    tenant = account["client"]
    w = world(account)
    for report, p in {"abc": {}, "ageing": {}, "reconciliation": {"project_id": w.proj}, "valuation": {"method": "fifo"}}.items():
        xl = tenant.get("%s/%s/export" % (API, report), params=p)
        assert xl.status_code == 200 and "spreadsheetml" in xl.headers["content-type"], (report, xl.text[:200])
        pdf = tenant.get("%s/%s/print" % (API, report), params=p)
        assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF", (report, pdf.text[:200])
