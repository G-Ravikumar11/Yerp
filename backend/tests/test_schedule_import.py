"""Importing a programme sheet into Planning & Progress: a block heading, its activities, each one's own stages, read
by their own indentation - the signal a real Excel file carries that a flattened PDF text dump does not.
"""
import io

import openpyxl
from openpyxl.styles import Alignment, Font

from test_jobs_and_orders import staff
from test_measurement_and_ra_bills import placed_order


def job(tenant, name="Nuzividu houses"):
    return tenant.post("/api/jobs", json={"name": name, "customer_name": "APTIDCO"}).json()


def sheet_bytes(rows):
    """rows: (sno, text, start, end, indent, bold)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Project - demo", None, None, None])
    ws.append(["SL.NO", "ITEMES / ACTIVITY", "Start", "End"])
    for sno, text, start, end, indent, bold in rows:
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=sno)
        c = ws.cell(row=r, column=2, value=text)
        c.alignment = Alignment(indent=indent)
        c.font = Font(bold=bold)
        ws.cell(row=r, column=3, value=start)
        ws.cell(row=r, column=4, value=end)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_sheet(tenant, job_id, rows):
    return tenant.post("/api/jobs/%d/schedule/import" % job_id, files={"file": ("sch.xlsx", sheet_bytes(rows))})


def confirm(tenant, job_id, rows):
    return tenant.post("/api/jobs/%d/schedule/import/confirm" % job_id, json={"rows": rows})


# --- the simple 2-level sheet: a block heading, then its activities side by side -------------------------------------

TWO_LEVEL = [
    ("1.00", "300 SFT Block Type - A", "14-Oct-26", "16-Feb-27", 0, True),
    ("1.01", "A-16 Third floor slab", "14-Oct-26", "12-Nov-26", 0, False),
    ("1.16", "A-16 Terrace floor slab", "10-Nov-26", "28-Nov-26", 0, False),
    ("1.17", "A-18 Head room", "29-Nov-26", "9-Dec-26", 0, False),
]


def test_a_block_heading_is_skipped_and_its_activities_read_side_by_side(tenant):
    j = job(tenant)
    out = import_sheet(tenant, j["id"], TWO_LEVEL)
    assert out.status_code == 200, out.text
    body = out.json()
    assert body["warnings"] == []
    assert [r["kind"] for r in body["rows"]] == ["activity", "activity", "activity"]
    assert [r["text"] for r in body["rows"]] == ["A-16 Third floor slab", "A-16 Terrace floor slab", "A-18 Head room"]
    assert body["rows"][0]["start"] == "2026-10-14" and body["rows"][0]["finish"] == "2026-11-12"


def test_confirming_creates_chained_activities_with_no_tasks(tenant):
    j = job(tenant)
    rows = import_sheet(tenant, j["id"], TWO_LEVEL).json()["rows"]
    done = confirm(tenant, j["id"], rows)
    assert done.status_code == 200, done.text
    assert "3 activities with 0 tasks" in done.json()["message"]
    acts = done.json()["schedule"]["activities"]
    assert [a["name"] for a in acts] == ["A-16 Third floor slab", "A-16 Terrace floor slab", "A-18 Head room"]
    assert acts[1]["depends_on_id"] == acts[0]["id"] and acts[2]["depends_on_id"] == acts[1]["id"]


# --- the 3-level sheet: a block, its activities, and each one's own construction stages as tasks ----------------------

THREE_LEVEL = [
    ("1.00", "300 SFT Block Type - A", "14-Oct-26", "13-Feb-27", 0, True),
    ("1.01", "A-16 Third floor slab", "14-Oct-26", "27-Nov-26", 0, False),
    ("1.12", "Reinforcement", "14-Oct-26", "19-Oct-26", 1, False),
    ("1.13", "Shuttering", "19-Oct-26", "23-Nov-26", 1, False),
    ("1.14", "Concreting", "23-Nov-26", "24-Nov-26", 1, False),
    ("1.15", "De-shuttering", "25-Nov-26", "27-Nov-26", 1, False),
    ("1.16", "A-16 Terrace floor slab", "25-Nov-26", "8-Dec-26", 0, False),
    ("1.12", "Reinforcement", "25-Nov-26", "29-Nov-26", 1, False),
    ("1.13", "Shuttering", "29-Nov-26", "4-Dec-26", 1, False),
]


def test_stages_indented_under_an_activity_become_its_tasks(tenant):
    j = job(tenant)
    out = import_sheet(tenant, j["id"], THREE_LEVEL)
    assert out.status_code == 200, out.text
    rows = out.json()["rows"]
    kinds = [(r["kind"], r["text"]) for r in rows]
    assert kinds == [
        ("activity", "A-16 Third floor slab"),
        ("task", "Reinforcement"), ("task", "Shuttering"), ("task", "Concreting"), ("task", "De-shuttering"),
        ("activity", "A-16 Terrace floor slab"), ("task", "Reinforcement"), ("task", "Shuttering"),
    ]
    assert rows[1]["under"] == 0 and rows[6]["under"] == 5


def test_confirming_the_3_level_sheet_makes_real_tickable_tasks(tenant):
    j = job(tenant)
    rows = import_sheet(tenant, j["id"], THREE_LEVEL).json()["rows"]
    done = confirm(tenant, j["id"], rows)
    assert done.status_code == 200 and "2 activities with 6 tasks" in done.json()["message"]
    acts = done.json()["schedule"]["activities"]
    slab1 = [a for a in acts if a["name"] == "A-16 Third floor slab"][0]
    detail = tenant.get("/api/schedule/activities/%d" % slab1["id"]).json()["activity"]
    assert [t["text"] for t in detail["tasks"]] == ["Reinforcement", "Shuttering", "Concreting", "De-shuttering"]
    assert detail["tasks"][0]["date"] == "2026-10-14"
    tenant.put("/api/schedule/tasks/%d" % detail["tasks"][0]["id"], json={"done": True})
    assert tenant.get("/api/schedule/activities/%d" % slab1["id"]).json()["activity"]["actual_percent"] == 25


# --- the 4-level sheet: a block, a unit, its tasks, and a task's own sub-works folded into it --------------------------

FOUR_LEVEL = [
    ("A", "365 SFT Block Type - B", "15-Nov-25", "28-Jan-26", 0, True),
    ("1", "B - 24", "15-Nov-25", "28-Jan-26", 1, False),
    ("1.1", "Cupboard work", "15-Nov-25", "25-Nov-25", 2, False),
    ("1.2", "Internal Painting", "15-Nov-25", "13-Jan-26", 2, False),
    ("", "Internal Putty", "15-Nov-25", "12-Dec-25", 3, False),
    ("", "Primer", "13-Dec-25", "25-Dec-25", 3, False),
    ("", "Internal Paint", "26-Dec-25", "13-Jan-26", 3, False),
    ("1.3", "Plumbing - External piping", "10-Dec-25", "17-Dec-25", 2, False),
]


def test_a_tasks_own_sub_works_are_folded_into_its_text(tenant):
    j = job(tenant)
    out = import_sheet(tenant, j["id"], FOUR_LEVEL)
    assert out.status_code == 200, out.text
    rows = out.json()["rows"]
    assert [(r["kind"], r["text"].split(";")[0]) for r in rows] == [
        ("activity", "B - 24"), ("task", "Cupboard work"), ("task", "Internal Painting"), ("task", "Plumbing - External piping")]
    painting = rows[2]
    assert "Internal Putty (2025-11-15 to 2025-12-12)" in painting["text"]
    assert "Primer (2025-12-13 to 2025-12-25)" in painting["text"]
    assert "Internal Paint (2025-12-26 to 2026-01-13)" in painting["text"]
    assert painting["start"] == "2025-11-15" and painting["finish"] == "2026-01-13", "the task keeps its own, correct dates"


# --- what a bad or thin sheet says -------------------------------------------------------------------------------------

def test_a_row_with_no_readable_date_is_left_out_and_said_so(tenant):
    j = job(tenant)
    out = import_sheet(tenant, j["id"], TWO_LEVEL + [("1.18", "A row with no dates", "", "", 0, False)])
    assert out.status_code == 200
    body = out.json()
    assert len(body["rows"]) == 3 and len(body["warnings"]) == 1 and "no readable start and end" in body["warnings"][0]


def test_a_sheet_with_no_date_columns_is_refused_in_words(tenant):
    j = job(tenant)
    wb = openpyxl.Workbook()
    wb.active.append(["Description", "Quantity"])
    wb.active.append(["Excavation", 500])
    buf = io.BytesIO()
    wb.save(buf)
    out = tenant.post("/api/jobs/%d/schedule/import" % j["id"], files={"file": ("x.xlsx", buf.getvalue())})
    assert out.status_code == 400 and "Start and End date columns" in out.json()["detail"]


def test_confirming_nothing_is_refused(tenant):
    j = job(tenant)
    assert confirm(tenant, j["id"], []).status_code == 400


def test_dates_in_slash_form_are_read_too(tenant):
    j = job(tenant)
    rows = [("", "Internal Back filling in G.F corridor", "4/15/2026", "4/17/2026", 0, False)]
    out = import_sheet(tenant, j["id"], rows)
    assert out.status_code == 200, out.text
    assert out.json()["rows"][0]["start"] == "2026-04-15" and out.json()["rows"][0]["finish"] == "2026-04-17"


def test_one_company_cannot_import_into_anothers_project(tenant, second_tenant):
    theirs = job(second_tenant, "Rival site")
    out = import_sheet(tenant, theirs["id"], TWO_LEVEL)
    assert out.status_code == 404


def test_a_task_numbered_x_point_zero_is_not_mistaken_for_a_block(tenant):
    """2.0 Electrical - Switches & Sockets is an ordinary task continuing the 1.1, 1.2... numbering, not a block
    marker like 1.00 / 2.00 - the single trailing zero must not be read the same as the double one."""
    j = job(tenant)
    rows = [
        ("A", "365 SFT Block Type - B", "15-Nov-25", "28-Jan-26", 0, True),
        ("1", "B - 24", "15-Nov-25", "28-Jan-26", 1, False),
        ("1.9", "Electrical - Wiring", "14-Jan-26", "26-Jan-26", 2, False),
        ("2.0", "Electrical - Switches & Sockets", "16-Jan-26", "24-Jan-26", 2, False),
        ("2.1", "Electrical - Panel board installation", "25-Jan-26", "28-Jan-26", 2, False),
    ]
    out = import_sheet(tenant, j["id"], rows)
    assert out.status_code == 200, out.text
    kinds = [(r["kind"], r["text"]) for r in out.json()["rows"]]
    assert kinds == [
        ("activity", "B - 24"), ("task", "Electrical - Wiring"),
        ("task", "Electrical - Switches & Sockets"), ("task", "Electrical - Panel board installation"),
    ]
