"""Importing items from a sheet (preview first, then commit) and keeping the categories tidy."""
import io

import openpyxl

import main
from app import models
from conftest import make_employee
from test_site_pricing import make_item, site_with

HEAD = ["Item code", "Item name", "What is it?", "Category", "Subcategory", "Unit", "Grade or specification", "Usual purchase rate", "Usual selling rate", "GST %"]


def sheet(rows, head=HEAD):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(head)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def preview(tenant, rows, head=HEAD):
    res = tenant.post("/api/items/import/preview", files={"file": ("items.xlsx", sheet(rows, head))})
    assert res.status_code == 200, res.text
    return res.json()


def by_line(body):
    return {r["line"]: r for r in body["rows"]}


def to_commit(body, overrides=None):
    overrides = overrides or {}
    return {"rows": [{"line": r["line"], "action": overrides.get(r["line"], r["action"]), "values": r["values"]} for r in body["rows"]]}


def test_the_template_downloads_and_reads_back_as_it_is(tenant):
    res = tenant.get("/api/items/import/template")
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    rows = [[c.value for c in r] for r in wb.active.iter_rows()]
    head_row = next(r for r in rows if r and r[0] == "Item code")
    assert "Item name" in head_row and "GST %" in head_row
    again = tenant.post("/api/items/import/preview", files={"file": ("t.xlsx", res.content)}).json()
    assert again["summary"]["error"] == 0 and again["summary"]["new"] == 2


def test_each_row_is_marked_new_update_duplicate_or_error_with_the_reason(tenant):
    make_item(tenant, "OLD-1", "Steel TMT bar 12mm", units_of_measure="Kgs", item_category="Cement")
    body = preview(tenant, [
        ["", "River sand", "Material", "Sand & Aggregate", "", "cum", "Zone II", 1850, "", 5],                    # 2  new, code to be issued
        ["NEW-1", "Brick masonry CM 1:6", "Work item", "Masonry", "", "cum", "", "", 6100, 18],                    # 3  new with its own code
        ["OLD-1", "Steel TMT bar 12mm", "", "", "", "Kgs", "", 68, "", ""],                                       # 4  update: a rate is new
        ["OLD-1", "Steel TMT bar 12mm", "", "", "", "Kgs", "", 68, "", ""],                                       # 5  repeats line 4
        ["", "steel  TMT bar 12MM", "", "", "", "Kgs", "", "", "", ""],                                           # 6  same name as an existing item
        ["", "", "", "", "", "cum", "", "", "", ""],                                                              # 7  no name
        ["BAD-1", "Odd unit", "", "", "", "furlong", "", "", "", ""],                                             # 8  unknown unit
        ["BAD-2", "Odd category", "", "Not set up", "", "cum", "", "", "", ""],                                   # 9
        ["BAD-3", "Negative", "", "", "", "cum", "", -5, "", ""],                                                 # 10
        ["BAD-4", "Odd gst", "", "", "", "cum", "", "", "", 7],                                                   # 11
        ["HAS SPACE", "Spaced code", "", "", "", "cum", "", "", "", ""],                                          # 12
    ])
    got = by_line(body)
    assert got[2]["status"] == "New" and "next free code" in got[2]["reason"] and got[2]["values"]["kind"] == "RM"
    assert got[3]["status"] == "New" and got[3]["values"]["kind"] == "FG"
    assert got[4]["status"] == "Update" and "purchase_rate" in got[4]["changes"] and "Steel TMT bar 12mm" in got[4]["reason"]
    assert got[5]["status"] == "Duplicate" and "Line 4" in got[5]["reason"] and got[5]["action"] == "skip"
    assert got[6]["status"] == "Duplicate" and "OLD-1" in got[6]["reason"] and got[6]["action"] == "skip"
    assert "no name" in got[7]["reason"] and got[7]["status"] == "Error"
    assert "furlong" in got[8]["reason"] and "not a unit" in got[8]["reason"]
    assert "not set up" in got[9]["reason"]
    assert "0 or more" in got[10]["reason"]
    assert "0, 5, 12, 18 or 28" in got[11]["reason"]
    assert "space" in got[12]["reason"]
    assert body["summary"] == {"total": 11, "new": 2, "update": 1, "duplicate": 2, "error": 6}
    with main.SessionLocal() as db:
        assert db.query(models.DBItem).count() == 1, "a preview saves nothing"


def test_commit_adds_updates_and_skips_and_lists_what_it_left_out(tenant):
    make_item(tenant, "OLD-1", "Steel TMT bar 12mm", units_of_measure="Kgs", item_category="Cement", specification="Fe550D", remarks="keep me")
    body = preview(tenant, [
        ["", "River sand", "Material", "Sand & Aggregate", "", "cum", "Zone II", 1850, "", 5],
        ["OLD-1", "Steel TMT bar 12mm", "", "", "", "Kgs", "", 68, "", ""],
        ["BAD-1", "Odd unit", "", "", "", "furlong", "", "", "", ""],
        ["OLD-1", "Steel TMT bar 12mm", "", "", "", "Kgs", "", 68, "", ""],
    ])
    done = tenant.post("/api/items/import/commit", json=to_commit(body)).json()
    assert (done["created"], done["updated"], done["skipped"]) == (1, 1, 2)
    assert {e["line"] for e in done["errors"]} == {4, 5} and "furlong" in done["errors"][0]["reason"]
    items = {i["item_name"]: i for i in tenant.get("/api/items", params={"status": "all"}).json()["items"]}
    assert items["River sand"]["item_code"] and items["River sand"]["purchase_rate"] == 1850 and items["River sand"]["item_tax_type"] == "5%"
    steel = items["Steel TMT bar 12mm"]
    assert steel["purchase_rate"] == 68
    assert steel["specification"] == "Fe550D" and steel["remarks"] == "keep me" and steel["item_category"] == "Cement", "a blank cell never wipes what is held"


def test_importing_the_same_sheet_twice_does_not_make_the_items_twice(tenant):
    rows = [["PCC-M15", "PCC M15 in foundation", "Work item", "Concrete (PCC / RCC)", "", "cum", "M15", "", "", 18],
            ["RCC-M25", "RCC M25 in raft", "Work item", "Concrete (PCC / RCC)", "", "cum", "M25", "", "", 18]]
    first = tenant.post("/api/items/import/commit", json=to_commit(preview(tenant, rows))).json()
    assert first["created"] == 2
    again_preview = preview(tenant, rows)
    assert again_preview["summary"]["duplicate"] == 2 and again_preview["summary"]["new"] == 0
    second = tenant.post("/api/items/import/commit", json=to_commit(again_preview)).json()
    assert (second["created"], second["updated"]) == (0, 0)
    with main.SessionLocal() as db:
        assert db.query(models.DBItem).filter(models.DBItem.item_code.in_(["PCC-M15", "RCC-M25"])).count() == 2


def test_a_row_fixed_on_screen_is_checked_again_and_a_forged_status_is_not_believed(tenant):
    make_item(tenant, "TAKEN-1", "Existing")
    body = preview(tenant, [["TAKEN-1", "Existing", "", "", "", "cum", "", "", "", ""], ["X-1", "Odd unit", "", "", "", "furlong", "", "", "", ""]])
    fixed = to_commit(body)
    fixed["rows"][1]["values"]["units_of_measure"] = "cum"
    checked = tenant.post("/api/items/import/check", json=fixed).json()
    assert [r["status"] for r in checked["rows"]] == ["Duplicate", "New"]
    forged = {"rows": [{"line": 2, "action": "import", "values": {"item_code": "F-1", "item_name": "Bad", "units_of_measure": "furlong"}}]}
    done = tenant.post("/api/items/import/commit", json=forged).json()
    assert done["created"] == 0 and done["skipped"] == 1


def test_a_same_name_duplicate_can_be_added_anyway_when_the_person_says_so(tenant):
    make_item(tenant, "SAME-1", "Plywood 19mm", units_of_measure="Nos")
    body = preview(tenant, [["", "Plywood 19mm", "", "", "", "Nos", "", "", "", ""]])
    assert body["rows"][0]["status"] == "Duplicate" and body["rows"][0]["action"] == "skip"
    skipped = tenant.post("/api/items/import/commit", json=to_commit(body)).json()
    assert skipped["created"] == 0
    forced = tenant.post("/api/items/import/commit", json=to_commit(body, {2: "import"})).json()
    assert forced["created"] == 1


def test_a_sheet_with_no_name_column_is_refused_in_words(tenant):
    res = tenant.post("/api/items/import/preview", files={"file": ("x.xlsx", sheet([["a"]], head=["Colour"]))})
    assert res.status_code == 400 and "Item name" in res.json()["detail"]


def test_only_those_who_manage_items_may_import(tenant, portal):
    emp = make_employee(tenant, password="EmpPass123")
    tenant.put("/api/employees/%d/permissions" % emp["id"], json={"permission_role": "staff", "permissions": ["self.service"]})
    portal.post("/api/employee/auth/login", json={"email": emp["email"], "password": "EmpPass123"})
    assert portal.post("/api/items/import/preview", files={"file": ("x.xlsx", sheet([]))}).status_code == 403
    assert portal.post("/api/items/import/commit", json={"rows": [{"line": 2, "values": {}}]}).status_code == 403
    assert portal.get("/api/items/import/template").status_code == 200


# --- categories ----------------------------------------------------------------------------------------------------------

def cats(tenant):
    return {c["name"]: c for c in tenant.get("/api/item-categories").json()["categories"]}


def test_categories_show_how_many_items_each_holds(tenant):
    make_item(tenant, "C-1", "One", item_category="Cement")
    make_item(tenant, "C-2", "Two", item_category="Cement")
    make_item(tenant, "S-1", "Three", item_category="Plastering")
    got = cats(tenant)
    assert (got["Cement"]["items"], got["Plastering"]["items"], got["Masonry"]["items"]) == (2, 1, 0)


def test_merging_moves_the_items_and_the_subcategories_and_removes_the_old_one(tenant):
    tenant.post("/api/item-categories", json={"name": "Glazing"})
    glazing = cats(tenant)["Glazing"]
    tenant.post("/api/item-categories", json={"name": "Toughened", "parent_id": glazing["id"]})
    make_item(tenant, "G-1", "Glass", item_category="Glazing", item_subcategory="Toughened")
    make_item(tenant, "O-1", "Other", item_category="Other BOQ Work")
    target = cats(tenant)["Other BOQ Work"]
    done = tenant.post("/api/item-categories/%d/merge" % glazing["id"], json={"into_id": target["id"]})
    assert done.status_code == 200 and done.json()["moved"] == 1
    assert "Glazing" not in cats(tenant)
    glass = tenant.get("/api/items", params={"q": "Glass"}).json()["items"][0]
    assert (glass["item_category"], glass["item_subcategory"]) == ("Other BOQ Work", "Toughened")
    assert [s["name"] for s in cats(tenant)["Other BOQ Work"]["subcategories"]] == ["Toughened"]
    assert cats(tenant)["Other BOQ Work"]["items"] == 2


def test_merge_rules(tenant):
    cement, steel = cats(tenant)["Cement"], cats(tenant)["Steel & Reinforcement"]
    assert tenant.post("/api/item-categories/%d/merge" % cement["id"], json={"into_id": cement["id"]}).status_code == 422
    assert tenant.post("/api/item-categories/%d/merge" % cement["id"], json={"into_id": 999999}).status_code == 404
    tenant.post("/api/item-categories", json={"name": "OPC", "parent_id": cement["id"]})
    sub = cats(tenant)["Cement"]["subcategories"][0]
    assert tenant.post("/api/item-categories/%d/merge" % sub["id"], json={"into_id": steel["id"]}).status_code == 422, "a subcategory does not merge into a category"


def test_a_category_that_holds_items_cannot_be_deleted_until_they_move(tenant):
    make_item(tenant, "D-1", "Held", item_category="Painting")
    painting = cats(tenant)["Painting"]
    refused = tenant.delete("/api/item-categories/%d" % painting["id"])
    assert refused.status_code == 409 and "1 item is filed under 'Painting'" in refused.json()["detail"] and "Merge" in refused.json()["detail"]
    assert "Painting" in cats(tenant)
    other = cats(tenant)["Plumbing"]
    tenant.post("/api/item-categories/%d/merge" % painting["id"], json={"into_id": other["id"]})
    assert "Painting" not in cats(tenant)
    empty = cats(tenant)["Electrical"]
    assert tenant.delete("/api/item-categories/%d" % empty["id"]).status_code == 200
    assert "Electrical" not in cats(tenant)


def test_deleting_a_category_with_empty_subcategories_removes_them_too(tenant):
    tenant.post("/api/item-categories", json={"name": "Temp"})
    temp = cats(tenant)["Temp"]
    tenant.post("/api/item-categories", json={"name": "Sub", "parent_id": temp["id"]})
    assert tenant.delete("/api/item-categories/%d" % temp["id"]).status_code == 200
    with main.SessionLocal() as db:
        assert db.query(models.DBItemCategory).filter(models.DBItemCategory.name.in_(["Temp", "Sub"])).count() == 0


def test_category_changes_are_for_those_who_manage_items_and_only_for_their_own_company(tenant, second_tenant, portal):
    cement = cats(tenant)["Cement"]
    assert second_tenant.delete("/api/item-categories/%d" % cement["id"]).status_code == 404
    emp = make_employee(tenant, password="EmpPass123")
    tenant.put("/api/employees/%d/permissions" % emp["id"], json={"permission_role": "staff", "permissions": ["self.service"]})
    portal.post("/api/employee/auth/login", json={"email": emp["email"], "password": "EmpPass123"})
    assert portal.delete("/api/item-categories/%d" % cement["id"]).status_code == 403
    assert portal.post("/api/item-categories/%d/merge" % cement["id"], json={"into_id": cement["id"]}).status_code == 403
    assert site_with  # the shared helper module stays imported for the suites that follow
