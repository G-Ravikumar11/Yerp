"""A payslip downloaded as PDF or Excel carries every column of the office's sheet."""
import io

from conftest import make_employee

HEADS = ["Employee No.", "Location .", "Name.", "Department.", "Bank Name.", "Designation.", "Bank Acc.No.", "Total Days.",
         "Father Name.", "LOP .", "Salary Per Month.", "Balance Leaves.", "DOJ.", "Work Days.", "Aadhar No.",
         "Effective Work days.", "UAN No", "ESI No", "Earnings", "Deductions",
         "Basic", "HRA", "Site Allowance", "Special Allownace", "Mobile Allowance", "Transport Allowance", "Other Allowance",
         "PF", "ESI", "Prof Tax", "Salary Advance", "TDS Deduction", "Loss Of Pay", "Other Deductions",
         "Total Earnings : Rs.", "Total Deductions : Rs.", "Net Pay :", "In Words :"]


def _payslip(tenant):
    emp = make_employee(tenant, job_title="Project Manager", salary=65000)
    res = tenant.post("/api/payslips", json={
        "employee_id": emp["id"], "period_start": "2026-09-01", "period_end": "2026-09-30", "pay_date": "2026-09-30",
        "basic_salary": 60000, "allowances": 5000, "retirement": 3000})
    assert res.status_code == 200, res.text
    return res.json()["id"]


def test_the_excel_download_has_every_column_and_the_figures(tenant):
    import openpyxl
    pid = _payslip(tenant)
    slip = tenant.get("/api/payslips/%d" % pid).json()
    res = tenant.get("/api/payslips/%d/xlsx" % pid)
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"] and ".xlsx" in res.headers["content-disposition"]
    sheet = openpyxl.load_workbook(io.BytesIO(res.content)).active
    text = [str(c.value) for row in sheet.iter_rows() for c in row if c.value not in (None, "")]
    for head in HEADS:
        assert head in text, head
    assert "Payslip For The Month of Sep-26" in text
    figures = [c.value for row in sheet.iter_rows() for c in row]
    assert 65000.0 in figures and float(slip["net_pay"]) in figures and float(slip["total_deductions"]) in figures
    assert any(t.startswith("Rupees") and t.endswith("Only") for t in text)


def test_the_pdf_download_has_every_column(tenant):
    res = tenant.get("/api/payslips/%d/pdf" % _payslip(tenant))
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf" and res.content[:4] == b"%PDF"
    from pypdf import PdfReader
    text = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(res.content)).pages)
    for head in HEADS:
        assert head.strip(" :") in text.replace("\n", " "), head


def test_a_downloads_belongs_to_the_company_that_owns_it(tenant, second_tenant):
    pid = _payslip(tenant)
    assert second_tenant.get("/api/payslips/%d/pdf" % pid).status_code == 404
    assert second_tenant.get("/api/payslips/%d/xlsx" % pid).status_code == 404


def _two_payslips(tenant):
    ids = [_payslip(tenant)]
    other = make_employee(tenant, job_title="Engineer", salary=40000)
    res = tenant.post("/api/payslips", json={
        "employee_id": other["id"], "period_start": "2026-09-01", "period_end": "2026-09-30", "pay_date": "2026-09-30",
        "basic_salary": 40000, "allowances": 2000})
    assert res.status_code == 200, res.text
    return ids + [res.json()["id"]]


def test_the_register_has_a_row_for_each_payslip_and_every_column(tenant):
    import openpyxl
    from app.documents.payslip import REGISTER_COLUMNS
    ids = _two_payslips(tenant)
    res = tenant.get("/api/payslips/register.xlsx", params={"ids": ",".join(map(str, ids))})
    assert res.status_code == 200 and "Payroll_register_2026-09" in res.headers["content-disposition"]
    sheet = openpyxl.load_workbook(io.BytesIO(res.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    header = next(i for i, r in enumerate(rows) if r and r[0] == "Payslip No.")
    assert list(rows[header]) == REGISTER_COLUMNS
    body = rows[header + 1:-1]
    assert len(body) == 2 and rows[-1][0] == "Total"
    net = REGISTER_COLUMNS.index("Net Pay")
    slips = [tenant.get("/api/payslips/%d" % i).json() for i in ids]
    assert round(sum(r[net] for r in body), 2) == round(sum(s["net_pay"] for s in slips), 2) == rows[-1][net]


def test_the_register_takes_only_the_payslips_asked_for_and_only_ours(tenant, second_tenant):
    import openpyxl
    a, b = _two_payslips(tenant)
    res = tenant.get("/api/payslips/register.xlsx", params={"ids": str(a)})
    sheet = openpyxl.load_workbook(io.BytesIO(res.content)).active
    assert len([r for r in sheet.iter_rows(values_only=True) if r[0] and str(r[0]).startswith("PS-")]) == 1
    assert second_tenant.get("/api/payslips/register.xlsx", params={"ids": "%d,%d" % (a, b)}).status_code == 404
    assert tenant.get("/api/payslips/register.xlsx", params={"ids": "x"}).status_code == 400
